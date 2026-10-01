# -*- coding: utf-8 -*-
"""
Встроенный редактор Word (.docx): правка текста прямо в LegalHelper, без отдельного Word.

Как устроено, чтобы ничего не терялось:
* абзацы и таблицы, которые не меняли, при сохранении записываются обратно байт в байт (вместе с картинками,
  нумерацией, сносками, полями и закладками);
* изменённый абзац собирается заново, но с прежними свойствами абзаца (отступы, интервалы, нумерация, стиль)
  и прежними свойствами каждого куска текста; картинки, поля и сноски внутри него — «значки» ■, которые
  сохраняются как были, если их не удалить;
* в таблицах правится текст ячеек (таблицы с объединёнными ячейками — только в Word);
* поля страницы, колонтитулы и стили документа не трогаются;
* перед каждым сохранением прежняя версия файла кладётся в резервные копии.
"""
import copy
import datetime as dt
import os
import shutil

from PySide6.QtCore import Qt, QTimer
from PySide6.QtGui import (QAction, QColor, QFont, QKeySequence, QTextBlockFormat, QTextCharFormat, QTextCursor,
                           QTextDocument, QTextFormat, QTextTableFormat, QTextFrameFormat, QTextLength)
from PySide6.QtWidgets import (QComboBox, QDialog, QFontComboBox, QHBoxLayout, QInputDialog, QLabel,
                               QMessageBox, QTextEdit, QToolBar, QVBoxLayout, QWidget)

from docx import Document
from docx.oxml.ns import qn

P_ORIG = QTextFormat.UserProperty + 1      # блок: номер исходного элемента (абзаца) документа
P_AUTO = QTextFormat.UserProperty + 2      # блок, который Qt добавил сам (рядом с таблицей), — не абзац Word
T_ORIG = QTextFormat.UserProperty + 3      # таблица: номер исходной таблицы
R_RPR = QTextFormat.UserProperty + 4       # кусок текста: номер исходных свойств текста (rPr)
R_OBJ = QTextFormat.UserProperty + 5       # «значок» ■: номер исходного объекта (картинка, поле, сноска…)
R_B0 = QTextFormat.UserProperty + 6        # как было при открытии: жирный/курсив/подчёркнутый/размер/шрифт
R_I0 = QTextFormat.UserProperty + 7
R_U0 = QTextFormat.UserProperty + 8
R_S0 = QTextFormat.UserProperty + 9
R_F0 = QTextFormat.UserProperty + 10

OBJ = "■"                                  # так в тексте выглядит картинка, поле или сноска
LINE = " "                            # перенос строки внутри абзаца (Shift+Enter)
OBJ_RUN_TAGS = {qn(t) for t in ("w:drawing", "w:pict", "w:object", "w:fldChar", "w:instrText", "w:footnoteReference",
                                "w:endnoteReference", "w:sym", "w:ptab", "w:commentReference", "w:ruby")}
OBJ_PAR_TAGS = {qn(t) for t in ("w:fldSimple", "w:sdt", "w:smartTag", "w:ins", "w:del", "w:moveTo", "w:moveFrom",
                                "w:customXml")} | {"{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath",
                                                   "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathPara"}
SKIP_PAR_TAGS = {qn(t) for t in ("w:pPr", "w:bookmarkStart", "w:bookmarkEnd", "w:proofErr", "w:permStart",
                                 "w:permEnd", "w:commentRangeStart", "w:commentRangeEnd")}
ALIGN = {"left": Qt.AlignLeft, "start": Qt.AlignLeft, "center": Qt.AlignHCenter, "right": Qt.AlignRight,
         "end": Qt.AlignRight, "both": Qt.AlignJustify, "distribute": Qt.AlignJustify}
PX = 96 / 72.0                             # пунктов → точек экрана
DEFAULT_FONT = "Times New Roman"           # шрифт юридических документов: везде, где документ не задал другой


def _iv(x):
    """Значение перечисления Qt (выравнивание, жирность) — обычным числом."""
    return int(getattr(x, "value", x))


BOLD_MIN = _iv(QFont.DemiBold)
ALIGN_BACK = {_iv(Qt.AlignLeft): "left", _iv(Qt.AlignHCenter): "center", _iv(Qt.AlignRight): "right",
              _iv(Qt.AlignJustify): "both"}


def _w(el, tag):
    return el.find(qn(tag))


def _val(el, attr="w:val"):
    return None if el is None else el.get(qn(attr))


def _onoff(el):
    """<w:b/> → True, <w:b w:val="0"/> → False, нет элемента → None."""
    if el is None:
        return None
    v = _val(el)
    return v not in ("0", "false", "off")


def backups_dir():
    base = os.environ.get("LOCALAPPDATA") or os.path.join(os.path.expanduser("~"), ".cache")
    d = os.path.join(base, "PDFMaster", "word_backups")
    os.makedirs(d, exist_ok=True)
    return d


def backup(path, keep=60):
    """Прежнюю версию файла — в резервные копии (последние keep штук)."""
    try:
        stamp = dt.datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        stem, ext = os.path.splitext(os.path.basename(path))
        shutil.copy2(path, os.path.join(backups_dir(), f"{stem} — {stamp}{ext}"))
        files = sorted((os.path.join(backups_dir(), f) for f in os.listdir(backups_dir())), key=os.path.getmtime)
        for f in files[:-keep]:
            os.remove(f)
    except OSError:
        pass


class DocxModel:
    """Загрузка .docx в QTextDocument и сборка .docx обратно из QTextDocument."""

    def __init__(self, path):
        self.path = path
        self.doc = Document(path)
        self.body = self.doc.element.body
        self.orig = []           # копии исходных элементов тела (абзацы, таблицы, прочее)
        self.sig = {}            # номер элемента → «отпечаток» при открытии (для проверки «не меняли»)
        self.rprs = []           # исходные свойства кусков текста
        self.objs = []           # исходные объекты внутри абзацев (картинки, поля…)
        self.complex_tables = 0
        st = self.doc.styles
        self.theme_fonts = self._read_theme_fonts()
        try:
            nf = st["Normal"].font
            self.base_family = nf.name or DEFAULT_FONT
            self.base_size = nf.size.pt if nf.size else 14.0
        except KeyError:
            self.base_family, self.base_size = DEFAULT_FONT, 14.0
        defaults = self.doc.styles.element.find(qn("w:docDefaults"))
        if defaults is not None:
            rpr = defaults.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
            if rpr is not None:
                fonts = rpr.find(qn("w:rFonts"))
                if fonts is not None and not self.doc.styles["Normal"].font.name:
                    self.base_family = self._font_name(fonts) or self.base_family
                sz = rpr.find(qn("w:sz"))
                if sz is not None and not self.doc.styles["Normal"].font.size:
                    try:
                        self.base_size = int(_val(sz)) / 2
                    except (TypeError, ValueError):
                        pass

    def _read_theme_fonts(self):
        """Шрифты темы оформления (Word пишет не «Times New Roman», а «основной шрифт темы»)."""
        out = {}
        try:
            for part in self.doc.part.package.iter_parts():
                if "/theme/" in str(part.partname):
                    import re
                    xml = part.blob.decode("utf-8", "replace")
                    for kind in ("major", "minor"):
                        m = re.search(rf"<a:{kind}Font>\s*<a:latin typeface=\"([^\"]*)\"", xml)
                        if m and m.group(1):
                            out[kind] = m.group(1)
                    break
        except Exception:
            pass
        return out

    def _font_name(self, rfonts):
        """Имя шрифта из <w:rFonts>: прямое или из темы; None — не задано."""
        if rfonts is None:
            return None
        name = rfonts.get(qn("w:ascii")) or rfonts.get(qn("w:hAnsi")) or rfonts.get(qn("w:cs"))
        if name:
            return name
        theme = rfonts.get(qn("w:asciiTheme")) or rfonts.get(qn("w:hAnsiTheme"))
        if theme:
            return self.theme_fonts.get("major" if theme.startswith("major") else "minor")
        return None

    # ------------------------------------------------------------------ свойства из Word
    def _style_font(self, style_id, kind):
        """Жирность/курсив/размер/шрифт из стиля абзаца (с наследованием)."""
        seen = 0
        st = self.doc.styles.element.get_by_id(style_id) if style_id else None
        while st is not None and seen < 10:
            rpr = st.find(qn("w:rPr"))
            if rpr is not None:
                if kind in ("b", "i", "u"):
                    el = rpr.find(qn("w:" + kind))
                    if el is not None:
                        return (_val(el) not in ("none", "0", "false")) if kind == "u" else _onoff(el)
                elif kind == "sz":
                    el = rpr.find(qn("w:sz"))
                    if el is not None and _val(el):
                        return int(_val(el)) / 2
                elif kind == "font":
                    name = self._font_name(rpr.find(qn("w:rFonts")))
                    if name:
                        return name
            based = st.find(qn("w:basedOn"))
            st = self.doc.styles.element.get_by_id(_val(based)) if based is not None else None
            seen += 1
        return None

    def _run_format(self, rpr, pstyle):
        def get(kind):
            if rpr is not None:
                el = rpr.find(qn("w:" + kind))
                if el is not None:
                    return (_val(el) not in ("none", "0", "false")) if kind == "u" else _onoff(el)
            return self._style_font(pstyle, kind)
        b, i, u = bool(get("b")), bool(get("i")), bool(get("u"))
        size, family = None, None
        if rpr is not None:
            sz = rpr.find(qn("w:sz"))
            if sz is not None and _val(sz):
                try:
                    size = int(_val(sz)) / 2
                except ValueError:
                    pass
            family = self._font_name(rpr.find(qn("w:rFonts")))
        size = size or self._style_font(pstyle, "sz") or self.base_size
        family = family or self._style_font(pstyle, "font") or self.base_family
        return b, i, u, float(size), family

    def _char_fmt(self, b, i, u, size, family, rpr_idx):
        cf = QTextCharFormat()
        cf.setFontWeight(QFont.Bold if b else QFont.Normal)
        cf.setFontItalic(i)
        cf.setFontUnderline(u)
        cf.setFontPointSize(size)
        cf.setFontFamilies([family, DEFAULT_FONT] if family != DEFAULT_FONT else [family])
        cf.setProperty(R_RPR, rpr_idx)
        cf.setProperty(R_B0, b)
        cf.setProperty(R_I0, i)
        cf.setProperty(R_U0, u)
        cf.setProperty(R_S0, size)
        cf.setProperty(R_F0, family)
        return cf

    def _block_fmt(self, p_el, idx):
        bf = QTextBlockFormat()
        bf.setProperty(P_ORIG, idx)
        ppr = p_el.find(qn("w:pPr"))
        pstyle = _val(ppr.find(qn("w:pStyle"))) if ppr is not None and ppr.find(qn("w:pStyle")) is not None else None
        if ppr is not None:
            jc = ppr.find(qn("w:jc"))
            if jc is not None and _val(jc) in ALIGN:
                bf.setAlignment(ALIGN[_val(jc)])
            ind = ppr.find(qn("w:ind"))
            if ind is not None:
                def tw(a):
                    v = ind.get(qn("w:" + a))
                    try:
                        return int(v) / 20.0 * PX if v else None
                    except ValueError:
                        return None
                left = tw("left") or tw("start")
                if left:
                    bf.setLeftMargin(left)
                right = tw("right") or tw("end")
                if right:
                    bf.setRightMargin(right)
                first, hang = tw("firstLine"), tw("hanging")
                if first:
                    bf.setTextIndent(first)
                elif hang:
                    bf.setTextIndent(-hang)
            sp = ppr.find(qn("w:spacing"))
            if sp is not None:
                for attr, setter in (("before", bf.setTopMargin), ("after", bf.setBottomMargin)):
                    v = sp.get(qn("w:" + attr))
                    if v and v.lstrip("-").isdigit():
                        setter(int(v) / 20.0 * PX)
                line, rule = sp.get(qn("w:line")), sp.get(qn("w:lineRule")) or "auto"
                if line and line.isdigit() and rule == "auto":
                    bf.setLineHeight(int(line) / 240.0 * 100, 1)        # 1 = ProportionalHeight
        return bf, pstyle

    # ------------------------------------------------------------------ загрузка
    def fill(self, qdoc):
        qdoc.clear()
        qdoc.setDefaultFont(QFont(self.base_family, int(round(self.base_size))))
        cur = QTextCursor(qdoc)
        state = {"first": True}            # первый блок документа уже есть — его не добавляем, а используем
        for child in list(self.body.iterchildren()):
            tag = child.tag
            if tag == qn("w:sectPr"):
                continue
            idx = len(self.orig)
            self.orig.append(copy.deepcopy(child))
            if tag == qn("w:p"):
                self._put_par(cur, child, idx, state)
            elif tag == qn("w:tbl") and self._simple_table(child):
                self._put_table(qdoc, cur, child, idx, state)
            else:                           # прочее (содержимое-элемент, сложная таблица…) — одним значком
                if tag == qn("w:tbl"):
                    self.complex_tables += 1
                self._put_opaque(cur, idx, state, "▦ Таблица со сложной структурой — правится в Word"
                                 if tag == qn("w:tbl") else "▦ Особый элемент Word — правится в Word")
        # отпечатки «как было» — для проверки при сохранении
        for kind, idx, sig, _ref in self.signatures(qdoc):
            self.sig.setdefault((kind, idx), sig)

    def _new_block(self, cur, bf, state):
        if state["first"]:
            cur.setBlockFormat(bf)
            state["first"] = False
        else:
            cur.insertBlock(bf)

    def _put_par(self, cur, p_el, idx, state):
        bf, pstyle = self._block_fmt(p_el, idx)
        self._new_block(cur, bf, state)
        cur.setBlockCharFormat(self._char_fmt(*self._run_format(None, pstyle), -1))
        for ch in p_el:
            t = ch.tag
            if t in SKIP_PAR_TAGS:
                continue
            if t == qn("w:r"):
                self._put_run(cur, ch, pstyle)
            elif t == qn("w:hyperlink"):
                for r in ch.iter(qn("w:r")):
                    self._put_run(cur, r, pstyle)
            elif t in OBJ_PAR_TAGS:
                self._put_obj(cur, ch, pstyle)

    def _put_run(self, cur, r, pstyle):
        rpr = r.find(qn("w:rPr"))
        if any(c.tag in OBJ_RUN_TAGS for c in r):
            return self._put_obj(cur, r, pstyle)
        self.rprs.append(copy.deepcopy(rpr) if rpr is not None else None)
        cf = self._char_fmt(*self._run_format(rpr, pstyle), len(self.rprs) - 1)
        text = []
        for c in r:
            t = c.tag
            if t == qn("w:t"):
                text.append(c.text or "")
            elif t == qn("w:tab"):
                text.append("\t")
            elif t in (qn("w:br"), qn("w:cr")):
                text.append(LINE)
            elif t == qn("w:noBreakHyphen"):
                text.append("-")
        if text:
            cur.insertText("".join(text), cf)

    def _put_obj(self, cur, el, pstyle):
        self.objs.append(copy.deepcopy(el))
        rpr = el.find(qn("w:rPr")) if el.tag == qn("w:r") else None
        cf = self._char_fmt(*self._run_format(rpr, pstyle), -1)
        cf.setProperty(R_OBJ, len(self.objs) - 1)
        cf.setForeground(QColor("#8e8e93"))
        cf.setToolTip("Картинка, поле или сноска из Word — сохранится как есть. Удалите значок, чтобы убрать её.")
        cur.insertText(OBJ, cf)

    def _put_opaque(self, cur, idx, state, text):
        bf = QTextBlockFormat()
        bf.setProperty(P_ORIG, idx)
        self._new_block(cur, bf, state)
        cf = QTextCharFormat()
        cf.setForeground(QColor("#8e8e93"))
        cf.setFontItalic(True)
        cf.setProperty(R_OBJ, -2)            # особый: весь элемент целиком
        cur.insertText(text, cf)

    @staticmethod
    def _simple_table(tbl):
        """Таблица без объединённых ячеек и вложенных таблиц — её текст можно править здесь."""
        rows = tbl.findall(qn("w:tr"))
        if not rows:
            return False
        ncols = None
        for tr in rows:
            tcs = tr.findall(qn("w:tc"))
            if ncols is None:
                ncols = len(tcs)
            if len(tcs) != ncols or ncols == 0:
                return False
            for tc in tcs:
                pr = tc.find(qn("w:tcPr"))
                if pr is not None and (pr.find(qn("w:gridSpan")) is not None or pr.find(qn("w:vMerge")) is not None
                                       or pr.find(qn("w:hMerge")) is not None):
                    return False
                if tc.find(".//" + qn("w:tbl")) is not None:
                    return False
        return True

    def _put_table(self, qdoc, cur, tbl, idx, state):
        rows = tbl.findall(qn("w:tr"))
        ncols = len(rows[0].findall(qn("w:tc")))
        if state["first"]:                  # таблица в самом начале: перед ней Qt оставит пустой служебный блок
            bf = QTextBlockFormat()
            bf.setProperty(P_AUTO, True)
            cur.setBlockFormat(bf)
            state["first"] = False
        tf = QTextTableFormat()
        tf.setBorder(0.6)
        tf.setCellPadding(4)
        tf.setCellSpacing(0)
        tf.setBorderCollapse(True)
        tf.setWidth(QTextLength(QTextLength.PercentageLength, 100))
        tf.setProperty(T_ORIG, idx)
        table = cur.insertTable(len(rows), ncols, tf)
        for r, tr in enumerate(rows):
            for c, tc in enumerate(tr.findall(qn("w:tc"))):
                cc = table.cellAt(r, c).firstCursorPosition()
                first = True
                for p in tc.findall(qn("w:p")):
                    if not first:
                        cc.insertBlock()
                    first = False
                    pstyle = None
                    ppr = p.find(qn("w:pPr"))
                    if ppr is not None and ppr.find(qn("w:jc")) is not None and _val(ppr.find(qn("w:jc"))) in ALIGN:
                        bf = cc.blockFormat()
                        bf.setAlignment(ALIGN[_val(ppr.find(qn("w:jc")))])
                        cc.setBlockFormat(bf)
                    for run in p.iter(qn("w:r")):
                        rpr = run.find(qn("w:rPr"))
                        b, i, u, size, fam = self._run_format(rpr, pstyle)
                        cf = self._char_fmt(b, i, u, size, fam, -1)
                        txt = "".join((x.text or "") if x.tag == qn("w:t") else "\t" if x.tag == qn("w:tab")
                                      else LINE if x.tag in (qn("w:br"), qn("w:cr")) else "" for x in run)
                        if txt:
                            cc.insertText(txt, cf)
        # после таблицы Qt сам добавляет пустой блок — следующий абзац займёт его
        cur.movePosition(QTextCursor.End)
        bf = QTextBlockFormat()
        bf.setProperty(P_AUTO, True)
        cur.setBlockFormat(bf)
        state["first"] = True

    # ------------------------------------------------------------------ отпечатки
    @staticmethod
    def _block_sig(block):
        bf = block.blockFormat()
        frags = []
        it = block.begin()
        while not it.atEnd():
            f = it.fragment()
            if f.isValid():
                cf = f.charFormat()
                frags.append((f.text(), cf.fontWeight() >= BOLD_MIN, cf.fontItalic(), cf.fontUnderline(),
                              round(cf.fontPointSize(), 1), tuple(cf.fontFamilies() or []),
                              cf.property(R_OBJ), cf.property(R_RPR)))
            it += 1
        return (_iv(bf.alignment()) & 0x0F, round(bf.leftMargin(), 1), round(bf.textIndent(), 1), tuple(frags))

    @staticmethod
    def _table_cells(table):
        out = []
        for r in range(table.rows()):
            row = []
            for c in range(table.columns()):
                cell = table.cellAt(r, c)
                cur = QTextCursor(table.document())
                cur.setPosition(cell.firstPosition())
                cur.setPosition(cell.lastPosition(), QTextCursor.KeepAnchor)
                row.append(cur.selectedText().replace(" ", "\n"))
            out.append(tuple(row))
        return tuple(out)

    def signatures(self, qdoc):
        """[(вид, номер исходного элемента, отпечаток)] по порядку документа."""
        out, seen_tables = [], set()
        block = qdoc.begin()
        while block.isValid():
            cur = QTextCursor(block)
            table = cur.currentTable()
            if table is not None:
                key = table.firstPosition()
                if key not in seen_tables:
                    seen_tables.add(key)
                    idx = table.format().property(T_ORIG)
                    out.append(("t", idx if idx is not None else -1, self._table_cells(table), table))
                block = qdoc.findBlock(table.lastPosition() + 1)
                continue
            bf = block.blockFormat()
            idx = bf.property(P_ORIG)
            if bf.property(P_AUTO) and not block.text():
                out.append(("auto", -1, None, block))
            else:
                out.append(("p", idx if idx is not None else -1, self._block_sig(block), block))
            block = block.next()
        return out

    # ------------------------------------------------------------------ сохранение
    def build(self, qdoc):
        """Собрать новое тело документа из текста в редакторе. Возвращает число изменённых элементов."""
        items = self.signatures(qdoc)
        new, changed, used = [], 0, set()
        for kind, idx, sig, ref in items:
            if kind == "auto":
                continue
            if kind == "t":
                orig = self.orig[idx] if 0 <= idx < len(self.orig) else None
                if orig is not None and self.sig.get(("t", idx)) == sig and idx not in used:
                    new.append(copy.deepcopy(orig))
                else:
                    new.append(self._rebuild_table(orig, sig))
                    changed += 1
                used.add(idx)
                continue
            orig = self.orig[idx] if 0 <= idx < len(self.orig) else None
            block = ref
            if orig is not None and orig.tag != qn("w:p"):        # особый элемент целиком
                if idx not in used and block.text():
                    new.append(copy.deepcopy(orig))
                used.add(idx)
                continue
            if orig is not None and self.sig.get(("p", idx)) == sig and idx not in used:
                new.append(copy.deepcopy(orig))
            else:
                new.append(self._rebuild_par(orig, block, keep_sect=idx not in used))
                changed += 1
            used.add(idx)
        changed += sum(1 for i in range(len(self.orig)) if i not in used)    # удалённые
        sect = self.body.find(qn("w:sectPr"))
        for ch in list(self.body):
            if ch is not sect:
                self.body.remove(ch)
        for el in new:
            if sect is not None:
                sect.addprevious(el)
            else:
                self.body.append(el)
        return changed

    def _rebuild_par(self, orig, block, keep_sect=True):
        from lxml import etree
        p = etree.Element(qn("w:p"))
        ppr = orig.find(qn("w:pPr")) if orig is not None else None
        ppr = copy.deepcopy(ppr) if ppr is not None else None
        bf = block.blockFormat()
        cur_al = _iv(bf.alignment()) & 0x0F or _iv(Qt.AlignLeft)
        al = ALIGN_BACK.get(cur_al)
        if ppr is None and al not in (None, "left"):
            ppr = etree.Element(qn("w:pPr"))
        if ppr is not None:
            if not keep_sect:                        # разрыв раздела — только у исходного абзаца
                for s in ppr.findall(qn("w:sectPr")):
                    ppr.remove(s)
            jc = ppr.find(qn("w:jc"))
            orig_al = _val(jc) if jc is not None else None
            if al and _iv(ALIGN.get(orig_al or "left", Qt.AlignLeft)) != cur_al:
                if jc is None:
                    jc = etree.Element(qn("w:jc"))
                    _insert_ppr_child(ppr, jc)
                jc.set(qn("w:val"), al)
            p.append(ppr)
        it = block.begin()
        while not it.atEnd():
            f = it.fragment()
            if f.isValid():
                self._emit_fragment(p, f.text(), f.charFormat())
            it += 1
        return p

    def _emit_fragment(self, p, text, cf):
        from lxml import etree
        obj = cf.property(R_OBJ)
        if obj is not None and obj >= 0:
            for ch in text:                     # каждый ■ — свой исходный объект
                if ch == OBJ and obj < len(self.objs):
                    p.append(copy.deepcopy(self.objs[obj]))
            return
        if not text:
            return
        r = etree.SubElement(p, qn("w:r"))
        ridx = cf.property(R_RPR)
        rpr = copy.deepcopy(self.rprs[ridx]) if ridx is not None and 0 <= ridx < len(self.rprs) and \
            self.rprs[ridx] is not None else None
        b, i, u = cf.fontWeight() >= BOLD_MIN, cf.fontItalic(), cf.fontUnderline()
        size = round(cf.fontPointSize(), 1) if cf.fontPointSize() > 0 else None
        fams = cf.fontFamilies() or []
        fam = fams[0] if fams else None
        changes = []
        if b != bool(cf.property(R_B0)):
            changes.append(("w:b", None if b else "0"))
        if i != bool(cf.property(R_I0)):
            changes.append(("w:i", None if i else "0"))
        if u != bool(cf.property(R_U0)):
            changes.append(("w:u", "single" if u else "none"))
        s0 = cf.property(R_S0)
        if size and (s0 is None or abs(float(s0) - size) > 0.05):
            changes.append(("w:sz", str(int(round(size * 2)))))
        if fam and fam != cf.property(R_F0):
            changes.append(("w:rFonts", fam))
        if changes and rpr is None:
            rpr = etree.Element(qn("w:rPr"))
        for tag, val in changes:
            for old in rpr.findall(qn(tag)):
                rpr.remove(old)
            el = etree.Element(qn(tag))
            if tag == "w:rFonts":
                for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                    el.set(qn(a), val)
            elif val is not None:
                el.set(qn("w:val"), val)
            _insert_rpr_child(rpr, el)
            if tag == "w:sz":
                cs = etree.Element(qn("w:szCs"))
                cs.set(qn("w:val"), val)
                for old in rpr.findall(qn("w:szCs")):
                    rpr.remove(old)
                _insert_rpr_child(rpr, cs)
        if rpr is not None:
            r.append(rpr)
        buf = []

        def flush():
            if buf:
                t = etree.SubElement(r, qn("w:t"))
                t.text = "".join(buf)
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                buf.clear()
        for ch in text:
            if ch == "\t":
                flush()
                etree.SubElement(r, qn("w:tab"))
            elif ch == LINE:
                flush()
                etree.SubElement(r, qn("w:br"))
            elif ch == OBJ:
                continue
            else:
                buf.append(ch)
        flush()

    def _rebuild_table(self, orig, cells):
        """Таблица с изменённым текстом: копия исходной, в ячейках — новый текст (с прежним оформлением)."""
        from lxml import etree
        tbl = copy.deepcopy(orig)
        for r, tr in enumerate(tbl.findall(qn("w:tr"))):
            for c, tc in enumerate(tr.findall(qn("w:tc"))):
                if r >= len(cells) or c >= len(cells[r]):
                    continue
                lines = cells[r][c].split("\n")
                ps = tc.findall(qn("w:p"))
                old = [("".join(t.text or "" for t in p.iter(qn("w:t")))) for p in ps]
                if old == lines:
                    continue
                proto = ps[0] if ps else None
                ppr = proto.find(qn("w:pPr")) if proto is not None else None
                first_r = proto.find(qn("w:r")) if proto is not None else None
                rpr = first_r.find(qn("w:rPr")) if first_r is not None else None
                for p in ps:
                    tc.remove(p)
                for line in lines:
                    p = etree.SubElement(tc, qn("w:p"))
                    if ppr is not None:
                        p.append(copy.deepcopy(ppr))
                    if line:
                        rr = etree.SubElement(p, qn("w:r"))
                        if rpr is not None:
                            rr.append(copy.deepcopy(rpr))
                        parts = line.split(LINE)
                        for k, part in enumerate(parts):
                            if k:
                                etree.SubElement(rr, qn("w:br"))
                            t = etree.SubElement(rr, qn("w:t"))
                            t.text = part
                            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        return tbl

    def save(self, path=None):
        path = path or self.path
        tmp = path + ".lh_tmp"
        self._forget_word_stats()
        self.doc.save(tmp)
        os.replace(tmp, path)

    def _forget_word_stats(self):
        """Число страниц, записанное Word, после нашей правки уже неверно — убрать его (см. word_page_count)."""
        try:
            for rel in self.doc.part.package.iter_parts():
                if str(rel.partname).endswith("/docProps/app.xml"):
                    import re
                    xml = rel.blob.decode("utf-8", "replace")
                    xml = re.sub(r"<Pages>\d+</Pages>", "", xml)
                    rel._blob = xml.encode("utf-8")
        except Exception:
            pass


# порядок элементов внутри pPr/rPr важен для Word — вставляем на «своё» место
_PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
              "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
              "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid",
              "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
              "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
_RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
              "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w",
              "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign",
              "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]


def _insert_ordered(parent, el, order):
    name = el.tag.split("}")[1]
    pos = order.index(name) if name in order else len(order)
    for i, ch in enumerate(parent):
        cname = ch.tag.split("}")[1] if "}" in ch.tag else ch.tag
        if cname in order and order.index(cname) > pos:
            parent.insert(i, el)
            return
    parent.append(el)


def _insert_ppr_child(ppr, el):
    _insert_ordered(ppr, el, _PPR_ORDER)


def _insert_rpr_child(rpr, el):
    _insert_ordered(rpr, el, _RPR_ORDER)


# =============================================================================
#  Окно редактора
# =============================================================================
class WordEditor(QDialog):
    def __init__(self, main, path):
        super().__init__(main)
        self.main, self.path = main, path
        self.setWindowTitle(f"{os.path.basename(path)} — правка")
        self.setWindowFlag(Qt.WindowMaximizeButtonHint, True)
        scr = main.screen().availableGeometry() if main is not None and main.screen() else None
        self.resize(min(1240, scr.width() - 40) if scr else 1240, min(900, scr.height() - 60) if scr else 860)
        T = getattr(main, "T", None) or {}
        v = QVBoxLayout(self)
        v.setContentsMargins(0, 0, 0, 0)
        v.setSpacing(0)
        tb = QToolBar()
        tb.setObjectName("wordbar")
        v.addWidget(tb)
        self.tb = tb
        self.info = QLabel()
        self.info.setObjectName("hint")
        self.info.setWordWrap(True)
        self.info.setContentsMargins(14, 6, 14, 6)
        v.addWidget(self.info)
        host = QWidget()
        host.setObjectName("wordhost")
        hl = QHBoxLayout(host)
        hl.setContentsMargins(0, 14, 0, 0)
        self.ed = QTextEdit()
        self.ed.setObjectName("wordpage")
        self.ed.setAcceptRichText(False)               # вставка из буфера — простым текстом с текущим оформлением
        self.ed.setFrameShape(QTextEdit.NoFrame)
        hl.addStretch(1)
        hl.addWidget(self.ed, 0)
        hl.addStretch(1)
        v.addWidget(host, 1)
        host.setStyleSheet(f"QWidget#wordhost {{ background: {T.get('pages', '#e9e9ee')}; }}"
                           "QTextEdit#wordpage { background: white; color: black; border: none; }")
        self._build_toolbar()
        self.load()
        self.ed.document().modificationChanged.connect(self._sync_title)
        self.ed.cursorPositionChanged.connect(self._sync_buttons)

    # ------------------------------------------------------------------ панель
    def _act(self, text, tip, fn, key=None, checkable=False):
        a = QAction(text, self)
        a.setToolTip(tip + (f"  ({QKeySequence(key).toString()})" if key else ""))
        if key:
            a.setShortcut(key)
        a.setCheckable(checkable)
        a.triggered.connect(fn)
        self.tb.addAction(a)
        return a

    def _build_toolbar(self):
        self.a_save = self._act("💾  Сохранить", "Сохранить в тот же файл Word; страницы в PDF дела обновятся сами",
                                self.save, QKeySequence.Save)
        self.tb.widgetForAction(self.a_save).setObjectName("tbprimary")
        self.tb.addSeparator()
        self._act("↶", "Отменить", lambda: self.ed.undo(), QKeySequence.Undo)
        self._act("↷", "Повторить", lambda: self.ed.redo(), QKeySequence.Redo)
        self.tb.addSeparator()
        self.font_box = QFontComboBox()
        self.font_box.setMaximumWidth(160)
        self.font_box.setCurrentFont(QFont(DEFAULT_FONT))
        self.font_box.setToolTip("Шрифт")
        self.font_box.currentFontChanged.connect(lambda f: self._merge(lambda cf: cf.setFontFamilies([f.family()])))
        self.tb.addWidget(self.font_box)
        self.size_box = QComboBox()
        self.size_box.setEditable(True)
        self.size_box.addItems([str(s) for s in (8, 9, 10, 11, 12, 13, 14, 16, 18, 20, 24, 28, 36)])
        self.size_box.setToolTip("Размер шрифта")
        self.size_box.setFixedWidth(64)
        self.size_box.textActivated.connect(self._set_size)
        self.tb.addWidget(self.size_box)
        self.a_b = self._act("Ж", "Жирный", lambda on: self._merge(lambda cf: cf.setFontWeight(
            QFont.Bold if on else QFont.Normal)), QKeySequence.Bold, True)
        self.a_i = self._act("К", "Курсив", lambda on: self._merge(lambda cf: cf.setFontItalic(on)),
                             QKeySequence.Italic, True)
        self.a_u = self._act("Ч", "Подчёркнутый", lambda on: self._merge(lambda cf: cf.setFontUnderline(on)),
                             QKeySequence.Underline, True)
        for a, st in ((self.a_b, "font-weight:700"), (self.a_i, "font-style:italic"),
                      (self.a_u, "text-decoration:underline")):
            self.tb.widgetForAction(a).setStyleSheet(st)
        self.tb.addSeparator()
        self.a_al = []
        for text, tip, al in (("⯇", "По левому краю", Qt.AlignLeft), ("≡", "По центру", Qt.AlignHCenter),
                              ("⯈", "По правому краю", Qt.AlignRight), ("☰", "По ширине", Qt.AlignJustify)):
            a = self._act(text, tip, lambda _=False, al=al: self._align(al), checkable=True)
            self.a_al.append((a, al))
        self.tb.addSeparator()
        self._act("🔍", "Найти", self.find, QKeySequence.Find)
        self.tb.addSeparator()
        self._act("В Word", "Для сложного оформления: открыть этот файл в Word "
                            "(изменения после сохранения в Word подтянутся сами)", self.open_in_word)

    def _merge(self, fn):
        cf = QTextCharFormat()
        fn(cf)
        cur = self.ed.textCursor()
        if not cur.hasSelection():
            self.ed.mergeCurrentCharFormat(cf)
        else:
            cur.mergeCharFormat(cf)
        self.ed.setFocus()

    def _set_size(self, text):
        try:
            size = float(text.replace(",", "."))
        except ValueError:
            return
        if 4 <= size <= 96:
            self._merge(lambda cf: cf.setFontPointSize(size))

    def _align(self, al):
        self.ed.setAlignment(al)
        self._sync_buttons()

    def _sync_buttons(self):
        cf = self.ed.currentCharFormat()
        for a, val in ((self.a_b, cf.fontWeight() >= BOLD_MIN), (self.a_i, cf.fontItalic()),
                       (self.a_u, cf.fontUnderline())):
            a.blockSignals(True)
            a.setChecked(val)
            a.blockSignals(False)
        al = _iv(self.ed.alignment()) & 0x0F or _iv(Qt.AlignLeft)
        for a, v in self.a_al:
            a.setChecked(al == _iv(v))
        self.size_box.blockSignals(True)
        size = cf.fontPointSize() if cf.fontPointSize() > 0 else self.ed.document().defaultFont().pointSizeF()
        self.size_box.setCurrentText(f"{size:g}")
        self.size_box.blockSignals(False)
        fams = cf.fontFamilies() or [self.ed.document().defaultFont().family()]
        if fams:
            self.font_box.blockSignals(True)
            self.font_box.setCurrentFont(QFont(fams[0]))
            self.font_box.blockSignals(False)

    def _sync_title(self, modified):
        self.setWindowTitle(("● " if modified else "") + f"{os.path.basename(self.path)} — правка")

    # ------------------------------------------------------------------ файл
    def load(self, keep_pos=None):
        self.model = DocxModel(self.path)
        doc = QTextDocument(self)
        doc.setUndoRedoEnabled(False)
        self.model.fill(doc)
        doc.setUndoRedoEnabled(True)
        sec = self.model.doc.sections[0] if self.model.doc.sections else None
        width = 794
        if sec is not None and sec.page_width:
            width = int(sec.page_width.pt * PX)
            fmt = doc.rootFrame().frameFormat()
            fmt.setLeftMargin((sec.left_margin.pt if sec.left_margin else 85) * PX)
            fmt.setRightMargin((sec.right_margin.pt if sec.right_margin else 42) * PX)
            fmt.setTopMargin((sec.top_margin.pt if sec.top_margin else 56) * PX)
            fmt.setBottomMargin((sec.bottom_margin.pt if sec.bottom_margin else 56) * PX)
            doc.rootFrame().setFrameFormat(fmt)
        self.ed.setDocument(doc)
        self.ed.setFixedWidth(min(width + 24, 1400))
        doc.setModified(False)
        notes = []
        if self.model.objs:
            notes.append(f"■ — картинки, поля или сноски ({len(self.model.objs)}): сохранятся как есть")
        if self.model.complex_tables:
            notes.append("таблицы с объединёнными ячейками правятся в Word")
        self.info.setText("Правьте текст как в Word. «Сохранить» (Ctrl+S) — запишет в тот же файл, а страницы в PDF "
                          "дела обновятся сами. Колонтитулы, поля страницы и нетронутые абзацы не меняются."
                          + ("  " + "; ".join(notes).capitalize() + "." if notes else ""))
        if keep_pos is not None:
            cur = self.ed.textCursor()
            cur.setPosition(min(keep_pos, doc.characterCount() - 1))
            self.ed.setTextCursor(cur)
            self.ed.ensureCursorVisible()
        self._sync_title(False)
        self._sync_buttons()

    def save(self):
        doc = self.ed.document()
        if not doc.isModified():
            self.main.toast("Изменений нет") if hasattr(self.main, "toast") else None
            return True
        pos = self.ed.textCursor().position()
        try:
            self.model.build(doc)
            backup(self.path)
            self.model.save()
        except PermissionError:
            QMessageBox.warning(self, "Сохранение", f"Файл «{os.path.basename(self.path)}» сейчас открыт в Word или "
                                "другой программе. Закройте его там и нажмите «Сохранить» ещё раз.")
            return False
        except Exception as e:
            QMessageBox.critical(self, "Сохранение", f"Не удалось сохранить: {e}\n\nФайл не изменён.")
            return False
        self.load(pos)                           # заново из файла: следующее сохранение сравнивает уже с ним
        if hasattr(self.main, "on_word_saved"):
            QTimer.singleShot(0, lambda: self.main.on_word_saved(self.path))
        return True

    def find(self):
        text, ok = QInputDialog.getText(self, "Найти", "Что найти:", text=getattr(self, "_last_find", ""))
        if not ok or not text:
            return
        self._last_find = text
        if not self.ed.find(text):
            cur = self.ed.textCursor()
            cur.movePosition(QTextCursor.Start)
            self.ed.setTextCursor(cur)
            if not self.ed.find(text):
                QMessageBox.information(self, "Найти", f"«{text}» не найдено.")

    def open_in_word(self):
        if self.ed.document().isModified():
            if not self._ask_save():
                return
        self.main.open_in_app(self.path)
        self.close()

    def _ask_save(self):
        r = QMessageBox.question(self, "Правка документа", f"Сохранить изменения в «{os.path.basename(self.path)}»?",
                                 QMessageBox.Save | QMessageBox.Discard | QMessageBox.Cancel, QMessageBox.Save)
        if r == QMessageBox.Save:
            return self.save()
        return r == QMessageBox.Discard

    def closeEvent(self, e):
        if self.ed.document().isModified() and not self._ask_save():
            e.ignore()
            return
        e.accept()


def can_edit(path):
    return bool(path) and path.lower().endswith(".docx") and os.path.isfile(path)
