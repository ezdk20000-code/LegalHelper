# -*- coding: utf-8 -*-
"""
Юридические функции LegalHelper (без интерфейса): расчёты сроков, госпошлины,
процентов и неустоек; подготовка пакета документов в суд и на почту;
нумерация листов дела; штампы заверения; опись вложения ф. 107;
проверка перед подачей; обезличивание; сравнение редакций; шаблоны документов.
"""
import os
import re
import math
import html
import difflib
import calendar
import datetime as dt
from pathlib import Path

import pymupdf as fitz

import pdf_core as C
import legal_data as D

MM = C.MM
A4 = fitz.paper_rect("a4")


# ============================================================================
#  Общие помощники
# ============================================================================
def money(x, cents=True):
    """1234567.8 -> '1 234 567,80'."""
    s = f"{abs(x):,.2f}" if cents else f"{abs(round(x)):,.0f}"
    s = s.replace(",", " ").replace(".", ",")
    return ("-" if x < 0 else "") + s


def ddmmyyyy(d):
    return d.strftime("%d.%m.%Y")


def parse_date(s):
    s = str(s).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Неверная дата: «{s}». Формат: ДД.ММ.ГГГГ")


_UNITS = ["", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
_UNITS_F = ["", "одна", "две"] + _UNITS[3:]
_TEENS = ["десять", "одиннадцать", "двенадцать", "тринадцать", "четырнадцать", "пятнадцать",
          "шестнадцать", "семнадцать", "восемнадцать", "девятнадцать"]
_TENS = ["", "", "двадцать", "тридцать", "сорок", "пятьдесят", "шестьдесят", "семьдесят", "восемьдесят",
         "девяносто"]
_HUNDREDS = ["", "сто", "двести", "триста", "четыреста", "пятьсот", "шестьсот", "семьсот", "восемьсот",
             "девятьсот"]


def plural(n, one, few, many):
    n = abs(n) % 100
    if 11 <= n <= 19:
        return many
    n %= 10
    return one if n == 1 else few if 2 <= n <= 4 else many


def num_words(n, feminine=False):
    """Число прописью: 3 -> 'три', 251 -> 'двести пятьдесят один'."""
    n = int(n)
    if n == 0:
        return "ноль"
    parts = []
    groups = [("", "", "", False), ("тысяча", "тысячи", "тысяч", True),
              ("миллион", "миллиона", "миллионов", False), ("миллиард", "миллиарда", "миллиардов", False)]
    k = 0
    while n > 0 and k < len(groups):
        g = n % 1000
        n //= 1000
        if g:
            one, few, many, fem = groups[k]
            fem = fem or (k == 0 and feminine)
            w = [_HUNDREDS[g // 100]]
            t = g % 100
            if 10 <= t <= 19:
                w.append(_TEENS[t - 10])
            else:
                w.append(_TENS[t // 10])
                w.append((_UNITS_F if fem else _UNITS)[t % 10])
            if one:
                w.append(plural(g, one, few, many))
            parts.insert(0, " ".join(x for x in w if x))
        k += 1
    return " ".join(parts)


def cap(s):
    return s[:1].upper() + s[1:]


def sheets_word(n):
    return plural(n, "лист", "листа", "листов")


def year_days(y):
    return 366 if calendar.isleap(y) else 365


# ============================================================================
#  Процессуальные сроки
# ============================================================================
def next_workday(d):
    while not D.is_workday(d):
        d += dt.timedelta(days=1)
    return d


def add_months(d, months):
    m = d.month - 1 + months
    y = d.year + m // 12
    m = m % 12 + 1
    day = min(d.day, calendar.monthrange(y, m)[1])
    return dt.date(y, m, day)


def compute_deadline(event_date, n, unit):
    """Последний день срока. unit: wd — рабочие дни, cd — календарные дни, m — месяцы, y — годы.
    Течение срока начинается на следующий день после события; если последний день нерабочий —
    срок переносится на ближайший рабочий день (ст. 114 АПК, ст. 108 ГПК, ст. 93 КАС, ст. 193 ГК).
    Возвращает (дата, [пояснения])."""
    notes = []
    if unit == "wd":
        d, left = event_date, int(n)
        while left > 0:
            d += dt.timedelta(days=1)
            if D.is_workday(d):
                left -= 1
        notes.append("В сроки, исчисляемые днями, не включаются нерабочие дни "
                     "(ч. 3 ст. 113 АПК, ч. 3 ст. 107 ГПК, ч. 2 ст. 92 КАС).")
    else:
        if unit == "cd":
            d = event_date + dt.timedelta(days=int(n))
        elif unit == "m":
            d = add_months(event_date, int(n))
        elif unit == "y":
            d = add_months(event_date, int(n) * 12)
        else:
            raise ValueError(unit)
        if not D.is_workday(d):
            nd = next_workday(d)
            notes.append(f"{ddmmyyyy(d)} — нерабочий день, окончание срока переносится на "
                         f"{ddmmyyyy(nd)}.")
            d = nd
    for y in {event_date.year, d.year}:
        if not D.year_calendar(y)[2]:
            notes.append(f"Производственный календарь на {y} год ещё не утверждён — учтены только "
                         "праздники по ст. 112 ТК РФ. Проверьте переносы выходных.")
    return d, notes


def days_between(a, b):
    """(календарных дней, рабочих дней) между датами включительно."""
    if b < a:
        a, b = b, a
    cal = (b - a).days + 1
    work = sum(1 for i in range(cal) if D.is_workday(a + dt.timedelta(days=i)))
    return cal, work


# ============================================================================
#  Госпошлина
# ============================================================================
def _scale(amount, scale, maximum):
    for top, base, pct, over in scale:
        if top is None or amount <= top:
            v = base + (amount - over) * pct / 100 if pct else base
            rule = f"{money(base, False)} ₽" + (f" + {str(pct).replace('.', ',')}% от суммы свыше "
                                                 f"{money(over, False)} ₽" if pct else "")
            if maximum and v > maximum:
                return maximum, rule + f", но не более {money(maximum, False)} ₽"
            return v, rule
    raise ValueError


def state_duty(court, amount=0.0, fixed_index=None, person="org", order=False):
    """court: 'soj' (ст. 333.19 НК) | 'arb' (ст. 333.21 НК).
    fixed_index — номер фиксированного вида из legal_data; иначе имущественный иск на amount.
    Возвращает (сумма, пояснение)."""
    if fixed_index is not None:
        table = D.DUTY_SOJ_FIXED if court == "soj" else D.DUTY_ARB_FIXED
        name, (fl, ul) = table[fixed_index]
        v = fl if person == "fl" else ul
        return float(v), f"{name}: {'физлицо' if person == 'fl' else 'организация'} — {money(v, False)} ₽"
    art = "ст. 333.19 НК РФ" if court == "soj" else "ст. 333.21 НК РФ"
    if court == "soj":
        v, rule = _scale(amount, D.DUTY_SOJ_SCALE, D.DUTY_SOJ_MAX)
    else:
        v, rule = _scale(amount, D.DUTY_ARB_SCALE, D.DUTY_ARB_MAX)
    v = math.floor(v + 0.5)   # до полного рубля: 50 коп. и более — вверх (п. 6 ст. 52 НК)
    expl = f"Цена иска {money(amount)} ₽. Шкала ({art}): {rule}. Пошлина: {money(v, False)} ₽."
    if order:
        half = math.ceil(v / 2)
        if court == "arb":
            half = max(half, 8000)
            expl += f" Судебный приказ — 50%, но не менее 8 000 ₽: {money(half, False)} ₽."
        else:
            expl += f" Судебный приказ — 50%: {money(half, False)} ₽."
        v = half
    return float(v), expl


# ============================================================================
#  Проценты по ст. 395 ГК и неустойка
# ============================================================================
def key_rate_on(d):
    rate = None
    for s, r in D.KEY_RATES:
        if dt.date.fromisoformat(s) <= d:
            rate = r
        else:
            break
    return rate


def _breakpoints(start, end, moves, use_rates, exclude_moratorium):
    """Даты, с которых меняется что-то одно: ставка, долг, граница года, мораторий."""
    pts = {start, end + dt.timedelta(days=1)}
    for y in range(start.year + 1, end.year + 1):
        pts.add(dt.date(y, 1, 1))
    if use_rates:
        for s, _ in D.KEY_RATES:
            d = dt.date.fromisoformat(s)
            if start < d <= end:
                pts.add(d)
    for d, _ in moves:
        if start < d <= end:
            pts.add(d)
    if exclude_moratorium:
        m0, m1 = (dt.date.fromisoformat(x) for x in D.MORATORIUM_2022)
        for d in (m0, m1 + dt.timedelta(days=1)):
            if start < d <= end:
                pts.add(d)
    return sorted(pts)


def interest_calc(principal, start, end, payments=(), additions=(), mode="395", pct_per_day=0.1,
                  rate_divisor=300, fixed_rate_date=None, exclude_moratorium=False, cap_principal=False):
    """Расчёт процентов/неустойки.
    mode: '395' — ключевая ставка ЦБ (ст. 395 ГК), годовая, по периодам;
          'pct' — неустойка pct_per_day % в день;
          'frac' — неустойка 1/rate_divisor ключевой ставки в день (по периодам или на fixed_rate_date).
    payments — [(дата, сумма)] оплаты: долг уменьшается со следующего дня после оплаты
    (день оплаты входит в период просрочки — п. 48 ПП ВС РФ № 7);
    additions — [(дата, сумма)] увеличение долга: с указанной даты.
    Возвращает dict(rows=[...], total, notes)."""
    if end < start:
        raise ValueError("Дата окончания раньше даты начала")
    notes = []
    moves = [(d + dt.timedelta(days=1), -float(a)) for d, a in payments] + \
            [(d, float(a)) for d, a in additions]
    use_rates = mode in ("395", "frac") and not fixed_rate_date
    if mode == "395" and start < dt.date.fromisoformat(D.RATE_395_FROM):
        notes.append("До 01.08.2016 проценты по ст. 395 ГК определялись иначе (средние ставки по "
                     "округам, ставка рефинансирования) — для этого периода расчёт условный.")
    pts = _breakpoints(start, end, moves, use_rates, exclude_moratorium)
    m0, m1 = (dt.date.fromisoformat(x) for x in D.MORATORIUM_2022)
    debt = float(principal)
    for d, a in moves:
        if d <= start:
            debt += a
    rows, total = [], 0.0
    fixed_rate = key_rate_on(fixed_rate_date) if fixed_rate_date else None
    for i in range(len(pts) - 1):
        a, b = pts[i], pts[i + 1] - dt.timedelta(days=1)
        for d, amt in moves:
            if d == a and d > start:
                debt += amt
        if debt <= 0:
            debt = max(debt, 0.0)
            continue
        days = (b - a).days + 1
        if exclude_moratorium and m0 <= a <= m1:
            rows.append(dict(start=a, end=b, days=days, debt=debt, rate=None, basis="мораторий",
                             amount=0.0))
            continue
        if mode == "395":
            rate = key_rate_on(a)
            yd = year_days(a.year)
            amt = debt * rate / 100 * days / yd
            basis = f"{str(rate).replace('.', ',')}% / {yd}"
        elif mode == "frac":
            rate = fixed_rate if fixed_rate is not None else key_rate_on(a)
            amt = debt * rate / 100 / rate_divisor * days
            basis = f"1/{rate_divisor} × {str(rate).replace('.', ',')}%"
        else:
            rate = pct_per_day
            amt = debt * pct_per_day / 100 * days
            basis = f"{str(pct_per_day).replace('.', ',')}% в день"
        amt = round(amt, 2)
        total += amt
        rows.append(dict(start=a, end=b, days=days, debt=debt, rate=rate, basis=basis, amount=amt))
    if cap_principal and total > principal:
        notes.append(f"Сумма ограничена размером основного долга ({money(principal)} ₽).")
        total = float(principal)
    if exclude_moratorium:
        notes.append("Период моратория 01.04.2022–01.10.2022 исключён (ПП РФ от 28.03.2022 № 497).")
    return dict(rows=rows, total=round(total, 2), notes=notes)


def calc_table(res, mode):
    """Строки таблицы для вывода: заголовки и значения."""
    head = ["Период", "Дней", "Долг, ₽", "Ставка", "Сумма, ₽"]
    body = []
    for r in res["rows"]:
        body.append([f"{ddmmyyyy(r['start'])} – {ddmmyyyy(r['end'])}", str(r["days"]), money(r["debt"]),
                     r["basis"], money(r["amount"])])
    return head, body


# ============================================================================
#  Таблица в Word / HTML
# ============================================================================
def table_html(title, head, body, footer_lines=()):
    h = [f"<h3>{html.escape(title)}</h3>", "<table border='1' cellspacing='0' cellpadding='4' "
         "style='border-collapse:collapse;font-size:10pt'>",
         "<tr>" + "".join(f"<th>{html.escape(x)}</th>" for x in head) + "</tr>"]
    for row in body:
        h.append("<tr>" + "".join(f"<td>{html.escape(str(x))}</td>" for x in row) + "</tr>")
    h.append("</table>")
    for line in footer_lines:
        h.append(f"<p>{html.escape(line)}</p>")
    return "\n".join(h)


def table_docx(path, title, head, body, footer_lines=()):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    d.add_paragraph().add_run(title).bold = True
    t = d.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    for i, x in enumerate(head):
        t.rows[0].cells[i].text = x
        for r in t.rows[0].cells[i].paragraphs[0].runs:
            r.bold = True
    for row in body:
        cells = t.add_row().cells
        for i, x in enumerate(row):
            cells[i].text = str(x)
    for line in footer_lines:
        d.add_paragraph(line)
    d.save(path)
    return path


def table_tsv(head, body, footer_lines=()):
    return "\n".join(["\t".join(head)] + ["\t".join(map(str, r)) for r in body] + list(footer_lines))


def html_to_pdf_file(html_text, out, title=""):
    """Простой HTML (таблицы, абзацы) -> PDF A4 через fitz.Story."""
    css = ("body{font-family:sans-serif;font-size:10pt} td,th{padding:3px;border:1px solid #888} "
           "table{border-collapse:collapse} h3{font-size:12pt} ins{color:#0a7a28;background:#e3f6e8;"
           "text-decoration:underline} del{color:#b3261e;background:#fde7e5;text-decoration:line-through}")
    ff = C.font_file()
    arch = None
    if ff:
        arch = fitz.Archive(os.path.dirname(ff))
        css = f"@font-face{{font-family:ru;src:url({os.path.basename(ff)});}} " + \
              css.replace("sans-serif", "ru")
    story = fitz.Story(html=html_text, user_css=css, archive=arch)
    writer = fitz.DocumentWriter(out)
    rect = A4 + (40, 40, -40, -40)
    more = True
    while more:
        dev = writer.begin_page(A4)
        more, _ = story.place(rect)
        story.draw(dev)
        writer.end_page()
    writer.close()
    return out


# ============================================================================
#  Подготовка пакета документов (суд / почта)
# ============================================================================
def clean_filename(s, maxlen=150):
    s = re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", s).strip().strip(".")
    s = re.sub(r"\s+", " ", s)
    return s[:maxlen] or "Документ"


def count_sheets(pages, duplex=False):
    return math.ceil(pages / 2) if duplex else pages


def attachments_text(items, start=1, skip_first=True):
    """Перечень приложений: items — [dict(title, sheets, copies)]."""
    lines = []
    src = items[1:] if skip_first else items
    for i, it in enumerate(src, start):
        cp = int(it.get("copies") or 1)
        lines.append(f"{i}. {it['title']} — на {it['sheets']} л. в {cp} экз.")
    return lines


def _fit_size(doc, path, limit_mb, progress=None):
    """Сохранить doc в path, при необходимости сжимая до limit_mb. Возвращает (размер, уровень)."""
    C.save_pdf(doc, path, garbage=4)
    size = os.path.getsize(path)
    if not limit_mb or size <= limit_mb * 1024 * 1024:
        return size, None
    for level in list(C.COMPRESS_LEVELS)[1:]:
        size = C.compress_document(doc, path, level)
        if size <= limit_mb * 1024 * 1024:
            return size, level
    return size, "превышен"


def _append_list_page(doc, title, lines):
    page = doc.new_page(width=A4.width, height=A4.height)
    text = title + "\n\n" + "\n".join(lines)
    C.write_text(page, fitz.Rect(70, 60, A4.width - 50, A4.height - 50), text, fontsize=12,
                 font_name="Times New Roman")


def build_package(items, portal, out_dir, number_prefix=True, add_list=False, list_title="Приложения:",
                  combined=True, split_letters=True, progress=None, password_cb=None):
    """items — [dict(path, title, sheets, copies)]; первый — основной документ (иск, жалоба, письмо).
    portal — ключ legal_data.PORTALS. Возвращает dict(files=[(путь, стр., байт)], warnings, list_lines)."""
    P = D.PORTALS[portal]
    os.makedirs(out_dir, exist_ok=True)
    warnings, files = [], []
    docs = []
    total = len(items)
    for k, it in enumerate(items):
        d = C.open_as_pdf(it["path"], password_cb)
        if d is None:
            raise ValueError(f"Не удалось открыть: {it['path']}")
        docs.append(d)
        it.setdefault("sheets", d.page_count)
        it.setdefault("copies", 1)
        it.setdefault("title", Path(it["path"]).stem)
        if progress:
            progress(k + 1, total * 2)
    list_lines = attachments_text(items)
    if add_list and len(items) > 1:
        _append_list_page(docs[0], list_title, list_lines)
        items[0]["sheets"] = items[0]["sheets"] + 1

    if P["a4_only"]:
        for idx, d in enumerate(docs):
            bad = [i for i in range(d.page_count) if not _is_a4(d[i])]
            if bad:
                docs[idx] = C.resize_pages(d, bad, "A4 (210×297 мм)", "Книжная")
                warnings.append(f"Страницы не формата A4 приведены к A4 ({len(bad)} шт.).")

    if P["one_file_per_doc"]:
        for k, (it, d) in enumerate(zip(items, docs)):
            name = it["title"]
            if P["name_with_sheets"]:
                name += f" {it['sheets']} л"
            if number_prefix:
                name = f"{k + 1:02d}. {name}"
            path = os.path.join(out_dir, clean_filename(name) + ".pdf")
            size, lvl = _fit_size(d, path, P["max_file_mb"])
            if lvl == "превышен":
                # делим на части
                os.remove(path)
                parts = _split_by_size(d, P["max_file_mb"])
                for j, (a, b) in enumerate(parts, 1):
                    sub = C.subset_doc(d, list(range(a, b + 1)))
                    pth = os.path.join(out_dir, clean_filename(f"{name} (часть {j})") + ".pdf")
                    sz, _ = _fit_size(sub, pth, P["max_file_mb"])
                    files.append((pth, sub.page_count, sz))
                warnings.append(f"«{it['title']}» больше {P['max_file_mb']} МБ даже после сжатия — "
                                f"разделён на {len(parts)} части.")
            else:
                if lvl:
                    warnings.append(f"«{it['title']}» сжат ({lvl}), чтобы уложиться в {P['max_file_mb']} МБ.")
                files.append((path, d.page_count, size))
            if progress:
                progress(total + k + 1, total * 2)
    else:
        big = fitz.open()
        toc = []
        for it, d in zip(items, docs):
            toc.append([1, it["title"], big.page_count + 1])
            big.insert_pdf(d)
        big.set_toc(toc)
        chunks = [(0, big.page_count - 1)]
        if P["max_pages"] and big.page_count > P["max_pages"]:
            if split_letters:
                chunks = _split_by_docs(docs, P["max_pages"])
                warnings.append(f"Всего {big.page_count} стр. — больше {P['max_pages']}. Пакет разбит на "
                                f"{len(chunks)} письма (документы не разрываются, если это возможно).")
            else:
                warnings.append(f"Всего {big.page_count} стр. — больше допустимых {P['max_pages']}.")
        for j, (a, b) in enumerate(chunks, 1):
            sub = C.subset_doc(big, list(range(a, b + 1))) if len(chunks) > 1 else big
            nm = "Письмо" + (f" {j}" if len(chunks) > 1 else "")
            path = os.path.join(out_dir, nm + ".pdf")
            size, lvl = _fit_size(sub, path, P["max_total_mb"])
            if lvl == "превышен":
                warnings.append(f"{nm}: {size / 1048576:.1f} МБ — больше {P['max_total_mb']} МБ даже после "
                                "сжатия. Уменьшите количество документов.")
            elif lvl:
                warnings.append(f"{nm} сжато ({lvl}), чтобы уложиться в {P['max_total_mb']} МБ.")
            files.append((path, sub.page_count, size))
        combined = False
    if combined and len(docs) > 1:
        big = fitz.open()
        toc = []
        for it, d in zip(items, docs):
            toc.append([1, it["title"], big.page_count + 1])
            big.insert_pdf(d)
        big.set_toc(toc)
        path = os.path.join(out_dir, "_Весь пакет одним файлом (для себя).pdf")
        C.save_pdf(big, path, garbage=4)
    if P["max_total_mb"] and P["one_file_per_doc"]:
        tot = sum(f[2] for f in files)
        if tot > P["max_total_mb"] * 1048576:
            warnings.append(f"Общий размер {tot / 1048576:.1f} МБ больше {P['max_total_mb']} МБ.")
    return dict(files=files, warnings=warnings, list_lines=list_lines)


def _is_a4(page, tol=3):
    w, h = sorted((page.rect.width / MM, page.rect.height / MM))
    return abs(w - 210) <= tol and abs(h - 297) <= tol


def _split_by_size(doc, limit_mb):
    per_page = max(1, len(doc.tobytes(garbage=3, deflate=True)) / max(1, doc.page_count))
    n = max(1, int(limit_mb * 1048576 * 0.9 // per_page))
    return [(a, min(a + n, doc.page_count) - 1) for a in range(0, doc.page_count, n)]


def _split_by_docs(docs, max_pages):
    """Группировать документы по письмам не более max_pages страниц (большой документ режется)."""
    chunks, cur_start, cur, pos = [], 0, 0, 0
    for d in docs:
        n = d.page_count
        if cur and cur + n > max_pages:
            chunks.append((cur_start, pos - 1))
            cur_start, cur = pos, 0
        while n > max_pages:
            chunks.append((pos, pos + max_pages - 1))
            pos += max_pages
            n -= max_pages
            cur_start = pos
        cur += n
        pos += n
    if cur:
        chunks.append((cur_start, pos - 1))
    return chunks


# ============================================================================
#  Проверка перед подачей
# ============================================================================
def preflight(doc, portal=None, file_size=None):
    """Возвращает список (уровень, страница|None, текст). Уровни: 'error', 'warn', 'info', 'ok'."""
    P = D.PORTALS.get(portal) if portal else None
    out = []
    n = doc.page_count
    if file_size is None:
        file_size = len(doc.tobytes(garbage=3, deflate=True))
    mb = file_size / 1048576
    limit = (P or {}).get("max_file_mb") or (P or {}).get("max_total_mb")
    if limit and mb > limit:
        out.append(("error", None, f"Размер файла {mb:.1f} МБ — больше допустимых {limit} МБ. "
                                   "Используйте «Сжать PDF» или «Пакет в суд» (сожмёт автоматически)."))
    else:
        out.append(("ok", None, f"Размер файла: {mb:.1f} МБ" + (f" (лимит {limit} МБ)" if limit else "")))
    if P and P.get("max_pages") and n > P["max_pages"]:
        out.append(("error", None, f"Страниц: {n} — больше допустимых {P['max_pages']}."))
    blank, rotated, scans, lowdpi, not_a4, landscape = [], [], [], [], [], []
    for i in range(n):
        p = doc[i]
        pix = p.get_pixmap(dpi=18, colorspace=fitz.csGRAY, alpha=False)
        s = pix.samples
        if s:
            mean = sum(s) / len(s)
            dark = sum(1 for v in s if v < 200)
            if mean > 247 and dark <= max(2, len(s) // 2000):
                blank.append(i + 1)
        if p.rotation:
            rotated.append(i + 1)
        if p.rect.width > p.rect.height:
            landscape.append(i + 1)
        if not _is_a4(p):
            not_a4.append(i + 1)
        txt = p.get_text().strip()
        imgs = p.get_image_info()
        if not txt and imgs:
            scans.append(i + 1)
        for im in imgs:
            bw = abs(im["bbox"][2] - im["bbox"][0]) / 72
            if bw > 2 and im.get("width"):
                dpi = im["width"] / bw
                if dpi < 150:
                    lowdpi.append(i + 1)
                    break

    def rng(lst):
        return C_ranges(lst)
    if blank:
        out.append(("warn", blank[0], f"Похоже на пустые страницы: {rng(blank)}."))
    if scans:
        out.append(("warn" if P else "info", scans[0],
                    f"Страницы без текстового слоя (скан): {rng(scans)}. Рекомендуется OCR — суды просят "
                    "документы «с возможностью копирования текста»."))
    if lowdpi:
        out.append(("warn", lowdpi[0], f"Низкое разрешение скана (< 150 dpi): {rng(lowdpi)}. "
                                       "Требуется 200–300 dpi."))
    if rotated:
        out.append(("info", rotated[0], f"Повёрнутые страницы: {rng(rotated)} — проверьте ориентацию."))
    if P and P.get("a4_only") and not_a4:
        out.append(("error", not_a4[0], f"Не формат A4: {rng(not_a4)}. Для почты — только A4 "
                                        "(инструмент «Размер страниц»)."))
    elif not_a4:
        out.append(("info", not_a4[0], f"Не формат A4: {rng(not_a4)}."))
    if landscape and not rotated:
        out.append(("info", landscape[0], f"Альбомные страницы: {rng(landscape)}."))
    if doc.is_encrypted or doc.permissions != -1 and doc.permissions & fitz.PDF_PERM_COPY == 0:
        out.append(("warn", None, "Документ защищён или запрещено копирование — суд может не принять. "
                                  "Используйте «Снять пароль»."))
    if not any(l in ("error", "warn") for l, _, _ in out):
        out.append(("ok", None, "Серьёзных проблем не найдено."))
    return out


def C_ranges(nums):
    """[1,2,3,5] -> '1–3, 5'."""
    nums = sorted(set(nums))
    res, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        res.append(str(nums[i]) if i == j else f"{nums[i]}–{nums[j]}")
        i = j + 1
    s = ", ".join(res)
    return s if len(s) < 120 else s[:117] + "…"


# ============================================================================
#  Нумерация листов дела, тома
# ============================================================================
SHEET_POSITIONS = ["Вверху справа", "Внизу справа", "Вверху по центру", "Внизу по центру"]


def number_sheets(doc, start=1, per_volume=250, position="Вверху справа", fontsize=12, margin_mm=8,
                  duplex=False, restart_each_volume=True, pages=None, progress=None):
    """Проставить номера листов. duplex — номер только на лицевой стороне (каждая 2-я страница — оборот).
    Возвращает список томов [(первая стр., последняя стр., первый лист, последний лист)]."""
    pages = list(range(doc.page_count)) if pages is None else list(pages)
    m = margin_mm * MM
    align = 2 if "справа" in position else 1
    volumes, sheet, vol_start_page, vol_first_sheet, count_in_vol = [], start, pages[0] if pages else 0, start, 0
    for k, pno in enumerate(pages):
        front = (k % 2 == 0) if duplex else True
        if not front:
            continue
        if per_volume and count_in_vol == per_volume:
            volumes.append((vol_start_page, pno - 1, vol_first_sheet, sheet - 1))
            vol_start_page = pno
            if restart_each_volume:
                sheet = 1
            vol_first_sheet, count_in_vol = sheet, 0
        page = doc[pno]
        vr = page.rect
        h = fontsize * 1.6
        y0 = vr.y0 + m if "Вверху" in position else vr.y1 - m - h
        rect = fitz.Rect(vr.x0 + m, y0, vr.x1 - m, y0 + h)
        C.write_text(page, rect, str(sheet), fontsize, (0, 0, 0), "Times New Roman", align, extend=False)
        sheet += 1
        count_in_vol += 1
        if progress:
            progress(k + 1, len(pages))
    if pages:
        volumes.append((vol_start_page, pages[-1], vol_first_sheet, sheet - 1))
    return volumes


def volume_cover(title, vol_no, sheets_from, sheets_to, extra=""):
    """Титульный лист тома (A4)."""
    d = fitz.open()
    p = d.new_page(width=A4.width, height=A4.height)
    txt = f"{title}\n\nТОМ № {vol_no}\n\nлисты {sheets_from}–{sheets_to}\n\n{extra}".strip()
    C.write_text(p, fitz.Rect(60, 250, A4.width - 60, A4.height - 60), txt, 20, (0, 0, 0),
                 "Times New Roman", align=1)
    return d


# ============================================================================
#  Штамп заверения
# ============================================================================
STAMP_POSITIONS = ["Внизу справа", "Внизу слева", "Вверху справа", "Внизу по центру"]
STAMP_MODES = ["На каждой странице", "Только на последней", "На первой и последней", "Только на первой"]


def stamp_pages(n, mode):
    if n == 0:
        return []
    if mode.startswith("Только на последней"):
        return [n - 1]
    if mode.startswith("На первой и последней"):
        return sorted({0, n - 1})
    if mode.startswith("Только на первой"):
        return [0]
    return list(range(n))


def certify_stamp(doc, pages, lines, position="Внизу справа", fontsize=10, width_mm=75, margin_mm=10,
                  frame=True, color=(0.05, 0.1, 0.55), sig_png=None, progress=None):
    """Нанести штамп из строк (напр. «Копия верна», должность, ФИО, дата)."""
    text = "\n".join(l for l in lines if l.strip())
    n_lines = max(1, text.count("\n") + 1)
    w = width_mm * MM
    h = n_lines * fontsize * 1.35 + 10 + (28 if sig_png else 0)
    m = margin_mm * MM
    for k, pno in enumerate(pages):
        page = doc[pno]
        vr = page.rect
        if "справа" in position:
            x0 = vr.x1 - m - w
        elif "слева" in position:
            x0 = vr.x0 + m
        else:
            x0 = (vr.x0 + vr.x1 - w) / 2
        y0 = vr.y0 + m if "Вверху" in position else vr.y1 - m - h
        box = fitz.Rect(x0, y0, x0 + w, y0 + h)
        if frame:
            page.draw_rect(C.vis_to_shape_rect(page, box), color=color, width=1.1)
        C.write_text(page, box + (6, 5, -6, -(28 if sig_png else 0)), text, fontsize, color,
                     "Times New Roman", 0, extend=False)
        if sig_png:
            sr = fitz.Rect(box.x1 - 90, box.y1 - 32, box.x1 - 6, box.y1 - 3)
            page.insert_image(C.vis_to_shape_rect(page, sr), stream=sig_png, keep_proportion=True,
                              rotate=page.rotation)
        if progress:
            progress(k + 1, len(pages))


def bound_stamp_lines(sheets, who="", date=""):
    """«Прошито, пронумеровано …» для последнего листа."""
    words = num_words(sheets, feminine=False)
    return [f"Прошито, пронумеровано", f"и скреплено печатью",
            f"{sheets} ({words}) {sheets_word(sheets)}", who, date]


# ============================================================================
#  Опись вложения ф. 107
# ============================================================================
F107_KINDS = ["ценное письмо", "ценную бандероль", "ценную посылку", "ценное письмо 1 класса",
              "ценную бандероль 1 класса"]


# Бланк ф. 107 — официальный бланк Почты России (pochta.ru): альбомный A4, два одинаковых экземпляра
# рядом, 14 строк. Пустой бланк лежит в forms/f107_blank.pdf, данные впечатываются в те же места
# тем же шрифтом (Arial Bold 9 пт), что и на бланках с pochta.ru.
F107_ROWS = 14
F107_COPY_DX = 417.8                       # сдвиг второго экземпляра вправо
F107_ROW_Y = (136.1, 151.1, 165.5, 180.0, 194.4, 208.9, 223.4, 237.8, 252.3, 266.7, 281.2, 295.6, 310.1,
              324.5, 339.4)                # линии между 14 строками таблицы (замерены по отрисовке бланка)
F107_COLS = (50.5, 77.4, 258.0, 299.6, 376.4)   # № | наименование | кол-во | ценность
F107_TEXT_TOP = 138.4                     # верх текста 1-й строки; дальше — ровный шаг, как у генератора pochta.ru
F107_PITCH = 14.5
F107_TOTAL = (344.6, 279.2, 338.7)         # верх текста итога, центры «кол-во» и «ценность»
F107_SENDER = (51.0, 281.1, 389.3, 401.8)  # x начала, x конца линий, линии 1-й и 2-й строки
F107_SPI = (52.6, 14.3, 13.0, 88.6, 101.6)  # x первой клетки, шаг, ширина, верх, низ (14 клеток)
F107_SIZE = 9


def _forms_dir():
    return os.path.join(C._app_dir(), "forms")


def _f107_font():
    """Arial Bold, как на бланках Почты; без него — метрически совместимый Liberation Sans Bold."""
    for p in (os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts", "arialbd.ttf"),
              os.path.join(_forms_dir(), "LiberationSans-Bold.ttf")):
        if os.path.exists(p):
            return p
    raise FileNotFoundError("Не найден шрифт для описи (forms/LiberationSans-Bold.ttf).")


class _F107Writer:
    def __init__(self, page, fontfile):
        self.p = page
        self.ff = fontfile
        self.font = fitz.Font(fontfile=fontfile)

    def width(self, s, size=F107_SIZE):
        return self.font.text_length(s, size)

    def text(self, x, baseline, s, size=F107_SIZE):
        if s:
            self.p.insert_text((x, baseline), s, fontsize=size, fontname="F107", fontfile=self.ff,
                               color=(0, 0, 0))

    def centered(self, x0, x1, top, bottom, s, size=F107_SIZE):
        """По центру ячейки — как на бланке pochta.ru."""
        base = (top + bottom) / 2 + (self.font.ascender + self.font.descender) * size / 2
        self.text((x0 + x1) / 2 - self.width(s, size) / 2, base, s, size)

    def at(self, cx, text_top, s, size=F107_SIZE):
        """Центр по x, верх текста по y (координаты сняты с бланка pochta.ru)."""
        self.text(cx - self.width(s, size) / 2, text_top + self.font.ascender * size, s, size)

    def left(self, x0, x1, top, bottom, s, size=F107_SIZE):
        """Слева в ячейке; длинный текст — мельче или в две строки, но не за границу."""
        room = x1 - x0 - 3
        for sz in (size, 8.5, 8, 7.5, 7):
            if self.width(s, sz) <= room:
                base = (top + bottom) / 2 + (self.font.ascender + self.font.descender) * sz / 2
                return self.text(x0, base, s, sz)
        sz, lh = 6.0, 6.6                     # две строки мелким шрифтом внутри строки таблицы
        lines = _wrap_to(self, s, room, sz, 2)
        first = top + 0.6 + self.font.ascender * sz
        if len(lines) == 1:
            first += lh / 2
        for i, line in enumerate(lines):
            self.text(x0, first + i * lh, line, sz)


def _wrap_full(w, s, room, size):
    """Разбить на строки по ширине room без потерь: слово длиннее строки режется по буквам."""
    lines, cur = [], ""
    for wd in s.split():
        while w.width(wd, size) > room:              # очень длинное слово (номер, адрес сайта…)
            k = len(wd)
            while k > 1 and w.width(wd[:k], size) > room:
                k -= 1
            if cur:
                lines.append(cur)
                cur = ""
            lines.append(wd[:k])
            wd = wd[k:]
        t = (cur + " " + wd).strip()
        if not cur or w.width(t, size) <= room:
            cur = t
        else:
            lines.append(cur)
            cur = wd
    if cur:
        lines.append(cur)
    return lines or [""]


def _wrap_to(w, s, room, size, max_lines):
    words, lines, cur = s.split(), [], ""
    for wd in words:
        t = (cur + " " + wd).strip()
        if w.width(t, size) <= room or not cur:
            cur = t
        else:
            lines.append(cur)
            cur = wd
    if cur:
        lines.append(cur)
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        while lines[-1] and w.width(lines[-1] + "…", size) > room:
            lines[-1] = lines[-1][:-1]
        lines[-1] += "…"
    return lines


def _f107_value(v):
    if v in (None, ""):
        return "—"
    v = float(v)
    return money(v, cents=abs(v - round(v)) > 0.004)


F107_TWO = 6.4                             # шрифт двух строк внутри одной строки бланка
F107_TWO_LH = 6.6


def _f107_layout(items, ff):
    """Раскладка пунктов по листам бланка: [[(строка, строк, №, [текст], шрифт, вид, кол-во, ценность)]].
    вид: «one» — одна строка; «two» — две строки мельче внутри одной строки бланка; «merge» — длинный
    пункт в объединённой ячейке из нескольких строк. Пункт не разрывается между листами, если он сам не
    длиннее листа."""
    probe = _F107Writer(fitz.open().new_page(), ff)
    room = F107_COLS[2] - (F107_COLS[1] + 2.8) - 3
    row_h = (F107_ROW_Y[-1] - F107_ROW_Y[0]) / F107_ROWS
    sheets, cur, free = [], [], F107_ROWS
    fits = [next((sz for sz in (F107_SIZE, 8.5, 8, 7.5, 7) if probe.width(n, sz) <= room), None)
            for n, _q, _v in items]
    one = min([sz for sz in fits if sz] or [F107_SIZE])      # однострочные — одним размером, чтобы опись была ровной
    for no, (name, q, v) in enumerate(items, 1):
        size = one if fits[no - 1] else None
        if size:
            chunks = [(1, "one", [name], size)]
        else:
            two = _wrap_full(probe, name, room, F107_TWO)
            if len(two) <= 2:
                chunks = [(1, "two", two, F107_TWO)]
            else:
                lines = _wrap_full(probe, name, room, 8)
                lh = 8 * 1.18
                per_sheet = max(1, int((F107_ROWS * row_h - 3) // lh))
                chunks = []
                for i in range(0, len(lines), per_sheet):
                    part = lines[i:i + per_sheet]
                    chunks.append((min(F107_ROWS, math.ceil((len(part) * lh + 3) / row_h)), "merge", part, 8))
        for k, (rows, kind, part, sz) in enumerate(chunks):
            if rows > free:
                sheets.append(cur)
                cur, free = [], F107_ROWS
            cur.append((F107_ROWS - free, rows, no if k == 0 else None, part, sz, kind,
                        q if k == 0 else None, v if k == 0 else None))
            free -= rows
    sheets.append(cur)
    return sheets


def f107_sheets(items):
    """Сколько листов займёт опись (для подсказки в окне)."""
    items = [(str(n).strip(), int(q or 0), v) for n, q, v in items if str(n).strip()]
    return len(_f107_layout(items, _f107_font())) if items else 0


_F107_SHORT = [
    (r"\bобществ\w* с ограниченной ответственностью\b", "ООО"), (r"\bпубличн\w* акционерн\w* обществ\w*\b", "ПАО"),
    (r"\bакционерн\w* обществ\w*\b", "АО"), (r"\bиндивидуальн\w* предпринимател\w*\b", "ИП"),
    (r"\bРоссийской Федерации\b", "РФ"), (r"\bгосударственной пошлины\b", "госпошлины"),
    (r"\bплатёжного поручения\b|\bплатежного поручения\b", "п/п"),
    (r"\bза пользование чужими денежными средствами\b", "по ст. 395 ГК РФ"),
    (r"\bс доказательств\w* (?:её |его |их )?направления\b", "с док-вом направления"),
    (r"\b(доверенност\w*) представителя\b", r"\1"),
    (r"\b(выписк\w*) из Единого государственного реестра юридических лиц\b", r"\1 из ЕГРЮЛ"),
    (r"(\d{1,2}\.\d{1,2}\.\d{4})\s*(?:г\.|года)", r"\1"), (r"\s+№\s+", " № "), (r"\s{2,}", " "),
]


def f107_shorten(name):
    """Короче, но без потери смысла: ООО, РФ, госпошлина, п/п, «по ст. 395 ГК РФ», даты без «г.»."""
    s = str(name)
    for rx, rep in _F107_SHORT:
        s = re.sub(rx, rep, s, flags=re.I)
    s = s.strip()
    return s[:1].upper() + s[1:] if str(name)[:1].isupper() else s



def f107_pdf(items, sender="", spi="", out=None, **_old):
    """Опись вложения ф. 107 на официальном бланке Почты России.
    items — [(наименование, количество, ценность руб. или None)]; sender — отправитель (ФИО или организация);
    spi — номер почтового идентификатора (14 цифр, необязательно). На листе 14 строк; длинное наименование
    занимает несколько строк подряд. Не поместилось — следующий лист, нумерация сквозная, итог — на последнем."""
    blank = os.path.join(_forms_dir(), "f107_blank.pdf")
    if not os.path.exists(blank):
        raise FileNotFoundError("Не найден бланк описи (forms/f107_blank.pdf).")
    ff = _f107_font()
    items = [(str(n).strip(), int(q or 0), v) for n, q, v in items if str(n).strip()]
    tot_q = sum(q for _n, q, _v in items)
    vals = [float(v) for _n, _q, v in items if v not in (None, "")]
    tot_v = _f107_value(sum(vals)) if vals else "0"
    spi = "".join(ch for ch in str(spi or "") if not ch.isspace())[:14]
    sheets = _f107_layout(items, ff)
    doc = fitz.open()
    tpl = fitz.open(blank)
    for n, placed in enumerate(sheets):
        doc.insert_pdf(tpl)
        page = doc[-1]
        w = _F107Writer(page, ff)
        last = n == len(sheets) - 1
        for dx in (0, F107_COPY_DX):
            c = [x + dx for x in F107_COLS]
            for row, rows, no, lines, size, kind, q, v in placed:
                top, bot = F107_ROW_Y[row], F107_ROW_Y[row + rows]
                if rows > 1:                  # одна ячейка на весь пункт: убрать линии между её строками
                    for y in F107_ROW_Y[row + 1:row + rows]:
                        for x0, x1 in zip(c, c[1:]):
                            page.draw_rect(fitz.Rect(x0 + 0.36, y - 0.6, x1 - 0.36, y + 0.6), color=None,
                                           fill=(1, 1, 1), overlay=True)
                if kind == "two":             # две строки мельче — внутри одной строки бланка
                    ty = F107_TEXT_TOP + row * F107_PITCH
                    if no is not None:
                        w.at((c[0] + c[1]) / 2, ty, str(no))
                        w.at((c[2] + c[3]) / 2, ty, str(q))
                        w.at((c[3] + c[4]) / 2, ty, _f107_value(v))
                    y0 = (top + bot) / 2 - len(lines) * F107_TWO_LH / 2
                    for j, line in enumerate(lines):
                        lt = y0 + j * F107_TWO_LH
                        base = lt + F107_TWO_LH / 2 + (w.font.ascender + w.font.descender) * size / 2
                        w.text(c[1] + 2.8, base, line, size)
                    continue
                if rows == 1:
                    ty = F107_TEXT_TOP + row * F107_PITCH
                    if no is not None:
                        w.at((c[0] + c[1]) / 2, ty, str(no))
                        w.at((c[2] + c[3]) / 2, ty, str(q))
                        w.at((c[3] + c[4]) / 2, ty, _f107_value(v))
                    base = (top + bot) / 2 + (w.font.ascender + w.font.descender) * size / 2
                    w.text(c[1] + 2.8, base, lines[0], size)
                    continue
                if no is not None:            # номер, количество и ценность — по центру объединённой ячейки
                    w.centered(c[0], c[1], top, bot, str(no))
                    w.centered(c[2], c[3], top, bot, str(q))
                    w.centered(c[3], c[4], top, bot, _f107_value(v))
                lh = size * 1.18
                y0 = (top + bot) / 2 - len(lines) * lh / 2
                for j, line in enumerate(lines):
                    lt = y0 + j * lh
                    base = lt + lh / 2 + (w.font.ascender + w.font.descender) * size / 2
                    w.text(c[1] + 2.8, base, line, size)
            if last:
                ty, cq, cv = F107_TOTAL
                w.at(cq + dx, ty, str(tot_q))
                w.at(cv + dx, ty, tot_v)
            x0, x1, l1, l2 = F107_SENDER
            lines = _wrap_to(w, sender or "", x1 - x0 - 4, F107_SIZE, 2) if sender else []
            for j, line in enumerate(lines):
                w.text(x0 + dx, (l1, l2)[j] - 1.2, line)
            sx, step, cw, top, bot = F107_SPI
            for k, ch in enumerate(spi):
                w.centered(sx + dx + k * step, sx + dx + k * step + cw, top, bot, ch)
    doc.set_metadata({"title": "Опись вложения ф. 107", "creator": "LegalHelper"})
    if out:
        C.save_pdf(doc, out)
    return doc


# ============================================================================
#  Обезличивание (152-ФЗ)
# ============================================================================
_SURN = r"[А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?"
_SURN_END = (r"[А-ЯЁ][а-яё]*(?:ов|ова|ев|ева|ёв|ёва|ин|ина|ын|ына|ский|ская|цкий|цкая|ской|ко|енко|ук|юк|"
             r"ян|дзе|швили|их|ых|ер|ман|ич|ец|ай|ий|ая|ен|ова)")
_PATR = r"[А-ЯЁ][а-яё]+(?:вич|вна|ична|инична|ич|оглы|кызы)"
ANON_PATTERNS = {
    "fio": ("ФИО полностью (Иванов Иван Иванович)",
            [rf"\b{_SURN}\s+[А-ЯЁ][а-яё]+\s+{_PATR}\b", rf"\b[А-ЯЁ][а-яё]+\s+{_PATR}\s+{_SURN}\b"], "[ФИО]"),
    "fio_short": ("ФИО с инициалами (Иванов И.И., И.И. Иванов)",
                  [rf"\b{_SURN}\s+[А-ЯЁ]\.\s?[А-ЯЁ]\.", rf"\b[А-ЯЁ]\.\s?[А-ЯЁ]\.\s?{_SURN_END}\b"], "[ФИО]"),
    "name_patr": ("Имя и отчество (Иван Иванович)", [rf"\b[А-ЯЁ][а-яё]+\s+{_PATR}\b"], "[ИО]"),
    "birth": ("Даты рождения", [r"(?i:дата рождения|года рождения|г\.\s?р\.|родил(?:ся|ась))[:\s]*(\d{2}\.\d{2}\.\d{4})",
                                r"(\d{2}\.\d{2}\.\d{4})\s*(?:г\.\s?р\.|года рождения)"], "[дата]"),
    "passport": ("Паспорт (серия, номер)", [r"\b\d{2}\s?\d{2}\s?№?\s?\d{6}\b"], "[паспорт]"),
    "snils": ("СНИЛС", [r"\b\d{3}-\d{3}-\d{3}[\s-]\d{2}\b"], "[СНИЛС]"),
    "inn": ("ИНН, ОГРН", [r"ИНН[\s:№]*(\d{12}|\d{10})\b", r"ОГРН(?:ИП)?[\s:№]*(\d{15}|\d{13})\b"], "[ИНН]"),
    "account": ("Банковские счета и карты", [r"\b\d{20}\b", r"\b(?:\d{4}[\s-]?){3}\d{4}\b"], "[счёт]"),
    "phone": ("Телефоны", [C.PATTERNS["phone"]], "[тел.]"),
    "email": ("E-mail", [C.PATTERNS["email"]], "[e-mail]"),
    "address": ("Адреса (ул., д., кв.)",
                [r"(?:ул\.|улица|пр-кт|пр-т|проспект|пер\.|переулок|б-р|бульвар|ш\.|шоссе|мкр\.?|микрорайон)"
                 r"\s*[А-ЯЁ0-9][^\n,;]{1,40}(?:,\s*(?:д\.|дом)\s*[\dА-Яа-я/\-]+)?"
                 r"(?:,\s*(?:корп\.|стр\.)\s*[\dА-Яа-я]+)?(?:,\s*(?:кв\.|квартира|оф\.|офис)\s*\d+)?"], "[адрес]"),
    "car": ("Госномера авто (А123ВС 777)", [r"\b[АВЕКМНОРСТУХ]\s?\d{3}\s?[АВЕКМНОРСТУХ]{2}\s?\d{2,3}\b"], "[госномер]"),
}
ANON_MODES = ["Чёрная плашка", "Метка вместо данных ([ФИО], [адрес]…)", "Белая заливка"]


def find_personal(doc, keys, custom=(), pages=None):
    """Найти персональные данные. Возвращает {pno: [(текст, метка)]}."""
    pages = range(doc.page_count) if pages is None else pages
    found = {}
    for pno in pages:
        text = doc[pno].get_text()
        hits = []
        for k in keys:
            label, pats, tag = ANON_PATTERNS[k]
            for pat in pats:
                for m in re.finditer(pat, text):
                    s = m.group(1) if m.groups() else m.group(0)
                    hits.append((s.strip(), tag))
        for w in custom:
            if w.strip() and w.strip() in text:
                hits.append((w.strip(), "[скрыто]"))
        if hits:
            found[pno] = list(dict.fromkeys(hits))
    return found


def anonymize(doc, keys, custom=(), mode=ANON_MODES[0], clear_meta=True, pages=None, progress=None):
    """Безвозвратно удалить найденные данные. Возвращает число скрытых фрагментов."""
    found = find_personal(doc, keys, custom, pages)
    total = 0
    labels_font = C._font_kwargs("Arial")
    for k, (pno, hits) in enumerate(found.items()):
        page = doc[pno]
        marks = []
        for s, tag in hits:
            for q in page.search_for(s, quads=True):
                r = q.rect
                fill = (0, 0, 0) if mode.startswith("Чёрная") else (1, 1, 1)
                page.add_redact_annot(q, fill=fill)
                marks.append((r, tag))
                total += 1
        page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS)
        # вложенные совпадения (например, «Иван Иванович» внутри полного ФИО) — одна метка на место
        keep = []
        for r, tag in sorted(marks, key=lambda m: -m[0].width * m[0].height):
            big = fitz.Rect(r)
            if not any(fitz.Rect(k[0]) + (-1, -1, 1, 1) & big == big for k in keep):
                keep.append((r, tag))
        marks = keep
        if mode.startswith("Метка"):
            for r, tag in marks:
                fs = max(5, min(10, r.height * 0.75))
                try:
                    page.insert_textbox(r + (0, -1, 40, 2), tag, fontsize=fs, color=(0.35, 0.35, 0.35),
                                        rotate=page.rotation, **labels_font)
                except Exception:
                    pass
        if progress:
            progress(k + 1, len(found))
    if clear_meta:
        doc.set_metadata({})
        try:
            doc.del_xml_metadata()
        except Exception:
            pass
    return total


# ============================================================================
#  Сравнение редакций (правки в стиле «было / стало»)
# ============================================================================
def extract_text_any(path):
    ext = Path(path).suffix.lower()
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        paras = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                paras.append(" | ".join(c.text for c in row.cells))
        return "\n".join(paras)
    if ext in (".txt", ".md"):
        return Path(path).read_text(encoding="utf-8", errors="replace")
    d = C.open_as_pdf(path)
    return "\n".join(p.get_text() for p in d)


def _paragraphs(text):
    """Склеить строки PDF в абзацы (перенос строки внутри предложения убираем)."""
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    out, cur = [], ""
    for line in text.splitlines():
        s = line.strip()
        if not s:
            if cur:
                out.append(cur)
                cur = ""
            continue
        if cur and (re.search(r"[.:;!?»)]$", cur) and (s[:1].isupper() or re.match(r"^\d+[.)]", s))):
            out.append(cur)
            cur = s
        else:
            cur = (cur + " " + s).strip()
    if cur:
        out.append(cur)
    return out


def _tokens(s):
    return re.findall(r"\s+|[\w\-]+|[^\w\s]", s)


def redline(old_text, new_text):
    """HTML-правки и статистика."""
    a, b = _paragraphs(old_text), _paragraphs(new_text)
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    parts, st = [], dict(added=0, deleted=0, changed=0)
    esc = html.escape
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            for p in a[i1:i2]:
                parts.append(f"<p class='same'>{esc(p)}</p>")
        elif op == "delete":
            for p in a[i1:i2]:
                parts.append(f"<p><del>{esc(p)}</del></p>")
                st["deleted"] += len(p.split())
        elif op == "insert":
            for p in b[j1:j2]:
                parts.append(f"<p><ins>{esc(p)}</ins></p>")
                st["added"] += len(p.split())
        else:
            st["changed"] += 1
            ta, tb = _tokens("\n".join(a[i1:i2])), _tokens("\n".join(b[j1:j2]))
            wm = difflib.SequenceMatcher(None, ta, tb, autojunk=False)
            buf = []
            for o, x1, x2, y1, y2 in wm.get_opcodes():
                sa, sb = "".join(ta[x1:x2]), "".join(tb[y1:y2])
                if o == "equal":
                    buf.append(esc(sa))
                else:
                    if sa.strip():
                        buf.append(f"<del>{esc(sa)}</del>")
                        st["deleted"] += len(sa.split())
                    if sb.strip():
                        buf.append(f"<ins>{esc(sb)}</ins>")
                        st["added"] += len(sb.split())
            parts.append("<p>" + "".join(buf).replace("\n", "<br/>") + "</p>")
    return "\n".join(parts), st


def compare_versions(old_path, new_path, out_base, only_changes=False):
    """Создать отчёт сравнения .html и .pdf. Возвращает (html_path, pdf_path, статистика)."""
    body, st = redline(extract_text_any(old_path), extract_text_any(new_path))
    if only_changes:
        body = "\n".join(l for l in body.splitlines() if "class='same'" not in l)
    head = (f"<h3>Сравнение редакций</h3><p>Было: <b>{html.escape(Path(old_path).name)}</b><br/>"
            f"Стало: <b>{html.escape(Path(new_path).name)}</b><br/>"
            f"Изменённых фрагментов: {st['changed']}; добавлено слов: {st['added']}; удалено слов: "
            f"{st['deleted']}</p><p><del>удалено</del> &nbsp; <ins>добавлено</ins></p><hr/>")
    page = ("<!DOCTYPE html><html><head><meta charset='utf-8'><title>Сравнение редакций</title><style>"
            "body{font-family:'Times New Roman',serif;font-size:13pt;max-width:900px;margin:30px auto;"
            "line-height:1.45} p.same{color:#333} ins{color:#0a7a28;background:#e3f6e8;text-decoration:underline}"
            "del{color:#b3261e;background:#fde7e5}</style></head><body>" + head + body + "</body></html>")
    hp = out_base + ".html"
    Path(hp).write_text(page, encoding="utf-8")
    pp = out_base + ".pdf"
    try:
        html_to_pdf_file(head + body, pp)
    except Exception:
        pp = None
    return hp, pp, st


# ============================================================================
#  Поиск по документам дела, цитаты
# ============================================================================
SEARCH_CACHE = None          # путь к кэшу текста (задаёт программа: папка данных/search_cache.sqlite)


def _cache_con():
    import sqlite3
    if not SEARCH_CACHE:
        return None
    con = sqlite3.connect(SEARCH_CACHE, timeout=5)
    con.execute("CREATE TABLE IF NOT EXISTS texts(path TEXT PRIMARY KEY, size INTEGER, mtime REAL, units TEXT)")
    return con


def file_text_units(p, con=None):
    """Текст файла кусками [(метка, текст)]: PDF — по страницам, Word (.docx) — по абзацам без конвертации
    через Office. Результат кэшируется по (размер, время изменения) — повторный поиск почти мгновенный."""
    import json as _json
    st = os.stat(p)
    if con is not None:
        r = con.execute("SELECT size, mtime, units FROM texts WHERE path=?", (p,)).fetchone()
        if r and r[0] == st.st_size and abs(r[1] - st.st_mtime) < 0.001:
            return [tuple(u) for u in _json.loads(r[2])]
    ext = Path(p).suffix.lower()
    units = []
    if ext == ".docx":
        import docx
        d = docx.Document(p)
        n = 0
        for par in _iter_paragraphs(d):
            if par.text.strip():
                n += 1
                units.append((f"абз. {n}", par.text))
    elif ext in (".txt", ".md", ".csv"):
        units = [("", Path(p).read_text(encoding="utf-8", errors="replace"))]
    else:
        d = fitz.open(p) if ext == ".pdf" else C.open_as_pdf(p)
        units = [(str(i + 1), page.get_text()) for i, page in enumerate(d)]
    if con is not None:
        con.execute("INSERT OR REPLACE INTO texts(path,size,mtime,units) VALUES (?,?,?,?)",
                    (p, st.st_size, st.st_mtime, _json.dumps(units, ensure_ascii=False)))
        con.commit()
    return units


def search_files(paths, query, case_sensitive=False, progress=None):
    """Возвращает (результаты [(путь, стр./абзац, фрагмент)], [пути со сканами без текста])."""
    res, no_text = [], []
    q = query if case_sensitive else query.lower()
    con = None
    try:
        con = _cache_con()
    except Exception:
        con = None
    try:
        for k, p in enumerate(paths):
            if progress:
                progress(k, len(paths))
            try:
                units = file_text_units(p, con)
            except Exception:
                continue
            if units and not any(t.strip() for _l, t in units):
                no_text.append(p)
                continue
            for label, t in units:
                tt = t if case_sensitive else t.lower()
                pos = tt.find(q)
                while pos >= 0:
                    a, b = max(0, pos - 60), min(len(t), pos + len(q) + 60)
                    res.append((p, label, re.sub(r"\s+", " ", t[a:b]).strip()))
                    pos = tt.find(q, pos + len(q))
                    if len(res) > 2000:
                        return res, no_text
    finally:
        if con is not None:
            con.close()
    return res, no_text


def quote_from_rect(page, rect):
    t = page.get_textbox(rect)
    return re.sub(r"[ \t]+", " ", re.sub(r"-\n(\w)", r"\1", t)).strip()


# ============================================================================
#  Шаблоны DOCX
# ============================================================================
PLACEHOLDER = re.compile(r"\{([^{}\n]{1,60})\}")


def _iter_paragraphs(d):
    def cells(tbls):
        for t in tbls:
            for row in t.rows:
                for c in row.cells:
                    yield from c.paragraphs
                    yield from cells(c.tables)
    yield from d.paragraphs
    yield from cells(d.tables)
    for s in d.sections:
        for part in (s.header, s.footer):
            yield from part.paragraphs
            yield from cells(part.tables)


def template_fields(path):
    import docx
    d = docx.Document(path)
    names = []
    for p in _iter_paragraphs(d):
        names += PLACEHOLDER.findall(p.text)
    return list(dict.fromkeys(n.strip() for n in names))


BLOCK_PREFIXES = ("Сторона_", "Пункты_", "Блок_")     # поля-блоки: занимают весь абзац, могут быть пустыми


def _strip_bullet(line):
    # «1) текст», «2. текст», «- текст» → «текст»; даты вроде «01.02.2026 — …» не трогаем
    return re.sub(r"^\s*(?:\d{1,3}[.)](?=\s)|[-–—•*](?=\s))\s*", "", line).strip()


def _expand_blocks(d, values):
    """Абзац, состоящий только из {Поля}: многострочное значение — отдельными абзацами с тем же оформлением
    («Пункты_…» — с нумерацией 1), 2)…); пустое значение поля-блока — абзац убирается."""
    import copy
    for p in list(_iter_paragraphs(d)):
        m = re.fullmatch(r"\s*\{([^{}\n]{1,60})\}\s*", p.text)
        if not m:
            continue
        name = m.group(1).strip()
        if name not in values:
            continue
        val = str(values[name] or "")
        lines = [x.rstrip() for x in val.split("\n")]
        if name.startswith("Пункты_"):
            lines = [f"{i}) {_strip_bullet(x)}" for i, x in enumerate([x for x in lines if x.strip()], 1)]
            if not lines:                        # пустой раздел плана — строки, чтобы дописать от руки
                lines = ["_" * 58, "_" * 58]
        else:
            lines = [x for x in lines if x.strip()] if name.startswith(BLOCK_PREFIXES) else lines
        if not lines:
            if name.startswith(BLOCK_PREFIXES):
                p._p.getparent().remove(p._p)
            continue
        runs = p.runs
        runs[0].text = lines[0]
        for r in runs[1:]:
            r.text = ""
        anchor = p._p
        for line in lines[1:]:
            el = copy.deepcopy(p._p)
            anchor.addnext(el)
            anchor = el
            from docx.text.paragraph import Paragraph
            q = Paragraph(el, p._parent)
            q.paragraph_format.space_before = 0          # отступ сверху — только у первой строки блока
            q.runs[0].text = line
            for r in q.runs[1:]:
                r.text = ""


def fill_template(path, values, out):
    """Заменить {Поле} значениями. Форматирование берётся у первого фрагмента абзаца с полем."""
    import docx
    d = docx.Document(path)
    _expand_blocks(d, values)
    for p in _iter_paragraphs(d):
        full = p.text
        if "{" not in full:
            continue
        new = PLACEHOLDER.sub(lambda m: str(values.get(m.group(1).strip(), m.group(0))), full)
        if new == full:
            continue
        # стараемся заменить внутри отдельных фрагментов, чтобы сохранить оформление
        done = True
        for r in p.runs:
            if "{" in r.text or "}" in r.text:
                t = PLACEHOLDER.sub(lambda m: str(values.get(m.group(1).strip(), m.group(0))), r.text)
                if PLACEHOLDER.search(t) or t.count("{") != t.count("}"):
                    done = False
                    break
                r.text = t
        if not done or p.text != new:
            runs = p.runs
            if runs:
                runs[0].text = new
                for r in runs[1:]:
                    r.text = ""
    d.save(out)
    return out


def ru_date_words(d):
    months = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа", "сентября",
              "октября", "ноября", "декабря"]
    return f"«{d.day:02d}» {months[d.month - 1]} {d.year} г."


def act_docx(path, case, entries, rate_default=0.0, executor="", client=""):
    """Акт об оказанной юридической помощи по записям учёта времени."""
    import docx
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("АКТ\nоб оказанной юридической помощи").bold = True
    d.add_paragraph(f"г. ____________\t\t\t\t\t\t{ru_date_words(dt.date.today())}")
    d.add_paragraph(f"Доверитель: {client or case.get('client', '')}")
    d.add_paragraph(f"Исполнитель: {executor}")
    if case.get("number") or case.get("court"):
        d.add_paragraph(f"Дело: {case.get('number', '')} {case.get('court', '')}".strip())
    d.add_paragraph("Исполнитель оказал, а Доверитель принял следующую юридическую помощь:")
    t = d.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for i, h in enumerate(["№", "Дата", "Содержание работы", "Часы", "Сумма, ₽"]):
        t.rows[0].cells[i].text = h
    total_h, total_s = 0.0, 0.0
    for i, e in enumerate(entries, 1):
        hours = float(e.get("hours") or 0)
        rate = float(e.get("rate") or rate_default or 0)
        s = float(e.get("amount") or 0) or hours * rate
        total_h += hours
        total_s += s
        c = t.add_row().cells
        c[0].text, c[1].text, c[2].text = str(i), e.get("date", ""), e.get("description", "")
        c[3].text, c[4].text = (f"{hours:g}" if hours else ""), money(s)
    c = t.add_row().cells
    c[2].text, c[3].text, c[4].text = "Итого", f"{total_h:g}", money(total_s)
    rub = int(total_s)
    kop = int(round((total_s - rub) * 100))
    d.add_paragraph(f"Итого: {money(total_s)} ₽ ({num_words(rub)} {plural(rub, 'рубль', 'рубля', 'рублей')} "
                    f"{kop:02d} коп.). Доверитель претензий по объёму, качеству и срокам не имеет.")
    d.add_paragraph("\nИсполнитель: ____________ /____________/\t\tДоверитель: ____________ /____________/")
    d.save(path)
    return total_h, total_s


def add_toc(doc, entries):
    """entries — [(заголовок, номер первой страницы 1-based)] — добавить закладки верхнего уровня."""
    toc = doc.get_toc(simple=True) or []
    toc += [[1, t, p] for t, p in entries]
    toc.sort(key=lambda x: x[2])
    doc.set_toc(toc)
