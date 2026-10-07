# -*- coding: utf-8 -*-
"""Вкладки карточки дела: «Подача» (чек-лист комплекта документов), «Карта дела» (Excalidraw),
«Нормы права» (дерево применяемых актов). Всё сохраняется в базе дел сразу при изменении."""
import os
import sys
import html
import uuid
from pathlib import Path

import pymupdf as fitz
from PySide6.QtCore import Qt, QTimer, QUrl, QSize, QEventLoop
from PySide6.QtGui import QImage, QPixmap, QColor, QFont, QDesktopServices, QBrush
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QFormLayout, QLabel, QPushButton, QToolButton, QComboBox, QLineEdit,
    QPlainTextEdit, QTableWidget, QTableWidgetItem, QHeaderView, QAbstractItemView, QSpinBox, QSplitter,
    QScrollArea, QFrame, QMenu, QInputDialog, QMessageBox, QFileDialog, QTreeWidget, QTreeWidgetItem,
    QProgressBar, QApplication, QStackedWidget, QDialog, QListWidget, QListWidgetItem, QCompleter, QDialogButtonBox)

import pdf_core as C
import laws_auto as LA

U = None          # legal_ui (задаётся при импорте из legal_ui)


def db():
    return U.db()


# =============================================================================
#  Типовые перечни приложений
# =============================================================================
TEMPLATES = {
    "Иск в арбитражный суд (ст. 125–126 АПК)": [
        "Исковое заявление",
        "Уведомление о вручении / документ о направлении копий иска и приложений ответчику и третьим лицам",
        "Документ об уплате государственной пошлины",
        "Документы, подтверждающие обстоятельства, на которых основаны требования",
        "Расчёт взыскиваемой суммы",
        "Выписка из ЕГРЮЛ/ЕГРИП на истца и ответчика (не ранее 30 дней до подачи)",
        "Документы, подтверждающие соблюдение претензионного порядка",
        "Доверенность или иной документ о полномочиях на подписание иска",
    ],
    "Иск в суд общей юрисдикции (ст. 131–132 ГПК)": [
        "Исковое заявление",
        "Документ, подтверждающий направление другим лицам копий иска и приложений",
        "Документ об уплате государственной пошлины",
        "Доверенность и документ о высшем юридическом образовании представителя",
        "Документы, подтверждающие обстоятельства, на которых основаны требования",
        "Документы о соблюдении досудебного порядка (если он предусмотрен)",
        "Расчёт взыскиваемой (оспариваемой) суммы",
    ],
    "Административный иск (ст. 125–126 КАС)": [
        "Административное исковое заявление",
        "Документ о вручении другим лицам копий иска и приложений",
        "Документ об уплате государственной пошлины",
        "Доверенность и документ о высшем юридическом образовании представителя",
        "Документы, подтверждающие обстоятельства, на которых основаны требования",
    ],
    "Отзыв на иск (ст. 131 АПК)": [
        "Отзыв на исковое заявление",
        "Документ о направлении отзыва и приложений другим лицам",
        "Документы, подтверждающие возражения",
        "Доверенность",
    ],
    "Апелляционная жалоба (ст. 260 АПК)": [
        "Апелляционная жалоба",
        "Документ об уплате государственной пошлины",
        "Документ о направлении копий жалобы другим лицам",
        "Доверенность",
    ],
    "Апелляционная жалоба (ст. 322 ГПК)": [
        "Апелляционная жалоба",
        "Документ об уплате государственной пошлины",
        "Документ о направлении копий жалобы и приложений другим лицам",
        "Доверенность и документ о высшем юридическом образовании представителя",
    ],
    "Претензия": [
        "Претензия",
        "Документы, подтверждающие требования",
        "Расчёт суммы требований",
        "Доверенность",
    ],
}

COMMON_ACTS = [
    "Гражданский кодекс РФ", "Арбитражный процессуальный кодекс РФ", "Гражданский процессуальный кодекс РФ",
    "Кодекс административного судопроизводства РФ", "Налоговый кодекс РФ", "Трудовой кодекс РФ",
    "Жилищный кодекс РФ", "Кодекс РФ об административных правонарушениях",
    "Закон РФ от 07.02.1992 № 2300-1 «О защите прав потребителей»",
    "Постановление Пленума ВС РФ от 28.06.2012 № 17 (защита прав потребителей)",
    "Постановление Пленума ВС РФ от 24.03.2016 № 7 (ответственность за нарушение обязательств)",
    "Постановление Пленума ВС РФ от 23.06.2015 № 25 (общие положения ГК РФ)",
    "Постановление Пленума ВС РФ от 21.01.2016 № 1 (судебные издержки)",
    "Обзор судебной практики ВС РФ", "Определение ВС РФ", "Постановление Конституционного Суда РФ",
]
KIND_NAMES = {"group": "Направление / тема", "act": "Нормативный акт / постановление", "norm": "Статья, пункт, позиция"}
KIND_GLYPH = {"group": "▣", "act": "▤", "norm": "§"}


def _btn(text, slot, primary=False, tip=None):
    b = QPushButton(text)
    if primary:
        b.setObjectName("primary")
    if tip:
        b.setToolTip(tip)
    b.clicked.connect(slot)
    return b


def _page_count(path, _cache={}):
    try:
        st = os.stat(path)
    except OSError:
        return None
    key = (path, st.st_mtime, st.st_size)
    if key not in _cache:
        n = None
        try:
            if path.lower().endswith(".pdf") or Path(path).suffix.lower() in (".xps", ".epub", ".fb2"):
                with fitz.open(path) as d:
                    n = d.page_count
            elif Path(path).suffix.lower() in (".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp", ".gif"):
                n = 1
        except Exception:
            n = None
        _cache[key] = n
    return _cache[key]


# =============================================================================
#  Просмотр документа (страницы по ширине окна)
# =============================================================================
class PdfPreview(QScrollArea):
    def __init__(self):
        super().__init__()
        self.setWidgetResizable(True)
        self.setFrameShape(QFrame.NoFrame)
        self.setObjectName("canvasArea")
        self.host = QWidget()
        self.lay = QVBoxLayout(self.host)
        self.lay.setAlignment(Qt.AlignHCenter | Qt.AlignTop)
        self.lay.setSpacing(12)
        self.lay.setContentsMargins(16, 16, 16, 16)
        self.setWidget(self.host)
        self.doc = None
        self.labels = []
        self.next = 0
        self.timer = QTimer(self)
        self.timer.setInterval(0)
        self.timer.timeout.connect(self._render_next)
        self.rtimer = QTimer(self)
        self.rtimer.setSingleShot(True)
        self.rtimer.setInterval(250)
        self.rtimer.timeout.connect(self._rerender)
        self.message("Выберите документ в списке справа")

    def _clear(self):
        self.timer.stop()
        while self.lay.count():
            w = self.lay.takeAt(0).widget()
            if w:
                w.hide()
                w.setParent(None)
                w.deleteLater()
        self.labels = []

    def message(self, text):
        self._clear()
        if self.doc is not None:
            self.doc.close()
            self.doc = None
        lab = QLabel(text)
        lab.setObjectName("hint")
        lab.setStyleSheet(f"color: {U.M.T['muted']}; font-size: 11pt;")
        lab.setAlignment(Qt.AlignCenter)
        lab.setWordWrap(True)
        self.lay.addWidget(lab)

    def set_file(self, path):
        if not path:
            return self.message("Файл ещё не прикреплён.\nНажмите «Прикрепить файл…» или перетащите файл на строку.")
        if not os.path.exists(path):
            return self.message(f"Файл не найден:\n{path}")
        try:
            doc = C.open_as_pdf(path, None)
        except Exception as e:
            return self.message(f"Не удалось показать файл:\n{e}")
        if doc is None:
            return self.message("Файл защищён паролем — откройте его в редакторе.")
        self.message("")
        self._clear()
        self.doc = doc
        n = min(doc.page_count, 150)
        for _ in range(n):
            lab = QLabel()
            lab.setAlignment(Qt.AlignCenter)
            self.lay.addWidget(lab)
            self.labels.append(lab)
        if doc.page_count > n:
            more = QLabel(f"… ещё {doc.page_count - n} стр. — откройте документ в редакторе")
            more.setStyleSheet("color: white;")
            self.lay.addWidget(more)
        self.next = 0
        self.timer.start()

    def _render_next(self):
        if not self.doc or self.next >= len(self.labels):
            self.timer.stop()
            return
        i = self.next
        self.next += 1
        page = self.doc[i]
        w = max(200, self.viewport().width() - 60)
        dpr = self.devicePixelRatioF()
        zoom = w / page.rect.width * dpr
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        img = QImage(pix.samples, pix.width, pix.height, pix.stride, QImage.Format_RGB888).copy()
        pm = QPixmap.fromImage(img)
        pm.setDevicePixelRatio(dpr)
        self.labels[i].setPixmap(pm)

    def _rerender(self):
        if self.doc:
            self.next = 0
            self.timer.start()

    def resizeEvent(self, e):
        super().resizeEvent(e)
        if self.doc:
            self.rtimer.start()


# =============================================================================
#  Вкладка «Подача»
# =============================================================================
class ChecklistTable(QTableWidget):
    def __init__(self, tab):
        super().__init__(0, 5)
        self.tab = tab
        self.setAcceptDrops(True)
        self.viewport().setAcceptDrops(True)
        self.setDragDropMode(QAbstractItemView.DropOnly)

    def dragEnterEvent(self, e):
        if e.mimeData().hasUrls():
            e.acceptProposedAction()

    def dragMoveEvent(self, e):
        if e.mimeData().hasUrls():
            e.acceptProposedAction()

    def dropEvent(self, e):
        paths = [u.toLocalFile() for u in e.mimeData().urls() if u.isLocalFile()]
        if not paths:
            return
        row = self.rowAt(e.position().toPoint().y())
        e.acceptProposedAction()
        self.tab.dropped(paths, row)


class SubmissionTab(QWidget):
    COLS = ["✓", "Документ", "Листов", "Экз.", "Файл"]

    def __init__(self, main):
        super().__init__()
        self.main = main
        self.cid = None
        self.pid = None
        self._loading = False
        v = QVBoxLayout(self)
        v.setContentsMargins(0, 6, 0, 0)
        top = QHBoxLayout()
        top.addWidget(QLabel("Комплект:"))
        self.packs = QComboBox()
        self.packs.setMinimumWidth(260)
        self.packs.currentIndexChanged.connect(self.on_pack)
        top.addWidget(self.packs)
        mb = QToolButton()
        mb.setText("⋯")
        mb.setToolTip("Новый комплект, переименовать, удалить")
        mb.setPopupMode(QToolButton.InstantPopup)
        m = QMenu(mb)
        m.addAction("Новый комплект…", self.new_pack)
        m.addAction("Переименовать…", self.rename_pack)
        m.addSeparator()
        m.addAction("Удалить комплект", self.delete_pack)
        mb.setMenu(m)
        top.addWidget(mb)
        top.addWidget(U.HelpButton("submission"))
        top.addStretch(1)
        self.progress = QProgressBar()
        self.progress.setFixedWidth(160)
        self.progress.setTextVisible(False)
        self.progress.setFixedHeight(8)
        self.prog_lbl = QLabel()
        top.addWidget(self.prog_lbl)
        top.addWidget(self.progress)
        top.addWidget(_btn("Собрать пакет…", self.build, primary=True,
                           tip="Открыть «Пакет в суд / на почту» с документами этого комплекта"))
        v.addLayout(top)

        split = QSplitter()
        split.setHandleWidth(1)
        # слева — просмотр
        left = QWidget()
        lv = QVBoxLayout(left)
        lv.setContentsMargins(0, 0, 0, 0)
        hl = QHBoxLayout()
        self.pv_title = QLabel("")
        self.pv_title.setObjectName("subtitle")
        hl.addWidget(self.pv_title, 1)
        hl.addWidget(_btn("Открыть в редакторе", self.open_in_editor))
        lv.addLayout(hl)
        self.preview = PdfPreview()
        lv.addWidget(self.preview, 1)
        split.addWidget(left)
        # справа — чек-лист
        right = QFrame()
        right.setObjectName("card")
        rv = QVBoxLayout(right)
        rv.setContentsMargins(12, 12, 12, 12)
        rv.addWidget(QLabel("<b>Документы комплекта</b> — первая строка основной документ"))
        self.t = ChecklistTable(self)
        self.t.setHorizontalHeaderLabels(self.COLS)
        h = self.t.horizontalHeader()
        h.setSectionResizeMode(1, QHeaderView.Stretch)
        for c, wdt in ((0, 34), (2, 62), (3, 58), (4, 110)):
            h.setSectionResizeMode(c, QHeaderView.Fixed)
            h.resizeSection(c, wdt)
        self.t.verticalHeader().setVisible(False)
        self.t.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.t.setSelectionMode(QAbstractItemView.SingleSelection)
        self.t.setWordWrap(True)
        self.t.itemChanged.connect(self.on_item_changed)
        self.t.currentCellChanged.connect(lambda r, *_: self.show_row(r))
        self.t.cellDoubleClicked.connect(lambda r, c: self.attach(r) if c == 4 else None)
        rv.addWidget(self.t, 1)
        b1 = QHBoxLayout()
        b1.addWidget(_btn("+ Файлы…", self.add_files, tip="Добавить документы из файлов (можно перетащить в список)"))
        b1.addWidget(_btn("+ Пункт", self.add_item, tip="Пункт без файла — то, что ещё нужно подготовить"))
        tb = QPushButton("Типовой перечень")
        tm = QMenu(tb)
        for name in TEMPLATES:
            tm.addAction(name, lambda n=name: self.apply_template(n))
        tb.setMenu(tm)
        b1.addWidget(tb)
        b1.addStretch(1)
        rv.addLayout(b1)
        b2 = QHBoxLayout()
        b2.addWidget(_btn("Прикрепить файл…", lambda: self.attach(self.t.currentRow())))
        b2.addWidget(_btn("▲", lambda: self.move(-1), tip="Выше"))
        b2.addWidget(_btn("▼", lambda: self.move(1), tip="Ниже"))
        b2.addWidget(_btn("Удалить", self.delete_item))
        b2.addStretch(1)
        rv.addLayout(b2)
        self.hint = QLabel("Отмечайте ✓ по мере готовности. Всё сохраняется автоматически.")
        self.hint.setObjectName("hint")
        self.hint.setWordWrap(True)
        rv.addWidget(self.hint)
        split.addWidget(right)
        split.setStretchFactor(0, 3)
        split.setStretchFactor(1, 2)
        split.setSizes([620, 520])
        v.addWidget(split, 1)
        self.stack_empty = None

    # ---------------------------------------------------------------- загрузка
    def set_case(self, cid):
        self.cid = cid
        packs = db().packs(cid) if cid else []
        if cid and not packs:
            db().add_pack(cid, "Исковое заявление")
            packs = db().packs(cid)
        self._loading = True
        self.packs.clear()
        for p in packs:
            self.packs.addItem(p["title"], p["id"])
        last = int(U.M.settings().value(f"pack_{cid}", 0) or 0)
        idx = max(0, self.packs.findData(last))
        self._loading = False
        self.packs.setCurrentIndex(idx)
        self.on_pack()

    def on_pack(self, *_):
        if self._loading:
            return
        self.pid = self.packs.currentData()
        if self.cid and self.pid:
            U.M.settings().setValue(f"pack_{self.cid}", self.pid)
        self.reload()

    def rows(self):
        return db().pack_items(self.pid) if self.pid else []

    def reload(self, select=None):
        items = self.rows()
        cur = select if select is not None else self.t.currentRow()
        self._loading = True
        self.t.setRowCount(0)
        bold = QFont()
        bold.setBold(True)
        for r, it in enumerate(items):
            self.t.insertRow(r)
            chk = QTableWidgetItem()
            chk.setFlags(Qt.ItemIsEnabled | Qt.ItemIsUserCheckable | Qt.ItemIsSelectable)
            chk.setCheckState(Qt.Checked if it["done"] else Qt.Unchecked)
            chk.setData(Qt.UserRole, it["id"])
            self.t.setItem(r, 0, chk)
            ti = QTableWidgetItem(it["title"])
            ti.setToolTip(it["title"] + ("\n\nОсновной документ" if r == 0 else ""))
            if r == 0:
                ti.setFont(bold)
            self.t.setItem(r, 1, ti)
            path = it["path"]
            n = _page_count(path) if path else None
            si = QTableWidgetItem(str(n) if n else "—")
            si.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable)
            si.setTextAlignment(Qt.AlignCenter)
            self.t.setItem(r, 2, si)
            cp = QSpinBox()
            cp.setRange(1, 50)
            cp.setValue(it["copies"] or 1)
            cp.valueChanged.connect(lambda val, iid=it["id"]: db().update_pack_item(iid, copies=val))
            cp.setFixedHeight(30)
            box = QWidget()
            bl = QHBoxLayout(box)
            bl.setContentsMargins(4, 0, 4, 0)
            bl.addWidget(cp)
            box.setProperty("spin", cp)
            self.t.setCellWidget(r, 3, box)
            if not path:
                fi = QTableWidgetItem("＋ прикрепить")
                fi.setForeground(QBrush(QColor("#8e8e93")))
                fi.setToolTip("Дважды щёлкните, чтобы прикрепить файл")
            elif not os.path.exists(path):
                fi = QTableWidgetItem("⚠ не найден")
                fi.setForeground(QBrush(QColor("#ff3b30")))
                fi.setToolTip(path)
            else:
                fi = QTableWidgetItem(Path(path).name)
                fi.setToolTip(path)
            fi.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable)
            self.t.setItem(r, 4, fi)
        self.t.resizeRowsToContents()
        self._loading = False
        self.update_progress(items)
        if items:
            self.t.setCurrentCell(min(max(cur, 0), len(items) - 1), 1)
            self.show_row(self.t.currentRow())
        else:
            self.preview.message("Комплект пуст.\nДобавьте основной документ (иск, жалобу) и приложения —\n"
                                 "кнопкой «+ Файлы…», перетаскиванием или «Типовой перечень».")
            self.pv_title.setText("")

    def update_progress(self, items=None):
        items = self.rows() if items is None else items
        done = sum(1 for i in items if i["done"])
        self.progress.setMaximum(max(1, len(items)))
        self.progress.setValue(done)
        self.prog_lbl.setText(f"Готово {done} из {len(items)}  " if items else "")

    def item_id(self, row):
        it = self.t.item(row, 0)
        return it.data(Qt.UserRole) if it else None

    def item_row(self, row):
        iid = self.item_id(row)
        return next((i for i in self.rows() if i["id"] == iid), None)

    def show_row(self, row):
        it = self.item_row(row)
        if not it:
            return
        self.pv_title.setText(it["title"])
        self.preview.set_file(it["path"])

    # ---------------------------------------------------------------- правка
    def on_item_changed(self, item):
        if self._loading:
            return
        iid = self.item_id(item.row())
        if item.column() == 0:
            db().update_pack_item(iid, done=1 if item.checkState() == Qt.Checked else 0)
            self.update_progress()
        elif item.column() == 1:
            db().update_pack_item(iid, title=item.text().strip())

    def _title_from(self, path):
        return Path(path).stem.replace("_", " ").strip()

    def add_paths(self, paths, at_row=-1):
        if not self.pid:
            return
        for p in paths:
            db().add_pack_item(self.pid, self._title_from(p), p)
            try:
                if not any(d["path"] == p for d in db().docs(self.cid)):
                    db().add_doc(self.cid, p, self._title_from(p))
            except Exception:
                pass
        self.reload(select=len(self.rows()) - 1)

    def dropped(self, paths, row):
        it = self.item_row(row) if row >= 0 else None
        if it and not it["path"] and len(paths) == 1:
            db().update_pack_item(it["id"], path=paths[0])
            self.reload(select=row)
        else:
            self.add_paths(paths)

    def start_dir(self):
        c = db().case(self.cid) or {}
        return c.get("folder") or self.main.default_dir()

    def add_files(self):
        paths, _ = QFileDialog.getOpenFileNames(self, "Документы комплекта", self.start_dir(), U.M.OPEN_FILTER)
        if paths:
            self.add_paths(paths)

    def add_item(self):
        name, ok = QInputDialog.getText(self, "Новый пункт", "Название документа (как в перечне приложений):")
        if ok and name.strip():
            db().add_pack_item(self.pid, name.strip())
            self.reload(select=len(self.rows()) - 1)

    def apply_template(self, name):
        existing = {i["title"].lower() for i in self.rows()}
        added = 0
        for title in TEMPLATES[name]:
            if title.lower() not in existing:
                db().add_pack_item(self.pid, title)
                added += 1
        self.reload()
        self.main.msg(f"Добавлено пунктов: {added}. Лишнее удалите, названия можно поправить двойным щелчком.")

    def attach(self, row):
        it = self.item_row(row)
        if not it:
            return
        p, _ = QFileDialog.getOpenFileName(self, f"Файл для «{it['title']}»", self.start_dir(), U.M.OPEN_FILTER)
        if p:
            db().update_pack_item(it["id"], path=p)
            if it["title"] in ("", "Новый документ"):
                db().update_pack_item(it["id"], title=self._title_from(p))
            self.reload(select=row)

    def move(self, step):
        r = self.t.currentRow()
        ids = [i["id"] for i in self.rows()]
        n = r + step
        if r < 0 or not (0 <= n < len(ids)):
            return
        ids[r], ids[n] = ids[n], ids[r]
        db().reorder_pack_items(ids)
        self.reload(select=n)

    def delete_item(self):
        it = self.item_row(self.t.currentRow())
        if it and QMessageBox.question(self, U.M.APP_NAME, f"Убрать из комплекта «{it['title']}»?\n"
                                       "Сам файл на диске не удаляется.") == QMessageBox.Yes:
            snap = db().snapshot(("pack_items", "id=?", (it["id"],)))
            db().delete_pack_item(it["id"])
            self.reload()

            def undo():
                db().restore(snap)
                self.reload()
            self.main.toast_undo(f"«{it['title']}» убрано из комплекта", undo)

    def open_in_editor(self):
        it = self.item_row(self.t.currentRow())
        if it and it["path"]:
            self.main.open_external(it["path"], 0)

    # ---------------------------------------------------------------- комплекты
    def new_pack(self):
        name, ok = QInputDialog.getItem(self, "Новый комплект", "Что подаём:",
                                        ["Исковое заявление", "Отзыв на иск", "Апелляционная жалоба",
                                         "Кассационная жалоба", "Ходатайство", "Претензия"], 0, True)
        if ok and name.strip():
            pid = db().add_pack(self.cid, name.strip())
            U.M.settings().setValue(f"pack_{self.cid}", pid)
            self.set_case(self.cid)

    def rename_pack(self):
        if not self.pid:
            return
        name, ok = QInputDialog.getText(self, "Переименовать", "Название комплекта:",
                                        text=self.packs.currentText())
        if ok and name.strip():
            db().update_pack(self.pid, title=name.strip())
            self.set_case(self.cid)

    def delete_pack(self):
        if self.pid and QMessageBox.question(self, U.M.APP_NAME, f"Удалить комплект «{self.packs.currentText()}»?\n"
                                             "Файлы на диске не удаляются.") == QMessageBox.Yes:
            pid, cid, name = self.pid, self.cid, self.packs.currentText()
            snap = db().snapshot(("packs", "id=?", (pid,)), ("pack_items", "pack_id=?", (pid,)))
            db().delete_pack(pid)
            self.set_case(cid)

            def undo():
                db().restore(snap)
                U.M.settings().setValue(f"pack_{cid}", pid)
                if self.cid == cid:
                    self.set_case(cid)
            self.main.toast_undo(f"Комплект «{name}» удалён", undo)

    def build(self):
        items = [i for i in self.rows() if i["path"] and os.path.exists(i["path"])]
        missing = [i["title"] for i in self.rows() if not i["path"] or not os.path.exists(i["path"])]
        not_done = [i["title"] for i in self.rows() if not i["done"]]
        if not items:
            return QMessageBox.information(self, U.M.APP_NAME, "В комплекте нет прикреплённых файлов.")
        warn = []
        if missing:
            warn.append("Без файла (в пакет не попадут):\n• " + "\n• ".join(missing[:10]))
        if not_done:
            warn.append("Не отмечены как готовые:\n• " + "\n• ".join(not_done[:10]))
        if warn and QMessageBox.question(self, U.M.APP_NAME, "\n\n".join(warn) + "\n\nВсё равно собрать пакет?") \
                != QMessageBox.Yes:
            return
        dlg = U.PackageDialog(self.main)
        for it in items:
            dlg.add_path(it["path"], it["title"])
            dlg.t.cellWidget(dlg.t.rowCount() - 1, 2).setValue(it["copies"] or 1)
        c = db().case(self.cid) or {}
        if c.get("folder"):
            dlg.out.setText(os.path.join(c["folder"], f"Пакет — {self.packs.currentText()}"))
        dlg.exec()


def assemble_pdf(items, out, progress=None):
    """Сложить документы комплекта (PDF, Word, картинки…) в один PDF с закладками по документам.
    items — [(id, название, путь)]. Возвращает ({id: (первая, последняя страница)}, [не удалось])."""
    doc = fitz.open()
    toc, ranges, failed = [], {}, []
    for n, (iid, title, path) in enumerate(items):
        if progress:
            progress(n, len(items))
        try:
            d = C.open_as_pdf(path, None)
        except Exception as e:
            failed.append(f"{title}: {e}")
            continue
        if d is None or not d.page_count:
            failed.append(f"{title}: не удалось открыть (пароль или пустой файл)")
            continue
        start = doc.page_count
        doc.insert_pdf(d)
        gid, nm = fitz.get_pdf_str(uuid.uuid4().hex[:10]), fitz.get_pdf_str(title)
        for i in range(start, doc.page_count):       # страницы документа сворачиваются в рабочей области
            doc.xref_set_key(doc[i].xref, "LHGroup", gid)
            doc.xref_set_key(doc[i].xref, "LHName", nm)
            doc.xref_set_key(doc[i].xref, "LHSrc", fitz.get_pdf_str(os.path.abspath(path)))
        ranges[iid] = (start, doc.page_count - 1)
        toc.append([1, title, start + 1])
    if not doc.page_count:
        raise ValueError("Ни один документ комплекта не удалось открыть.")
    doc.set_toc(toc)
    doc.set_metadata({"title": os.path.splitext(os.path.basename(out))[0], "creator": "LegalHelper"})
    C.save_pdf(doc, out)
    return ranges, failed


class KitPanel(SubmissionTab):
    """Комплект для подачи — во вкладке «Документы», слева от рабочей области: список документов комплекта,
    «Собрать в один PDF» (открывается справа) и сверка по пунктам: щелчок — к первой странице документа."""
    COLS = ["✓", "Документ", "Стр.", "Экз."]

    def __init__(self, main):
        QWidget.__init__(self)
        self.main = main
        self.cid = None
        self.pid = None
        self._loading = False
        self.ranges = {}
        self.kit_path = None
        v = QVBoxLayout(self)
        v.setContentsMargins(0, 4, 0, 0)
        v.setSpacing(8)
        top = QHBoxLayout()
        self.packs = QComboBox()
        self.packs.currentIndexChanged.connect(self.on_pack)
        top.addWidget(self.packs, 1)
        mb = QToolButton()
        mb.setText("⋯")
        mb.setObjectName("moretabs")
        mb.setToolTip("Новый комплект, переименовать, удалить")
        mb.setPopupMode(QToolButton.InstantPopup)
        m = QMenu(mb)
        m.addAction("Новый комплект…", self.new_pack)
        m.addAction("Переименовать…", self.rename_pack)
        m.addSeparator()
        m.addAction("Удалить комплект", self.delete_pack)
        mb.setMenu(m)
        top.addWidget(mb)
        top.addWidget(U.HelpButton("submission"))
        v.addLayout(top)
        pr = QHBoxLayout()
        self.prog_lbl = QLabel()
        self.prog_lbl.setObjectName("hint")
        pr.addWidget(self.prog_lbl)
        self.progress = QProgressBar()
        self.progress.setTextVisible(False)
        self.progress.setFixedHeight(6)
        pr.addWidget(self.progress, 1)
        v.addLayout(pr)
        self.t = ChecklistTable(self)
        self.t.setColumnCount(len(self.COLS))
        self.t.setObjectName("docs")
        self.t.setHorizontalHeaderLabels(self.COLS)
        h = self.t.horizontalHeader()
        h.setSectionResizeMode(1, QHeaderView.Stretch)
        for c, wdt in ((0, 34), (2, 70), (3, 58)):
            h.setSectionResizeMode(c, QHeaderView.Fixed)
            h.resizeSection(c, wdt)
        self.t.verticalHeader().setVisible(False)
        self.t.setShowGrid(False)
        self.t.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.t.setSelectionMode(QAbstractItemView.SingleSelection)
        self.t.setWordWrap(True)
        self.t.setEditTriggers(QAbstractItemView.EditKeyPressed)
        self.t.itemChanged.connect(self.on_item_changed)
        self.t.cellClicked.connect(lambda r, c: self.show_row(r))
        self.t.cellDoubleClicked.connect(self.on_double)
        v.addWidget(self.t, 1)
        b1 = QHBoxLayout()
        for text, fn, tip in (("+ Из дела…", self.add_from_case, "Выбрать документы из списка дела"),
                              ("+ Файл…", self.add_files, "Добавить файлы с компьютера"),
                              ("+ Пункт", self.add_item, "Пункт без файла — то, что ещё нужно подготовить")):
            b = _btn(text, fn, tip=tip)
            b.setObjectName("compact")
            b1.addWidget(b)
        tb = QPushButton("Перечень ▾")
        tb.setObjectName("compact")
        tb.setToolTip("Типовой перечень приложений по АПК / ГПК / КАС")
        tm = QMenu(tb)
        for name in TEMPLATES:
            tm.addAction(name, lambda n=name: self.apply_template(n))
        tb.setMenu(tm)
        b1.addWidget(tb)
        v.addLayout(b1)
        b2 = QHBoxLayout()
        for text, fn, tip in (("▲", lambda: self.move(-1), "Выше"), ("▼", lambda: self.move(1), "Ниже"),
                              ("Прикрепить файл…", lambda: self.attach(self.t.currentRow()), "Файл для выбранного пункта"),
                              ("Удалить", self.delete_item, "Убрать пункт из комплекта (файл не удаляется)")):
            b = _btn(text, fn, tip=tip)
            b.setObjectName("compact")
            b2.addWidget(b)
        b2.addStretch(1)
        v.addLayout(b2)
        self.b_assemble = _btn("📑  Собрать в один PDF и проверить", self.assemble, primary=True,
                               tip="Все документы комплекта — одним PDF справа, по порядку, с закладками")
        v.addWidget(self.b_assemble)
        b3 = QHBoxLayout()
        b = _btn("📨 Пакет для подачи…", self.build, tip="Файлы под «Мой арбитр», ГАС «Правосудие», Почту России")
        b.setObjectName("compact")
        b3.addWidget(b)
        b = _btn("📮 Опись ф. 107", self.to_f107, tip="Опись вложения по документам комплекта")
        b.setObjectName("compact")
        b3.addWidget(b)
        v.addLayout(b3)
        self.hint = QLabel("Перетащите документы из списка дела сюда. Первая строка — основной документ. "
                           "Соберите комплект в один PDF — он откроется справа; щёлкайте по пунктам, "
                           "сверяйте страницы и отмечайте ✓.")
        self.hint.setObjectName("hint")
        self.hint.setWordWrap(True)
        v.addWidget(self.hint)

    def reload(self, select=None):
        items = self.rows()
        cur = select if select is not None else self.t.currentRow()
        self._loading = True
        self.t.setRowCount(0)
        bold = QFont()
        bold.setBold(True)
        for r, it in enumerate(items):
            self.t.insertRow(r)
            chk = QTableWidgetItem()
            chk.setFlags(Qt.ItemIsEnabled | Qt.ItemIsUserCheckable | Qt.ItemIsSelectable)
            chk.setCheckState(Qt.Checked if it["done"] else Qt.Unchecked)
            chk.setData(Qt.UserRole, it["id"])
            chk.setToolTip("Проверено")
            self.t.setItem(r, 0, chk)
            path = it["path"]
            ok = bool(path) and os.path.exists(path)
            ti = QTableWidgetItem(it["title"] + ("" if ok else ("  ⚠ файл не найден" if path else "  ＋ файл")))
            ti.setData(Qt.UserRole, it["title"])
            ti.setToolTip((path or "Файл не прикреплён — «Прикрепить файл…» или перетащите файл на строку") +
                          ("\n\nОсновной документ" if r == 0 else ""))
            if r == 0:
                ti.setFont(bold)
            if not ok:
                ti.setForeground(QBrush(QColor(U.M.T["danger"] if path else U.M.T["muted"])))
            self.t.setItem(r, 1, ti)
            rng = self.ranges.get(it["id"])
            n = _page_count(path) if ok else None
            text = (f"{rng[0] + 1}–{rng[1] + 1}" if rng and rng[1] > rng[0] else f"{rng[0] + 1}") if rng else \
                (f"{n} л." if n else "—")
            si = QTableWidgetItem(text)
            si.setToolTip("Страницы в собранном PDF" if rng else "Листов в документе")
            si.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable)
            si.setTextAlignment(Qt.AlignCenter)
            self.t.setItem(r, 2, si)
            cp = QSpinBox()
            cp.setRange(1, 50)
            cp.setValue(it["copies"] or 1)
            cp.valueChanged.connect(lambda val, iid=it["id"]: db().update_pack_item(iid, copies=val))
            self.t.setCellWidget(r, 3, cp)
        self.t.resizeRowsToContents()
        self._loading = False
        self.update_progress(items)
        if items:
            self.t.setCurrentCell(min(max(cur, 0), len(items) - 1), 1)

    def on_item_changed(self, item):
        if self._loading:
            return
        iid = self.item_id(item.row())
        if item.column() == 0:
            db().update_pack_item(iid, done=1 if item.checkState() == Qt.Checked else 0)
            self.update_progress()
        elif item.column() == 1:
            db().update_pack_item(iid, title=item.text().split("  ⚠")[0].split("  ＋")[0].strip())

    def update_progress(self, items=None):
        items = self.rows() if items is None else items
        done = sum(1 for i in items if i["done"])
        self.progress.setMaximum(max(1, len(items)))
        self.progress.setValue(done)
        self.prog_lbl.setText(f"Проверено {done} из {len(items)}" if items else "Комплект пуст")

    def set_case(self, cid):
        if cid != self.cid:
            self.ranges, self.kit_path = {}, None
        super().set_case(cid)

    def on_pack(self, *_):
        self.ranges, self.kit_path = {}, None
        super().on_pack()

    def show_row(self, row):
        """Щелчок по пункту: если комплект собран и открыт — к первой странице этого документа."""
        it = self.item_row(row)
        if not it:
            return
        rng = self.ranges.get(it["id"])
        if rng and self.kit_path and self.main.ws[self.main.cur_ws].get("kit"):
            self.main.goto_page(rng[0])

    def on_double(self, r, c):
        if c == 1:
            it = self.item_row(r)
            if it and it["path"] and os.path.exists(it["path"]):
                self.main.open_external(it["path"], 0, self.cid)
            elif it:
                self.attach(r)

    def add_from_case(self):
        docs = [d for d in db().docs(self.cid) if d["path"]]
        if not docs:
            return QMessageBox.information(self, U.M.APP_NAME, "В деле пока нет документов.")
        have = {i["path"] for i in self.rows()}
        dlg = QDialog(self)
        dlg.setWindowTitle("Документы в комплект")
        dlg.resize(460, 480)
        dv = QVBoxLayout(dlg)
        dv.addWidget(QLabel("Отметьте документы, которые войдут в комплект:"))
        from PySide6.QtWidgets import QListWidget, QListWidgetItem
        lst = QListWidget()
        for d in docs:
            li = QListWidgetItem(f"{d.get('icon') or ''}  {d['title']}".strip())
            li.setData(Qt.UserRole, d["path"])
            li.setFlags(li.flags() | Qt.ItemIsUserCheckable)
            li.setCheckState(Qt.Unchecked)
            if d["path"] in have:
                li.setFlags(li.flags() & ~Qt.ItemIsEnabled)
                li.setText(li.text() + "  — уже в комплекте")
            lst.addItem(li)
        dv.addWidget(lst, 1)
        row = QHBoxLayout()
        row.addStretch(1)
        c = QPushButton("Отмена")
        c.clicked.connect(dlg.reject)
        ok = QPushButton("Добавить")
        ok.setObjectName("primary")
        ok.clicked.connect(dlg.accept)
        row.addWidget(c)
        row.addWidget(ok)
        dv.addLayout(row)
        if not dlg.exec():
            return
        paths = [lst.item(i).data(Qt.UserRole) for i in range(lst.count())
                 if lst.item(i).checkState() == Qt.Checked]
        if paths:
            titles = {d["path"]: d["title"] for d in docs}
            for p in paths:
                db().add_pack_item(self.pid, titles.get(p) or self._title_from(p), p)
            self.reload(select=len(self.rows()) - 1)

    def assemble(self):
        rows = self.rows()
        items = [(i["id"], i["title"], i["path"]) for i in rows if i["path"] and os.path.exists(i["path"])]
        if not items:
            return QMessageBox.information(self, U.M.APP_NAME, "В комплекте нет документов с файлами. "
                                           "Перетащите документы из списка дела или нажмите «+ Из дела…».")
        import tempfile
        name = U.L.clean_filename(f"Комплект — {self.packs.currentText()}")
        out = os.path.join(tempfile.mkdtemp(prefix="lh_kit_"), name + ".pdf")   # не мусорим в папке дела
        res = self.main.run("Собираю комплект в один PDF…", assemble_pdf, items, out)
        if res is U.M.FAILED:
            return
        self.ranges, failed = res
        self.kit_path = out
        self.main.open_kit(out, name, self.cid)
        self.reload(select=0)
        skipped = [i["title"] for i in rows if not i["path"] or not os.path.exists(i["path"])]
        msg = []
        if skipped:
            msg.append("Без файла (не вошли):\n• " + "\n• ".join(skipped[:10]))
        if failed:
            msg.append("Не удалось добавить:\n• " + "\n• ".join(failed[:10]))
        if msg:
            QMessageBox.warning(self, U.M.APP_NAME, "\n\n".join(msg))
        else:
            self.main.msg(f"Комплект собран: {os.path.basename(out)}. Щёлкайте по пунктам слева — "
                          "откроется начало каждого документа. Нужен файл — «Сохранить».", 10000)

    def to_f107(self):
        rows = [(f"{i['title']}" + (f" на {_page_count(i['path'])} л." if i["path"] and os.path.exists(i["path"])
                                   and _page_count(i["path"]) else ""), i["copies"] or 1, 1) for i in self.rows()]
        if not rows:
            return QMessageBox.information(self, U.M.APP_NAME, "Комплект пуст.")
        U.F107Dialog(self.main, rows).exec()


# =============================================================================
#  Вкладка «Карта дела» (Excalidraw)
# =============================================================================
def excalidraw_dir():
    return os.path.join(getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__))), "excalidraw")


class BoardTab(QWidget):
    def __init__(self, main):
        super().__init__()
        self.main = main
        self.cid = None
        self.view = None
        self.ready = False
        self.pending = None
        self.loaded = None                 # дело, чья карта сейчас открыта в редакторе
        v = QVBoxLayout(self)
        v.setContentsMargins(0, 6, 0, 0)
        top = QHBoxLayout()
        cap = QLabel("Интеллект-карта: ход дела, позиции сторон, доказательства, риски")
        cap.setWordWrap(True)                          # длинные подписи переносятся, а не раздвигают окно
        cap.setMinimumWidth(10)
        top.addWidget(cap, 1)
        top.addWidget(U.HelpButton("board"))
        self.state = QLabel("")
        self.state.setObjectName("hint")
        self.state.setWordWrap(True)
        self.state.setMinimumWidth(10)
        top.addWidget(self.state)
        top.addWidget(_btn("Новая карта из шаблона", self.new_map))
        self.b_win = _btn("⛶  На весь экран", self.fullscreen, tip="Карта на весь экран (F11). Esc или F11 — "
                                                                  "вернуть во вкладку.")
        top.addWidget(self.b_win)
        v.addLayout(top)
        self.host = QVBoxLayout()
        v.addLayout(self.host, 1)
        self.placeholder = QLabel("Карта загружается…")
        self.placeholder.setAlignment(Qt.AlignCenter)
        self.placeholder.setObjectName("hint")
        self.host.addWidget(self.placeholder)
        self.timer = QTimer(self)
        self.timer.setInterval(2000)
        self.timer.timeout.connect(self.autosave)
        # встроенный браузер занимает сотни мегабайт памяти: если карту не открывали несколько минут,
        # он сохраняет карту и закрывается; при возврате на вкладку открывается снова сам
        self.release_timer = QTimer(self)
        self.release_timer.setSingleShot(True)
        self.release_timer.setInterval(self.RELEASE_MS)
        self.release_timer.timeout.connect(self.release)

    RELEASE_MS = 3 * 60 * 1000

    def hideEvent(self, e):
        super().hideEvent(e)
        if self.view is not None:
            self.release_timer.start()

    def release(self):
        """Освободить память встроенного браузера (карта сначала сохраняется)."""
        if self.view is None or self.isVisible() or getattr(self, "win", None) is not None:
            return
        self.flush()
        self.timer.stop()
        view, self.view = self.view, None
        self.ready = False
        self.pending = self.cid
        self.loaded = None
        try:
            view.loadFinished.disconnect(self.on_loaded)
        except (RuntimeError, TypeError):
            pass
        view.setParent(None)
        view.deleteLater()
        self.placeholder.setText("Карта загружается…")
        self.placeholder.show()

    def _ensure_view(self):
        if self.view is not None:
            return True
        try:
            from PySide6.QtWebEngineWidgets import QWebEngineView
            from PySide6.QtWebEngineCore import QWebEngineSettings
        except Exception as e:
            self.placeholder.setText(f"Компонент для карт недоступен в этой сборке:\n{e}")
            return False
        idx = os.path.join(excalidraw_dir(), "index.html")
        if not os.path.exists(idx):
            self.placeholder.setText("Не найдены файлы редактора карт (папка excalidraw).")
            return False
        self.view = QWebEngineView()
        s = self.view.settings()
        s.setAttribute(QWebEngineSettings.LocalContentCanAccessFileUrls, True)
        s.setAttribute(QWebEngineSettings.LocalContentCanAccessRemoteUrls, False)
        prof = self.view.page().profile()
        try:
            prof.downloadRequested.connect(self.on_download)
        except Exception:
            pass
        self.view.loadFinished.connect(self.on_loaded)
        self.placeholder.hide()
        self.host.addWidget(self.view, 1)
        self.view.load(QUrl.fromLocalFile(idx))
        return True

    def on_loaded(self, ok):
        self.ready = ok
        if ok and self.pending is not None:
            cid = self.pending
            self.pending = None
            self._load(cid)

    def set_case(self, cid):
        if cid == self.cid and self.view is not None:
            return
        self.flush()
        self.cid = cid
        if not cid or not self.isVisible():
            self.pending = cid
            return
        self.activate()

    def showEvent(self, e):
        super().showEvent(e)
        self.release_timer.stop()
        QTimer.singleShot(60, self.activate)       # сначала показать «Карта загружается…»

    def activate(self):
        """Вызывается, когда вкладка становится видимой. Пока карту не видно, встроенный браузер
        не создаётся — он тяжёлый и раньше замедлял запуск программы."""
        if not self.isVisible():
            return
        if not self._ensure_view():
            return
        cid = self.cid
        if self.ready:
            if cid != self.loaded:            # та же карта уже открыта — не перечитывать (правки и масштаб целы)
                self._load(cid)
        else:
            self.pending = cid

    def _load(self, cid):
        if not cid:
            return
        scene = db().board(cid) or ""
        title = (db().case(cid) or {}).get("title", "Дело")
        theme = U.M.T.get("name", "light")
        js = f"PM.load({_js(scene)}, {_js(theme)}, {_js(title)})"
        self.view.page().runJavaScript(js)
        self.loaded = cid
        self.state.setText("Сохранено" if scene else "Новая карта — сохранится автоматически")
        self.timer.start()

    def autosave(self, cb=None):
        if not (self.view and self.ready and self.cid):
            if cb:
                cb()
            return
        cid = self.cid

        def got(res):
            if res and db().case(cid):           # дело могли удалить — карту «в пустоту» не пишем
                db().save_board(cid, res)
                self.state.setText("Сохранено ✓")
            if cb:
                cb()
        self.view.page().runJavaScript("PM.takeIfDirty()", 0, got)

    def flush(self, timeout_ms=1500):
        """Синхронно сохранить текущую карту (перед сменой дела и при выходе)."""
        if not (self.view and self.ready and self.cid):
            return
        loop = QEventLoop()
        self.autosave(loop.quit)
        QTimer.singleShot(timeout_ms, loop.quit)
        loop.exec()

    def new_map(self):
        if not (self.view and self.ready):
            return
        if QMessageBox.question(self, U.M.APP_NAME, "Заменить текущую карту шаблоном?\n"
                                "(Отменить можно сочетанием Ctrl+Z на карте.)") != QMessageBox.Yes:
            return
        title = (db().case(self.cid) or {}).get("title", "Дело")
        self.view.page().runJavaScript(f"PM.newMindmap({_js(title)})")

    def set_theme(self, name):
        if self.view and self.ready:
            self.view.page().runJavaScript(f"PM.setTheme({_js(name)})")

    def fullscreen(self):
        """Карта на весь экран; «Вернуть во вкладку», Esc или F11 — обратно во вкладку."""
        if not self._ensure_view():
            return
        if getattr(self, "win", None) is not None:
            self.win.raise_()
            self.win.activateWindow()
            return
        tab = self

        class BoardWindow(QWidget):
            def closeEvent(self, e):
                tab.autosave()
                lay = self.layout()
                lay.removeWidget(tab.view)
                tab.host.addWidget(tab.view, 1)
                tab.view.show()
                tab.b_win.setEnabled(True)
                tab.win = None
                if not tab.isVisible():
                    tab.release_timer.start()
                e.accept()

        from PySide6.QtGui import QShortcut, QKeySequence
        self.win = BoardWindow()
        self.win.setAttribute(Qt.WA_DeleteOnClose)
        title = (db().case(self.cid) or {}).get("title", "Дело") if self.cid else "Дело"
        self.win.setWindowTitle(f"Карта дела — {title}")
        self.win.setWindowIcon(self.main.windowIcon())
        lay = QVBoxLayout(self.win)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(0)
        bar = QFrame()
        bar.setObjectName("boardbar")
        bar.setStyleSheet("QFrame#boardbar { background: palette(window); border-bottom: 1px solid rgba(128,128,128,.25); }")
        bh = QHBoxLayout(bar)
        bh.setContentsMargins(14, 6, 8, 6)
        cap = QLabel(f"🗺️  Карта дела — {title}")
        f = cap.font()
        f.setBold(True)
        cap.setFont(f)
        bh.addWidget(cap)
        bh.addStretch(1)
        hint = QLabel("Esc или F11 — вернуть во вкладку")
        hint.setObjectName("hint")
        bh.addWidget(hint)
        back = _btn("✕  Вернуть во вкладку", self.win.close, primary=True)
        bh.addWidget(back)
        lay.addWidget(bar)
        self.host.removeWidget(self.view)
        lay.addWidget(self.view, 1)
        for key in ("Esc", "F11"):
            sc = QShortcut(QKeySequence(key), self.win)
            sc.setContext(Qt.WindowShortcut)
            sc.activated.connect(self.win.close)
        self.b_win.setEnabled(False)
        self.win.showFullScreen()
        self.view.setFocus()

    def shutdown(self):
        """Перед выходом: остановить автосохранение (сама карта уже сохранена в flush)."""
        self.timer.stop()
        if getattr(self, "win", None) is not None:
            self.win.close()

    def on_download(self, item):
        name = item.downloadFileName() if hasattr(item, "downloadFileName") else "карта"
        p, _ = QFileDialog.getSaveFileName(self, "Сохранить", os.path.join(self.main.default_dir(), name))
        if not p:
            item.cancel()
            return
        item.setDownloadDirectory(os.path.dirname(p))
        item.setDownloadFileName(os.path.basename(p))
        item.accept()


def _js(s):
    import json
    return json.dumps(s if s is not None else "")


# =============================================================================
#  Вкладка «Нормы права»
# =============================================================================
class LawTree(QTreeWidget):
    def __init__(self, tab):
        super().__init__()
        self.tab = tab
        self.setHeaderHidden(True)
        self.setDragDropMode(QAbstractItemView.InternalMove)
        self.setSelectionMode(QAbstractItemView.SingleSelection)
        self.setIndentation(18)
        self.setObjectName("lawtree")
        self.setUniformRowHeights(False)

    def dropEvent(self, e):
        super().dropEvent(e)
        self.tab.save_structure()


class LawsTab(QWidget):
    def __init__(self, main):
        super().__init__()
        self.main = main
        self.cid = None
        self.cur = None
        self._loading = False
        v = QVBoxLayout(self)
        v.setContentsMargins(0, 6, 0, 0)
        top = QHBoxLayout()
        self.filter = QLineEdit()
        self.filter.setPlaceholderText("Поиск по нормам…")
        self.filter.textChanged.connect(self.apply_filter)
        top.addWidget(self.filter, 1)
        top.addWidget(U.HelpButton("laws"))
        ex = QPushButton("Экспорт")
        em = QMenu(ex)
        em.addAction("Скопировать списком (для иска)", self.copy_list)
        em.addAction("Сохранить в Word…", self.to_word)
        em.addSeparator()
        em.addAction("Взять из другого дела…", self.import_from_case)
        ex.setMenu(em)
        top.addWidget(ex)
        v.addLayout(top)
        # быстрое добавление: ссылка как в иске → акт и статья сами
        quick = QHBoxLayout()
        self.quick = QLineEdit()
        self.quick.setPlaceholderText("➕  Добавить нормы: напишите или вставьте «п. 1 ст. 395 ГК РФ, ст. 110 АПК РФ» и нажмите Enter")
        self.quick.returnPressed.connect(self.quick_add)
        self._completer_model = None
        quick.addWidget(self.quick, 1)
        quick.addWidget(_btn("Добавить", self.quick_add, primary=True))
        quick.addWidget(_btn("🔎 Найти в документах дела", self.find_in_docs,
                             tip="Программа просмотрит иск, решения и другие документы дела и найдёт все ссылки на статьи"))
        quick.addWidget(_btn("📚 Мои нормы", self.pick_library,
                             tip="Нормы, которые вы уже добавляли в другие дела (с текстом и комментарием), и частые статьи"))
        v.addLayout(quick)
        split = QSplitter()
        split.setHandleWidth(1)
        left = QWidget()
        lv = QVBoxLayout(left)
        lv.setContentsMargins(0, 0, 0, 0)
        self.tree = LawTree(self)
        self.tree.currentItemChanged.connect(self.on_current)
        lv.addWidget(self.tree, 1)
        b = QHBoxLayout()
        b.addWidget(_btn("+ Направление", lambda: self.add("group"), tip="Раздел: например, «Неустойка», «Моральный вред»"))
        b.addWidget(_btn("+ Акт", lambda: self.add("act"), tip="Закон, кодекс, постановление Пленума, определение"))
        b.addWidget(_btn("+ Статья / пункт", lambda: self.add("norm"), tip="Конкретная статья, пункт, правовая позиция"))
        lv.addLayout(b)
        b2 = QHBoxLayout()
        b2.addWidget(_btn("▲", lambda: self.move(-1)))
        b2.addWidget(_btn("▼", lambda: self.move(1)))
        b2.addWidget(_btn("Удалить", self.delete))
        b2.addStretch(1)
        lv.addLayout(b2)
        split.addWidget(left)
        # редактор
        right = QFrame()
        right.setObjectName("card")
        f = QFormLayout(right)
        f.setContentsMargins(14, 14, 14, 14)
        self.kind = QComboBox()
        for k, n in KIND_NAMES.items():
            self.kind.addItem(f"{KIND_GLYPH[k]}  {n}", k)
        self.title = QComboBox()
        self.title.setEditable(True)
        self.title.setInsertPolicy(QComboBox.NoInsert)
        self.title.lineEdit().setPlaceholderText("Напр.: ст. 13, п. 6 — штраф 50%")
        self.url = QLineEdit()
        self.url.setPlaceholderText("Ссылка (КонсультантПлюс, Гарант, pravo.gov.ru, kad.arbitr.ru …)")
        urow = QHBoxLayout()
        urow.addWidget(self.url, 1)
        urow.addWidget(_btn("Открыть", lambda: self.url.text().strip() and
                            QDesktopServices.openUrl(QUrl.fromUserInput(self.url.text().strip()))))
        self.body = QPlainTextEdit()
        self.body.setPlaceholderText("Текст нормы или правовой позиции (вставьте из КонсультантПлюс/Гарант)")
        self.note = QPlainTextEdit()
        self.note.setPlaceholderText("Как применяю в деле: к какому доводу, какие факты подпадают, риски толкования")
        self.note.setMaximumHeight(140)
        f.addRow("Тип", self.kind)
        f.addRow("Название", self.title)
        f.addRow("Ссылка", urow)
        f.addRow("Текст", self.body)
        f.addRow("Применение", self.note)
        self.editor = right
        split.addWidget(right)
        split.setSizes([430, 620])
        v.addWidget(split, 1)
        self.save_timer = QTimer(self)
        self.save_timer.setSingleShot(True)
        self.save_timer.setInterval(400)
        self.save_timer.timeout.connect(self.save_current)
        for sig in (self.title.currentTextChanged, self.url.textChanged, self.body.textChanged,
                    self.note.textChanged):
            sig.connect(self.schedule)
        self.kind.currentIndexChanged.connect(self.on_kind)
        right.setEnabled(False)

    # ---------------------------------------------------------------- загрузка
    def set_case(self, cid):
        self.save_now()
        self.cid = cid
        self.reload()

    def reload(self, select=None):
        self._loading = True
        self.tree.clear()
        rows = db().laws(self.cid) if self.cid else []
        by_parent = {}
        for r in rows:
            by_parent.setdefault(r["parent_id"], []).append(r)
        items = {}

        def build(parent_item, pid):
            for r in sorted(by_parent.get(pid, []), key=lambda x: (x["pos"], x["id"])):
                it = QTreeWidgetItem([self.label(r)])
                it.setData(0, Qt.UserRole, r["id"])
                it.setData(0, Qt.UserRole + 1, r["kind"])
                it.setToolTip(0, (r["body"] or "")[:600])
                self.style_item(it, r["kind"])
                if parent_item is None:
                    self.tree.addTopLevelItem(it)
                else:
                    parent_item.addChild(it)
                items[r["id"]] = it
                build(it, r["id"])
        build(None, 0)
        self.tree.expandAll()
        self._loading = False
        target = items.get(select) if select else None
        if target is None and self.tree.topLevelItemCount():
            target = self.tree.topLevelItem(0)
        if target:
            self.tree.setCurrentItem(target)
        else:
            self.on_current(None)
        self.apply_filter()

    @staticmethod
    def label(r):
        t = r["title"] or {"group": "Новое направление", "act": "Новый акт", "norm": "Новая статья"}[r["kind"]]
        return f"{KIND_GLYPH.get(r['kind'], '')}  {t}"

    @staticmethod
    def style_item(it, kind):
        f = it.font(0)
        if kind == "group":
            f.setBold(True)
            f.setPointSizeF(f.pointSizeF() + 0.5)
        elif kind == "act":
            f.setWeight(QFont.DemiBold)
        it.setFont(0, f)

    def on_current(self, it, _prev=None):
        if self._loading:
            return
        self.save_now()
        self.cur = it.data(0, Qt.UserRole) if it else None
        r = next((x for x in db().laws(self.cid) if x["id"] == self.cur), None) if self.cur else None
        self._loading = True
        self.editor.setEnabled(bool(r))
        self.title.clear()
        if r:
            self.kind.setCurrentIndex(max(0, self.kind.findData(r["kind"])))
            if r["kind"] == "act":
                self.title.addItems(COMMON_ACTS)
            self.title.setCurrentText(r["title"])
            self.url.setText(r["url"])
            self.body.setPlainText(r["body"])
            self.note.setPlainText(r["note"])
        else:
            for w in (self.url,):
                w.clear()
            self.body.clear()
            self.note.clear()
        self._loading = False

    def on_kind(self):
        if self._loading or not self.cur:
            return
        k = self.kind.currentData()
        db().update_law(self.cur, kind=k)
        it = self.tree.currentItem()
        if it:
            it.setData(0, Qt.UserRole + 1, k)
            self.style_item(it, k)
            it.setText(0, self.label(dict(kind=k, title=self.title.currentText().strip())))
        self._loading = True
        t = self.title.currentText()
        self.title.clear()
        if k == "act":
            self.title.addItems(COMMON_ACTS)
        self.title.setCurrentText(t)
        self._loading = False

    def schedule(self, *_):
        if not self._loading and self.cur:
            self.save_timer.start()
            it = self.tree.currentItem()
            if it:
                it.setText(0, self.label(dict(kind=self.kind.currentData(), title=self.title.currentText().strip())))

    def save_current(self):
        if not self.cur:
            return
        db().update_law(self.cur, title=self.title.currentText().strip(), url=self.url.text().strip(),
                        body=self.body.toPlainText(), note=self.note.toPlainText())

    def save_now(self):
        if self.save_timer.isActive():
            self.save_timer.stop()
            self.save_current()

    # ---------------------------------------------------------------- структура
    def add(self, kind):
        if not self.cid:
            return
        self.save_now()
        it = self.tree.currentItem()
        parent = 0
        if it is not None:
            ck = it.data(0, Qt.UserRole + 1)
            cid_ = it.data(0, Qt.UserRole)
            order = ["group", "act", "norm"]
            if order.index(kind) > order.index(ck):          # внутрь выбранного
                parent = cid_
            elif kind == ck or order.index(kind) < order.index(ck):
                p = it.parent()                               # рядом / на уровень выше
                while p is not None and order.index(p.data(0, Qt.UserRole + 1)) >= order.index(kind):
                    p = p.parent()
                parent = p.data(0, Qt.UserRole) if p is not None else 0
        title = ""
        if kind == "norm" and it is None:
            title = ""
        nid = db().add_law(self.cid, parent, kind, title)
        self.reload(select=nid)
        self.title.setFocus()

    def save_structure(self):
        def walk(item, parent_id):
            for i in range(item.childCount()):
                ch = item.child(i)
                db().update_law(ch.data(0, Qt.UserRole), parent_id=parent_id, pos=i)
                walk(ch, ch.data(0, Qt.UserRole))
        root = self.tree.invisibleRootItem()
        walk(root, 0)

    def move(self, step):
        it = self.tree.currentItem()
        if not it:
            return
        par = it.parent() or self.tree.invisibleRootItem()
        i = par.indexOfChild(it)
        n = i + step
        if not (0 <= n < par.childCount()):
            return
        par.takeChild(i)
        par.insertChild(n, it)
        it.setExpanded(True)
        self.tree.setCurrentItem(it)
        self.save_structure()

    def delete(self):
        it = self.tree.currentItem()
        if not it:
            return
        extra = " вместе со всем вложенным" if it.childCount() else ""
        if QMessageBox.question(self, U.M.APP_NAME, f"Удалить «{it.text(0).strip()}»{extra}?") == QMessageBox.Yes:
            self.save_timer.stop()
            self.save_now()
            lid, cid = it.data(0, Qt.UserRole), self.cid
            name = next((r["title"] for r in db().laws(cid) if r["id"] == lid), "") or it.text(0).strip()
            ids = db().law_subtree(lid)
            snap = db().snapshot(("laws", f"id IN ({','.join('?' * len(ids))})", ids))
            db().delete_law(lid)
            self.cur = None
            self.reload()

            def undo():
                db().restore(snap)
                if self.cid == cid:
                    self.reload(select=lid)
            self.main.toast_undo(f"«{name}» удалено", undo)

    def apply_filter(self, *_):
        q = self.filter.text().strip().lower()
        rows = {r["id"]: r for r in db().laws(self.cid)} if self.cid else {}

        def visit(item):
            r = rows.get(item.data(0, Qt.UserRole), {})
            hit = not q or any(q in (r.get(k) or "").lower() for k in ("title", "body", "note"))
            child_hit = False
            for i in range(item.childCount()):
                child_hit = visit(item.child(i)) or child_hit
            item.setHidden(not (hit or child_hit))
            return hit or child_hit
        for i in range(self.tree.topLevelItemCount()):
            visit(self.tree.topLevelItem(i))

    # ---------------------------------------------------------------- экспорт
    def structured(self):
        rows = db().laws(self.cid)
        by_parent = {}
        for r in rows:
            by_parent.setdefault(r["parent_id"], []).append(r)
        out = []

        def walk(pid, depth):
            for r in sorted(by_parent.get(pid, []), key=lambda x: (x["pos"], x["id"])):
                out.append((depth, r))
                walk(r["id"], depth + 1)
        walk(0, 0)
        return out

    def copy_list(self):
        from PySide6.QtCore import QMimeData
        items = self.structured()
        if not items:
            return
        plain, htm = [], []
        for depth, r in items:
            t = r["title"] or ""
            ind = "    " * depth
            if r["kind"] == "group":
                plain.append(f"{ind}{t.upper()}")
                htm.append(f"<p style='margin:10px 0 4px'><b>{html.escape(t)}</b></p>")
            elif r["kind"] == "act":
                plain.append(f"{ind}{t}")
                htm.append(f"<p style='margin:6px 0 2px {depth * 18}px'><b>{html.escape(t)}</b></p>")
            else:
                body = (r["body"] or "").strip()
                line = f"{t}" + (f": «{body}»" if body else "")
                plain.append(f"{ind}— {line}")
                htm.append(f"<p style='margin:2px 0 2px {depth * 18}px'>— <i>{html.escape(t)}</i>"
                           + (f": «{html.escape(body)}»" if body else "") + "</p>")
            if r["note"].strip():
                plain.append(f"{ind}    Применение: {r['note'].strip()}")
                htm.append(f"<p style='margin:0 0 4px {depth * 18 + 18}px;color:#555'>"
                           f"Применение: {html.escape(r['note'].strip())}</p>")
        md = QMimeData()
        md.setText("\n".join(plain))
        md.setHtml("".join(htm))
        QApplication.clipboard().setMimeData(md)
        self.main.msg("Нормы скопированы — вставьте в документ (Ctrl+V)")

    def to_word(self):
        items = self.structured()
        if not items:
            return
        c = db().case(self.cid) or {}
        p, _ = QFileDialog.getSaveFileName(self, "Сохранить в Word",
                                           os.path.join(c.get("folder") or self.main.default_dir(),
                                                        "Нормы права — " + (c.get("title") or "дело") + ".docx"),
                                           "Word (*.docx)")
        if not p:
            return
        import docx
        from docx.shared import Pt, Cm
        d = docx.Document()
        d.styles["Normal"].font.name = "Times New Roman"
        d.styles["Normal"].font.size = Pt(12)
        d.add_heading(f"Применимые нормы права: {c.get('title', '')}", level=1)
        for depth, r in items:
            if r["kind"] == "group":
                d.add_heading(r["title"] or "", level=2)
                continue
            para = d.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.8 * depth)
            run = para.add_run(r["title"] or "")
            run.bold = r["kind"] == "act"
            run.italic = r["kind"] == "norm"
            if r["body"].strip():
                q = d.add_paragraph(f"«{r['body'].strip()}»")
                q.paragraph_format.left_indent = Cm(0.8 * depth + 0.6)
            if r["note"].strip():
                n = d.add_paragraph()
                n.paragraph_format.left_indent = Cm(0.8 * depth + 0.6)
                rr = n.add_run("Применение: " + r["note"].strip())
                rr.italic = True
            if r["url"].strip():
                u = d.add_paragraph(r["url"].strip())
                u.paragraph_format.left_indent = Cm(0.8 * depth + 0.6)
        d.save(p)
        QDesktopServices.openUrl(QUrl.fromLocalFile(p))

    # ---------------------------------------------------------------- автоматическое добавление
    def _library(self):
        return LA.library(db())

    def _refresh_completer(self):
        lib = self._library()
        self._lib = {f"{e['label']} {LA.short_act(e['act'])}": e for e in lib if e["act"]}
        labels = []
        for k, e in self._lib.items():
            about = e.get("about") or (e["body"][:60].replace("\n", " ") if e["body"] else "")
            uses = f" · в {e['uses']} {U.L.plural(e['uses'], 'деле', 'делах', 'делах')}" if e["uses"] else ""
            labels.append(k + (f" — {about}" if about else "") + uses)
        self._lib_labels = dict(zip(labels, self._lib.values()))
        comp = QCompleter(labels, self.quick)
        comp.setCaseSensitivity(Qt.CaseInsensitive)
        comp.setFilterMode(Qt.MatchContains)
        comp.setMaxVisibleItems(12)
        comp.activated.connect(lambda text: QTimer.singleShot(0, lambda: self._add_from_label(text)))
        self.quick.setCompleter(comp)

    def showEvent(self, e):
        super().showEvent(e)
        QTimer.singleShot(0, self._refresh_completer)

    def _add_from_label(self, text):
        e = self._lib_labels.get(text)
        if e:
            self.quick.clear()
            self.add_entries([e])

    def quick_add(self):
        text = self.quick.text().strip()
        if not text or not self.cid:
            return
        if text in getattr(self, "_lib_labels", {}):
            return self._add_from_label(text)
        cites = LA.parse(text)
        if not cites:
            QMessageBox.information(self, "Нормы права", "Не нашёл ссылок на статьи. Пишите так же, как в иске: "
                                    "«п. 1 ст. 395 ГК РФ», «ст. 110 АПК РФ», «ст. 13 Закона о защите прав "
                                    "потребителей», «п. 75 постановления Пленума ВС РФ от 24.03.2016 № 7».")
            return
        self.quick.clear()
        self.add_entries([{"act": c["act"], "label": c["label"]} for c in cites])

    def _target_group(self):
        """Куда класть акты: в выбранное направление (или направление выбранной нормы), иначе — в корень."""
        it = self.tree.currentItem()
        while it is not None and it.data(0, Qt.UserRole + 1) != "group":
            it = it.parent()
        return it.data(0, Qt.UserRole) if it is not None else 0

    def add_entries(self, entries):
        """Добавить нормы: акт находится или создаётся сам, повторы не добавляются. Текст и комментарий
        подставляются из «моих норм» (если эту статью уже добавляли в другое дело)."""
        if not self.cid:
            return
        self.save_now()
        lib = {LA.norm_key(e["act"], e["label"]): e for e in self._library()}
        group = self._target_group()
        added, last = 0, None
        for e in entries:
            act, label = LA.canon_act(e["act"]), LA.canon_label(e["label"])
            ka, kl = LA.norm_key(act, label)
            rows = db().laws(self.cid)
            acts = [r for r in rows if r["kind"] == "act" and LA.canon_act(r["title"]).lower() == ka]
            act_row = next((r for r in acts if r["parent_id"] == group), None) or (acts[0] if acts else None)
            if act_row is None:
                aid = db().add_law(self.cid, group, "act", act, "", "", LA.url_for(act))
            else:
                aid = act_row["id"]
            dup = next((r["id"] for r in rows if r["kind"] == "norm" and r["parent_id"] in {x["id"] for x in acts}
                        and LA.canon_label(r["title"]).lower() == kl), None)
            if dup:
                last = last or dup
                continue
            known = lib.get((ka, kl)) or {}
            body = e.get("body") or known.get("body", "")
            note = e.get("note") or known.get("note", "")
            url = e.get("url") or known.get("url") or LA.url_for(act, label)
            last = db().add_law(self.cid, aid, "norm", label, body, note, url)
            added += 1
        self.reload(select=last)
        self._refresh_completer()
        if added:
            self.main.toast(f"✓  Добавлено норм: {added}" + ("" if added == len(entries) else
                                                              f" (ещё {len(entries) - added} уже были)"))
        else:
            self.main.toast("Эти нормы уже есть в деле")

    def find_in_docs(self):
        if not self.cid:
            return
        docs = [d for d in db().docs(self.cid) if d["path"] and os.path.exists(d["path"])]
        if not docs:
            return QMessageBox.information(self, "Нормы права", "В деле пока нет документов — добавьте иск, решения "
                                           "и другие документы во вкладке «Документы».")
        pages = []
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            for d in docs:
                name = d["title"] or Path(d["path"]).name
                ext = Path(d["path"]).suffix.lower()
                try:
                    if ext == ".pdf":
                        with fitz.open(d["path"]) as doc:
                            for i, pg in enumerate(doc):
                                pages.append((name, i + 1, pg.get_text()))
                    elif ext in (".docx", ".txt", ".md"):
                        pages.append((name, None, U.L.extract_text_any(d["path"])))
                    else:
                        conv = C.cached_pdf(d["path"]) if hasattr(C, "cached_pdf") else None
                        if conv and os.path.exists(conv):
                            with fitz.open(conv) as doc:
                                for i, pg in enumerate(doc):
                                    pages.append((name, i + 1, pg.get_text()))
                except Exception as ex:
                    U.M.log_error(f"Поиск норм в {name}", ex)
            found = LA.find_in_pages(pages)
        finally:
            QApplication.restoreOverrideCursor()
        if not found:
            return QMessageBox.information(self, "Нормы права", "В текстах документов дела ссылок на статьи не нашлось. "
                                           "Сканы без распознанного текста программа прочитать не может — "
                                           "сначала распознайте текст (OCR).")
        have = self._have()
        d = NormPicker(self, "Нормы в документах дела",
                       f"Найдено ссылок: {len(found)}. Отметьте нужные — акты и статьи создадутся сами.",
                       [dict(c, where_text=_where(c["where"]),
                             exists=LA.norm_key(c["act"], c["label"]) in have) for c in found])
        if d.exec() == QDialog.Accepted:
            self.add_entries(d.chosen())

    def _have(self):
        """Пары (акт, статья), которые уже есть в деле — в нижнем регистре."""
        if not self.cid:
            return set()
        rows = db().laws(self.cid)
        acts = {r["id"]: r["title"] for r in rows if r["kind"] == "act"}
        return {LA.norm_key(acts[r["parent_id"]], r["title"]) for r in rows
                if r["kind"] == "norm" and r["parent_id"] in acts}

    def pick_library(self):
        lib = [e for e in self._library() if e["act"]]
        if not lib:
            return QMessageBox.information(self, "Мои нормы", "Пока пусто.")
        have = self._have()
        d = NormPicker(self, "Мои нормы", "Нормы из ваших дел (с текстом и комментарием) и частые статьи. "
                       "Отметьте нужные.", [dict(e, where_text=(f"в {e['uses']} {U.L.plural(e['uses'], 'деле', 'делах', 'делах')}"
                                                              if e["uses"] else e.get("about", "")),
                                                  exists=LA.norm_key(e["act"], e["label"]) in have)
                                            for e in lib], checked=False, search=True)
        if d.exec() == QDialog.Accepted:
            self.add_entries(d.chosen())

    def import_from_case(self):
        others = [c for c in db().cases() if c["id"] != self.cid]
        others += [c for c in db().cases(archived=True) if c["id"] != self.cid]
        others = [c for c in others if db().laws(c["id"])]
        if not others:
            return QMessageBox.information(self, U.M.APP_NAME, "В других делах пока нет сохранённых норм.")
        labels = [c["title"] + (f"  ·  {c['number']}" if c["number"] else "") for c in others]
        s, ok = QInputDialog.getItem(self, "Взять нормы из другого дела", "Дело:", labels, 0, False)
        if ok:
            n = db().copy_laws(others[labels.index(s)]["id"], self.cid)
            self.reload()
            self.main.msg(f"Скопировано разделов: {n}")


def _where(where):
    parts = []
    for doc, page in where[:4]:
        parts.append(f"{doc}" + (f", стр. {page}" if page else ""))
    more = f" и ещё {len(where) - 4}" if len(where) > 4 else ""
    return "; ".join(parts) + more


class NormPicker(QDialog):
    """Список найденных или своих норм с галочками: «Добавить отмеченные»."""

    def __init__(self, parent, title, note, entries, checked=True, search=False):
        super().__init__(parent)
        self.setWindowTitle(title)
        self.resize(720, 560)
        self.entries = entries
        v = QVBoxLayout(self)
        lab = QLabel(note)
        lab.setWordWrap(True)
        lab.setObjectName("hint")
        v.addWidget(lab)
        if search:
            self.q = QLineEdit()
            self.q.setPlaceholderText("Поиск: 395, неустойка, АПК…")
            self.q.textChanged.connect(self._filter)
            v.addWidget(self.q)
        self.list = QListWidget()
        self.list.setWordWrap(True)
        self.list.setStyleSheet("QListWidget::item { padding: 7px 6px; }")
        for idx, e in enumerate(entries):
            extra = [x for x in (e.get("where_text"), "уже есть в деле" if e.get("exists") else "") if x]
            it = QListWidgetItem(f"{e['label']} {LA.short_act(e['act'])}" + ("   —   " + " · ".join(extra) if extra else ""))
            if e.get("exists"):
                it.setForeground(QBrush(QColor(140, 140, 150)))
            it.setFlags(it.flags() | Qt.ItemIsUserCheckable)
            it.setCheckState(Qt.Checked if checked and not e.get("exists") else Qt.Unchecked)
            it.setToolTip((e.get("body") or e.get("raw") or "")[:500])
            it.setData(Qt.UserRole, idx)
            self.list.addItem(it)
        v.addWidget(self.list, 1)
        row = QHBoxLayout()
        row.addWidget(_btn("Отметить все", lambda: self._all(True)))
        row.addWidget(_btn("Снять все", lambda: self._all(False)))
        row.addStretch(1)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Ok).setText("Добавить отмеченные")
        bb.button(QDialogButtonBox.Cancel).setText("Отмена")
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        row.addWidget(bb)
        v.addLayout(row)

    def _filter(self, text):
        words = text.lower().split()
        for i in range(self.list.count()):
            it = self.list.item(i)
            e = self.entries[it.data(Qt.UserRole)]
            hay = f"{e['label']} {e['act']} {LA.short_act(e['act'])} {e.get('about', '')} {e.get('body', '')}".lower()
            it.setHidden(not all(w in hay for w in words))

    def _all(self, on):
        for i in range(self.list.count()):
            if not self.list.item(i).isHidden():
                self.list.item(i).setCheckState(Qt.Checked if on else Qt.Unchecked)

    def chosen(self):
        return [self.entries[self.list.item(i).data(Qt.UserRole)] for i in range(self.list.count())
                if self.list.item(i).checkState() == Qt.Checked]
