# -*- coding: utf-8 -*-
"""
PDF Мастер — ядро: все операции над PDF без графического интерфейса.
Основано на PyMuPDF (MuPDF).
"""
import os
import io
import re
import sys
import glob
import shutil
import difflib
import tempfile
import subprocess
from pathlib import Path
from html import escape

import pymupdf as fitz

IS_WIN = sys.platform.startswith("win")
MM = 72 / 25.4                      # пунктов в миллиметре
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".bmp", ".gif", ".tif", ".tiff", ".webp", ".jfif", ".pnm"}
OFFICE_EXT = {".doc", ".docx", ".rtf", ".odt", ".xls", ".xlsx", ".ods", ".csv",
              ".ppt", ".pptx", ".odp", ".pps", ".ppsx"}
HTML_EXT = {".html", ".htm", ".mht", ".mhtml"}
MUPDF_EXT = {".xps", ".oxps", ".epub", ".fb2", ".cbz", ".svg", ".txt"}

_NO_WINDOW = 0x08000000 if IS_WIN else 0   # CREATE_NO_WINDOW для subprocess


class Cancelled(Exception):
    """Операция отменена пользователем."""


def _tick(progress, i, total):
    if progress:
        progress(i, total)


# ----------------------------------------------------------------------------
#  Шрифты (для кириллицы нужен настоящий TTF-шрифт)
# ----------------------------------------------------------------------------
FONT_CANDIDATES = {
    "Arial": ["arial.ttf", "LiberationSans-Regular.ttf", "DejaVuSans.ttf", "Arial.ttf"],
    "Arial жирный": ["arialbd.ttf", "LiberationSans-Bold.ttf", "DejaVuSans-Bold.ttf"],
    "Times New Roman": ["times.ttf", "LiberationSerif-Regular.ttf", "DejaVuSerif.ttf"],
    "Times New Roman жирный": ["timesbd.ttf", "LiberationSerif-Bold.ttf", "DejaVuSerif-Bold.ttf"],
    "Calibri": ["calibri.ttf"],
    "Verdana": ["verdana.ttf"],
    "Georgia": ["georgia.ttf"],
    "Courier New": ["cour.ttf", "LiberationMono-Regular.ttf", "DejaVuSansMono.ttf"],
    "Segoe Script (рукописный)": ["segoesc.ttf"],
}
_FONT_DIRS = [
    os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts"),
    os.path.expanduser(r"~\AppData\Local\Microsoft\Windows\Fonts"),
    "/usr/share/fonts", "/usr/local/share/fonts", os.path.expanduser("~/.fonts"),
    "/Library/Fonts", "/System/Library/Fonts/Supplemental",
]
_font_cache = None


def available_fonts():
    """{название: путь к .ttf} для шрифтов, найденных в системе."""
    global _font_cache
    if _font_cache is not None:
        return _font_cache
    index = {}
    for d in _FONT_DIRS:
        if os.path.isdir(d):
            for root, _dirs, files in os.walk(d):
                for f in files:
                    index.setdefault(f.lower(), os.path.join(root, f))
    res = {}
    for name, cands in FONT_CANDIDATES.items():
        for c in cands:
            p = index.get(c.lower())
            if p:
                res[name] = p
                break
    _font_cache = res
    return res


def font_file(name=None):
    fonts = available_fonts()
    if name and name in fonts:
        return fonts[name]
    for n in ("Arial", "Times New Roman", "Calibri", "Verdana"):
        if n in fonts:
            return fonts[n]
    return next(iter(fonts.values()), None)


def _font_kwargs(name=None):
    ff = font_file(name)
    if ff:
        return {"fontfile": ff, "fontname": "F" + re.sub(r"[^A-Za-z0-9]", "", Path(ff).stem)}
    return {"fontname": "helv"}


def _font_obj(name=None):
    ff = font_file(name)
    return fitz.Font(fontfile=ff) if ff else fitz.Font("helv")


# ----------------------------------------------------------------------------
#  Координаты. «Видимые» координаты — то, что видит пользователь (с учётом
#  поворота страницы). MuPDF хочет «неповёрнутые» координаты.
# ----------------------------------------------------------------------------
def vis_to_unrot_rect(page, vis):
    """Для текста, аннотаций, поиска, скрытия данных, полей форм."""
    return (fitz.Rect(vis) * page.derotation_matrix).normalize()


def vis_to_unrot_point(page, pt):
    return fitz.Point(pt) * page.derotation_matrix


def _shape_offset(page):
    # Особенность MuPDF: для повёрнутых страниц с обрезкой рисование фигур
    # и картинок смещено на величину CropBox.
    if page.rotation == 0:
        return fitz.Point(0, 0)
    cb, mb = page.cropbox, page.mediabox
    return fitz.Point(cb.x0, -(mb.y1 - cb.y1))


def vis_to_shape_rect(page, vis):
    """Для draw_* и insert_image."""
    o = _shape_offset(page)
    return vis_to_unrot_rect(page, vis) + (o.x, o.y, o.x, o.y)


def set_visible_crop(page, vis):
    """Обрезать страницу до видимого прямоугольника vis."""
    r = vis_to_unrot_rect(page, vis)
    cb = page.cropbox
    r = fitz.Rect(r.x0 + cb.x0, r.y0 + cb.y0, r.x1 + cb.x0, r.y1 + cb.y0) & page.mediabox
    if r.is_empty or r.width < 5 or r.height < 5:
        raise ValueError("Слишком маленькая область обрезки")
    page.set_cropbox(r)


# ----------------------------------------------------------------------------
#  Текст
# ----------------------------------------------------------------------------
def write_text(page, vis_rect, text, fontsize=12, color=(0, 0, 0), font_name=None,
               align=0, opacity=1.0, extend=True):
    """Вписать текст в видимый прямоугольник. Возвращает фактический кегль."""
    vis = fitz.Rect(vis_rect)
    if extend:  # вниз до края страницы — чтобы длинный текст не пропадал
        vis.y1 = max(vis.y1, page.rect.y1 - 2)
    if vis.width < fontsize:
        vis.x1 = min(page.rect.x1, vis.x0 + max(fontsize * 4, 60))
    r = vis_to_unrot_rect(page, vis)
    kw = dict(color=color, align=align, rotate=page.rotation, **_font_kwargs(font_name))
    if opacity < 1:
        kw["fill_opacity"] = opacity
    fs = float(fontsize)
    while fs >= 3:
        try:
            rc = page.insert_textbox(r, text, fontsize=fs, **kw)
        except TypeError:
            kw.pop("fill_opacity", None)
            rc = page.insert_textbox(r, text, fontsize=fs, **kw)
        if rc >= 0:
            return fs
        fs = fs * 0.9 if fs > 6 else fs - 1
    return None


def text_style_in_rect(page, unrot):
    """Примерный кегль и цвет текста в области (для «Изменить текст»)."""
    try:
        d = page.get_text("dict", clip=unrot)
        for b in d.get("blocks", []):
            for line in b.get("lines", []):
                for s in line.get("spans", []):
                    if s.get("text", "").strip():
                        c = s.get("color", 0)
                        rgb = ((c >> 16) & 255, (c >> 8) & 255, c & 255)
                        return float(s.get("size", 12)), tuple(x / 255 for x in rgb), s.get("font", "")
    except Exception:
        pass
    return 12.0, (0, 0, 0), ""


# ----------------------------------------------------------------------------
#  Диапазоны страниц
# ----------------------------------------------------------------------------
def parse_ranges(text, n):
    """'1-3, 5, 8-' -> [[0,1,2],[4],[7..n-1]] (номера с нуля)."""
    text = re.sub(r"\s*[-–—]\s*", "-", (text or "").strip())
    groups = []
    for part in re.split(r"[,;\s]+", text):
        if not part:
            continue
        m = re.fullmatch(r"(\d*)-(\d*)", part)
        if m:
            a = int(m.group(1)) if m.group(1) else 1
            b = int(m.group(2)) if m.group(2) else n
        elif part.isdigit():
            a = b = int(part)
        else:
            raise ValueError(f"Непонятный диапазон: «{part}»")
        if a < 1 or b < 1 or a > n or b > n:
            raise ValueError(f"Страницы «{part}» нет в документе (всего {n})")
        step = 1 if b >= a else -1
        groups.append(list(range(a - 1, b - 1 + step, step)))
    return groups


def parse_pages(text, n):
    """Пусто = все страницы. Возвращает отсортированный список без повторов."""
    if not (text or "").strip() or text.strip().lower() in ("все", "all", "*"):
        return list(range(n))
    s = set()
    for g in parse_ranges(text, n):
        s.update(g)
    return sorted(s)


# ----------------------------------------------------------------------------
#  Открытие файлов и конвертация в PDF
# ----------------------------------------------------------------------------
def image_bytes(path):
    """Байты картинки, пригодные для MuPDF, и её размер в пикселях."""
    data = Path(path).read_bytes()
    try:
        pix = fitz.Pixmap(data)
        return data, pix.width, pix.height
    except Exception:
        pass
    # запасной вариант — через Qt (webp и прочее)
    from PySide6.QtGui import QImage
    from PySide6.QtCore import QBuffer, QByteArray, QIODevice
    img = QImage(str(path))
    if img.isNull():
        raise ValueError(f"Не удалось прочитать изображение: {path}")
    ba = QByteArray()
    buf = QBuffer(ba)
    buf.open(QIODevice.WriteOnly)
    img.save(buf, "PNG")
    return bytes(ba), img.width(), img.height()


# Форматы бумаги: название -> (ширина, высота) в мм, книжная ориентация
PAPER_SIZES = {
    "A3 (297×420 мм)": (297, 420),
    "A4 (210×297 мм)": (210, 297),
    "A5 (148×210 мм)": (148, 210),
    "A6 (105×148 мм)": (105, 148),
    "A2 (420×594 мм)": (420, 594),
    "A1 (594×841 мм)": (594, 841),
    "A0 (841×1189 мм)": (841, 1189),
    "B4 (250×353 мм)": (250, 353),
    "B5 (176×250 мм)": (176, 250),
    "Letter (216×279 мм)": (215.9, 279.4),
    "Legal (216×356 мм)": (215.9, 355.6),
    "Tabloid (279×432 мм)": (279.4, 431.8),
    "Конверт C5 (162×229 мм)": (162, 229),
    "Конверт DL (110×220 мм)": (110, 220),
}
CUSTOM_SIZE = "Свой размер…"
ORIENTATIONS = ["Автоматически", "Книжная", "Альбомная"]
FIT_MODES = ["Вписать, сохраняя пропорции", "Растянуть на весь лист", "Без масштабирования (по центру)"]

IMAGE_SIZE = "По размеру изображения"
PAGE_MODES = list(PAPER_SIZES) + [IMAGE_SIZE, CUSTOM_SIZE]


def paper_points(size_name, orientation="Книжная", like_w=None, like_h=None, custom_mm=(210, 297)):
    """Размер листа в пунктах. «Автоматически» — ориентация как у образца (like_w × like_h)."""
    w, h = custom_mm if size_name == CUSTOM_SIZE else PAPER_SIZES[size_name]
    w, h = w * MM, h * MM
    if orientation == "Альбомная":
        w, h = max(w, h), min(w, h)
    elif orientation == "Книжная":
        w, h = min(w, h), max(w, h)
    elif like_w and like_h and (like_w > like_h) != (w > h):   # автоматически
        w, h = h, w
    return w, h


def images_to_pdf(paths, page_mode="A4 (210×297 мм)", margin_mm=0, progress=None,
                  orientation="Автоматически", custom_mm=(210, 297)):
    # совместимость со старыми названиями
    legacy = {"A4 (ориентация по картинке)": ("A4 (210×297 мм)", "Автоматически"),
              "A4 книжная": ("A4 (210×297 мм)", "Книжная"),
              "A4 альбомная": ("A4 (210×297 мм)", "Альбомная"),
              "Letter": ("Letter (216×279 мм)", orientation)}
    if page_mode in legacy:
        page_mode, orientation = legacy[page_mode]
    out = fitz.open()
    m = margin_mm * MM
    for k, p in enumerate(paths):
        data, w, h = image_bytes(p)
        if page_mode.startswith("По размеру"):
            scale = 0.75                         # 96 dpi -> пункты
            if max(w, h) * scale > 3000:          # не больше ~1 м
                scale = 3000 / max(w, h)
            pw, ph = w * scale + 2 * m, h * scale + 2 * m
        else:
            pw, ph = paper_points(page_mode, orientation, w, h, custom_mm)
        page = out.new_page(width=pw, height=ph)
        page.insert_image(fitz.Rect(m, m, pw - m, ph - m), stream=data, keep_proportion=True)
        _tick(progress, k + 1, len(paths))
    return out


def page_size_label(page):
    """Человекочитаемый размер страницы: «A4 книжная (210×297 мм)»."""
    w, h = page.rect.width / MM, page.rect.height / MM
    orient = "альбомная" if w > h else "книжная"
    a, b = sorted((w, h))
    for name, (pw, ph) in PAPER_SIZES.items():
        if abs(a - pw) < 3 and abs(b - ph) < 3:
            return f"{name.split(' (')[0]} {orient} ({w:.0f}×{h:.0f} мм)"
    return f"{w:.0f}×{h:.0f} мм"


def _normalize_page(doc, page):
    """Сделать видимую область = всему листу без поворота (обрезка и поворот «впечатываются»)."""
    try:
        kind, val = doc.xref_get_key(page.xref, "CropBox")
        if kind == "array":
            nums = [float(v) for v in val.strip("[]").split()]
            if len(nums) == 4:
                page.set_mediabox(fitz.Rect(nums))
    except Exception:
        pass
    if page.rotation:
        try:
            page.remove_rotation()
        except Exception:
            pass


def resize_pages(doc, pages, size_name, orientation="Автоматически", fit=FIT_MODES[0],
                 margin_mm=0, custom_mm=(210, 297), progress=None):
    """Привести страницы к формату бумаги. Возвращает новый документ.
    Содержимое масштабируется (текст остаётся векторным и доступным для поиска).
    Аннотации и поля форм на изменяемых страницах предварительно «вшиваются»."""
    pages = set(pages if pages is not None else range(doc.page_count))
    src = fitz.open("pdf", doc.tobytes())
    try:
        src.bake(annots=True, widgets=True)
    except Exception:
        pass
    out = fitz.open()
    out.set_metadata(doc.metadata or {})
    m = margin_mm * MM
    total = src.page_count
    for i in range(total):
        sp = src[i]
        if i in pages:
            _normalize_page(src, sp)
        if i not in pages:
            out.insert_pdf(doc, from_page=i, to_page=i)       # без изменений, с аннотациями
        else:
            vw, vh = sp.rect.width, sp.rect.height              # видимый размер (с учётом поворота)
            pw, ph = paper_points(size_name, orientation, vw, vh, custom_mm)
            np_ = out.new_page(width=pw, height=ph)
            box = fitz.Rect(m, m, pw - m, ph - m)
            if fit.startswith("Растянуть"):
                np_.show_pdf_page(box, src, i, keep_proportion=False)
            elif fit.startswith("Без"):
                cx, cy = (box.x0 + box.x1) / 2, (box.y0 + box.y1) / 2
                r = fitz.Rect(cx - vw / 2, cy - vh / 2, cx + vw / 2, cy + vh / 2)
                np_.show_pdf_page(r, src, i)
            else:
                np_.show_pdf_page(box, src, i, keep_proportion=True)
        _tick(progress, i + 1, total)
    # закладки: номера страниц не меняются
    try:
        toc = doc.get_toc(simple=False)
        if toc:
            out.set_toc(toc)
    except Exception:
        pass
    return out


def open_as_pdf(path, password_cb=None):
    """Открыть любой поддерживаемый файл как PDF-документ в памяти.
    password_cb(name) -> пароль или None. Возвращает fitz.Document или None."""
    ext = Path(path).suffix.lower()
    if ext in IMAGE_EXT:
        return images_to_pdf([path])
    if ext in OFFICE_EXT:
        tmp = os.path.join(tempfile.mkdtemp(prefix="pdfm_"), Path(path).stem + ".pdf")
        office_to_pdf(path, tmp)
        return fitz.open("pdf", Path(tmp).read_bytes())
    if ext in HTML_EXT:
        tmp = os.path.join(tempfile.mkdtemp(prefix="pdfm_"), Path(path).stem + ".pdf")
        html_to_pdf(path, tmp)
        return fitz.open("pdf", Path(tmp).read_bytes())
    doc = fitz.open(path)
    if doc.needs_pass:
        while True:
            pw = password_cb(Path(path).name) if password_cb else None
            if pw is None:
                return None
            if doc.authenticate(pw):
                break
    if not doc.is_pdf:
        return fitz.open("pdf", doc.convert_to_pdf())
    # копия в памяти, уже без шифрования
    return fitz.open("pdf", doc.tobytes(encryption=fitz.PDF_ENCRYPT_NONE))


def save_pdf(doc, path, garbage=3):
    tmp = str(path) + ".tmp_pdfm"
    doc.save(tmp, garbage=garbage, deflate=True, encryption=fitz.PDF_ENCRYPT_NONE)
    os.replace(tmp, path)


# ----------------------------------------------------------------------------
#  Разделение, извлечение
# ----------------------------------------------------------------------------
def subset_doc(doc, pages):
    new = fitz.open()
    for p in pages:
        new.insert_pdf(doc, from_page=p, to_page=p)
    return new


def split_document(doc, mode, value, outdir, base, progress=None):
    n = doc.page_count
    if mode == "each":
        groups = [[i] for i in range(n)]
    elif mode == "every":
        k = max(1, int(value))
        groups = [list(range(i, min(i + k, n))) for i in range(0, n, k)]
    else:
        groups = parse_ranges(value, n)
        if not groups:
            raise ValueError("Укажите диапазоны, например: 1-3, 4-10, 11-")
    os.makedirs(outdir, exist_ok=True)
    paths = []
    for gi, g in enumerate(groups):
        new = subset_doc(doc, g)
        name = f"{base}_стр{g[0] + 1}.pdf" if len(g) == 1 else f"{base}_стр{g[0] + 1}-{g[-1] + 1}.pdf"
        path = os.path.join(outdir, name)
        new.save(path, garbage=3, deflate=True)
        new.close()
        paths.append(path)
        _tick(progress, gi + 1, len(groups))
    return paths


# ----------------------------------------------------------------------------
#  Сжатие и восстановление
# ----------------------------------------------------------------------------
COMPRESS_LEVELS = {
    "Без потери качества": None,
    "Рекомендуемое (150 dpi)": (170, 150, 75),
    "Сильное (100 dpi)": (120, 100, 60),
    "Максимальное (72 dpi)": (90, 72, 45),
}


def _manual_image_compress(doc, dpi_target, quality, progress=None):
    done = set()
    for pno, page in enumerate(doc):
        for img in page.get_images(full=True):
            xref, smask = img[0], img[1]
            if xref in done or smask:
                continue
            done.add(xref)
            try:
                pix = fitz.Pixmap(doc, xref)
                if pix.alpha:
                    continue
                if pix.n > 3:
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                bbox = page.get_image_bbox(img)
                if not bbox.is_empty and bbox.width > 0:
                    dpi = pix.width / (bbox.width / 72)
                    while dpi > dpi_target * 1.4 and pix.width > 64:
                        pix.shrink(1)
                        dpi /= 2
                data = pix.tobytes("jpeg", jpg_quality=quality)
                if len(data) < len(doc.xref_stream_raw(xref) or b"") * 0.95:
                    page.replace_image(xref, stream=data)
            except Exception:
                continue
        _tick(progress, pno + 1, doc.page_count)


def compress_document(doc, path, level, progress=None):
    d = fitz.open("pdf", doc.tobytes())
    params = COMPRESS_LEVELS.get(level)
    if params:
        th, tg, q = params
        try:
            d.rewrite_images(dpi_threshold=th, dpi_target=tg, quality=q)
        except Exception:
            _manual_image_compress(d, tg, q, progress)
    try:
        d.subset_fonts()
    except Exception:
        pass
    kw = dict(garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True)
    try:
        d.save(path, use_objstms=1, **kw)
    except TypeError:
        d.save(path, **kw)
    d.close()
    return os.path.getsize(path)


def repair_file(path):
    """MuPDF сам чинит повреждённую структуру при открытии."""
    try:
        d = fitz.open(path)
    except Exception:
        d = fitz.open(stream=Path(path).read_bytes(), filetype="pdf")
    if d.page_count == 0:
        raise ValueError("Не удалось восстановить ни одной страницы")
    good = fitz.open()
    bad = 0
    for i in range(d.page_count):
        try:
            good.insert_pdf(d, from_page=i, to_page=i)
        except Exception:
            bad += 1
    data = good.tobytes(garbage=4, deflate=True, clean=True)
    return fitz.open("pdf", data), bad


# ----------------------------------------------------------------------------
#  Водяной знак, номера страниц, обрезка полей
# ----------------------------------------------------------------------------
def _center(r):
    return fitz.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)


def add_text_watermark(doc, pages, text, fontsize=60, color=(0.6, 0.6, 0.6), opacity=0.3,
                       angle=45, tile=False, font_name=None, progress=None):
    font = _font_obj(font_name)
    for k, pno in enumerate(pages):
        page = doc[pno]
        vr = page.rect
        lines = text.split("\n")
        widths = [font.text_length(t, fontsize) for t in lines]
        wmax = max(widths) if widths else 0
        lh = fontsize * 1.2
        centers = []
        if tile:
            sx, sy = wmax + fontsize * 3, lh * len(lines) + fontsize * 4
            row = 0
            y = vr.y0
            while y < vr.y1 + sy:
                x = vr.x0 + (sx / 2 if row % 2 else 0)
                while x < vr.x1 + sx:
                    centers.append(fitz.Point(x, y))
                    x += sx
                y += sy
                row += 1
        else:
            centers = [_center(vr)]
        for cv in centers:
            c = vis_to_unrot_point(page, cv)
            tw = fitz.TextWriter(page.rect, opacity=opacity, color=color)
            y0 = c.y - lh * (len(lines) - 1) / 2 + fontsize * 0.35
            for i, t in enumerate(lines):
                tw.append(fitz.Point(c.x - widths[i] / 2, y0 + i * lh), t, font=font, fontsize=fontsize)
            tw.write_text(page, morph=(c, fitz.Matrix(angle + page.rotation)))
        _tick(progress, k + 1, len(pages))


def add_image_stamp(doc, pages, png_bytes, rel_width=0.4, position="По центру", margin_mm=10,
                    progress=None):
    pix = fitz.Pixmap(png_bytes)
    ar = pix.height / max(1, pix.width)
    m = margin_mm * MM
    for k, pno in enumerate(pages):
        page = doc[pno]
        vr = page.rect
        w = vr.width * rel_width
        h = w * ar
        if h > vr.height * 0.9:
            h = vr.height * 0.9
            w = h / ar
        c = _center(vr)
        x, y = c.x - w / 2, c.y - h / 2
        if "Вверху" in position:
            y = vr.y0 + m
        if "Внизу" in position:
            y = vr.y1 - m - h
        if "слева" in position:
            x = vr.x0 + m
        if "справа" in position:
            x = vr.x1 - m - w
        rect = vis_to_shape_rect(page, fitz.Rect(x, y, x + w, y + h))
        page.insert_image(rect, stream=png_bytes, keep_proportion=True, overlay=True,
                          rotate=page.rotation)
        _tick(progress, k + 1, len(pages))


NUMBER_POSITIONS = ["Внизу по центру", "Внизу справа", "Внизу слева",
                    "Вверху по центру", "Вверху справа", "Вверху слева"]
NUMBER_FORMATS = ["{n}", "{n} / {total}", "Страница {n}", "Страница {n} из {total}", "- {n} -"]


def add_page_numbers(doc, pages, fmt="{n}", position=NUMBER_POSITIONS[0], fontsize=11,
                     margin_mm=10, start=1, color=(0, 0, 0), font_name=None, progress=None):
    m = margin_mm * MM
    total = len(pages) + start - 1
    align = 1 if "центру" in position else (2 if "справа" in position else 0)
    for k, pno in enumerate(pages):
        page = doc[pno]
        vr = page.rect
        h = fontsize * 1.8
        y0 = vr.y1 - m - h if "Внизу" in position else vr.y0 + m
        rect = fitz.Rect(vr.x0 + m, y0, vr.x1 - m, y0 + h)
        txt = fmt.replace("{n}", str(start + k)).replace("{total}", str(total))
        write_text(page, rect, txt, fontsize, color, font_name, align, extend=False)
        _tick(progress, k + 1, len(pages))


def crop_margins(doc, pages, left, top, right, bottom, progress=None):
    for k, pno in enumerate(pages):
        page = doc[pno]
        vr = page.rect
        set_visible_crop(page, fitz.Rect(vr.x0 + left, vr.y0 + top, vr.x1 - right, vr.y1 - bottom))
        _tick(progress, k + 1, len(pages))


def auto_crop(doc, pages, pad_mm=5, progress=None):
    """Обрезать пустые поля по содержимому."""
    pad = pad_mm * MM
    for k, pno in enumerate(pages):
        page = doc[pno]
        pix = page.get_pixmap(dpi=50, alpha=False, colorspace=fitz.csGRAY)
        w, h, s = pix.width, pix.height, pix.samples
        xs, ys = [], []
        for y in range(h):
            row = s[y * pix.stride:y * pix.stride + w]
            if min(row) < 235:
                ys.append(y)
                xs_row = [x for x in range(w) if row[x] < 235]
                xs.append(xs_row[0])
                xs.append(xs_row[-1])
        if ys:
            sc = page.rect.width / w
            vis = fitz.Rect(min(xs) * sc - pad, min(ys) * sc - pad,
                            (max(xs) + 1) * sc + pad, (max(ys) + 1) * sc + pad) & page.rect
            try:
                set_visible_crop(page, vis)
            except ValueError:
                pass
        _tick(progress, k + 1, len(pages))


# ----------------------------------------------------------------------------
#  Поиск: выделить / скрыть
# ----------------------------------------------------------------------------
PATTERNS = {
    "email": r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+",
    "phone": r"(?:\+7|8)[\s\-(]*\d{3}[\s\-)]*\d{3}[\s\-]*\d{2}[\s\-]*\d{2}",
    "cards": r"\b(?:\d{4}[\s-]?){3}\d{4}\b",
    "passport": r"\b\d{2}\s?\d{2}\s?\d{6}\b",
}


def find_and_mark(doc, pages, words, action, patterns=(), color=(1, 1, 0), progress=None):
    """action: 'redact' | 'whiteout' | 'highlight' | 'underline' | 'strike'.
    Возвращает число найденных фрагментов."""
    count = 0
    for k, pno in enumerate(pages):
        page = doc[pno]
        needles = [w for w in words if w.strip()]
        if patterns:
            text = page.get_text()
            for pkey in patterns:
                needles += [m.group(0) for m in re.finditer(PATTERNS[pkey], text)]
        quads = []
        for w in dict.fromkeys(needles):
            quads += page.search_for(w, quads=True)
        count += len(quads)
        if not quads:
            _tick(progress, k + 1, len(pages))
            continue
        if action in ("redact", "whiteout"):
            fill = (0, 0, 0) if action == "redact" else (1, 1, 1)
            for q in quads:
                page.add_redact_annot(q, fill=fill)
            page.apply_redactions()
        else:
            fn = {"highlight": page.add_highlight_annot, "underline": page.add_underline_annot,
                  "strike": page.add_strikeout_annot}[action]
            a = fn(quads)
            if action == "highlight":
                a.set_colors(stroke=color)
                a.update()
        _tick(progress, k + 1, len(pages))
    return count


# ----------------------------------------------------------------------------
#  Экспорт из PDF
# ----------------------------------------------------------------------------
def pdf_to_images(doc, pages, outdir, base, dpi=150, fmt="jpg", quality=90, progress=None):
    os.makedirs(outdir, exist_ok=True)
    paths = []
    width = len(str(doc.page_count))
    for k, pno in enumerate(pages):
        pix = doc[pno].get_pixmap(dpi=dpi, alpha=False)
        path = os.path.join(outdir, f"{base}_{str(pno + 1).zfill(width)}.{fmt}")
        if fmt == "jpg":
            Path(path).write_bytes(pix.tobytes("jpeg", jpg_quality=quality))
        else:
            pix.save(path)
        paths.append(path)
        _tick(progress, k + 1, len(pages))
    return paths


def extract_images(doc, pages, outdir, base, progress=None):
    os.makedirs(outdir, exist_ok=True)
    seen, paths = set(), []
    for k, pno in enumerate(pages):
        for img in doc[pno].get_images(full=True):
            xref = img[0]
            if xref in seen:
                continue
            seen.add(xref)
            try:
                info = doc.extract_image(xref)
            except Exception:
                continue
            if not info or info.get("width", 0) < 16 or info.get("height", 0) < 16:
                continue
            path = os.path.join(outdir, f"{base}_стр{pno + 1}_{xref}.{info['ext']}")
            Path(path).write_bytes(info["image"])
            paths.append(path)
        _tick(progress, k + 1, len(pages))
    return paths


def pdf_to_text(doc, path, progress=None):
    parts = []
    for i, page in enumerate(doc):
        parts.append(f"===== Страница {i + 1} =====\n{page.get_text('text', sort=True)}")
        _tick(progress, i + 1, doc.page_count)
    Path(path).write_text("\n".join(parts), encoding="utf-8")


def pdf_to_word(doc, path, progress=None):
    from pdf2docx import Converter
    tmpdir = tempfile.mkdtemp(prefix="pdfm_")
    src = os.path.join(tmpdir, "src.pdf")
    doc.save(src, garbage=3, deflate=True)
    _tick(progress, 0, 0)
    cv = Converter(src)
    try:
        cv.convert(path, multi_processing=False)
    finally:
        cv.close()
    shutil.rmtree(tmpdir, ignore_errors=True)


def _cell(v):
    if v is None:
        return ""
    s = str(v).strip()
    t = s.replace("\u00a0", "").replace(" ", "")
    if re.fullmatch(r"-?\d+([.,]\d+)?", t):
        try:
            return float(t.replace(",", ".")) if re.search(r"[.,]", t) else int(t)
        except ValueError:
            pass
    return s


def _page_text_rows(page):
    """Строки текста, разбитые на колонки по крупным промежуткам."""
    words = page.get_text("words", sort=True)
    lines = {}
    for w in words:
        key = round((w[1] + w[3]) / 2 / 3)
        lines.setdefault(key, []).append(w)
    rows = []
    for key in sorted(lines):
        ws = sorted(lines[key], key=lambda w: w[0])
        cells, cur, last_x1 = [], [], None
        for w in ws:
            h = max(1, w[3] - w[1])
            if last_x1 is not None and w[0] - last_x1 > h * 1.5:
                cells.append(" ".join(cur))
                cur = []
            cur.append(w[4])
            last_x1 = w[2]
        if cur:
            cells.append(" ".join(cur))
        rows.append([_cell(c) for c in cells])
    return rows


def pdf_to_excel(doc, path, progress=None):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    found = 0
    for i, page in enumerate(doc):
        try:
            tabs = page.find_tables()
            tables = list(tabs.tables) if tabs else []
        except Exception:
            tables = []
        for t_i, t in enumerate(tables):
            data = t.extract()
            if not data:
                continue
            ws = wb.create_sheet(f"Стр{i + 1} табл{t_i + 1}"[:31])
            for r_i, row in enumerate(data):
                ws.append([_cell(c) for c in row])
                if r_i == 0:
                    for c in ws[1]:
                        c.font = Font(bold=True)
            found += 1
        _tick(progress, i + 1, doc.page_count)
    if not found:  # таблиц нет — выгружаем текст по колонкам
        for i, page in enumerate(doc):
            ws = wb.create_sheet(f"Страница {i + 1}"[:31])
            for row in _page_text_rows(page):
                ws.append(row)
    for ws in wb.worksheets:
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(60, max(8, width + 2))
    wb.save(path)
    return found


def pdf_to_pptx(doc, path, dpi=150, progress=None):
    from pptx import Presentation
    from pptx.util import Emu
    EMU = 12700
    prs = Presentation()
    r0 = doc[0].rect
    prs.slide_width, prs.slide_height = Emu(int(r0.width * EMU)), Emu(int(r0.height * EMU))
    blank = prs.slide_layouts[6]
    sw, sh = prs.slide_width, prs.slide_height
    for i, page in enumerate(doc):
        pix = page.get_pixmap(dpi=dpi, alpha=False)
        slide = prs.slides.add_slide(blank)
        pr = page.rect
        sc = min(sw / (pr.width * EMU), sh / (pr.height * EMU))
        w, h = int(pr.width * EMU * sc), int(pr.height * EMU * sc)
        slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), (sw - w) // 2, (sh - h) // 2, w, h)
        txt = page.get_text().strip()
        if txt:
            slide.notes_slide.notes_text_frame.text = txt[:10000]
        _tick(progress, i + 1, doc.page_count)
    prs.save(path)


def _simple_markdown(doc, progress=None):
    sizes = {}
    for page in doc:
        for b in page.get_text("dict")["blocks"]:
            for l in b.get("lines", []):
                for s in l["spans"]:
                    sizes[round(s["size"])] = sizes.get(round(s["size"]), 0) + len(s["text"])
    body = max(sizes, key=sizes.get) if sizes else 11
    out = []
    for i, page in enumerate(doc):
        for b in page.get_text("dict")["blocks"]:
            if b.get("type") != 0:
                continue
            text = " ".join("".join(s["text"] for s in l["spans"]) for l in b["lines"]).strip()
            if not text:
                continue
            spans = [s for l in b["lines"] for s in l["spans"] if s["text"].strip()]
            size = max((s["size"] for s in spans), default=body)
            bold = all("bold" in s["font"].lower() or s["flags"] & 16 for s in spans) if spans else False
            if size >= body * 1.6:
                out.append("# " + text)
            elif size >= body * 1.3:
                out.append("## " + text)
            elif (size >= body * 1.1 or bold) and len(text) < 120:
                out.append("### " + text)
            elif re.match(r"^[•●▪\-–]\s*", text):
                out.append("- " + re.sub(r"^[•●▪\-–]\s*", "", text))
            else:
                out.append(text)
            out.append("")
        out.append("\n---\n")
        _tick(progress, i + 1, doc.page_count)
    return "\n".join(out)


def pdf_to_markdown(doc, path, progress=None):
    md = None
    try:
        import pymupdf4llm
        md = pymupdf4llm.to_markdown(doc)
    except Exception:
        md = None
    if not md or not md.strip():
        md = _simple_markdown(doc, progress)
    Path(path).write_text(md, encoding="utf-8")


# ----------------------------------------------------------------------------
#  Защита
# ----------------------------------------------------------------------------
def save_encrypted(doc, path, user_pw, owner_pw, allow_print=True, allow_copy=True,
                   allow_edit=False):
    perm = fitz.PDF_PERM_ACCESSIBILITY
    if allow_print:
        perm |= fitz.PDF_PERM_PRINT | fitz.PDF_PERM_PRINT_HQ
    if allow_copy:
        perm |= fitz.PDF_PERM_COPY
    if allow_edit:
        perm |= (fitz.PDF_PERM_MODIFY | fitz.PDF_PERM_ANNOTATE | fitz.PDF_PERM_FORM
                 | fitz.PDF_PERM_ASSEMBLE)
    doc.save(path, garbage=3, deflate=True, encryption=fitz.PDF_ENCRYPT_AES_256,
             owner_pw=owner_pw or user_pw, user_pw=user_pw, permissions=perm)


# ----------------------------------------------------------------------------
#  Сравнение
# ----------------------------------------------------------------------------
def _visual_diff(pa, pb):
    try:
        a = pa.get_pixmap(dpi=40, alpha=False, colorspace=fitz.csGRAY)
        b = pb.get_pixmap(dpi=40, alpha=False, colorspace=fitz.csGRAY)
        if (a.width, a.height) != (b.width, b.height):
            return 100.0
        sa, sb = a.samples, b.samples
        diff = sum(1 for x, y in zip(sa, sb) if abs(x - y) > 40)
        return 100.0 * diff / max(1, len(sa))
    except Exception:
        return 0.0


def compare_documents(a, b, out_html, name_a, name_b, progress=None):
    n = max(a.page_count, b.page_count)
    hd = difflib.HtmlDiff(wrapcolumn=70)
    rows, sections, changed = [], [], 0
    for i in range(n):
        ta = a[i].get_text(sort=True).splitlines() if i < a.page_count else []
        tb = b[i].get_text(sort=True).splitlines() if i < b.page_count else []
        vis = (_visual_diff(a[i], b[i]) if i < a.page_count and i < b.page_count else 100.0)
        text_diff = ta != tb
        status = "совпадает"
        if i >= a.page_count:
            status = "есть только во втором файле"
        elif i >= b.page_count:
            status = "есть только в первом файле"
        elif text_diff:
            status = "отличается текст"
        elif vis > 0.5:
            status = "отличается оформление/картинки"
        if status != "совпадает":
            changed += 1
        cls = "ok" if status == "совпадает" else "bad"
        rows.append(f"<tr class='{cls}'><td>{i + 1}</td><td>{status}</td><td>{vis:.1f}%</td></tr>")
        if text_diff:
            sections.append(f"<h2>Страница {i + 1}</h2>"
                            + hd.make_table(ta, tb, escape(name_a), escape(name_b), context=True, numlines=2))
        _tick(progress, i + 1, n)
    styles = getattr(difflib, "_styles", "")
    html = f"""<!doctype html><html lang="ru"><head><meta charset="utf-8">
<title>Сравнение PDF</title><style>{styles}
body{{font-family:Segoe UI,Arial,sans-serif;margin:24px;color:#222}}
table.sum{{border-collapse:collapse;margin:12px 0}} table.sum td,table.sum th{{border:1px solid #ccc;padding:4px 10px}}
tr.bad td{{background:#fde8e8}} tr.ok td{{background:#eaf7ea}} table.diff{{font-size:13px;margin-bottom:24px}}
</style></head><body><h1>Сравнение PDF</h1>
<p><b>1:</b> {escape(name_a)} ({a.page_count} стр.)<br><b>2:</b> {escape(name_b)} ({b.page_count} стр.)</p>
<p>Страниц с отличиями: <b>{changed}</b> из {n}</p>
<table class="sum"><tr><th>Стр.</th><th>Результат</th><th>Визуальная разница</th></tr>{''.join(rows)}</table>
{''.join(sections) or '<p>Текст документов совпадает.</p>'}
<p style="margin-top:20px">Легенда: <span class="diff_add">добавлено</span> · <span class="diff_chg">изменено</span> · <span class="diff_sub">удалено</span></p>
</body></html>"""
    Path(out_html).write_text(html, encoding="utf-8")
    return changed


# ----------------------------------------------------------------------------
#  Внешние программы: MS Office / LibreOffice, Edge / Chrome, Ghostscript
# ----------------------------------------------------------------------------
def _find_soffice():
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    cands = [r"C:\Program Files\LibreOffice\program\soffice.exe",
             r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
             "/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    return next((c for c in cands if os.path.isfile(c)), None)


def _office_com(src, out):
    import pythoncom
    import win32com.client
    ext = Path(src).suffix.lower()
    src, out = os.path.abspath(src), os.path.abspath(out)
    pythoncom.CoInitialize()
    try:
        if ext in (".doc", ".docx", ".rtf", ".odt"):
            app = win32com.client.DispatchEx("Word.Application")
            app.Visible = False
            app.DisplayAlerts = 0
            try:
                d = app.Documents.Open(src, ReadOnly=True, ConfirmConversions=False)
                d.ExportAsFixedFormat(out, 17)
                d.Close(False)
            finally:
                app.Quit()
        elif ext in (".xls", ".xlsx", ".ods", ".csv"):
            app = win32com.client.DispatchEx("Excel.Application")
            app.Visible = False
            app.DisplayAlerts = False
            try:
                wb = app.Workbooks.Open(src, ReadOnly=True)
                wb.ExportAsFixedFormat(0, out)
                wb.Close(False)
            finally:
                app.Quit()
        else:
            app = win32com.client.DispatchEx("PowerPoint.Application")
            try:
                p = app.Presentations.Open(src, ReadOnly=True, WithWindow=False)
                p.SaveAs(out, 32)
                p.Close()
            finally:
                app.Quit()
    finally:
        pythoncom.CoUninitialize()


def office_to_pdf(src, out):
    errors = []
    if IS_WIN:
        try:
            _office_com(src, out)
            if os.path.isfile(out):
                return out
        except Exception as e:
            errors.append(f"Microsoft Office: {e}")
    soffice = _find_soffice()
    if soffice:
        outdir = tempfile.mkdtemp(prefix="pdfm_lo_")
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", outdir, src],
                       capture_output=True, timeout=300, creationflags=_NO_WINDOW)
        res = os.path.join(outdir, Path(src).stem + ".pdf")
        if os.path.isfile(res):
            shutil.move(res, out)
            return out
        errors.append("LibreOffice не смог сконвертировать файл")
    raise RuntimeError(
        "Для конвертации Word/Excel/PowerPoint нужен установленный Microsoft Office "
        "или бесплатный LibreOffice (libreoffice.org).\n\n" + "\n".join(errors))


def _find_browser():
    cands = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
        os.path.expanduser(r"~\AppData\Local\Google\Chrome\Application\chrome.exe"),
        r"C:\Program Files\Yandex\YandexBrowser\Application\browser.exe",
        os.path.expanduser(r"~\AppData\Local\Yandex\YandexBrowser\Application\browser.exe"),
        "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    ]
    for c in cands:
        if os.path.isfile(c):
            return c
    for name in ("msedge", "google-chrome", "chromium", "chromium-browser", "chrome"):
        p = shutil.which(name)
        if p:
            return p
    return None


def html_to_pdf(src, out):
    """src — адрес сайта (https://...) или путь к .html файлу."""
    if os.path.isfile(src):
        url = Path(src).resolve().as_uri()
    else:
        url = src if re.match(r"^[a-z]+://", src, re.I) else "https://" + src
    browser = _find_browser()
    if browser:
        prof = tempfile.mkdtemp(prefix="pdfm_br_")
        cmd = [browser, "--headless=new", "--disable-gpu", "--no-first-run",
               f"--user-data-dir={prof}", "--no-pdf-header-footer",
               "--virtual-time-budget=8000", f"--print-to-pdf={os.path.abspath(out)}", url]
        subprocess.run(cmd, capture_output=True, timeout=180, creationflags=_NO_WINDOW)
        shutil.rmtree(prof, ignore_errors=True)
        if os.path.isfile(out) and os.path.getsize(out) > 0:
            return out
    if os.path.isfile(src):  # запасной вариант: встроенная вёрстка MuPDF (без JS/CSS-сложностей)
        html = Path(src).read_text(encoding="utf-8", errors="replace")
        story = fitz.Story(html=html, archive=os.path.dirname(os.path.abspath(src)))
        writer = fitz.DocumentWriter(out)
        mediabox = fitz.paper_rect("a4")
        where = mediabox + (36, 36, -36, -36)
        more = 1
        while more:
            dev = writer.begin_page(mediabox)
            more, _ = story.place(where)
            story.draw(dev)
            writer.end_page()
        writer.close()
        return out
    raise RuntimeError("Для HTML → PDF нужен браузер Microsoft Edge или Google Chrome.")


def _find_ghostscript():
    for n in ("gswin64c", "gswin32c", "gs"):
        p = shutil.which(n)
        if p:
            return p
    for pat in (r"C:\Program Files\gs\gs*\bin\gswin64c.exe", r"C:\Program Files (x86)\gs\gs*\bin\gswin32c.exe"):
        found = sorted(glob.glob(pat))
        if found:
            return found[-1]
    return None


def pdf_to_pdfa(doc, out):
    gs = _find_ghostscript()
    if not gs:
        raise RuntimeError("Для PDF/A нужен бесплатный Ghostscript (ghostscript.com/releases). "
                           "Установите его и повторите.")
    tmpdir = tempfile.mkdtemp(prefix="pdfm_")
    src = os.path.join(tmpdir, "src.pdf")
    doc.save(src, garbage=3, deflate=True)
    cmd = [gs, "-dPDFA=2", "-dBATCH", "-dNOPAUSE", "-dNOOUTERSAVE", "-dQUIET",
           "-sColorConversionStrategy=RGB", "-sDEVICE=pdfwrite", "-dPDFACompatibilityPolicy=1",
           f"-sOutputFile={out}", src]
    r = subprocess.run(cmd, capture_output=True, timeout=600, creationflags=_NO_WINDOW)
    shutil.rmtree(tmpdir, ignore_errors=True)
    if r.returncode != 0 or not os.path.isfile(out):
        raise RuntimeError("Ghostscript вернул ошибку:\n" + r.stderr.decode(errors="replace")[-1500:])


# ----------------------------------------------------------------------------
#  OCR (встроенный в MuPDF Tesseract, нужны только файлы языков *.traineddata)
# ----------------------------------------------------------------------------
def _app_dir():
    return getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))


def find_tessdata():
    cands = [os.environ.get("TESSDATA_PREFIX", ""),
             os.path.join(_app_dir(), "tessdata"),
             os.path.join(os.path.dirname(os.path.abspath(sys.argv[0])), "tessdata"),
             r"C:\Program Files\Tesseract-OCR\tessdata",
             r"C:\Program Files (x86)\Tesseract-OCR\tessdata",
             os.path.expanduser(r"~\AppData\Local\Programs\Tesseract-OCR\tessdata"),
             "/usr/share/tessdata", "/opt/homebrew/share/tessdata", "/usr/local/share/tessdata"]
    cands += glob.glob("/usr/share/tesseract-ocr/*/tessdata")
    for c in cands:
        if c and os.path.isdir(c) and glob.glob(os.path.join(c, "*.traineddata")):
            return c
    return None


def ocr_languages(tessdata):
    if not tessdata:
        return []
    return sorted(Path(f).name[:-len(".traineddata")] for f in glob.glob(os.path.join(tessdata, "*.traineddata"))
                  if not Path(f).name.startswith("osd"))


def ocr_document(doc, pages, lang="rus+eng", dpi=300, tessdata=None, only_without_text=True,
                 progress=None):
    """Возвращает новый документ: распознанные страницы = изображение + невидимый текст."""
    tessdata = tessdata or find_tessdata()
    if not tessdata:
        raise RuntimeError("Не найдены языковые файлы OCR (папка tessdata).")
    os.environ["TESSDATA_PREFIX"] = tessdata
    out = fitz.open()
    todo = set(pages)
    done = 0
    for i in range(doc.page_count):
        page = doc[i]
        if i in todo and not (only_without_text and len(page.get_text().strip()) > 20):
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            try:
                data = pix.pdfocr_tobytes(language=lang, tessdata=tessdata)
            except TypeError:
                data = pix.pdfocr_tobytes(language=lang)
            od = fitz.open("pdf", data)
            out.insert_pdf(od)
        else:
            out.insert_pdf(doc, from_page=i, to_page=i)
        done += 1
        _tick(progress, done, doc.page_count)
    return out
