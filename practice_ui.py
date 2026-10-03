# -*- coding: utf-8 -*-
"""
Вкладка дела «Судебная практика»: практика разложена по доводам.

Довод («Неустойка явно несоразмерна», «Подсудность») → под ним судебные акты, которые его подкрепляют.
У акта: вид, суд, дата, номер, номер дела, ссылка, позиция суда, цитата и как применяем в деле.
Быстрое добавление: вставить реквизиты так, как их пишут в иске, — «Определение ВС РФ от 12.03.2024
№ 305-ЭС23-1234» или ссылку на kad.arbitr.ru — вид, суд, дата и номер заполнятся сами (practice_auto).
"""
import html
import os

from PySide6.QtCore import Qt, QTimer, QUrl
from PySide6.QtGui import QDesktopServices, QFont
from PySide6.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QFormLayout, QLabel, QPushButton, QComboBox,
                               QLineEdit, QPlainTextEdit, QSplitter, QFrame, QMenu, QMessageBox, QFileDialog,
                               QTreeWidget, QTreeWidgetItem, QAbstractItemView, QApplication, QStackedWidget)

import practice_auto as PA

U = None          # legal_ui (задаётся в bind)
GLYPH = {"group": "▣", "act": "⚖"}


def bind(legal_ui):
    global U
    U = legal_ui


def db():
    return U.db()


def _btn(text, slot, primary=False, tip=None):
    b = QPushButton(text)
    if primary:
        b.setObjectName("primary")
    if tip:
        b.setToolTip(tip)
    b.clicked.connect(slot)
    return b


class PracticeTree(QTreeWidget):
    def __init__(self, tab):
        super().__init__()
        self.tab = tab
        self.setHeaderHidden(True)
        self.setDragDropMode(QAbstractItemView.InternalMove)
        self.setSelectionMode(QAbstractItemView.SingleSelection)
        self.setIndentation(18)
        self.setObjectName("lawtree")
        self.setWordWrap(True)
        self.setTextElideMode(Qt.ElideNone)
        self.setUniformRowHeights(False)

    def dropEvent(self, e):
        super().dropEvent(e)
        self.tab.save_structure()


class PracticeTab(QWidget):
    def __init__(self, main):
        super().__init__()
        self.main = main
        self.cid = None
        self.cur = None
        self._loading = False
        v = QVBoxLayout(self)
        v.setContentsMargins(0, 6, 0, 0)
        top = QHBoxLayout()
        self.filter = QLineEdit()
        self.filter.setPlaceholderText("Поиск по практике: суд, номер, слова из позиции…")
        self.filter.textChanged.connect(self.apply_filter)
        top.addWidget(self.filter, 1)
        top.addWidget(U.HelpButton("practice"))
        ex = QPushButton("Экспорт")
        em = QMenu(ex)
        em.addAction("Скопировать списком (для иска или отзыва)", self.copy_list)
        em.addAction("Сохранить в Word…", self.to_word)
        ex.setMenu(em)
        top.addWidget(ex)
        v.addLayout(top)
        quick = QHBoxLayout()
        self.quick = QLineEdit()
        self.quick.setPlaceholderText("➕  Вставьте реквизиты: «Определение ВС РФ от 12.03.2024 № 305-ЭС23-1234» "
                                      "или ссылку kad.arbitr.ru — и нажмите Enter")
        self.quick.returnPressed.connect(self.quick_add)
        quick.addWidget(self.quick, 1)
        quick.addWidget(_btn("Добавить", self.quick_add, primary=True))
        v.addLayout(quick)
        self.hint = QLabel("Вставьте реквизиты судебного акта в строку выше — вид акта, суд, дата и номер заполнятся "
                           "сами. Чтобы разложить практику по полочкам, создайте доводы кнопкой «+ Довод» "
                           "(например, «Неустойка явно несоразмерна») и добавляйте акты, выделив нужный довод.")
        self.hint.setObjectName("hint")
        self.hint.setWordWrap(True)
        v.addWidget(self.hint)
        split = QSplitter()
        split.setHandleWidth(1)
        left = QWidget()
        lv = QVBoxLayout(left)
        lv.setContentsMargins(0, 0, 0, 0)
        self.tree = PracticeTree(self)
        self.tree.currentItemChanged.connect(self.on_current)
        lv.addWidget(self.tree, 1)
        b = QHBoxLayout()
        b.addWidget(_btn("+ Довод", lambda: self.add("group"),
                         tip="Довод или вопрос: «Неустойка явно несоразмерна», «Подсудность», «Срок исковой давности»"))
        b.addWidget(_btn("+ Судебный акт", lambda: self.add("act"), tip="Пустая карточка акта — заполнить вручную"))
        b.addWidget(_btn("▲", lambda: self.move(-1)))
        b.addWidget(_btn("▼", lambda: self.move(1)))
        b.addWidget(_btn("Удалить", self.delete))
        b.addStretch(1)
        lv.addLayout(b)
        split.addWidget(left)
        # ---- справа: карточка довода или акта
        self.stack = QStackedWidget()
        g = QFrame()
        g.setObjectName("card")
        gf = QFormLayout(g)
        gf.setContentsMargins(14, 14, 14, 14)
        self.g_title = QLineEdit()
        self.g_title.setPlaceholderText("Напр.: Неустойка явно несоразмерна последствиям нарушения")
        self.g_note = QPlainTextEdit()
        self.g_note.setPlaceholderText("Суть довода, что нужно доказать, какие факты дела под него подходят")
        gf.addRow("Довод", self.g_title)
        gf.addRow("Суть", self.g_note)
        self.stack.addWidget(g)
        a = QFrame()
        a.setObjectName("card")
        af = QFormLayout(a)
        af.setContentsMargins(14, 14, 14, 14)
        self.preview = QLabel()
        self.preview.setWordWrap(True)
        self.preview.setTextInteractionFlags(Qt.TextSelectableByMouse)
        pf = self.preview.font()
        pf.setBold(True)
        pf.setPointSizeF(pf.pointSizeF() + 1)
        self.preview.setFont(pf)
        af.addRow(self.preview)
        self.a_kind = QComboBox()
        self.a_kind.setEditable(True)
        self.a_kind.addItems(PA.KINDS)
        self.a_court = QComboBox()
        self.a_court.setEditable(True)
        self.a_court.addItems(PA.COURTS)
        self.a_court.lineEdit().setPlaceholderText("Как в ссылке: «Верховного Суда РФ», «Арбитражного суда Московского округа»")
        row = QHBoxLayout()
        self.a_date = QLineEdit()
        self.a_date.setPlaceholderText("12.03.2024")
        self.a_date.setMaximumWidth(120)
        self.a_number = QLineEdit()
        self.a_number.setPlaceholderText("№ акта: 305-ЭС23-1234")
        self.a_case = QLineEdit()
        self.a_case.setPlaceholderText("№ дела: А40-12345/2023")
        row.addWidget(self.a_date)
        row.addWidget(self.a_number, 1)
        row.addWidget(self.a_case, 1)
        self.a_title = QLineEdit()
        self.a_title.setPlaceholderText("Необязательно — иначе название собирается из полей выше")
        urow = QHBoxLayout()
        self.a_url = QLineEdit()
        self.a_url.setPlaceholderText("Ссылка на текст акта (kad.arbitr.ru, vsrf.ru, КонсультантПлюс…)")
        urow.addWidget(self.a_url, 1)
        urow.addWidget(_btn("Открыть", self.open_url, tip="Открыть ссылку; если её нет — найти акт в КонсультантПлюс"))
        self.a_summary = QPlainTextEdit()
        self.a_summary.setPlaceholderText("Позиция суда коротко: что суд решил и почему это нам подходит")
        self.a_quote = QPlainTextEdit()
        self.a_quote.setPlaceholderText("Цитата из акта — дословно, для вставки в иск или отзыв")
        self.a_note = QPlainTextEdit()
        self.a_note.setPlaceholderText("Как применяем в деле: к какому доводу, чем наши обстоятельства похожи")
        self.a_note.setMaximumHeight(110)
        af.addRow("Вид акта", self.a_kind)
        af.addRow("Суд", self.a_court)
        af.addRow("Дата, номер", row)
        af.addRow("Ссылка", urow)
        af.addRow("Позиция суда", self.a_summary)
        af.addRow("Цитата", self.a_quote)
        af.addRow("Применение", self.a_note)
        af.addRow("Название", self.a_title)
        self.stack.addWidget(a)
        empty = QLabel("Выберите довод или судебный акт слева")
        empty.setAlignment(Qt.AlignCenter)
        empty.setObjectName("hint")
        self.stack.addWidget(empty)
        self.stack.setCurrentIndex(2)
        split.addWidget(self.stack)
        split.setSizes([480, 600])
        v.addWidget(split, 1)
        self.save_timer = QTimer(self)
        self.save_timer.setSingleShot(True)
        self.save_timer.setInterval(400)
        self.save_timer.timeout.connect(self.save_current)
        for sig in (self.g_title.textChanged, self.g_note.textChanged, self.a_kind.currentTextChanged,
                    self.a_court.currentTextChanged, self.a_date.textChanged, self.a_number.textChanged,
                    self.a_case.textChanged, self.a_title.textChanged, self.a_url.textChanged,
                    self.a_summary.textChanged, self.a_quote.textChanged, self.a_note.textChanged):
            sig.connect(self.schedule)

    # ---------------------------------------------------------------- загрузка
    def rows(self):
        return db().practice(self.cid) if self.cid else []

    def set_case(self, cid):
        self.save_now()
        self.cid = cid
        self.reload()

    def reload(self, select=None):
        self._loading = True
        self.tree.clear()
        rows = self.rows()
        by_parent = {}
        for r in rows:
            by_parent.setdefault(r["parent_id"], []).append(r)
        items = {}

        def build(parent_item, pid):
            for r in sorted(by_parent.get(pid, []), key=lambda x: (x["pos"], x["id"])):
                it = QTreeWidgetItem([self.label(r)])
                it.setData(0, Qt.UserRole, r["id"])
                it.setData(0, Qt.UserRole + 1, r["kind"])
                it.setToolTip(0, self.tip(r))
                if r["kind"] == "group":
                    f = it.font(0)
                    f.setBold(True)
                    f.setPointSizeF(f.pointSizeF() + 0.5)
                    it.setFont(0, f)
                    it.setFlags(it.flags() | Qt.ItemIsDropEnabled)
                else:
                    it.setFlags(it.flags() & ~Qt.ItemIsDropEnabled)
                if parent_item is None:
                    self.tree.addTopLevelItem(it)
                else:
                    parent_item.addChild(it)
                items[r["id"]] = it
                build(it, r["id"])
        build(None, 0)
        self.tree.expandAll()
        self._loading = False
        self.hint.setVisible(not rows)
        target = items.get(select) if select else None
        if target is None and self.tree.topLevelItemCount():
            target = self.tree.topLevelItem(0)
        if target:
            self.tree.setCurrentItem(target)
        else:
            self.on_current(None)
        self.apply_filter()

    @staticmethod
    def label(r):
        if r["kind"] == "group":
            return f"{GLYPH['group']}  {r['title'] or 'Новый довод'}"
        return f"{GLYPH['act']}  {PA.short_label(r) if (r.get('act_kind') or r.get('title') or r.get('court')) else 'Новый судебный акт'}"

    @staticmethod
    def tip(r):
        parts = [PA.label(r) if r.get("kind") == "act" else "", (r.get("summary") or "").strip(), (r.get("quote") or "").strip(), (r.get("note") or "").strip()]
        return "\n\n".join(p for p in parts if p)[:700]

    def on_current(self, it, _prev=None):
        if self._loading:
            return
        self.save_now()
        self.cur = it.data(0, Qt.UserRole) if it else None
        r = next((x for x in self.rows() if x["id"] == self.cur), None) if self.cur else None
        self._loading = True
        if r is None:
            self.stack.setCurrentIndex(2)
        elif r["kind"] == "group":
            self.stack.setCurrentIndex(0)
            self.g_title.setText(r["title"])
            self.g_note.setPlainText(r["note"])
        else:
            self.stack.setCurrentIndex(1)
            self.a_kind.setCurrentText(r["act_kind"])
            self.a_court.setCurrentText(r["court"])
            self.a_date.setText(r["act_date"])
            self.a_number.setText(r["number"])
            self.a_case.setText(r["case_no"])
            self.a_title.setText(r["title"])
            self.a_url.setText(r["url"])
            self.a_summary.setPlainText(r["summary"])
            self.a_quote.setPlainText(r["quote"])
            self.a_note.setPlainText(r["note"])
            self.preview.setText(PA.label(r))
        self._loading = False

    def _form(self):
        if self.stack.currentIndex() == 0:
            return dict(kind="group", title=self.g_title.text().strip(), note=self.g_note.toPlainText())
        return dict(kind="act", act_kind=self.a_kind.currentText().strip(), court=self.a_court.currentText().strip(),
                    act_date=self.a_date.text().strip(), number=self.a_number.text().strip().lstrip("№N ").strip(),
                    case_no=self.a_case.text().strip().lstrip("№N ").strip(), title=self.a_title.text().strip(),
                    url=self.a_url.text().strip(), summary=self.a_summary.toPlainText(),
                    quote=self.a_quote.toPlainText(), note=self.a_note.toPlainText())

    def schedule(self, *_):
        if self._loading or not self.cur:
            return
        self.save_timer.start()
        r = self._form()
        it = self.tree.currentItem()
        if it:
            it.setText(0, self.label(r))
        if r["kind"] == "act":
            self.preview.setText(PA.label(r))

    def save_current(self):
        if self.cur:
            r = self._form()
            r.pop("kind")
            db().update_practice(self.cur, **r)
            it = self.tree.currentItem()
            if it and it.data(0, Qt.UserRole) == self.cur:
                it.setToolTip(0, self.tip(r))

    def save_now(self):
        if self.save_timer.isActive():
            self.save_timer.stop()
            self.save_current()

    # ---------------------------------------------------------------- структура
    def _target_group(self):
        it = self.tree.currentItem()
        while it is not None and it.data(0, Qt.UserRole + 1) != "group":
            it = it.parent()
        return it.data(0, Qt.UserRole) if it is not None else 0

    def add(self, kind):
        if not self.cid:
            return
        self.save_now()
        parent = self._target_group() if kind == "act" else 0
        nid = db().add_practice(self.cid, parent, kind)
        self.reload(select=nid)
        (self.g_title if kind == "group" else self.a_kind).setFocus()

    def quick_add(self):
        text = self.quick.text().strip()
        if not text or not self.cid:
            return
        found = PA.parse(text)
        if not found:
            QMessageBox.information(self, "Судебная практика",
                                    "Не нашёл реквизитов судебного акта. Вставьте их так же, как в иске: "
                                    "«Определение ВС РФ от 12.03.2024 № 305-ЭС23-1234», «постановление Арбитражного "
                                    "суда Московского округа от 01.02.2024 по делу № А40-12345/2023» — или ссылку "
                                    "на акт (kad.arbitr.ru и др.). Карточку можно заполнить и вручную: «+ Судебный акт».")
            return
        self.save_now()
        have = {PA.key(r) for r in self.rows() if r["kind"] == "act"} - {""}
        group = self._target_group()
        added, last = 0, None
        for r in found:
            k = PA.key(r)
            if k and k in have:
                continue
            have.add(k)
            last = db().add_practice(self.cid, group, "act", **{f: r.get(f, "") for f in
                                                              ("title", "act_kind", "court", "act_date", "number",
                                                               "case_no", "url")})
            added += 1
        self.quick.clear()
        if last:
            self.reload(select=last)
            self.a_summary.setFocus()
        skipped = len(found) - added
        if added:
            self.main.toast(f"✓  Добавлено актов: {added}" + (f" (ещё {skipped} уже были)" if skipped else "")
                            + " — допишите позицию суда")
        else:
            self.main.toast("Эти акты уже есть в деле")

    def save_structure(self):
        """После перетаскивания: доводы — только наверху, акты — в доводе или наверху (не внутри акта)."""
        root = self.tree.invisibleRootItem()
        order = []
        for i in range(root.childCount()):
            top = root.child(i)
            if top.data(0, Qt.UserRole + 1) == "group":
                order.append((top, 0))
                for j in range(top.childCount()):
                    ch = top.child(j)
                    if ch.data(0, Qt.UserRole + 1) == "group":     # довод в доводе — вынести наверх, со своими актами
                        order.append((ch, 0))
                        for k in range(ch.childCount()):
                            order.append((ch.child(k), ch.data(0, Qt.UserRole)))
                    else:
                        order.append((ch, top.data(0, Qt.UserRole)))
                        for k in range(ch.childCount()):          # акт внутри акта — к его доводу
                            order.append((ch.child(k), top.data(0, Qt.UserRole)))
            else:
                order.append((top, 0))
                for k in range(top.childCount()):
                    order.append((top.child(k), 0))
        pos = {}
        for it, parent in order:
            n = pos.get(parent, 0)
            db().update_practice(it.data(0, Qt.UserRole), parent_id=parent, pos=n)
            pos[parent] = n + 1
        QTimer.singleShot(0, lambda c=self.cur: self.reload(select=c))

    def move(self, step):
        it = self.tree.currentItem()
        if not it:
            return
        par = it.parent() or self.tree.invisibleRootItem()
        i = par.indexOfChild(it)
        n = i + step
        if not (0 <= n < par.childCount()):
            return
        par.takeChild(i)
        par.insertChild(n, it)
        it.setExpanded(True)
        self.tree.setCurrentItem(it)
        self.save_structure()

    def delete(self):
        it = self.tree.currentItem()
        if not it:
            return
        pid, cid = it.data(0, Qt.UserRole), self.cid
        r = next((x for x in self.rows() if x["id"] == pid), None)
        if r is None:
            return
        name = r["title"] if r["kind"] == "group" else PA.label(r)
        extra = " вместе с его судебными актами" if it.childCount() else ""
        if QMessageBox.question(self, U.M.APP_NAME, f"Удалить «{name or 'без названия'}»{extra}?") != QMessageBox.Yes:
            return
        self.save_timer.stop()
        ids = db().practice_subtree(pid)
        snap = db().snapshot(("practice", f"id IN ({','.join('?' * len(ids))})", ids))
        db().delete_practice(pid)
        self.cur = None
        self.reload()

        def undo():
            db().restore(snap)
            if self.cid == cid:
                self.reload(select=pid)
        self.main.toast_undo(f"«{(name or 'без названия')[:60]}» удалено", undo)

    def open_url(self):
        u = self.a_url.text().strip()
        if not u:
            u = PA.search_url(self._form())
        QDesktopServices.openUrl(QUrl.fromUserInput(u))

    def apply_filter(self, *_):
        q = self.filter.text().strip().lower()
        rows = {r["id"]: r for r in self.rows()}

        def visit(item):
            r = rows.get(item.data(0, Qt.UserRole), {})
            hay = " ".join(str(r.get(k) or "") for k in ("title", "act_kind", "court", "act_date", "number",
                                                          "case_no", "summary", "quote", "note")).lower()
            hit = not q or all(w in hay for w in q.split())
            child_hit = False
            for i in range(item.childCount()):
                child_hit = visit(item.child(i)) or child_hit
            item.setHidden(not (hit or child_hit))
            return hit or child_hit
        for i in range(self.tree.topLevelItemCount()):
            visit(self.tree.topLevelItem(i))

    # ---------------------------------------------------------------- экспорт
    def structured(self):
        rows = self.rows()
        by_parent = {}
        for r in rows:
            by_parent.setdefault(r["parent_id"], []).append(r)
        out = []
        for r in sorted(by_parent.get(0, []), key=lambda x: (x["pos"], x["id"])):
            out.append(r)
            out += sorted(by_parent.get(r["id"], []), key=lambda x: (x["pos"], x["id"]))
        return out

    def copy_list(self):
        from PySide6.QtCore import QMimeData
        self.save_now()
        items = self.structured()
        if not items:
            return
        plain, htm = [], []
        for r in items:
            if r["kind"] == "group":
                plain.append(f"\n{(r['title'] or '').upper()}")
                htm.append(f"<p style='margin:12px 0 4px'><b>{html.escape(r['title'] or '')}</b></p>")
                continue
            ind = "    " if r["parent_id"] else ""
            line = f"{ind}— {PA.label(r)}."
            h = f"<p style='margin:4px 0 2px {18 if r['parent_id'] else 0}px'>— <b>{html.escape(PA.label(r))}</b>."
            if r["summary"].strip():
                line += f" {r['summary'].strip()}"
                h += f" {html.escape(r['summary'].strip())}"
            plain.append(line)
            htm.append(h + "</p>")
            if r["quote"].strip():
                plain.append(f"{ind}    «{r['quote'].strip()}»")
                htm.append(f"<p style='margin:0 0 4px {36 if r['parent_id'] else 18}px'><i>«{html.escape(r['quote'].strip())}»</i></p>")
        md = QMimeData()
        md.setText("\n".join(plain).strip())
        md.setHtml("".join(htm))
        QApplication.clipboard().setMimeData(md)
        self.main.toast("✓  Практика скопирована — вставьте в документ (Ctrl+V)")

    def to_word(self):
        self.save_now()
        items = self.structured()
        if not items:
            return
        c = db().case(self.cid) or {}
        p, _ = QFileDialog.getSaveFileName(self, "Сохранить в Word",
                                           os.path.join(c.get("folder") or self.main.default_dir(),
                                                        "Судебная практика — " + (c.get("title") or "дело") + ".docx"),
                                           "Word (*.docx)")
        if not p:
            return
        import docx
        from docx.shared import Pt, Cm
        d = docx.Document()
        d.styles["Normal"].font.name = "Times New Roman"
        d.styles["Normal"].font.size = Pt(12)
        d.add_heading(f"Судебная практика: {c.get('title', '')}", level=1)
        for r in items:
            if r["kind"] == "group":
                d.add_heading(r["title"] or "", level=2)
                if r["note"].strip():
                    d.add_paragraph(r["note"].strip()).runs[0].italic = True
                continue
            para = d.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.6 if r["parent_id"] else 0)
            para.add_run(PA.label(r)).bold = True
            if r["summary"].strip():
                para.add_run(". " + r["summary"].strip())
            if r["quote"].strip():
                q = d.add_paragraph(f"«{r['quote'].strip()}»")
                q.paragraph_format.left_indent = Cm(1.2)
                q.runs[0].italic = True
            if r["note"].strip():
                n = d.add_paragraph()
                n.paragraph_format.left_indent = Cm(1.2)
                n.add_run("Применение: ").bold = True
                n.add_run(r["note"].strip())
            if r["url"].strip():
                u = d.add_paragraph(r["url"].strip())
                u.paragraph_format.left_indent = Cm(1.2)
        try:
            d.save(p)
        except Exception as e:
            return QMessageBox.warning(self, U.M.APP_NAME, f"Не удалось сохранить файл (возможно, он открыт в Word):\n{e}")
        self.main.toast(f"✓  Сохранено: {os.path.basename(p)}")
        QDesktopServices.openUrl(QUrl.fromLocalFile(p))
