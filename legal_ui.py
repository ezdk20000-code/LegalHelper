# -*- coding: utf-8 -*-
"""Интерфейс юридических функций: калькуляторы, дела, пакет в суд, ф. 107, проверка, обезличивание,
поиск по делу, шаблоны, выписки, справка «?»."""
import os
import re
import json
import datetime as dt
import html
from pathlib import Path
from types import SimpleNamespace

import pymupdf as fitz
from PySide6.QtCore import Qt, QDate, QTimer, QUrl, Signal
from PySide6.QtGui import QDesktopServices, QColor, QFont
from PySide6.QtWidgets import (
    QWidget, QDialog, QVBoxLayout, QHBoxLayout, QFormLayout, QGridLayout, QLabel, QPushButton, QToolButton,
    QComboBox, QLineEdit, QSpinBox, QDoubleSpinBox, QDateEdit, QCheckBox, QTextBrowser, QPlainTextEdit,
    QTableWidget, QTableWidgetItem, QHeaderView, QAbstractItemView, QListWidget, QListWidgetItem,
    QTabWidget, QMessageBox, QFileDialog, QDialogButtonBox, QApplication, QSplitter, QFrame, QInputDialog,
    QRadioButton, QButtonGroup, QScrollArea, QSizePolicy, QMenu, QSystemTrayIcon, QStackedWidget)

import pdf_core as C
import legal_core as L
import legal_data as D
import cases as CS
import sys as _sys
import case_tabs as CT
import backup as BK
import casefile as CF
import anim
import folders_ui as FU
CT.U = _sys.modules[__name__]

M = None          # пространство имён главного модуля (bind)
_db = None


def bind(ns):
    global M
    M = SimpleNamespace(**ns)


def db():
    global _db
    if _db is None:
        _db = CS.CaseDB(os.path.join(M.data_dir(), "cases.sqlite"))
    return _db


def templates_dir():
    d = os.path.join(M.data_dir(), "Шаблоны")
    os.makedirs(d, exist_ok=True)
    return d


# =============================================================================
#  Справка «?»
# =============================================================================
HELP = {
    "submission": ("Подача: комплект документов",
        "Здесь вы собираете документы для подачи по делу — по одному, в удобном темпе, между делами."
        "<ul><li><b>Комплект</b> — то, что подаёте: иск, отзыв, жалоба. У дела может быть несколько комплектов "
        "(меню «⋯»).</li>"
        "<li>Первая строка — <b>основной документ</b> (иск, жалоба), дальше — приложения в нужном порядке "
        "(▲ ▼).</li>"
        "<li><b>«Типовой перечень»</b> добавит обычный список приложений по АПК, ГПК, КАС — останется "
        "прикрепить файлы и убрать лишнее.</li>"
        "<li>Файл прикрепляется кнопкой «Прикрепить файл…», двойным щелчком по колонке «Файл» или "
        "перетаскиванием файла на строку. Слева сразу виден сам документ.</li>"
        "<li>Отмечайте <b>✓</b> то, что проверено и готово. Сверху видно «Готово N из M».</li>"
        "<li>Всё сохраняется <b>автоматически</b>: закрыли программу — при следующем запуске откроется то же дело "
        "и тот же комплект.</li>"
        "<li><b>«Собрать пакет…»</b> передаёт документы в «Пакет в суд / на почту» — там выбирается площадка "
        "(«Мой арбитр», ГАС «Правосудие», Почта России) и формируются файлы и перечень приложений.</li></ul>"
        "Файлы не копируются — программа хранит ссылки на них. Если файл переместить, строка покажет "
        "«⚠ не найден» — прикрепите его заново."),
    "board": ("Карта дела",
        "Встроенный редактор <b>Excalidraw</b> (открытый проект, работает без интернета): схемы, "
        "интеллект-карты, стрелки, заметки, картинки."
        "<ul><li>Для нового дела создаётся шаблон: в центре дело, вокруг — «Факты и хронология», "
        "«Позиция доверителя», «Позиция оппонента», «Доказательства», «Риски», «Процесс и сроки».</li>"
        "<li>Двойной щелчок по пустому месту — текст; по фигуре — редактировать надпись. Стрелку можно "
        "привязать к фигурам — она будет двигаться вместе с ними.</li>"
        "<li>Клавиши: <b>R</b> — прямоугольник, <b>O</b> — эллипс, <b>A</b> — стрелка, <b>T</b> — текст, "
        "<b>Ctrl+Z</b> — отменить. Колесо мыши с Ctrl — масштаб, пробел + мышь — перемещение.</li>"
        "<li>Картинку (скриншот документа) можно вставить через Ctrl+V.</li>"
        "<li>Карта сохраняется <b>автоматически</b> каждые 2 секунды, отдельно для каждого дела.</li>"
        "<li>Меню ☰ на карте: экспорт в PNG/SVG, сохранение в файл .excalidraw, смена фона.</li></ul>"),
    "laws": ("Нормы права",
        "Список всех норм, которые вы применяете в деле, разложенный по полочкам."
        "<ul><li><b>Направление</b> — тема или довод: «Неустойка», «Моральный вред», «Штраф 50%», "
        "«Подсудность».</li>"
        "<li><b>Акт</b> — закон, кодекс, постановление Пленума, определение ВС (есть список частых).</li>"
        "<li><b>Статья / пункт</b> — конкретная норма или правовая позиция: название («п. 6 ст. 13»), текст "
        "и как вы её применяете в деле.</li></ul>"
        "Один и тот же акт можно добавить в разные направления — например, разные пункты Закона о защите "
        "прав потребителей под разные доводы. Порядок меняется кнопками ▲ ▼ или перетаскиванием мышью "
        "(в том числе в другое направление).<br><br>"
        "<b>Экспорт</b>: «Скопировать списком» — готовый текст для иска или позиции; «В Word» — отдельный "
        "документ; «Взять из другого дела» — перенести нормы из похожего дела. Всё сохраняется автоматически."),
    "package": ("Пакет в суд / на почту",
        "Собирает иск (жалобу, письмо) и приложения в пакет, который примет выбранная площадка."
        "<ul><li><b>«Мой арбитр» и ГАС «Правосудие»</b> — каждый документ отдельным PDF до 30 МБ, в имени "
        "файла название и число листов («02. Копия договора 6 л.pdf»). Большие файлы сжимаются, а если не "
        "помогает — делятся на части.</li>"
        "<li><b>Почта России, бумажная доставка</b> — один файл, все страницы A4, всего до 19 страниц и 20 МБ. "
        "Если страниц больше — пакет делится на несколько писем.</li>"
        "<li><b>Почта России, электронная доставка</b> — до 20 МБ.</li></ul>"
        "Первым в списке ставьте основной документ. Программа составит <b>перечень приложений</b> "
        "(«1. Копия договора — на 6 л. в 1 экз.») — его можно скопировать или добавить последней страницей "
        "в основной документ. Из этого же списка можно сразу сделать <b>опись ф. 107</b>."),
    "sheetnum": ("Нумерация листов дела",
        "Ставит номера листов в правом верхнем углу, как в материалах дела. "
        "<ul><li><b>Листов в томе</b> — по Инструкции по делопроизводству том не больше 250 листов; "
        "0 — без деления на тома.</li><li><b>Двусторонние листы</b> — номер ставится только на лицевой "
        "стороне (нечётные страницы).</li><li>Можно сохранить каждый том отдельным файлом с титульным листом.</li></ul>"),
    "certify": ("Штамп заверения",
        "Ставит штамп «Копия верна» с должностью, ФИО, датой — на всех страницах или только на последней. "
        "Можно добавить картинку подписи (PNG с прозрачным фоном) и надпись «Прошито, пронумеровано N листов» "
        "на последней странице. Наборы строк сохраняются как шаблоны."),
    "f107": ("Опись вложения ф. 107",
        "Официальный бланк Почты России — точно такой же, как на pochta.ru: альбомный лист A4, два одинаковых "
        "экземпляра рядом (разрезать по середине). Один вкладывается в отправление, второй с оттиском штемпеля "
        "остаётся у вас. На листе 14 строк; длинное название переносится на следующие строки целиком, без "
        "обрезки. Не поместилось — несколько листов со сквозной нумерацией, общий итог на последнем. Почтовый идентификатор можно вписать позже от руки. "
        "Если ценность не объявляется — оставьте графу пустой (будет прочерк)."),
    "preflight": ("Проверка перед подачей",
        "Находит то, из-за чего документы возвращают: превышение размера, пустые и перевёрнутые страницы, "
        "сканы без текстового слоя (нужен OCR), низкое разрешение (< 150 dpi вместо 200–300), не A4 для почты, "
        "защиту от копирования. Двойной щелчок по строке — перейти к странице."),
    "anonymize": ("Обезличивание (152-ФЗ)",
        "Находит персональные данные: ФИО, паспорта, СНИЛС, ИНН, счета и карты, телефоны, e-mail, адреса, "
        "даты рождения, госномера. Сначала показывает список найденного — снимите галочки с лишнего, затем "
        "данные удаляются <b>безвозвратно</b> (из текста, а не просто закрашиваются). Поиск эвристический: "
        "всегда просматривайте результат, особенно ФИО в падежах и адреса."),
    "compare_ed": ("Сравнение редакций",
        "Сравнивает две редакции договора или документа (PDF, DOCX, TXT) по словам и показывает правки "
        "в стиле «было / стало»: <del style='color:#b3261e'>удалённое</del> и "
        "<ins style='color:#0a7a28'>добавленное</ins>. Отчёт сохраняется в HTML (открывается в браузере, "
        "можно переслать) и PDF."),
    "case_search": ("Поиск по документам",
        "Ищет слово или фразу сразу во всех документах дела (или в выбранных файлах / папке). "
        "Сканы без текстового слоя найти нельзя — программа их перечислит, прогоните их через OCR."),
    "template": ("Шаблоны документов",
        "Готовые ходатайства и заявления по АПК РФ и ГПК РФ: ознакомление с материалами дела, отложение заседания, "
        "приобщение и истребование доказательств, участие по веб-конференции / ВКС, рассмотрение в отсутствие, "
        "исполнительный лист, копия судебного акта, судебные расходы; а также адвокатский запрос, претензия и "
        "возврат госпошлины."
        "<ul><li>Выберите дело — суд, номер дела, судья, доверитель и оппонент подставятся сами.</li>"
        "<li>Ваши ФИО и контакты задаются один раз в «Мои реквизиты».</li>"
        "<li>Остальные поля заполните в окне — введённое запоминается для этого дела.</li>"
        "<li>Готовый документ сохраняется в папку дела и сразу появляется в его документах. "
        "Можно сразу получить PDF и открыть его в программе.</li>"
        "<li><b>Свой шаблон</b> — обычный Word-файл, где вместо данных стоят поля в фигурных скобках: "
        "<code>{Суд}</code>, <code>{Номер_дела}</code>, <code>{Доверитель}</code>, <code>{Сумма_долга}</code>… "
        "Кнопка «+ Свой шаблон» добавит его в «Мои шаблоны». Встроенные шаблоны можно править в Word — "
        "программа их не перезапишет.</li></ul>"),
    "quote": ("Выписки",
        "Выделите рамкой фрагмент на странице — текст попадёт в выписки выбранного дела с пометкой «л. N». "
        "Потом выписки можно скопировать или выгрузить в Word для позиции по делу."),
    "calc_deadline": ("Процессуальные сроки",
        "Течение срока начинается на следующий день после события. Сроки в днях считаются в рабочих днях "
        "(АПК, ГПК, КАС), календарные — если так указано (сутки в УПК, претензионный срок). Если последний "
        "день нерабочий — срок переносится на следующий рабочий день. Учтены праздники и переносы по "
        "постановлениям Правительства на 2024–2027 годы. Результат можно добавить в дело — он появится "
        "в календаре с напоминанием."),
    "calc_duty": ("Госпошлина",
        "Размеры по ст. 333.19 (суды общей юрисдикции) и ст. 333.21 (арбитражные суды) НК РФ в редакции "
        "259-ФЗ — для заявлений, поданных после 08.09.2024. Пошлина округляется до полного рубля. "
        "Льготы (ст. 333.35–333.37 НК) и особые случаи не учитываются."),
    "calc_interest": ("Проценты и неустойка",
        "<b>Ст. 395 ГК</b> — ключевая ставка Банка России по периодам её действия, база — 365/366 дней. "
        "<b>Неустойка</b> — процент в день или доля ключевой ставки (1/300, 1/150, 1/130). "
        "Оплаты уменьшают долг со следующего дня (день оплаты включается в просрочку); увеличения долга — "
        "с указанной даты. Можно исключить мораторий 01.04.2022–01.10.2022 (ПП РФ № 497). "
        "Таблицу можно вставить в иск (Word) или сохранить в PDF."),
    "cases": ("Дела",
        "Карточки дел хранятся только на этом компьютере. В карточке — стороны, суд, стадия, сроки и "
        "заседания (с напоминаниями), документы, учёт времени и оплат с актом, выписки. "
        "Напоминания показываются при запуске программы и за день до события."),
    "ocr": ("OCR — распознавание текста",
        "Добавляет к сканам невидимый текстовый слой: документ выглядит так же, но в нём работает поиск "
        "и копирование. Суды просят документы «с возможностью копирования текста»."),
    "compress": ("Сжатие PDF",
        "Уменьшает размер за счёт картинок (сканов). «Рекомендуемое» — 150 dpi, текст остаётся чётким. "
        "Для подачи в суд обычно достаточно."),
    "pdfa": ("PDF/A", "Архивный формат для долгого хранения. Нужна бесплатная программа Ghostscript."),
    "redact": ("Скрыть данные",
        "Находит слова и шаблоны (e-mail, телефоны, карты, паспорта) и удаляет их безвозвратно. "
        "Для персональных данных удобнее «Обезличивание (152-ФЗ)»."),
    "split": ("Разделение PDF",
        "Каждая страница отдельным файлом, каждые N страниц, или по диапазонам: «1-3, 4-10, 11-»."),
    "protect": ("Пароль", "Шифрование AES-256. Можно запретить печать, копирование, изменение."),
    "compare": ("Сравнить PDF",
        "Техническое сравнение двух PDF по тексту и картинке страниц. Для договоров удобнее "
        "«Сравнение редакций» — там правки видны по словам."),
    "pagesize": ("Размер страниц", "Приводит страницы к A4 или другому формату. Текст остаётся текстом."),
    "forms": ("PDF-формы", "Заполнение полей интерактивных форм; «вшить» — сделать заполненное неизменяемым."),
}


def show_help(parent, key):
    title, text = HELP.get(key, ("Справка", "Нет справки"))
    HelpPopup(parent, title, text).exec()


class HelpPopup(QDialog):
    """Подсказка к функции (кнопка «?»): заголовок со значком, текст переносится по ширине окна и никогда
    не обрезается — длинный текст можно прокрутить, окно подстраивает высоту под содержимое."""
    WIDTH = 600

    def __init__(self, parent, title, text):
        super().__init__(parent)
        import modern_ui
        t = M.T if M else {"accent": "#007aff", "text": "#1c1c1e", "muted": "#8e8e93"}
        self.setWindowTitle(title)
        self.setObjectName("tooldlg")
        lay = QVBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(0)
        head = QHBoxLayout()
        head.setContentsMargins(24, 20, 24, 6)
        head.setSpacing(14)
        ic = QLabel()
        ic.setPixmap(modern_ui.emoji_tile("💡", t["accent"], 44))
        ic.setFixedSize(44, 44)
        head.addWidget(ic, 0, Qt.AlignTop)
        lt = QLabel(title)
        lt.setObjectName("dlgtitle")
        lt.setWordWrap(True)
        head.addWidget(lt, 1)
        lay.addLayout(head)
        view = QTextBrowser()
        view.setObjectName("helppop")
        view.setOpenExternalLinks(True)
        view.setFrameShape(QFrame.NoFrame)
        view.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        view.document().setDocumentMargin(0)
        view.setHtml(f"<html><body style='color:{t['text']}; font-size:10.5pt; line-height:150%'>{text}"
                     "</body></html>")
        self.view = view
        self._scr = (parent.screen() if parent else QApplication.primaryScreen()).availableGeometry()
        body = QHBoxLayout()
        body.setContentsMargins(24, 6, 24, 18)
        body.addWidget(view)
        lay.addLayout(body)
        foot = QFrame()
        foot.setObjectName("dlgfoot")
        row = QHBoxLayout(foot)
        row.setContentsMargins(24, 12, 24, 14)
        row.addStretch(1)
        ok = QPushButton("Понятно")
        ok.setObjectName("primary")
        ok.setDefault(True)
        ok.setMinimumWidth(110)
        ok.setCursor(Qt.PointingHandCursor)
        ok.clicked.connect(self.accept)
        row.addWidget(ok)
        lay.addWidget(foot)
        self.setFixedWidth(self.WIDTH)
        self._fit()

    def _fit(self):
        """Высота текста — по содержимому (после того как применены шрифты темы), не выше 62% экрана."""
        v = self.view
        v.ensurePolished()
        v.document().setDefaultFont(v.font())
        v.document().setTextWidth(self.WIDTH - 48 - 4)
        h = int(v.document().size().height()) + 12
        v.setFixedHeight(max(60, min(h, int(self._scr.height() * 0.62))))

    def showEvent(self, e):
        self._fit()
        super().showEvent(e)
        QTimer.singleShot(0, self._refit)

    def _refit(self):
        """Уже показанное окно: документ разложен по настоящей ширине — подогнать высоту точно."""
        v = self.view
        cap = int(self._scr.height() * 0.62)
        for _ in range(4):                       # пока есть что прокручивать — добавить высоты (до предела)
            extra = v.verticalScrollBar().maximum()
            if extra <= 0 or v.height() >= cap:
                break
            v.setFixedHeight(min(cap, v.height() + extra + 4))
            QApplication.processEvents()
        self.adjustSize()


class HelpButton(QToolButton):
    def __init__(self, key, parent=None):
        super().__init__(parent)
        self.key = key
        self.setText("?")
        self.setObjectName("help")
        self.setToolTip("Подробнее об этой функции")
        self.setCursor(Qt.PointingHandCursor)
        self.setFixedSize(22, 22)
        self.clicked.connect(lambda: show_help(self.window(), self.key))


def title_row(text, help_key=None, big=False):
    row = QHBoxLayout()
    lab = QLabel(text)
    lab.setObjectName("title" if big else "subtitle")
    row.addWidget(lab)
    if help_key:
        row.addWidget(HelpButton(help_key))
    row.addStretch(1)
    return row


# =============================================================================
#  Мелкие виджеты
# =============================================================================
def date_edit(d=None):
    w = QDateEdit()
    w.setCalendarPopup(True)
    w.setDisplayFormat("dd.MM.yyyy")
    d = d or dt.date.today()
    w.setDate(QDate(d.year, d.month, d.day))
    return w


def qdate(w):
    q = w.date()
    return dt.date(q.year(), q.month(), q.day())


def money_edit(v=0.0):
    w = QDoubleSpinBox()
    w.setRange(0, 1e13)
    w.setDecimals(2)
    w.setGroupSeparatorShown(True)
    w.setValue(v)
    w.setSuffix(" ₽")
    return w


WEEKDAYS = ["понедельник", "вторник", "среда", "четверг", "пятница", "суббота", "воскресенье"]


def pick_case(parent, title="Выберите дело"):
    items = db().cases()
    if not items:
        if QMessageBox.question(parent, M.APP_NAME, "Дел пока нет. Создать новое дело?") != QMessageBox.Yes:
            return None
        name, ok = QInputDialog.getText(parent, "Новое дело", "Название дела (например, «ООО Ромашка — взыскание долга»):")
        if not ok or not name.strip():
            return None
        return db().add_case(title=name.strip())
    labels = [f"{c['title']}" + (f"  ·  {c['number']}" if c["number"] else "") for c in items]
    s, ok = QInputDialog.getItem(parent, title, "Дело:", labels, 0, False)
    if not ok:
        return None
    return items[labels.index(s)]["id"]


class ResultView(QWidget):
    """Результат расчёта + экспорт (Word, PDF, буфер обмена)."""

    def __init__(self):
        super().__init__()
        lay = QVBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 0)
        self.view = QTextBrowser()
        self.view.setOpenExternalLinks(True)
        self.view.setObjectName("result")
        lay.addWidget(self.view, 1)
        row = QHBoxLayout()
        self.b_copy = QPushButton("Копировать")
        self.b_word = QPushButton("В Word…")
        self.b_pdf = QPushButton("В PDF…")
        for b in (self.b_copy, self.b_word, self.b_pdf):
            row.addWidget(b)
            b.setEnabled(False)
        row.addStretch(1)
        lay.addLayout(row)
        self.data = None
        self.b_copy.clicked.connect(self.copy)
        self.b_word.clicked.connect(self.to_word)
        self.b_pdf.clicked.connect(self.to_pdf)

    def set_html(self, html_text, data=None):
        self.view.setHtml(html_text)
        self.data = data
        for b in (self.b_copy, self.b_word, self.b_pdf):
            b.setEnabled(data is not None)

    def copy(self):
        from PySide6.QtCore import QMimeData
        title, head, body, foot = self.data
        md = QMimeData()
        md.setHtml(L.table_html(title, head, body, foot))
        md.setText(L.table_tsv(head, body, foot))
        QApplication.clipboard().setMimeData(md)
        QMessageBox.information(self, M.APP_NAME, "Таблица скопирована — вставьте её в Word (Ctrl+V).")

    def _ask(self, suffix, flt):
        p, _ = QFileDialog.getSaveFileName(self, "Сохранить", os.path.join(str(Path.home() / "Documents"),
                                           "Расчёт" + suffix), flt)
        if p and not p.lower().endswith(suffix):
            p += suffix
        return p

    def to_word(self):
        p = self._ask(".docx", "Word (*.docx)")
        if p:
            title, head, body, foot = self.data
            L.table_docx(p, title, head, body, foot)
            QDesktopServices.openUrl(QUrl.fromLocalFile(p))

    def to_pdf(self):
        p = self._ask(".pdf", "PDF (*.pdf)")
        if p:
            title, head, body, foot = self.data
            L.html_to_pdf_file(L.table_html(title, head, body, foot), p)
            QDesktopServices.openUrl(QUrl.fromLocalFile(p))


def _small_btn(text):
    b = QPushButton(text)
    b.setFixedWidth(38)
    b.setStyleSheet("padding: 4px 0px;")
    return b


def card(widget_or_layout):
    f = QFrame()
    f.setObjectName("card")
    lay = QVBoxLayout(f)
    lay.setContentsMargins(16, 14, 16, 14)
    if isinstance(widget_or_layout, QWidget):
        lay.addWidget(widget_or_layout)
    else:
        lay.addLayout(widget_or_layout)
    return f


# =============================================================================
#  Калькуляторы
# =============================================================================
class DeadlineTab(QWidget):
    def __init__(self):
        super().__init__()
        lay = QVBoxLayout(self)
        lay.addLayout(title_row("Процессуальный срок", "calc_deadline"))
        f = QFormLayout()
        self.preset = QComboBox()
        self.preset.addItems([p[0] for p in D.DEADLINE_PRESETS])
        self.start = date_edit()
        self.n = QSpinBox()
        self.n.setRange(1, 3650)
        self.unit = QComboBox()
        self.units = list(D.UNIT_NAMES)
        self.unit.addItems(list(D.UNIT_NAMES.values()))
        rown = QHBoxLayout()
        rown.addWidget(self.n)
        rown.addWidget(self.unit, 1)
        self.hint = QLabel()
        self.hint.setObjectName("hint")
        self.hint.setWordWrap(True)
        f.addRow("Вид срока", self.preset)
        f.addRow("", self.hint)
        f.addRow("Дата события", self.start)
        f.addRow("Длительность", rown)
        lay.addWidget(card(f))
        self.res = QLabel()
        self.res.setObjectName("bigresult")
        self.res.setTextInteractionFlags(Qt.TextSelectableByMouse)
        self.notes = QLabel()
        self.notes.setWordWrap(True)
        self.notes.setObjectName("hint")
        b_add = QPushButton("Добавить в дело…")
        b_add.clicked.connect(self.add_to_case)
        r = QVBoxLayout()
        r.addWidget(self.res)
        r.addWidget(self.notes)
        rb = QHBoxLayout()
        rb.addWidget(b_add)
        rb.addStretch(1)
        r.addLayout(rb)
        lay.addWidget(card(r))
        # между датами
        lay.addLayout(title_row("Дней между датами"))
        g = QHBoxLayout()
        self.d1, self.d2 = date_edit(), date_edit(dt.date.today() + dt.timedelta(days=30))
        g.addWidget(QLabel("с"))
        g.addWidget(self.d1)
        g.addWidget(QLabel("по"))
        g.addWidget(self.d2)
        self.between = QLabel()
        g.addWidget(self.between, 1)
        lay.addWidget(card(g))
        lay.addStretch(1)
        self.preset.currentIndexChanged.connect(self.on_preset)
        for w in (self.start, self.d1, self.d2):
            w.dateChanged.connect(self.calc)
        self.n.valueChanged.connect(self.calc)
        self.unit.currentIndexChanged.connect(self.calc)
        self.on_preset(0)

    def on_preset(self, i):
        name, n, unit, hint = D.DEADLINE_PRESETS[i]
        self.n.blockSignals(True)
        self.unit.blockSignals(True)
        self.n.setValue(n)
        self.unit.setCurrentIndex(self.units.index(unit))
        self.n.blockSignals(False)
        self.unit.blockSignals(False)
        self.hint.setText(hint)
        self.calc()

    def calc(self):
        d, notes = L.compute_deadline(qdate(self.start), self.n.value(), self.units[self.unit.currentIndex()])
        self.result = d
        self.res.setText(f"Последний день срока: <b>{L.ddmmyyyy(d)}</b> ({WEEKDAYS[d.weekday()]})")
        left = (d - dt.date.today()).days
        extra = f"Осталось дней: {left}." if left >= 0 else f"Срок истёк {-left} дн. назад."
        self.notes.setText(" ".join(notes + [extra]))
        cal, work = L.days_between(qdate(self.d1), qdate(self.d2))
        self.between.setText(f"   календарных: <b>{cal}</b>, рабочих: <b>{work}</b> (включительно)")

    def add_to_case(self):
        cid = pick_case(self)
        if not cid:
            return
        title = self.preset.currentText() if self.preset.currentIndex() else \
            f"Срок {self.n.value()} {self.unit.currentText()} от {L.ddmmyyyy(qdate(self.start))}"
        title, ok = QInputDialog.getText(self, "Срок в деле", "Название:", text=title.split(" — ")[0])
        if ok:
            db().add_event(cid, self.result, "Срок", title)
            QMessageBox.information(self, M.APP_NAME, f"Срок {L.ddmmyyyy(self.result)} добавлен в дело.")


class DutyTab(QWidget):
    def __init__(self):
        super().__init__()
        lay = QVBoxLayout(self)
        lay.addLayout(title_row("Государственная пошлина", "calc_duty"))
        f = QFormLayout()
        self.court = QComboBox()
        self.court.addItems(["Арбитражный суд (ст. 333.21 НК)", "Суд общей юрисдикции (ст. 333.19 НК)"])
        self.kind = QComboBox()
        self.person = QComboBox()
        self.person.addItems(["Организация", "Физическое лицо / ИП"])
        self.amount = money_edit(100000)
        f.addRow("Суд", self.court)
        f.addRow("Заявление", self.kind)
        f.addRow("Заявитель", self.person)
        f.addRow("Цена иска", self.amount)
        lay.addWidget(card(f))
        self.res = QLabel()
        self.res.setObjectName("bigresult")
        self.res.setTextInteractionFlags(Qt.TextSelectableByMouse)
        self.expl = QLabel()
        self.expl.setWordWrap(True)
        self.expl.setObjectName("hint")
        r = QVBoxLayout()
        r.addWidget(self.res)
        r.addWidget(self.expl)
        lay.addWidget(card(r))
        lay.addStretch(1)
        self.court.currentIndexChanged.connect(self.fill_kinds)
        self.kind.currentIndexChanged.connect(self.calc)
        self.person.currentIndexChanged.connect(self.calc)
        self.amount.valueChanged.connect(self.calc)
        self.fill_kinds()

    def fill_kinds(self):
        self.kind.blockSignals(True)
        self.kind.clear()
        fixed = D.DUTY_ARB_FIXED if self.court.currentIndex() == 0 else D.DUTY_SOJ_FIXED
        self.kind.addItems(["Имущественный иск (по цене иска)", "Судебный приказ"] + [n for n, _ in fixed])
        self.kind.blockSignals(False)
        self.calc()

    def calc(self):
        court = "arb" if self.court.currentIndex() == 0 else "soj"
        k = self.kind.currentIndex()
        person = "org" if self.person.currentIndex() == 0 else "fl"
        self.amount.setEnabled(k < 2)
        if k < 2:
            v, e = L.state_duty(court, self.amount.value(), order=(k == 1))
        else:
            v, e = L.state_duty(court, fixed_index=k - 2, person=person)
        self.res.setText(f"Госпошлина: <b>{L.money(v, False)} ₽</b>")
        self.expl.setText(e)


class InterestTab(QWidget):
    MODES = ["Проценты по ст. 395 ГК (ключевая ставка)", "Неустойка — % в день",
             "Неустойка — доля ключевой ставки (1/300, 1/150…)"]

    def __init__(self):
        super().__init__()
        lay = QVBoxLayout(self)
        lay.addLayout(title_row("Проценты и неустойка", "calc_interest"))
        top = QHBoxLayout()
        f = QFormLayout()
        self.mode = QComboBox()
        self.mode.addItems(self.MODES)
        self.principal = money_edit(100000)
        self.d_from = date_edit(dt.date.today() - dt.timedelta(days=180))
        self.d_to = date_edit()
        self.pct = QDoubleSpinBox()
        self.pct.setRange(0.0001, 100)
        self.pct.setDecimals(4)
        self.pct.setValue(0.1)
        self.pct.setSuffix(" % в день")
        self.div = QComboBox()
        self.div.addItems(["1/300", "1/150", "1/130", "1/360"])
        self.fixed = QCheckBox("Ставка на дату:")
        self.fixed_date = date_edit()
        rf = QHBoxLayout()
        rf.addWidget(self.fixed)
        rf.addWidget(self.fixed_date)
        rf.addStretch(1)
        self.morat = QCheckBox("Исключить мораторий 01.04–01.10.2022")
        self.capc = QCheckBox("Не больше суммы долга")
        f.addRow("Что считаем", self.mode)
        f.addRow("Сумма долга", self.principal)
        f.addRow("Просрочка с", self.d_from)
        f.addRow("по (включительно)", self.d_to)
        f.addRow("Размер", self.pct)
        f.addRow("Доля ставки", self.div)
        f.addRow("", rf)
        f.addRow("", self.morat)
        f.addRow("", self.capc)
        self.form = f
        top.addWidget(card(f), 3)
        # оплаты
        pv = QVBoxLayout()
        pv.addWidget(QLabel("Оплаты и увеличения долга"))
        self.moves = QTableWidget(0, 3)
        self.moves.setHorizontalHeaderLabels(["Дата", "Сумма, ₽", "Тип"])
        self.moves.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.moves.verticalHeader().setVisible(False)
        pv.addWidget(self.moves)
        rb = QHBoxLayout()
        b1, b2 = QPushButton("+ Оплата"), QPushButton("+ Долг")
        b3 = QPushButton("Удалить")
        b1.clicked.connect(lambda: self.add_move("Оплата"))
        b2.clicked.connect(lambda: self.add_move("Увеличение долга"))
        b3.clicked.connect(lambda: self.moves.removeRow(self.moves.currentRow()))
        for b in (b1, b2, b3):
            rb.addWidget(b)
        pv.addLayout(rb)
        top.addWidget(card(pv), 2)
        lay.addLayout(top)
        rowc = QHBoxLayout()
        b = QPushButton("Рассчитать")
        b.setObjectName("primary")
        b.clicked.connect(self.calc)
        self.total = QLabel()
        self.total.setObjectName("bigresult")
        rowc.addWidget(b)
        rowc.addSpacing(12)
        rowc.addWidget(self.total, 1)
        lay.addLayout(rowc)
        self.result = ResultView()
        lay.addWidget(self.result, 1)
        self.mode.currentIndexChanged.connect(self.on_mode)
        self.on_mode(0)

    def on_mode(self, i):
        for w, vis in ((self.pct, i == 1), (self.div, i == 2), (self.fixed, i == 2), (self.fixed_date, i == 2),
                       (self.capc, i > 0)):
            w.setVisible(vis)
            lab = self.form.labelForField(w)
            if lab:
                lab.setVisible(vis)

    def add_move(self, kind):
        r = self.moves.rowCount()
        self.moves.insertRow(r)
        de = date_edit()
        me = money_edit(0)
        self.moves.setCellWidget(r, 0, de)
        self.moves.setCellWidget(r, 1, me)
        it = QTableWidgetItem(kind)
        it.setFlags(Qt.ItemIsEnabled)
        self.moves.setItem(r, 2, it)

    def calc(self):
        pays, adds = [], []
        for r in range(self.moves.rowCount()):
            d = qdate(self.moves.cellWidget(r, 0))
            a = self.moves.cellWidget(r, 1).value()
            (pays if self.moves.item(r, 2).text() == "Оплата" else adds).append((d, a))
        m = self.mode.currentIndex()
        try:
            res = L.interest_calc(self.principal.value(), qdate(self.d_from), qdate(self.d_to), pays, adds,
                                  mode=["395", "pct", "frac"][m], pct_per_day=self.pct.value(),
                                  rate_divisor=int(self.div.currentText().split("/")[1]),
                                  fixed_rate_date=qdate(self.fixed_date) if (m == 2 and self.fixed.isChecked()) else None,
                                  exclude_moratorium=self.morat.isChecked(), cap_principal=self.capc.isChecked() and m > 0)
        except ValueError as e:
            QMessageBox.warning(self, M.APP_NAME, str(e))
            return
        head, body = L.calc_table(res, m)
        title = ["Расчёт процентов по ст. 395 ГК РФ", "Расчёт неустойки", "Расчёт неустойки"][m]
        title += f" на сумму {L.money(self.principal.value())} ₽ за период {L.ddmmyyyy(qdate(self.d_from))} – " \
                 f"{L.ddmmyyyy(qdate(self.d_to))}"
        foot = [f"Итого: {L.money(res['total'])} ₽"] + res["notes"]
        self.total.setText(f"Итого: <b>{L.money(res['total'])} ₽</b>")
        self.result.set_html(L.table_html(title, head, body, foot), (title, head, body, foot))


class CalcPage(QWidget):
    def __init__(self):
        super().__init__()
        lay = QVBoxLayout(self)
        lay.setContentsMargins(24, 18, 24, 18)
        lay.addLayout(title_row("Калькуляторы", big=True))
        self.tabs = tabs = QTabWidget()
        tabs.setObjectName("segmented")
        tabs.setDocumentMode(True)
        for w, name in ((DeadlineTab(), "Сроки"), (DutyTab(), "Госпошлина"), (InterestTab(), "Проценты и неустойка")):
            sc = QScrollArea()
            sc.setWidget(w)
            sc.setWidgetResizable(True)
            sc.setFrameShape(QFrame.NoFrame)
            tabs.addTab(sc, name)
        lay.addWidget(tabs, 1)
        info = QLabel(f"Данные (ставки ЦБ, календари, пошлины) — на {D.DATA_DATE}. Обновляются вместе с программой.")
        info.setObjectName("hint")
        lay.addWidget(info)


# =============================================================================
#  Дела
# =============================================================================
DOC_ICONS = ["📄", "📝", "📑", "⚖️", "🏛️", "📬", "✉️", "📦", "💰", "🧾", "📊", "📷",
             "📎", "📌", "⭐", "✅", "⏳", "❗", "🔒", "🗂️", "🤝", "🔍", "🖊️", "📅"]


def default_icon(path):
    ext = os.path.splitext(path or "")[1].lower()
    return {".pdf": "📄", ".doc": "📝", ".docx": "📝", ".rtf": "📝", ".odt": "📝", ".xls": "📊", ".xlsx": "📊",
            ".jpg": "📷", ".jpeg": "📷", ".png": "📷", ".tif": "📷", ".tiff": "📷"}.get(ext, "📎")


def fmt_sent_short(s):
    """Для узкого столбца: «23.09 14:35» в текущем году, иначе «23.09.25 14:35»."""
    try:
        d = dt.datetime.fromisoformat(s)
    except Exception:
        return s or ""
    return d.strftime("%d.%m %H:%M" if d.year == dt.date.today().year else "%d.%m.%y %H:%M")


def fmt_sent(s):
    if not s:
        return ""
    try:
        d = dt.datetime.fromisoformat(s)
        return d.strftime("%d.%m.%Y  %H:%M")
    except Exception:
        return s


class IconPicker(QMenu):
    """Сетка значков для документа."""

    def __init__(self, parent, on_pick):
        super().__init__(parent)
        from PySide6.QtWidgets import QWidgetAction
        box = QWidget()
        g = QGridLayout(box)
        g.setContentsMargins(8, 8, 8, 8)
        g.setSpacing(4)
        for i, ic in enumerate(DOC_ICONS):
            b = QToolButton()
            b.setText(ic)
            b.setObjectName("iconpick")
            b.setFixedSize(38, 38)
            b.setAutoRaise(True)
            f = b.font()
            f.setPointSize(15)
            b.setFont(f)
            b.clicked.connect(lambda _=False, ic=ic: (on_pick(ic), self.close()))
            g.addWidget(b, i // 6, i % 6)
        wa = QWidgetAction(self)
        wa.setDefaultWidget(box)
        self.addAction(wa)
        self.addSeparator()
        self.addAction("Значок по типу файла", lambda: on_pick(""))


class SentDialog(QDialog):
    """Дата и время отправки документа."""

    def __init__(self, parent, title, value):
        super().__init__(parent)
        from PySide6.QtWidgets import QDateTimeEdit
        from PySide6.QtCore import QDateTime
        self.setWindowTitle("Когда отправлен")
        self.result_value = None
        v = QVBoxLayout(self)
        lab = QLabel(f"«{title}»")
        lab.setObjectName("subtitle")
        lab.setWordWrap(True)
        v.addWidget(lab)
        self.ed = QDateTimeEdit()
        self.ed.setCalendarPopup(True)
        self.ed.setDisplayFormat("dd.MM.yyyy  HH:mm")
        cur = None
        try:
            cur = dt.datetime.fromisoformat(value) if value else None
        except Exception:
            cur = None
        cur = cur or dt.datetime.now().replace(second=0, microsecond=0)
        self.ed.setDateTime(QDateTime(QDate(cur.year, cur.month, cur.day),
                                      __import__("PySide6.QtCore", fromlist=["QTime"]).QTime(cur.hour, cur.minute)))
        row = QHBoxLayout()
        row.addWidget(self.ed, 1)
        now = QPushButton("Сейчас")
        now.clicked.connect(lambda: self.ed.setDateTime(QDateTime.currentDateTime()))
        row.addWidget(now)
        v.addLayout(row)
        bb = QHBoxLayout()
        clr = QPushButton("Не отправлен")
        clr.clicked.connect(self.clear)
        bb.addWidget(clr)
        bb.addStretch(1)
        cancel = QPushButton("Отмена")
        cancel.clicked.connect(self.reject)
        ok = QPushButton("Сохранить")
        ok.setObjectName("primary")
        ok.clicked.connect(self.save)
        bb.addWidget(cancel)
        bb.addWidget(ok)
        v.addLayout(bb)

    def clear(self):
        self.result_value = ""
        self.accept()

    def save(self):
        q = self.ed.dateTime()
        d, t = q.date(), q.time()
        self.result_value = dt.datetime(d.year(), d.month(), d.day(), t.hour(), t.minute()).isoformat(timespec="minutes")
        self.accept()


class DocsTable(QTableWidget):
    """Документы дела: значок, название (редактируется), отметка «отправлен» галочкой, файл.
    Щелчок — выделить; двойной щелчок или перетаскивание вправо, в рабочую область — открыть.
    Перетаскивание в список комплекта — добавить в комплект."""
    COLS = ["", "Документ", "Отправлен", "Файл"]

    def mimeData(self, items):
        from PySide6.QtCore import QMimeData
        md = QMimeData()
        paths = [p for p in self.paths(only_selected=True) if p and os.path.exists(p)]
        md.setUrls([QUrl.fromLocalFile(p) for p in paths])
        return md

    def mimeTypes(self):
        return ["text/uri-list"]

    def __init__(self, page):
        super().__init__(0, 4)
        self.page = page
        self._loading = False
        self.setObjectName("docs")
        self.setHorizontalHeaderLabels(self.COLS)
        self.verticalHeader().hide()
        self.setShowGrid(False)
        self.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.setEditTriggers(QAbstractItemView.EditKeyPressed)      # переименование — F2 или кнопкой «✎ Название»
        self.setDragEnabled(True)
        self.setDragDropMode(QAbstractItemView.DragOnly)
        self.setDefaultDropAction(Qt.CopyAction)
        self.setWordWrap(False)
        h = self.horizontalHeader()
        h.setSectionResizeMode(0, QHeaderView.Fixed)
        self.setColumnWidth(0, 46)
        h.setSectionResizeMode(1, QHeaderView.Stretch)
        h.setSectionResizeMode(2, QHeaderView.Fixed)
        self.setColumnWidth(2, 170)
        h.setSectionResizeMode(3, QHeaderView.Interactive)
        self.setColumnWidth(3, 220)
        self.verticalHeader().setDefaultSectionSize(40)
        self.cellClicked.connect(self.on_click)
        self.cellDoubleClicked.connect(self.on_double)
        self.itemChanged.connect(self.on_changed)
        self.setContextMenuPolicy(Qt.CustomContextMenu)
        self.customContextMenuRequested.connect(self.menu)

    # ---------------------------------------------------------------- данные
    def load(self, docs):
        self._loading = True
        self.setRowCount(0)
        main = self.page.main
        try:
            in_pdf = main.case_pdf_sources(self.page.cid) if self.page.cid and hasattr(main, "case_pdf_sources") else set()
        except Exception:
            in_pdf = set()
        for d in docs:
            r = self.rowCount()
            self.insertRow(r)
            exists = os.path.exists(d["path"])
            ic = QTableWidgetItem(d.get("icon") or default_icon(d["path"]))
            f = ic.font()
            f.setPointSize(15)
            ic.setFont(f)
            ic.setTextAlignment(Qt.AlignCenter)
            ic.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable)
            ic.setToolTip("Щёлкните, чтобы выбрать значок")
            title = QTableWidgetItem(d["title"])
            title.setToolTip("Двойной щелчок — переименовать")
            if os.path.normcase(os.path.abspath(d["path"])) in in_pdf:
                fb = title.font()
                fb.setBold(True)
                title.setFont(fb)
                title.setForeground(QColor(M.T["accent"]))
                title.setToolTip("Уже в PDF дела (справа). Двойной щелчок — перейти к нему. F2 — переименовать")
            else:
                title.setToolTip("Ещё не в PDF дела. Двойной щелчок или перетаскивание вправо — добавить. "
                                 "F2 — переименовать")
            title.setData(Qt.UserRole, d["path"])
            title.setData(Qt.UserRole + 1, d["id"])
            sent = QTableWidgetItem(fmt_sent_short(d.get("sent")) if d.get("sent") else "нет")
            sent.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable | Qt.ItemIsUserCheckable)
            sent.setCheckState(Qt.Checked if d.get("sent") else Qt.Unchecked)
            sent.setToolTip((f"Отправлен {fmt_sent(d.get('sent'))}. Снимите галочку, если не отправлен. "
                             if d.get("sent") else "Поставьте галочку — запишутся текущие дата и время. ") +
                            "Точное время — кнопка «📅 Отправка».")
            sent.setForeground(QColor(M.T["success"] if d.get("sent") else M.T["muted"]))
            fn = QTableWidgetItem(("⚠ " if not exists else "") + os.path.basename(d["path"]))
            fn.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable)
            fn.setToolTip(d["path"] + ("" if exists else "\nФайл не найден — возможно, перемещён или удалён"))
            fn.setForeground(QColor(M.T["danger"] if not exists else M.T["muted"]))
            for c, it in enumerate((ic, title, sent, fn)):
                self.setItem(r, c, it)
        self._loading = False

    def row_doc(self, r):
        it = self.item(r, 1)
        return (it.data(Qt.UserRole + 1), it.data(Qt.UserRole), it.text()) if it else (None, None, "")

    def selected_rows(self):
        return sorted({i.row() for i in self.selectedIndexes()})

    def keyPressEvent(self, e):
        if e.key() == Qt.Key_Delete and self.state() != QAbstractItemView.EditingState:
            self.page.del_doc()
            return
        if e.modifiers() & Qt.ControlModifier and e.key() in (Qt.Key_Up, Qt.Key_Down):
            self.move(-1 if e.key() == Qt.Key_Up else 1)
            return
        super().keyPressEvent(e)

    def paths(self, only_selected=False):
        rows = self.selected_rows() if only_selected else range(self.rowCount())
        return [self.row_doc(r)[1] for r in rows]

    # ---------------------------------------------------------------- правка
    def on_changed(self, item):
        if self._loading:
            return
        if item.column() == 2:                      # галочка «отправлен»
            did = self.row_doc(item.row())[0]
            on = item.checkState() == Qt.Checked
            db().update_doc(did, sent=dt.datetime.now().isoformat(timespec="minutes") if on else "")
            QTimer.singleShot(0, self.page.load_docs)
            return
        if item.column() != 1:
            return
        did = item.data(Qt.UserRole + 1)
        text = item.text().strip()
        if not text:
            self.page.load_docs()
            return
        db().update_doc(did, title=text)
        self.page.main.navigator.refresh() if hasattr(self.page.main, "navigator") else None

    def on_click(self, r, c):
        did, path, title = self.row_doc(r)
        if c == 0:
            rect = self.visualItemRect(self.item(r, 0))
            IconPicker(self, lambda ic: self.set_icon(did, ic)).exec(self.viewport().mapToGlobal(rect.bottomLeft()))

    def on_double(self, r, c):
        if c in (1, 3):
            path = self.row_doc(r)[1]
            if path:
                self.page.openFile.emit(path, 0)

    def set_icon(self, did, ic):
        db().update_doc(did, icon=ic)
        self.page.load_docs()

    def edit_sent(self, r):
        did, path, title = self.row_doc(r)
        cur = next((d.get("sent") for d in db().docs(self.page.cid) if d["id"] == did), "")
        dlg = SentDialog(self, title, cur)
        if dlg.exec() and dlg.result_value is not None:
            db().update_doc(did, sent=dlg.result_value)
            self.page.load_docs()

    def _current(self):
        rows = self.selected_rows()
        return rows[0] if len(rows) == 1 else None

    def rename_current(self):
        r = self._current()
        if r is not None:
            self.editItem(self.item(r, 1))

    def sent_current(self):
        r = self._current()
        if r is not None:
            self.edit_sent(r)

    def icon_current(self):
        r = self._current()
        if r is not None:
            self.on_click(r, 0)

    def move(self, step):
        rows = self.selected_rows()
        if len(rows) != 1:
            return
        r = rows[0]
        n = r + step
        if not 0 <= n < self.rowCount():
            return
        ids = [self.row_doc(i)[0] for i in range(self.rowCount())]
        ids[r], ids[n] = ids[n], ids[r]
        for pos, did in enumerate(ids):
            db().update_doc(did, pos=pos)
        self.page.load_docs()
        self.selectRow(n)

    def menu(self, pos):
        r = self.rowAt(pos.y())
        if r < 0:
            return
        did, path, title = self.row_doc(r)
        m = QMenu(self)
        if (path or "").lower().endswith((".docx", ".doc", ".rtf", ".odt")):
            m.addAction("✏️ Править здесь (документ Word)", lambda: self.page.main.edit_word(path))
        m.addAction("Показать в PDF дела (справа)", lambda: self.page.openFile.emit(path, 0))
        m.addAction("Открыть в своей программе (Word, Acrobat…)", lambda: self.page.main.open_in_app(path))
        m.addAction("Переименовать", lambda: self.editItem(self.item(r, 1)))
        m.addAction("Выбрать значок…", lambda: self.on_click(r, 0))
        m.addAction("Отметить отправку…", lambda: self.edit_sent(r))
        m.addAction("Отправлен сейчас", lambda: (db().update_doc(did, sent=dt.datetime.now().isoformat(timespec="minutes")),
                                                self.page.load_docs()))
        m.addSeparator()
        m.addAction("▲  Выше (Ctrl+↑)", lambda: (self.selectRow(r), self.move(-1)))
        m.addAction("▼  Ниже (Ctrl+↓)", lambda: (self.selectRow(r), self.move(1)))
        m.addAction("Показать в папке", lambda: self.page.main.show_in_folder(path))
        m.addSeparator()
        if r not in self.selected_rows():
            self.selectRow(r)
        m.addAction("🗑  Убрать из дела (Delete)", self.page.del_doc)
        m.addAction("Удалить файл с диска…", self.page.del_doc_files)
        m.exec(self.viewport().mapToGlobal(pos))


class EmptyState(QWidget):
    """Пустой экран, который подсказывает, что делать: крупный значок, пара слов и одна кнопка."""

    def __init__(self, icon, title, text="", button=None, on_click=None, parent=None):
        super().__init__(parent)
        v = QVBoxLayout(self)
        v.setContentsMargins(24, 24, 24, 24)
        v.addStretch(1)
        ic = QLabel(icon)
        ic.setAlignment(Qt.AlignCenter)
        f = ic.font()
        f.setPointSize(34)
        ic.setFont(f)
        v.addWidget(ic)
        t = QLabel(title)
        t.setObjectName("subtitle")
        t.setAlignment(Qt.AlignCenter)
        t.setWordWrap(True)
        tf = t.font()
        tf.setPointSize(tf.pointSize() + 3)
        tf.setBold(True)
        t.setFont(tf)
        v.addWidget(t)
        self.text = QLabel(text)
        self.text.setObjectName("hint")
        self.text.setAlignment(Qt.AlignCenter)
        self.text.setWordWrap(True)
        v.addWidget(self.text)
        self.button = None
        if button:
            v.addSpacing(6)
            b = QPushButton(button)
            b.setObjectName("primary")
            b.setCursor(Qt.PointingHandCursor)
            b.clicked.connect(on_click)
            h = QHBoxLayout()
            h.addStretch(1)
            h.addWidget(b)
            h.addStretch(1)
            v.addLayout(h)
            self.button = b
        v.addStretch(2)


class ActionCard(QPushButton):
    """Крупная кнопка-действие: значок, название, пояснение."""

    def __init__(self, icon, title, desc, slot, help_key=None):
        super().__init__()
        self.setObjectName("actioncard")
        self.setCursor(Qt.PointingHandCursor)
        self.setMinimumHeight(88)                # название и две строки пояснения помещаются целиком
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        h = QHBoxLayout(self)
        h.setContentsMargins(14, 10, 10, 10)
        h.setSpacing(12)
        ic = QLabel(icon)
        ic.setObjectName("cardicon")
        f = ic.font()
        f.setPointSize(17)
        ic.setFont(f)
        ic.setFixedSize(42, 42)                  # значок в скруглённом квадрате, как в «Настройках» iOS
        ic.setAlignment(Qt.AlignCenter)
        ic.setAttribute(Qt.WA_TransparentForMouseEvents)
        h.addWidget(ic)
        tv = QVBoxLayout()
        tv.setSpacing(1)
        t = QLabel(title)
        t.setObjectName("cardtitle")
        t.setWordWrap(True)
        t.setAttribute(Qt.WA_TransparentForMouseEvents)
        d = QLabel(desc)
        d.setObjectName("carddesc")
        d.setWordWrap(True)
        d.setAttribute(Qt.WA_TransparentForMouseEvents)
        tv.addWidget(t)
        tv.addWidget(d)
        h.addLayout(tv, 1)
        if help_key:
            h.addWidget(HelpButton(help_key), 0, Qt.AlignTop)
        self.clicked.connect(slot)
        anim.hover_lift(self)


def card_grid(cards, cols=3):
    g = QGridLayout()
    g.setSpacing(10)
    for i, c in enumerate(cards):
        g.addWidget(c, i // cols, i % cols)
    return g


def now():
    """Текущее время с поправкой на отставание/спешку часов компьютера (см. timecheck.py)."""
    return dt.datetime.now() + dt.timedelta(seconds=getattr(M, "CLOCK_OFFSET", 0.0) or 0.0)


def event_dt(e):
    """Дата и время события; без времени — 9:00 утра."""
    t = (e.get("time") or "").strip() or "09:00"
    try:
        hh, mm = (int(x) for x in t.split(":")[:2])
    except ValueError:
        hh, mm = 9, 0
    d = dt.date.fromisoformat(e["date"])
    return dt.datetime(d.year, d.month, d.day, hh, mm)


class ReminderDialog(QDialog):
    """Новое напоминание или правка: когда и о чём."""

    def __init__(self, parent, text="", when=None):
        super().__init__(parent)
        from PySide6.QtWidgets import QDateTimeEdit
        from PySide6.QtCore import QDateTime, QTime
        self.setWindowTitle("Напоминание")
        self.setMinimumWidth(440)
        v = QVBoxLayout(self)
        v.addLayout(title_row("Напоминание"))
        v.addWidget(QLabel("О чём напомнить"))
        self.text = QPlainTextEdit(text)
        self.text.setPlaceholderText("Например: позвонить доверителю, запросить выписку, подготовить отзыв")
        self.text.setFixedHeight(80)
        v.addWidget(self.text)
        v.addWidget(QLabel("Когда"))
        when = when or (now() + dt.timedelta(days=1)).replace(hour=9, minute=0, second=0, microsecond=0)
        self.ed = QDateTimeEdit()
        self.ed.setCalendarPopup(True)
        self.ed.setDisplayFormat("dd.MM.yyyy  HH:mm")
        self.ed.setDateTime(QDateTime(QDate(when.year, when.month, when.day), QTime(when.hour, when.minute)))
        row = QHBoxLayout()
        row.addWidget(self.ed, 1)
        for label, delta in (("Через час", dt.timedelta(hours=1)), ("Завтра 9:00", None)):
            b = QPushButton(label)
            b.clicked.connect(lambda _=False, d=delta: self._set(d))
            row.addWidget(b)
        v.addLayout(row)
        bb = QHBoxLayout()
        bb.addStretch(1)
        cancel = QPushButton("Отмена")
        cancel.clicked.connect(self.reject)
        ok = QPushButton("Сохранить")
        ok.setObjectName("primary")
        ok.clicked.connect(self._ok)
        bb.addWidget(cancel)
        bb.addWidget(ok)
        v.addLayout(bb)
        self.value = None

    def _set(self, delta):
        from PySide6.QtCore import QDateTime, QTime
        w = (now() + delta).replace(second=0, microsecond=0) if delta else \
            (now() + dt.timedelta(days=1)).replace(hour=9, minute=0, second=0, microsecond=0)
        self.ed.setDateTime(QDateTime(QDate(w.year, w.month, w.day), QTime(w.hour, w.minute)))

    def _ok(self):
        text = self.text.toPlainText().strip()
        if not text:
            QMessageBox.information(self, M.APP_NAME, "Напишите, о чём напомнить.")
            return
        q = self.ed.dateTime()
        d, t = q.date(), q.time()
        self.value = (dt.datetime(d.year(), d.month(), d.day(), t.hour(), t.minute()), text)
        self.accept()


def court_search_targets(case):
    """Куда можно пойти искать дело по номеру: [(подпись, адрес)]."""
    num = (case.get("number") or "").strip()
    court = (case.get("court") or "").lower()
    arbitr = "арбитраж" in court or num[:1] in ("А", "A") and "-" in num and "/" in num
    out = []
    if arbitr:
        out.append(("Картотека арбитражных дел (kad.arbitr.ru)", "https://kad.arbitr.ru/"))
    else:
        if "москв" in court:
            out.append(("Портал судов общей юрисдикции Москвы (mos-sud.ru)", "https://mos-sud.ru/search"))
        out.append(("ГАС «Правосудие» — поиск по всем судам", "https://bsr.sudrf.ru/bigs/portal.html"))
        out.append(("Картотека арбитражных дел (kad.arbitr.ru)", "https://kad.arbitr.ru/"))
    return out


def court_menu(main, cid, anchor):
    """Меню кнопки «Дело на сайте суда»: открыть сохранённую ссылку, найти по номеру, вставить ссылку."""
    c = db().case(cid) or {}
    url = (c.get("court_url") or "").strip()
    m = QMenu(anchor)
    if url:
        m.addAction("🌐 Открыть карточку дела", lambda: QDesktopServices.openUrl(QUrl(url)))
        m.addSeparator()
    num = (c.get("number") or "").strip()

    def search(link):
        if num:
            QApplication.clipboard().setText(num)
            main.statusBar().showMessage(f"Номер дела {num} скопирован — вставьте его в поиск на сайте (Ctrl+V)", 12000)
        QDesktopServices.openUrl(QUrl(link))
    for label, link in court_search_targets(c):
        m.addAction(("🔎 Найти по номеру: " if num else "🔎 ") + label, lambda l=link: search(l))

    def paste():
        clip = QApplication.clipboard().text().strip()
        from PySide6.QtWidgets import QInputDialog
        text, ok = QInputDialog.getText(main, "Ссылка на дело", "Откройте карточку дела на сайте суда, скопируйте адрес "
                                        "из строки браузера и вставьте сюда:",
                                        text=clip if clip.startswith("http") else url)
        if ok:
            db().update_case(cid, court_url=text.strip())
            main.cases_page.load_case() if main.cases_page.cid == cid else None
            main.overview.set_case(cid)
    m.addSeparator()
    m.addAction("🔗 Вставить ссылку на карточку дела…" if not url else "🔗 Изменить ссылку…", paste)
    if url:
        m.addAction("Убрать ссылку", lambda: (db().update_case(cid, court_url=""), main.overview.set_case(cid)))
    m.exec(anchor.mapToGlobal(anchor.rect().bottomLeft()))


INSTANCE_KEYS = ("court", "number", "judge", "court_url")


class InstanceCard(QFrame):
    """Одна инстанция: уровень суда (список) и полное название рядом, номер дела, судья, ссылка."""

    def __init__(self, editor, inst, index):
        super().__init__()
        self.editor, self.iid = editor, inst["id"]
        self.setObjectName("card")
        v = QVBoxLayout(self)
        v.setContentsMargins(14, 10, 14, 10)
        h = QHBoxLayout()
        n = QLabel(f"{index + 1}.")
        n.setObjectName("subtitle")
        h.addWidget(n)
        self.level = QComboBox()
        self.level.addItems(CS.LEVEL_NAMES)
        self.level.setCurrentText(inst["level"] or CS.guess_level(inst["court"], inst["number"]))
        self.level.setMinimumWidth(300)
        h.addWidget(self.level, 1)
        self.cur = QLabel("текущая" if editor.is_last(inst) else "")
        self.cur.setObjectName("hint")
        h.addWidget(self.cur)
        rm = QPushButton("Удалить")
        rm.setObjectName("compact")
        rm.setProperty("danger", True)
        rm.setToolTip("Удалить эту инстанцию")
        rm.clicked.connect(lambda: editor.remove(self.iid))
        h.addWidget(rm)
        v.addLayout(h)
        f = QFormLayout()
        f.setLabelAlignment(Qt.AlignRight | Qt.AlignVCenter)
        self.court = M.GrowEdit(inst["court"], "Полное название: «Девятый арбитражный апелляционный суд»")
        self.number = M.GrowEdit(inst["number"], "Номер дела в этом суде")
        self.judge = M.GrowEdit(inst["judge"], "ФИО судьи или состав")
        self.url = M.GrowEdit(inst["url"], "Ссылка на карточку дела на сайте суда")
        urow = QHBoxLayout()
        urow.addWidget(self.url, 1)
        ob = QPushButton("Открыть")
        ob.setObjectName("compact")
        ob.clicked.connect(lambda: self.url.text().strip() and QDesktopServices.openUrl(QUrl(self.url.text().strip())))
        urow.addWidget(ob)
        f.addRow("Суд", self.court)
        f.addRow("Номер дела", self.number)
        f.addRow("Судья", self.judge)
        f.addRow("Сайт суда", urow)
        v.addLayout(f)
        self.timer = QTimer(self)
        self.timer.setSingleShot(True)
        self.timer.setInterval(600)
        self.timer.timeout.connect(self.save)
        for e in (self.court, self.number, self.judge, self.url):
            e.textChanged.connect(self.timer.start)
        self.level.currentTextChanged.connect(lambda *_: self.save())

    def save(self):
        self.timer.stop()
        db().update_instance(self.iid, level=self.level.currentText(), court=self.court.text().strip(),
                             number=self.number.text().strip(), judge=self.judge.text().strip(),
                             url=self.url.text().strip())
        self.editor.changed()


class InstancesEditor(QWidget):
    """«Суды и инстанции» в сведениях о деле: путь дела от первой инстанции дальше.
    Последняя инстанция — текущая: её суд, номер и судья используются во всей программе."""

    def __init__(self, page):
        super().__init__()
        self.page = page
        self.cid = None
        self.cards = []
        v = QVBoxLayout(self)
        v.setContentsMargins(0, 4, 0, 8)
        h = QHBoxLayout()
        t = QLabel("Суды и инстанции")
        t.setObjectName("subtitle")
        h.addWidget(t)
        h.addStretch(1)
        self.b_next = QPushButton("→ Следующая инстанция")
        self.b_next.setObjectName("primary")
        self.b_next.setToolTip("Дело перешло дальше: апелляция, кассация, Верховный Суд…")
        self.b_next.clicked.connect(self.add_next)
        h.addWidget(self.b_next)
        v.addLayout(h)
        hint = QLabel("Выберите уровень суда и впишите полное название. Когда дело переходит в апелляцию, "
                      "кассацию и дальше — «→ Следующая инстанция»: программа предложит уровень, стадия дела "
                      "обновится сама. Последняя инстанция — текущая: её суд и номер подставляются в шаблоны.")
        hint.setObjectName("hint")
        hint.setWordWrap(True)
        v.addWidget(hint)
        self.box = QVBoxLayout()
        self.box.setSpacing(8)
        v.addLayout(self.box)

    def is_last(self, inst):
        return self._last_id == inst["id"]

    def set_case(self, cid):
        self.cid = cid
        if not cid:
            return
        db().ensure_instances(cid)
        if not db().instances(cid):
            db().add_instance(cid, "Районный (городской) суд")
        self.reload()

    def reload(self):
        for c in self.cards:
            c.timer.isActive() and c.save()
            self.box.removeWidget(c)
            c.deleteLater()
        self.cards = []
        inst = db().instances(self.cid)
        self._last_id = inst[-1]["id"] if inst else None
        for i, it in enumerate(inst):
            card = InstanceCard(self, it, i)
            self.box.addWidget(card)
            self.cards.append(card)
        M.polish_ui(self)

    def add_next(self):
        inst = db().instances(self.cid)
        for c in self.cards:
            c.timer.isActive() and c.save()
        cur = inst[-1]["level"] if inst else ""
        nxt = CS.next_level(cur) or cur or "Районный (городской) суд"
        db().add_instance(self.cid, nxt)
        self.reload()
        self.changed()
        if self.cards:
            self.cards[-1].court.setFocus()

    def remove(self, iid):
        inst = db().instances(self.cid)
        if len(inst) <= 1:
            QMessageBox.information(self, M.APP_NAME, "Это единственная инстанция — её можно только изменить.")
            return
        it = next((i for i in inst if i["id"] == iid), None)
        if QMessageBox.question(self, M.APP_NAME, f"Удалить инстанцию «{it['court'] or it['level']}»?") != QMessageBox.Yes:
            return
        db().delete_instance(iid)
        self.reload()
        self.changed()

    def changed(self):
        """Текущая инстанция поменялась — обновить стадию в форме, заголовок и обзор."""
        c = db().case(self.cid) or {}
        st = self.page.fields.get("stage")
        if st is not None and st.currentText() != c.get("stage", ""):
            st.blockSignals(True)
            st.setCurrentText(c.get("stage", ""))
            st.blockSignals(False)
        for card in self.cards:
            card.cur.setText("текущая" if card.iid == self._last_id else "")
        main = self.page.main
        if getattr(main, "overview", None) is not None and main.overview.cid == self.cid:
            main.overview.set_case(self.cid)


class PhoneSetupDialog(QDialog):
    """Пошаговая настройка «Дела на телефоне» через Яндекс Диск — с объяснением каждого действия."""

    def __init__(self, main):
        super().__init__(main)
        import phone_export as PE
        self.PE = PE
        self.main = main
        st = M.settings()
        self.setWindowTitle("📱 Дела на телефоне — Яндекс Диск")
        self.resize(640, 560)
        v = QVBoxLayout(self)
        v.setContentsMargins(22, 18, 22, 16)
        self.step_lbl = QLabel()
        self.step_lbl.setObjectName("hint")
        v.addWidget(self.step_lbl)
        self.pages = QStackedWidget()
        v.addWidget(self.pages, 1)

        def page(title, html_text):
            w = QWidget()
            lv = QVBoxLayout(w)
            lv.setContentsMargins(0, 4, 0, 0)
            t = QLabel(title)
            t.setObjectName("title")
            t.setWordWrap(True)
            lv.addWidget(t)
            b = QLabel(html_text)
            b.setWordWrap(True)
            b.setTextFormat(Qt.RichText)
            b.setOpenExternalLinks(True)
            b.setStyleSheet("font-size: 11.5pt;")
            lv.addWidget(b)
            self.pages.addWidget(w)
            return lv

        # 1 — что это
        page("Дела на телефоне",
             "Перед заседанием можно открыть дело на телефоне и спокойно всё прочитать: <b>памятку</b> (когда и где "
             "заседание, суд, судья, стороны, <b>главное к заседанию</b>), <b>сроки</b>, <b>выписки</b> и все "
             "<b>документы дела</b> — с оглавлением, по которому документы открываются одним нажатием.<br><br>"
             "Файлы попадают на телефон через ваш <b>Яндекс Диск</b>. Программа сама обновляет их после каждого "
             "сохранения дела.<br><br>"
             "<b>Логин и пароль Яндекса в LegalHelper вводить не нужно.</b> Вы один раз входите в официальную "
             "программу «Яндекс Диск» — так безопаснее: пароль знает только Яндекс, а документы клиентов лежат "
             "только в вашем личном облаке.<br><br>Нажмите «Далее» — покажу, что сделать, по шагам.")
        # 2 — компьютер
        lv = page("Шаг 1. Яндекс Диск на этом компьютере",
                  "<ol style='margin-left:-20px'>"
                  "<li style='margin-bottom:6px'>Нажмите кнопку <b>«Скачать Яндекс Диск»</b> ниже — откроется сайт Яндекса. "
                  "Нажмите там «Скачать» для Windows.</li>"
                  "<li style='margin-bottom:6px'>Откройте скачанный файл (обычно он в папке «Загрузки») и установите программу, "
                  "нажимая «Далее».</li>"
                  "<li style='margin-bottom:6px'>В окне Яндекс Диска <b>войдите в свой аккаунт Яндекса</b>: логин — это ваша почта "
                  "вида <i>имя@yandex.ru</i> (или номер телефона), пароль — тот же, что от Яндекс Почты. "
                  "Нет аккаунта — нажмите там «Создать ID».</li>"
                  "<li style='margin-bottom:6px'>После входа на компьютере появится папка «Яндекс Диск». Вернитесь сюда и нажмите "
                  "<b>«Проверить»</b>.</li></ol>")
        row = QHBoxLayout()
        b_dl = QPushButton("⬇  Скачать Яндекс Диск")
        b_dl.setObjectName("primary")
        b_dl.clicked.connect(lambda: QDesktopServices.openUrl(QUrl(PE.YANDEX_DOWNLOAD)))
        row.addWidget(b_dl)
        b_chk = QPushButton("🔄  Проверить")
        b_chk.clicked.connect(self.check)
        row.addWidget(b_chk)
        b_pick = QPushButton("Указать папку вручную…")
        b_pick.setToolTip("Если вы при установке выбрали для Яндекс Диска другую папку")
        b_pick.clicked.connect(self.pick)
        row.addWidget(b_pick)
        row.addStretch(1)
        lv.addLayout(row)
        self.status = QLabel()
        self.status.setWordWrap(True)
        self.status.setStyleSheet("font-size: 11.5pt; padding: 8px 0;")
        lv.addWidget(self.status)
        lv.addStretch(1)
        # 3 — телефон
        page("Шаг 2. Яндекс Диск на телефоне",
             "<ol style='margin-left:-20px'>"
             "<li style='margin-bottom:6px'>Установите на телефон приложение <b>«Яндекс Диск»</b>: на iPhone — из App Store, "
             "на Android — из Google Play или RuStore (в поиске магазина наберите «Яндекс Диск»).</li>"
             "<li style='margin-bottom:6px'>Откройте его и <b>войдите тем же логином и паролем</b>, что и на компьютере.</li>"
             "<li style='margin-bottom:6px'>Откройте папку <b>«LegalHelper — дела»</b> — в ней по файлу на каждое дело.</li>"
             "<li style='margin-bottom:6px'>Чтобы файлы открывались <b>без интернета</b> (в суде он бывает плохой): нажмите "
             "на три точки <b>⋮</b> рядом с папкой и выберите пункт про офлайн («Офлайн» или «Сделать доступным "
             "офлайн» — зависит от версии приложения).</li>"
             "<li style='margin-bottom:6px'>Нажмите на дело — откроется памятка. Листайте вниз; в «Документы дела» нажмите "
             "на документ — он откроется. Увеличить — двумя пальцами, как фото.</li></ol>")
        # 4 — какие дела
        lv = page("Шаг 3. Какие дела отправлять",
                  "Выберите, какие дела будут на телефоне. Изменить можно в любой момент: «☰ → Файл → 📱 Дела на "
                  "телефоне».")
        from PySide6.QtWidgets import QRadioButton
        self.r_all = QRadioButton("Все дела в работе (кроме архива) — обновляются сами")
        self.r_manual = QRadioButton("Только те, где я нажму «📱 На телефон» в «Обзоре» дела")
        (self.r_manual if st.value("phone/mode", "all") == "manual" else self.r_all).setChecked(True)
        for r in (self.r_all, self.r_manual):
            r.setStyleSheet("font-size: 11.5pt; padding: 4px 0;")
            lv.addWidget(r)
        self.c_on = QCheckBox("Отправлять дела на телефон")
        self.c_on.setChecked(True)
        self.c_on.setStyleSheet("font-size: 11.5pt; padding: 8px 0;")
        lv.addWidget(self.c_on)
        tipl = QLabel("Нажмите «Готово» — дела сразу отправятся. Через минуту-две они появятся на телефоне в "
                      "папке «LegalHelper — дела».")
        tipl.setWordWrap(True)
        tipl.setObjectName("hint")
        lv.addWidget(tipl)
        lv.addStretch(1)
        # кнопки
        nav = QHBoxLayout()
        self.b_back = QPushButton("← Назад")
        self.b_back.clicked.connect(lambda: self.go(-1))
        nav.addWidget(self.b_back)
        nav.addStretch(1)
        self.b_next = QPushButton("Далее →")
        self.b_next.setObjectName("primary")
        self.b_next.clicked.connect(lambda: self.go(1))
        nav.addWidget(self.b_next)
        v.addLayout(nav)
        self.base = st.value("phone/base", "") or ""
        if not (self.base and os.path.isdir(self.base)):
            self.base = PE.find_yandex_folder() or ""
        self.check(quiet=True)
        self.go(0)

    def check(self, quiet=False):
        if not (self.base and os.path.isdir(self.base)):
            self.base = self.PE.find_yandex_folder() or ""
        if self.base:
            self.status.setText(f"✅  <b>Яндекс Диск найден:</b> {html.escape(self.base)}<br>Можно нажимать «Далее».")
            self.status.setTextFormat(Qt.RichText)
        else:
            self.status.setTextFormat(Qt.RichText)
            self.status.setText("⏳  Папка Яндекс Диска пока не найдена. Установите программу и войдите в неё "
                                "(пункты 1–3), затем нажмите «Проверить». Если при установке вы выбрали другую "
                                "папку — «Указать папку вручную…».")
        self.update_nav()

    def pick(self):
        p = QFileDialog.getExistingDirectory(self, "Папка Яндекс Диска на компьютере", str(Path.home()))
        if p:
            self.base = p
            self.check()

    def go(self, step):
        i = max(0, min(self.pages.count() - 1, self.pages.currentIndex() + step))
        if step > 0 and self.pages.currentIndex() == 1 and not self.base:
            self.check()
            return
        if step > 0 and self.pages.currentIndex() == self.pages.count() - 1:
            return self.finish()
        self.pages.setCurrentIndex(i)
        self.update_nav()

    def update_nav(self):
        i = self.pages.currentIndex()
        n = self.pages.count()
        self.step_lbl.setText(f"{i + 1} из {n}")
        self.b_back.setVisible(i > 0)
        self.b_next.setText("Готово — отправить дела" if i == n - 1 else "Далее →")
        self.b_next.setEnabled(not (i == 1 and not self.base))

    def finish(self):
        st = M.settings()
        st.setValue("phone/base", self.base)
        st.setValue("phone/on", "1" if self.c_on.isChecked() else "0")
        st.setValue("phone/mode", "manual" if self.r_manual.isChecked() else "all")
        self.accept()


class OverviewTab(QWidget):
    """Обзор дела: главное о деле, ближайшие сроки, последние документы, быстрые действия."""

    def __init__(self, main):
        super().__init__()
        self.main = main
        self.cid = None
        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        sc = QScrollArea()
        sc.setWidgetResizable(True)
        sc.setFrameShape(QFrame.NoFrame)
        outer.addWidget(sc)
        body = QWidget()
        sc.setWidget(body)
        v = QVBoxLayout(body)
        v.setContentsMargins(2, 14, 8, 14)
        v.setSpacing(12)
        # сводка
        self.facts = QLabel()
        self.facts.setObjectName("facts")
        self.facts.setWordWrap(True)
        self.facts.setTextFormat(Qt.RichText)
        self.facts.setOpenExternalLinks(False)
        self.facts.linkActivated.connect(self._fact_link)
        v.addWidget(self.facts)
        crow = QHBoxLayout()
        self.b_court = QPushButton("🌐  Дело на сайте суда")
        self.b_court.setObjectName("primary")
        self.b_court.setCursor(Qt.PointingHandCursor)
        self.b_court.clicked.connect(self.court_clicked)
        crow.addWidget(self.b_court)
        self.b_court_more = QPushButton("▾")
        self.b_court_more.setToolTip("Найти по номеру, вставить или изменить ссылку")
        self.b_court_more.clicked.connect(lambda: court_menu(self.main, self.cid, self.b_court_more))
        crow.addWidget(self.b_court_more)
        self.court_hint = QLabel()
        self.court_hint.setObjectName("hint")
        self.court_hint.setWordWrap(True)             # не распирать окно в ширину на небольших экранах
        self.court_hint.setMinimumWidth(10)
        crow.addWidget(self.court_hint, 1)
        self.b_phone = QPushButton("📱  На телефон")
        self.b_phone.setObjectName("primary")
        self.b_phone.setCursor(Qt.PointingHandCursor)
        self.b_phone.setToolTip("Отправить дело на телефон через Яндекс Диск: памятка к заседанию и все документы")
        self.b_phone.clicked.connect(lambda: self.cid and self.main.phone_button(self.cid))
        crow.addWidget(self.b_phone)
        v.addLayout(crow)
        # главное к заседанию (коротко — это видно на телефоне первым)
        box = QFrame()
        box.setObjectName("card")
        bv = QVBoxLayout(box)
        bv.setContentsMargins(14, 12, 14, 10)
        bh = QHBoxLayout()
        lab = QLabel("⚖️  Главное к заседанию")
        lab.setObjectName("subtitle")
        bh.addWidget(lab)
        hint = QLabel("каждая мысль — с новой строки; «!» в начале — выделить красным. Видно на телефоне первым.")
        hint.setObjectName("hint")
        hint.setWordWrap(True)
        hint.setMinimumWidth(10)
        bh.addWidget(hint, 1)
        self.brief_state = QLabel()
        self.brief_state.setObjectName("hint")
        bh.addWidget(self.brief_state)
        bv.addLayout(bh)
        self.brief = QPlainTextEdit()
        self.brief.setObjectName("notes")
        self.brief.setPlaceholderText("Например:\nДолг подтверждён актом сверки (л. 9)\nПретензия 01.04 — ответа нет\n"
                                      "!Спросить: приобщить переписку о переносе сроков")
        self.brief.setFixedHeight(104)
        self.brief_timer = QTimer(self)
        self.brief_timer.setSingleShot(True)
        self.brief_timer.setInterval(700)
        self.brief_timer.timeout.connect(self.save_brief)
        self.brief.textChanged.connect(lambda: (self.brief_state.setText("…"), self.brief_timer.start()))
        bv.addWidget(self.brief)
        v.addWidget(box)
        # две колонки
        cols = QHBoxLayout()
        cols.setSpacing(14)
        for attr, title, btn_text, slot in (
                ("l_events", "Ближайшие сроки и заседания", "+ Срок или заседание", lambda: self.main.open_case_tab("events")),
                ("l_docs", "Последние документы", "Все документы", lambda: self.main.open_case_tab("docs"))):
            box = QFrame()
            box.setObjectName("card")
            bv = QVBoxLayout(box)
            bv.setContentsMargins(14, 12, 14, 10)
            lab = QLabel(title)
            lab.setObjectName("subtitle")
            bv.addWidget(lab)
            lst = QListWidget()
            lst.setObjectName("overlist")
            lst.setMinimumHeight(170)
            lst.setWordWrap(True)
            bv.addWidget(lst, 1)
            b = QPushButton(btn_text)
            b.clicked.connect(slot)
            bv.addWidget(b, 0, Qt.AlignLeft)
            setattr(self, attr, lst)
            cols.addWidget(box, 1)
        self.l_docs.itemClicked.connect(lambda it: it.data(Qt.UserRole) and self.main.open_external(
            it.data(Qt.UserRole), 0, self.cid))
        self.l_events.itemClicked.connect(lambda _it: self.main.open_case_tab("events"))
        v.addLayout(cols)
        # напоминания и заметки
        cols2 = QHBoxLayout()
        cols2.setSpacing(14)
        box = QFrame()
        box.setObjectName("card")
        bv = QVBoxLayout(box)
        bv.setContentsMargins(14, 12, 14, 10)
        lab = QLabel("⏰  Напоминания")
        lab.setObjectName("subtitle")
        bv.addWidget(lab)
        self.l_rem = QListWidget()
        self.l_rem.setObjectName("overlist")
        self.l_rem.setMinimumHeight(150)
        self.l_rem.setWordWrap(True)
        self.l_rem.itemDoubleClicked.connect(lambda it: self.edit_reminder(it.data(Qt.UserRole)))
        self.l_rem.setContextMenuPolicy(Qt.CustomContextMenu)
        self.l_rem.customContextMenuRequested.connect(self.reminder_menu)
        bv.addWidget(self.l_rem, 1)
        rr = QHBoxLayout()
        b = QPushButton("+ Напоминание")
        b.clicked.connect(lambda: self.edit_reminder(None))
        rr.addWidget(b)
        b = QPushButton("✓ Выполнено")
        b.clicked.connect(lambda: self._rem_action("done"))
        rr.addWidget(b)
        rr.addStretch(1)
        bv.addLayout(rr)
        cols2.addWidget(box, 1)
        box = QFrame()
        box.setObjectName("card")
        bv = QVBoxLayout(box)
        bv.setContentsMargins(14, 12, 14, 10)
        nh = QHBoxLayout()
        lab = QLabel("📝  Заметки и мысли")
        lab.setObjectName("subtitle")
        nh.addWidget(lab)
        nh.addStretch(1)
        self.notes_state = QLabel()
        self.notes_state.setObjectName("hint")
        nh.addWidget(self.notes_state)
        bv.addLayout(nh)
        self.notes = QPlainTextEdit()
        self.notes.setObjectName("notes")
        self.notes.setPlaceholderText("Мысли по делу, позиция, что спросить у доверителя, идеи для выступления… "
                                      "Сохраняется само.")
        self.notes.setMinimumHeight(150)
        self.notes_timer = QTimer(self)
        self.notes_timer.setSingleShot(True)
        self.notes_timer.setInterval(700)
        self.notes_timer.timeout.connect(self.save_notes)
        self.notes.textChanged.connect(self._notes_changed)
        bv.addWidget(self.notes, 1)
        cols2.addWidget(box, 1)
        v.addLayout(cols2)
        # быстрые действия
        lab = QLabel("Что сделать")
        lab.setObjectName("subtitle")
        v.addWidget(lab)
        m = self.main
        v.addLayout(card_grid([
            ActionCard("📨", "Подать в суд или отправить", "Собрать иск и приложения под «Мой арбитр», ГАС или Почту",
                       lambda: m.open_case_tab("prepare")),
            ActionCard("📝", "Документ по шаблону", "Ходатайство, заявление, запрос — с данными дела",
                       lambda: tool_template(m, self.cid)),
            ActionCard("📂", "Открыть документы", "Посмотреть и отредактировать PDF дела",
                       lambda: m.open_case_tab("docs")),
            ActionCard("⏱️", "Посчитать срок", "Апелляция, кассация, частная жалоба…", lambda: m.show_calc(0)),
            ActionCard("💰", "Проценты и пошлина", "Ст. 395 ГК, неустойка, госпошлина", lambda: m.show_calc(2)),
            ActionCard("🗺️", "Карта дела", "Схема: факты, позиции, доказательства, риски",
                       lambda: m.open_case_tab("board")),
        ]))
        v.addStretch(1)

    def _fact_link(self, href):
        if href.startswith("case:"):
            self.main.enter_case(int(href[5:]))
            self.main.open_case_tab("overview")
        elif href.startswith("folder:"):
            self.main.show_folder(int(href[7:]))
        elif href == "link":
            allc = db().cases(False) + db().cases(True)
            if len(allc) < 2:
                QMessageBox.information(self, "Связать дела", "Для связи нужно хотя бы два дела.")
                return
            d = FU.LinkDialog(self, allc, self.cid)
            if d.exec() == QDialog.Accepted:
                a, b, note = d.value()
                if a != b and db().add_link(a, b, note) is None:
                    self.main.toast("Эти дела уже связаны")
                self.set_case(self.cid)
        else:
            self.main.open_case_tab("info")

    def set_case(self, cid):
        if self.notes_timer.isActive() and self.cid:
            self.notes_timer.stop()
            self.save_notes()
        if self.brief_timer.isActive() and self.cid:
            self.brief_timer.stop()
            self.save_brief()
        self.cid = cid
        if not cid:
            return
        c = db().case(cid) or {}
        self.brief.blockSignals(True)
        self.brief.setPlainText(c.get("brief") or "")
        self.brief.blockSignals(False)
        self.brief_state.setText("")
        url = (c.get("court_url") or "").strip()
        self.b_court.setText("🌐  Дело на сайте суда" if url else "🔎  Найти дело на сайте суда")
        self.court_hint.setText("" if url else "Найдите карточку дела и сохраните ссылку (▾) — кнопка будет вести прямо в неё.")
        self.notes.blockSignals(True)
        self.notes.setPlainText(c.get("notes") or "")
        self.notes.blockSignals(False)
        self.notes_state.setText("")
        self.load_reminders()

        def fact(label, value):
            value = html.escape(value or "") or f'<span style="color:{M.T["muted"]}">не указано</span>'
            return f'<td style="padding:4px 22px 4px 0"><span style="color:{M.T["muted"]}">{label}</span><br>{value}</td>'
        rows = [[("Номер дела", c.get("number")), ("Суд", c.get("court")), ("Судья", c.get("judge"))],
                [("Доверитель", c.get("client")), ("Оппонент", c.get("opponent")), ("Стадия", c.get("stage"))]]
        t = "".join("<tr>" + "".join(fact(a, b) for a, b in r) + "</tr>" for r in rows)
        inst = db().instances(cid)
        if len(inst) > 1:                               # путь дела по инстанциям
            steps = " → ".join(f"{html.escape(i['court'] or i['level'])}"
                               + (f" <span style='color:{M.T['muted']}'>({html.escape(i['number'])})</span>"
                                  if i["number"] else "") for i in inst)
            t += f'<tr><td colspan="3" style="padding:6px 0 0 0"><span style="color:{M.T["muted"]}">Путь дела</span><br>{steps}</td></tr>'
        claim = html.escape(c.get("claim") or "")
        A = M.T["accent"]
        extra = ""
        f = db().folder(c.get("folder_id")) if c.get("folder_id") else None
        if f:
            extra += (f'<p style="margin-top:6px"><span style="color:{M.T["muted"]}">Папка</span> '
                      f'<a href="folder:{f["id"]}" style="color:{A}">📁 {html.escape(f["name"])}</a> '
                      f'<span style="color:{M.T["muted"]}">— сводка по всем делам доверителя</span></p>')
        links = db().links(cid)
        if links:
            items = []
            for l in links:
                o = db().case(l["other"]) or {}
                items.append(f'<a href="case:{l["other"]}" style="color:{A}">🔗 {html.escape(o.get("title", "?"))}</a>'
                             + (f' <span style="color:{M.T["muted"]}">({html.escape(l["note"])})</span>' if l["note"] else ""))
            extra += f'<p><span style="color:{M.T["muted"]}">Связанные дела</span><br>{"<br>".join(items)}</p>'
        self.facts.setText(f'<table>{t}</table>' + (f'<p style="margin-top:6px">{claim}</p>' if claim else "") + extra +
                           f'<p><a href="info" style="color:{A}">Изменить сведения о деле</a> &nbsp;·&nbsp; '
                           f'<a href="link" style="color:{A}">🔗 Связать с другим делом</a></p>')
        # сроки
        self.l_events.clear()
        today = dt.date.today()
        evs = [e for e in db().events(cid, include_done=False) if e["kind"] != CS.CaseDB.REMINDER]
        evs.sort(key=lambda e: (e["date"], e["time"] or ""))
        for e in evs[:7]:
            d = dt.date.fromisoformat(e["date"])
            left = (d - today).days
            when = "сегодня" if left == 0 else "завтра" if left == 1 else (
                f"просрочено {CS.ru(e['date'])}" if left < 0 else f"{CS.ru(e['date'])}, через {left} дн.")
            it = QListWidgetItem(f"{e['kind']}: {e['title']}\n{when}{(' в ' + e['time']) if e['time'] else ''}")
            if left <= 1:
                it.setForeground(QColor(M.T["danger"]))
            self.l_events.addItem(it)
        if not evs:
            it = QListWidgetItem("Сроков и заседаний пока нет")
            it.setFlags(Qt.NoItemFlags)
            self.l_events.addItem(it)
        # документы
        self.l_docs.clear()
        docs = db().docs(cid)
        docs_sorted = sorted(docs, key=lambda d: d.get("added") or "", reverse=True)
        for d in docs_sorted[:7]:
            sent = fmt_sent(d.get("sent"))
            it = QListWidgetItem(f"{d.get('icon') or default_icon(d['path'])}  {d['title']}" +
                                 (f"\n     отправлен {sent}" if sent else ""))
            it.setData(Qt.UserRole, d["path"])
            it.setToolTip(d["path"])
            self.l_docs.addItem(it)
        if not docs:
            it = QListWidgetItem("Документов пока нет — откройте вкладку «Документы» и добавьте файлы")
            it.setFlags(Qt.NoItemFlags)
            self.l_docs.addItem(it)


    # ---------------------------------------------------------------- сайт суда
    def court_clicked(self):
        c = db().case(self.cid) or {}
        url = (c.get("court_url") or "").strip()
        if url:
            QDesktopServices.openUrl(QUrl(url))
        else:
            court_menu(self.main, self.cid, self.b_court)

    # ---------------------------------------------------------------- заметки
    def _notes_changed(self):
        self.notes_state.setText("…")
        self.notes_timer.start()

    def save_brief(self):
        if not self.cid:
            return
        db().update_case(self.cid, brief=self.brief.toPlainText())
        self.brief_state.setText("сохранено")

    def save_notes(self):
        if not self.cid:
            return
        text = self.notes.toPlainText()
        db().update_case(self.cid, notes=text)
        cp = getattr(self.main, "cases_page", None)          # то же поле «Заметки» в сведениях о деле
        if cp is not None and cp.cid == self.cid and "notes" in cp.fields:
            f = cp.fields["notes"]
            f.blockSignals(True)
            f.setPlainText(text)
            f.blockSignals(False)
        self.notes_state.setText("сохранено ✓")

    # ---------------------------------------------------------------- напоминания
    def load_reminders(self):
        self.l_rem.clear()
        cur = now()
        for e in sorted(db().reminders(self.cid), key=event_dt):
            when = event_dt(e)
            days = (when.date() - cur.date()).days
            day = "сегодня" if days == 0 else "завтра" if days == 1 else when.strftime("%d.%m.%Y")
            it = QListWidgetItem(f"{e['title']}\n{day} в {when:%H:%M}" + ("  — время прошло" if when <= cur else ""))
            it.setData(Qt.UserRole, e["id"])
            if when <= cur:
                it.setForeground(QColor(M.T["danger"]))
            self.l_rem.addItem(it)
        if not self.l_rem.count():
            it = QListWidgetItem("Напоминаний нет. «+ Напоминание» — и программа напомнит в нужное время.")
            it.setFlags(Qt.NoItemFlags)
            self.l_rem.addItem(it)

    def edit_reminder(self, eid):
        if not self.cid:
            return
        e = next((x for x in db().reminders(self.cid, include_done=True) if x["id"] == eid), None) if eid else None
        dlg = ReminderDialog(self, e["title"] if e else "", event_dt(e) if e else None)
        if dlg.exec() and dlg.value:
            when, text = dlg.value
            if e:
                db().update_event(eid, date=when.date(), time=f"{when:%H:%M}", title=text, notified=0, done=0)
            else:
                db().add_event(self.cid, when.date(), CS.CaseDB.REMINDER, text, f"{when:%H:%M}")
            self.load_reminders()
            self.main.cases_page.reload_upcoming()

    def _rem_action(self, what, eid=None):
        if eid is None:
            it = self.l_rem.currentItem()
            eid = it.data(Qt.UserRole) if it else None
        if not eid:
            return
        if what == "done":
            db().update_event(eid, done=1)
        elif what == "delete":
            trash_event_with_undo(self.main, eid)
            return
        self.load_reminders()
        self.main.cases_page.reload_upcoming()

    def reminder_menu(self, pos):
        it = self.l_rem.itemAt(pos)
        eid = it.data(Qt.UserRole) if it else None
        if not eid:
            return
        m = QMenu(self)
        m.addAction("✓ Выполнено", lambda: self._rem_action("done", eid))
        m.addAction("Изменить…", lambda: self.edit_reminder(eid))
        m.addSeparator()
        m.addAction("Удалить", lambda: self._rem_action("delete", eid))
        m.exec(self.l_rem.viewport().mapToGlobal(pos))


class ReminderPopup(QDialog):
    """Окно «Пора!» поверх всех окон: готово / отложить / открыть дело."""

    def __init__(self, main, e):
        super().__init__(main)
        self.main, self.e = main, e
        self.setWindowTitle("Напоминание — " + M.APP_NAME)
        self.setWindowFlag(Qt.WindowStaysOnTopHint, True)
        self.setAttribute(Qt.WA_DeleteOnClose, True)
        self.setMinimumWidth(420)
        v = QVBoxLayout(self)
        v.setContentsMargins(22, 18, 22, 16)
        v.setSpacing(8)
        t = QLabel("⏰  Напоминание")
        t.setObjectName("title")
        v.addWidget(t)
        body = QLabel(html.escape(e["title"]).replace("\n", "<br>"))
        body.setWordWrap(True)
        body.setTextFormat(Qt.RichText)
        f = body.font()
        f.setPointSize(f.pointSize() + 2)
        body.setFont(f)
        v.addWidget(body)
        when = event_dt(e)
        sub = QLabel(f"{when:%d.%m.%Y %H:%M}" + (f" · {e['case_title']}" if e.get("case_title") else ""))
        sub.setObjectName("hint")
        v.addWidget(sub)
        row = QHBoxLayout()
        snooze = QPushButton("Отложить")
        sm = QMenu(snooze)
        for label, delta in (("на 10 минут", dt.timedelta(minutes=10)), ("на 1 час", dt.timedelta(hours=1)),
                             ("до завтра 9:00", None)):
            sm.addAction(label, lambda d=delta: self.snooze(d))
        snooze.setMenu(sm)
        row.addWidget(snooze)
        if e.get("case_id"):
            op = QPushButton("Открыть дело")
            op.clicked.connect(self.open_case)
            row.addWidget(op)
        row.addStretch(1)
        done = QPushButton("Готово")
        done.setObjectName("primary")
        done.clicked.connect(self.done_clicked)
        row.addWidget(done)
        v.addLayout(row)

    def showEvent(self, e):
        super().showEvent(e)
        import anim
        anim.window_fade(self, 0.0, 1.0, 260)

    def _refresh(self):
        try:
            self.main.overview.load_reminders()
            self.main.cases_page.reload_upcoming()
        except Exception:
            pass

    def snooze(self, delta):
        w = (now() + delta) if delta else (now() + dt.timedelta(days=1)).replace(hour=9, minute=0)
        db().update_event(self.e["id"], date=w.date(), time=f"{w:%H:%M}", notified=0)
        self._refresh()
        self.close()

    def done_clicked(self):
        db().update_event(self.e["id"], done=1)
        self._refresh()
        self.close()

    def open_case(self):
        self.main.showNormal()
        self.main.raise_()
        self.main.activateWindow()
        self.main.enter_case(self.e["case_id"])
        self.main.open_case_tab("overview")
        self.close()


TOOL_ICONS = {
    "package": "📨", "f107": "📮", "sheetnum": "🔢", "certify": "🖊️", "preflight": "✅", "anonymize": "🕶️",
    "compare_ed": "🔀", "board": "🗺️", "case_search": "🔍", "quote": "📌", "template": "📝",
    "calc_deadline": "⏱️", "calc_duty": "🏛️", "calc_interest": "💰",
    "merge": "➕", "organize": "🗂️", "split": "✂️", "rotate": "🔄", "delete": "🗑️", "extract": "📤",
    "pagesize": "📐", "blank": "📄", "reverse": "🔃", "edit": "✏️", "sign": "✍️", "watermark": "💧",
    "numbers": "#️⃣", "crop": "🖼️", "redact": "⬛", "highlight": "🖍️", "forms": "🧾", "meta": "🏷️",
    "compress": "🗜️", "ocr": "🔤", "repair": "🩹", "pdfa": "🗄️", "img2pdf": "🖼️", "word2pdf": "📘",
    "xls2pdf": "📗", "ppt2pdf": "📙", "html2pdf": "🌐", "pdf2word": "📘", "pdf2excel": "📗", "pdf2ppt": "📙",
    "pdf2jpg": "🖼️", "pdf2txt": "📃", "pdf2md": "⌨️", "protect": "🔒", "unlock": "🔓", "compare": "🆚",
}
# инструменты, которым не нужен открытый документ (создают новый или сами спрашивают файлы)
NO_DOC_TOOLS = {"merge", "img2pdf", "word2pdf", "xls2pdf", "ppt2pdf", "html2pdf", "repair", "unlock", "compare",
                "blank"}


def tool_callable(m, key):
    """Что делает карточка инструмента (вкладка «Создать документ», поиск Ctrl+K)."""
    special = {
        "template": lambda: tool_template(m, m.cases_page.cid),
        "package": lambda: m.tool_package(), "f107": lambda: m.tool_f107(),
        "board": lambda: m.open_case_tab("board"), "case_search": lambda: m.tool_case_search(),
        "compare_ed": lambda: m.tool_compare_ed(),
        "calc_deadline": lambda: m.show_calc(0), "calc_duty": lambda: m.show_calc(1),
        "calc_interest": lambda: m.show_calc(2),
    }
    if key in special:
        return special[key]
    if key in NO_DOC_TOOLS:
        return lambda: (m.open_case_tab("docs"), m.run_tool(key))
    return lambda: m.run_doc_tool(key)


class PrepareTab(QWidget):
    """Подготовить: все инструменты для документов дела — крупными карточками по разделам, с поиском
    и прокруткой. Комплект для подачи — во вкладке «Документы»."""
    COLS = 2

    def __init__(self, main, submission):
        super().__init__()
        self.main = main
        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        top = QHBoxLayout()
        top.setContentsMargins(2, 12, 10, 4)
        hint = QLabel("Инструменты работают с документом, открытым во вкладке «Документы».")
        hint.setObjectName("hint")
        top.addWidget(hint, 1)
        self.search = QLineEdit()
        self.search.setPlaceholderText("🔍  Найти инструмент: сжать, подпись, Word…")
        self.search.setClearButtonEnabled(True)
        self.search.setFixedWidth(300)
        self.search.textChanged.connect(self.filter)
        top.addWidget(self.search)
        outer.addLayout(top)
        sc = QScrollArea()
        sc.setWidgetResizable(True)
        sc.setFrameShape(QFrame.NoFrame)
        outer.addWidget(sc, 1)
        body = QWidget()
        sc.setWidget(body)
        v = QVBoxLayout(body)
        v.setContentsMargins(2, 6, 10, 14)
        v.setSpacing(8)
        self.sections = []                       # (заголовок, сетка-виджет, [(карточка, текст для поиска)])
        # сверху крупно — то, ради чего сюда приходят чаще всего: создать документ для суда
        hero = [("template", "📝", "По шаблону", "Иск, ходатайство, жалоба"),
                ("package", "📦", "Пакет в суд", "Иск и приложения одним PDF"),
                ("f107", "📮", "Опись ф. 107", "Для ценного письма")]
        hcards = []
        for key, icon, label, tip in hero:
            card = ActionCard(icon, label, tip, self._slot(key), key if key in HELP else None)
            hcards.append((card, f"{label} {tip} создать документ".lower()))
        hlab = QLabel("Создать документ")
        hlab.setObjectName("subtitle")
        hbox = QWidget()
        hg = card_grid([c for c, _t in hcards], cols=3)
        hg.setContentsMargins(0, 2, 0, 10)
        hbox.setLayout(hg)
        hbox.setProperty("cols", 3)
        hbox.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Maximum)
        v.addWidget(hlab)
        v.addWidget(hbox)
        self.sections.append((hlab, hbox, hcards))
        tools_lab = QLabel("Все инструменты для документов дела")
        tools_lab.setObjectName("hint")
        v.addSpacing(6)
        v.addWidget(tools_lab)
        self.tools_lab = tools_lab
        for cat, items in M.TOOLS:
            cards = []
            items = [it for it in items if it[0] not in {k for k, *_ in hero}]    # они уже наверху
            if not items:
                continue
            for key, label, tip in items:
                card = ActionCard(TOOL_ICONS.get(key, "•"), label, tip, self._slot(key),
                                  key if key in HELP else None)
                cards.append((card, f"{label} {tip} {cat}".lower()))
            lab = QLabel(cat)
            lab.setObjectName("subtitle")
            box = QWidget()
            grid = card_grid([c for c, _t in cards], cols=self.COLS)
            grid.setContentsMargins(0, 2, 0, 10)
            box.setLayout(grid)
            box.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Maximum)   # без пустых промежутков между рядами
            v.addWidget(lab)
            v.addWidget(box)
            self.sections.append((lab, box, cards))
        self.nothing = QLabel("Ничего не найдено")
        self.nothing.setObjectName("hint")
        self.nothing.hide()
        v.addWidget(self.nothing)
        kit = QPushButton("📦  Комплект для подачи — во вкладке «Документы» →")
        kit.setToolTip("Собрать документы комплекта в один PDF и сверить их")
        kit.clicked.connect(lambda: self.main.show_kit())
        v.addSpacing(6)
        v.addWidget(kit, 0, Qt.AlignLeft)
        v.addStretch(1)

    def _slot(self, key):
        return tool_callable(self.main, key)

    def filter(self, text):
        t = text.lower().strip()
        shown = 0
        for lab, box, cards in self.sections:
            grid = box.layout()
            vis = [c for c, hay in cards if not t or all(w in hay for w in t.split())]
            for c, _hay in cards:
                grid.removeWidget(c)
                c.setVisible(c in vis)
            cols = box.property("cols") or self.COLS
            for i, c in enumerate(vis):                   # переложить найденные плотно, без дыр
                grid.addWidget(c, i // cols, i % cols)
            lab.setVisible(bool(vis))
            box.setVisible(bool(vis))
            shown += len(vis)
        self.tools_lab.setVisible(not t)
        self.nothing.setVisible(shown == 0)


class CasesPage(QWidget):
    openFile = Signal(str, int)          # путь, страница
    buildPackage = Signal(list)          # пути

    def __init__(self, main):
        super().__init__()
        self.main = main
        self.cid = None
        self._loading = False
        self._restoring = False
        root = QHBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        split = QSplitter()
        root.addWidget(split)
        # ---- левая колонка
        left = QWidget()
        left.setObjectName("sidepanel")
        lv = QVBoxLayout(left)
        lv.setContentsMargins(14, 14, 14, 14)
        lv.addLayout(title_row("Дела", "cases", big=True))
        self.search = QLineEdit()
        self.search.setPlaceholderText("Фильтр дел: номер, доверитель, суд…")
        self.search.textChanged.connect(self.reload)
        lv.addWidget(self.search)
        self.list = QListWidget()
        self.list.setObjectName("caselist")
        self.list.currentItemChanged.connect(self.on_select)
        self.cur_fid = None                           # открытая папка (доверитель)
        FU.bind(_sys.modules[__name__])
        FU.setup_list(self)
        lv.addWidget(self.list, 1)
        rb = QHBoxLayout()
        b_new = QPushButton("+ Новое дело")
        b_new.setObjectName("primary")
        b_new.clicked.connect(lambda: self.new_case())
        b_fold = QPushButton("+ Папка")
        b_fold.setToolTip("Папка для дел одного доверителя (можно и подпапки). Дела перетаскиваются в папку мышью.")
        b_fold.clicked.connect(lambda: FU.new_folder(self, self.cur_fid if self._folder_selected() else None))
        self.b_folder = b_fold
        self.show_arch = QCheckBox("Архив")
        self.show_arch.toggled.connect(self.reload)
        rb.addWidget(b_new)
        rb.addWidget(b_fold)
        rb.addStretch(1)
        rb.addWidget(self.show_arch)
        lv.addLayout(rb)
        self.upcoming_lbl = QLabel("Ближайшие 14 дней")
        lv.addWidget(self.upcoming_lbl)
        self.upcoming_lbl.hide()                      # то же самое есть на «Главной» и в «Обзоре» дела
        self.upcoming = QListWidget()
        self.upcoming.setObjectName("upcoming")
        self.upcoming.setMaximumHeight(190)
        self.upcoming.setWordWrap(True)
        QTimer.singleShot(0, lambda: (self.reload(), self.restore_last()))
        self.upcoming.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.upcoming.itemDoubleClicked.connect(self.goto_event_case)
        lv.addWidget(self.upcoming)
        self.upcoming.hide()
        left.setMinimumWidth(300)
        self.left = left
        split.addWidget(left)
        # ---- карточка
        self.stack = QStackedWidget()
        empty = QLabel("Выберите дело слева или создайте новое")
        empty.setAlignment(Qt.AlignCenter)
        empty.setObjectName("hint")
        self.stack.addWidget(empty)
        self.card = QWidget()
        cv = QVBoxLayout(self.card)
        cv.setContentsMargins(22, 16, 22, 16)
        head = QHBoxLayout()
        self.h_title = QLabel()
        self.h_title.setObjectName("title")
        head.addWidget(self.h_title, 1)
        self.b_arch = QPushButton("В архив")
        self.b_arch.clicked.connect(self.toggle_archive)
        self.b_del = b_del = QPushButton("Удалить")
        b_del.clicked.connect(lambda: self.delete_case())
        b_del.hide()                                  # удаление — в «Ещё ▾»: не нажать случайно
        head.addWidget(self.b_arch)
        head.addWidget(b_del)
        cv.addLayout(head)
        self.tabs = QTabWidget()
        self.tabs.setObjectName("segmented")
        self.tabs.setDocumentMode(True)
        cv.addWidget(self.tabs, 1)
        self._build_info()
        self._build_events()
        self._build_docs()
        self._build_money()
        self._build_quotes()
        # новые вкладки: подача, карта дела, нормы права — сразу после «Сведения»
        self.events_tab = self.tabs.widget(1)
        self.sub_tab = CT.KitPanel(main)
        self.board_tab = CT.BoardTab(main)
        self.laws_tab = CT.LawsTab(main)
        self.tabs.insertTab(1, self.sub_tab, "Подача")
        self.tabs.insertTab(2, self.board_tab, "Карта дела")
        self.tabs.insertTab(3, self.laws_tab, "Нормы права")
        self.tabs.currentChanged.connect(self.on_tab)
        self.stack.addWidget(self.card)
        split.addWidget(self.stack)
        split.setStretchFactor(1, 1)
        split.setSizes([300, 900])
        self.save_timer = QTimer(self)
        self.save_timer.setSingleShot(True)
        self.save_timer.setInterval(600)
        self.save_timer.timeout.connect(self.save_info)
        self.reload()

    # ------------------------------------------------------------ вкладки
    def _build_info(self):
        w = QWidget()
        outer = QVBoxLayout(w)
        self.instances = InstancesEditor(self)
        outer.addWidget(self.instances)
        f = QFormLayout()
        outer.addLayout(f)
        outer.addStretch(1)
        self.fields = {}
        for key, label in CS.CASE_FIELDS:
            if key in INSTANCE_KEYS:
                continue                                # суд, номер, судья, ссылка — в блоке «Суды и инстанции»
            if key == "stage":
                e = QComboBox()
                e.setEditable(True)
                e.addItems(D.CASE_STAGES)
                e.currentTextChanged.connect(lambda *_: self.save_timer.start())
            elif key == "notes":
                e = QPlainTextEdit()
                e.setMinimumHeight(120)
                e.textChanged.connect(lambda *_: self.save_timer.start())
            elif key == "rate":
                e = money_edit(0)
                e.valueChanged.connect(lambda *_: self.save_timer.start())
            elif key == "folder":
                e = QLineEdit()
                e.textChanged.connect(lambda *_: self.save_timer.start())
                row = QHBoxLayout()
                row.addWidget(e, 1)
                bb = _small_btn("…")
                bb.setFixedWidth(34)
                bb.clicked.connect(self.pick_folder)
                row.addWidget(bb)
                bo = QPushButton("Открыть")
                bo.clicked.connect(lambda: self.fields["folder"].text() and
                                   QDesktopServices.openUrl(QUrl.fromLocalFile(self.fields["folder"].text())))
                row.addWidget(bo)
                self.fields[key] = e
                f.addRow(label, row)
                continue
            else:
                e = M.GrowEdit()                        # длинный текст переносится, поле растёт вниз
                e.textChanged.connect(lambda *_: self.save_timer.start())
            self.fields[key] = e
            f.addRow(label, e)
        sc = QScrollArea()
        sc.setWidget(w)
        sc.setWidgetResizable(True)
        sc.setFrameShape(QFrame.NoFrame)
        self.tabs.addTab(sc, "Сведения")

    def _table(self, headers):
        t = QTableWidget(0, len(headers))
        t.setHorizontalHeaderLabels(headers)
        t.verticalHeader().setVisible(False)
        t.setSelectionBehavior(QAbstractItemView.SelectRows)
        t.setEditTriggers(QAbstractItemView.NoEditTriggers)
        t.horizontalHeader().setStretchLastSection(True)
        t.setAlternatingRowColors(True)
        return t

    def _build_events(self):
        w = QWidget()
        v = QVBoxLayout(w)
        self.t_events = self._table(["Дата", "Время", "Вид", "Что", "Место", "Готово"])
        self.t_events.cellDoubleClicked.connect(self.edit_event)
        self.ev_stack = QStackedWidget()
        self.ev_stack.addWidget(self.t_events)
        self.ev_stack.addWidget(EmptyState("📅", "Сроков и заседаний пока нет",
                                           "Добавьте ближайшее заседание или процессуальный срок — программа "
                                           "заранее напомнит о нём.", "+ Заседание",
                                           lambda: self.add_event("Заседание")))
        v.addWidget(self.ev_stack, 1)
        r = QHBoxLayout()
        for text, fn in (("+ Заседание", lambda: self.add_event("Заседание")), ("+ Срок", lambda: self.add_event("Срок")),
                         ("+ Задача", lambda: self.add_event("Задача")), ("Отметить выполненным", self.toggle_done),
                         ("Удалить", self.del_event)):
            b = QPushButton(text)
            b.clicked.connect(fn)
            r.addWidget(b)
        r.addStretch(1)
        calc = QPushButton("⏱  Посчитать срок")
        calc.setToolTip("Калькулятор процессуальных сроков с учётом выходных и праздников; "
                        "результат можно сразу добавить в дело")
        calc.clicked.connect(lambda: self.main.show_calc(0))
        r.addWidget(calc)
        v.addLayout(r)
        self.tabs.addTab(w, "Сроки и заседания")

    def _build_docs(self):
        w = QWidget()
        v = QVBoxLayout(w)
        self.l_docs = DocsTable(self)
        self.l_docs.setToolTip("Синие — уже в PDF дела справа. Двойной щелчок — показать или добавить в него.\n"
                               "Правый щелчок — все действия: переименовать, отметить отправку, открыть в Word…")
        self.docs_stack = QStackedWidget()
        self.docs_stack.addWidget(self.l_docs)
        self.docs_empty = EmptyState("📄", "Документов пока нет",
                                     "Добавьте иск, договор, переписку — всё, что относится к делу. "
                                     "Можно просто перетащить файлы мышкой в окно программы. "
                                     "Они сразу соберутся в один «PDF дела» справа.",
                                     "+ Добавить файлы", lambda: self.add_docs())
        self.docs_stack.addWidget(self.docs_empty)
        v.addWidget(self.docs_stack, 1)
        tip = QLabel("Правый щелчок по документу — все действия")
        tip.setObjectName("hint")
        v.addWidget(tip)
        r = QHBoxLayout()
        b_add = QToolButton()
        b_add.setText("+ Добавить документ")
        b_add.setObjectName("primarytool")
        b_add.setToolTip("Файл с компьютера, документ по шаблону или пустой PDF. Добавленное сразу попадает в PDF дела")
        b_add.setPopupMode(QToolButton.InstantPopup)
        ma = QMenu(b_add)
        ma.addAction("Файлы с компьютера…", self.add_docs)
        ma.addAction("Документ по шаблону…", lambda: tool_template(self.main, self.cid))
        ma.addAction("Новый пустой PDF", lambda: self.main.new_doc())
        b_add.setMenu(ma)
        self.b_add_doc = b_add
        QTimer.singleShot(0, lambda: b_add.setMinimumWidth(b_add.sizeHint().width()))   # после стилей — не обрезать
        r.addWidget(b_add)
        r.addStretch(1)
        b_more = QToolButton()
        b_more.setText("Ещё ▾")
        b_more.setObjectName("moretabs")
        b_more.setPopupMode(QToolButton.InstantPopup)
        mm = QMenu(b_more)
        mm.addAction("📚 Собрать PDF дела из всех документов", lambda: self.main.build_case_pdf(self.cid))
        mm.addAction("📂 Показать PDF дела в папке", lambda: self.main.show_in_folder(self.main.case_pdf(self.cid)))
        mm.addSeparator()
        mm.addAction("Пакет в суд из выбранных", self.package_from_docs)
        mm.addAction("Поиск по документам", self.search_docs)
        mm.addSeparator()
        mm.addAction("📦 Собрать все файлы в папку дела", lambda: collect_case_files(self.main, self.cid))
        mm.addAction("📁 Открыть папку дела", lambda: open_case_folder(self.cid))
        a_copy = mm.addAction("Копировать добавляемые файлы в папку дела")
        a_copy.setCheckable(True)
        a_copy.setChecked(copy_docs_enabled())
        a_copy.toggled.connect(lambda on: M.settings().setValue("copy_docs", "1" if on else "0"))
        mm.addSeparator()
        mm.addAction("🗑  Убрать выбранные из дела", self.del_doc)
        mm.addAction("Удалить выбранные файлы с диска…", self.del_doc_files)
        mm.addAction("Убрать из списка файлы, которых больше нет", self.drop_missing_docs)
        b_more.setMenu(mm)
        r.addWidget(b_more)
        v.addLayout(r)
        self.docs_widget = w
        self.tabs.addTab(w, "Документы")

    def _build_money(self):
        w = QWidget()
        v = QVBoxLayout(w)
        v.addWidget(QLabel("Учёт времени и работ"))
        self.t_time = self._table(["Дата", "Часы", "Ставка", "Сумма, ₽", "Работа"])
        v.addWidget(self.t_time, 2)
        r = QHBoxLayout()
        for text, fn in (("+ Работа", self.add_time), ("Удалить", lambda: self._del_row(self.t_time, db().delete_time)),
                         ("Акт (Word)…", self.make_act)):
            b = QPushButton(text)
            b.clicked.connect(fn)
            r.addWidget(b)
        r.addStretch(1)
        v.addLayout(r)
        v.addWidget(QLabel("Оплаты"))
        self.t_pay = self._table(["Дата", "Сумма, ₽", "Комментарий"])
        v.addWidget(self.t_pay, 1)
        r2 = QHBoxLayout()
        for text, fn in (("+ Оплата", self.add_payment), ("Удалить", lambda: self._del_row(self.t_pay, db().delete_payment))):
            b = QPushButton(text)
            b.clicked.connect(fn)
            r2.addWidget(b)
        self.balance = QLabel()
        self.balance.setObjectName("subtitle")
        self.balance.setWordWrap(True)
        self.balance.setMinimumWidth(10)
        self.balance.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        r2.addSpacing(20)
        r2.addWidget(self.balance, 1)
        v.addLayout(r2)
        self.tabs.addTab(w, "Время и оплата")

    def _build_quotes(self):
        w = QWidget()
        v = QVBoxLayout(w)
        hint = QLabel("Выписки добавляются из редактора страницы: режим «Выписка» — выделите фрагмент рамкой.")
        hint.setObjectName("hint")
        hint.setWordWrap(True)
        v.addWidget(hint)
        self.l_quotes = QListWidget()
        self.l_quotes.setWordWrap(True)
        self.l_quotes.setSelectionMode(QAbstractItemView.ExtendedSelection)
        v.addWidget(self.l_quotes, 1)
        r = QHBoxLayout()
        for text, fn in (("Копировать", self.copy_quotes), ("В Word…", self.quotes_word), ("Удалить", self.del_quote)):
            b = QPushButton(text)
            b.clicked.connect(fn)
            r.addWidget(b)
        r.addStretch(1)
        v.addLayout(r)
        self.tabs.addTab(w, "Выписки")

    # ------------------------------------------------------------ список дел
    def reload(self, *_):
        cur = self.cid
        self.list.blockSignals(True)
        self.list.clear()
        cases = db().cases(self.show_arch.isChecked(), self.search.text())
        if not self.search.text().strip():            # обычный вид — папки доверителей и дела в них
            sel = FU.fill(self, cases, cur, self.cur_fid)
            if sel is not None:
                self.list.setCurrentItem(sel)
        else:                                         # поиск — плоский список, папка видна в подписи
            names = {f["id"]: f["name"] for f in db().folders()}
            for c in cases:
                sub = " · ".join(x for x in (c["number"], c["client"], c["stage"]) if x)
                if names.get(c.get("folder_id")):
                    sub = f"📁 {names[c['folder_id']]}" + (f" · {sub}" if sub else "")
                it = QListWidgetItem(c["title"] + (f"\n{sub}" if sub else ""))
                it.setData(Qt.UserRole, c["id"])
                self.list.addItem(it)
                if c["id"] == cur:
                    self.list.setCurrentItem(it)
        if not self.list.count():                     # пустой список — подсказать, что делать
            text = ("Ничего не найдено" if self.search.text().strip() else
                    "В архиве пусто" if self.show_arch.isChecked() else
                    "Дел пока нет.\nНажмите «+ Новое дело» ниже —\nпрограмма сама заведёт папку дела.")
            it = QListWidgetItem(text)
            it.setFlags(Qt.NoItemFlags)
            it.setTextAlignment(Qt.AlignCenter)
            it.setForeground(QColor(M.T["muted"]))
            self.list.addItem(it)
        self.list.blockSignals(False)
        self.reload_upcoming()
        if self.list.currentItem() is None or not self.list.currentItem().data(Qt.UserRole):
            self.cid = None
            self.stack.setCurrentIndex(0)

    def _folder_selected(self):
        it = self.list.currentItem()
        return it is not None and it.data(FU.FOLDER_ROLE) is not None

    def expand_to(self, cid):
        """Развернуть папки, в которых лежит дело (чтобы его строка была видна в списке)."""
        c = db().case(cid) or {}
        fid, seen, changed = c.get("folder_id"), set(), False
        while fid and fid not in seen:
            seen.add(fid)
            f = db().folder(fid)
            if not f:
                break
            if not f["expanded"]:
                db().update_folder(fid, expanded=1)
                changed = True
            fid = f["parent_id"]
        return changed

    def reload_upcoming(self):
        self.upcoming.clear()
        today = dt.date.today()
        for e in db().events(upcoming_days=14, include_done=False):
            d = dt.date.fromisoformat(e["date"])
            when = "сегодня" if d == today else "завтра" if d == today + dt.timedelta(days=1) else \
                ("просрочено " if d < today else "") + CS.ru(e["date"])
            it = QListWidgetItem(f"{when} {e['time']} — {e['kind']}: {e['title']}\n{e['case_title'] or ''}")
            it.setData(Qt.UserRole, e["case_id"])
            if d <= today:
                it.setForeground(QColor(M.T["danger"]))
            self.upcoming.addItem(it)
        if not self.upcoming.count():
            it = QListWidgetItem("Ничего не запланировано")
            it.setFlags(Qt.NoItemFlags)
            self.upcoming.addItem(it)

    def goto_event_case(self, it):
        cid = it.data(Qt.UserRole)
        if cid:
            self.select_case(cid)
            self.tabs.setCurrentWidget(self.events_tab)

    def select_case(self, cid):
        if self.expand_to(cid):
            self.reload()
        for i in range(self.list.count()):
            if self.list.item(i).data(Qt.UserRole) == cid:
                self.list.setCurrentRow(i)
                return
        self.show_arch.setChecked(not self.show_arch.isChecked())
        for i in range(self.list.count()):
            if self.list.item(i).data(Qt.UserRole) == cid:
                self.list.setCurrentRow(i)

    def on_tab(self, i):
        if self._restoring:
            return
        M.settings().setValue("case_tab", i)
        if self.tabs.widget(i) is self.board_tab:
            self.board_tab.activate()

    def flush(self):
        """Сохранить всё несохранённое (при смене дела и выходе)."""
        if self.save_timer.isActive():
            self.save_timer.stop()
            self.save_info()
        self.laws_tab.save_now()
        self.board_tab.flush()

    def restore_last(self):
        cid = int(M.settings().value("last_case", 0) or 0)
        if cid and db().case(cid):
            self.select_case(cid)
            self._restoring = True
            self.tabs.setCurrentIndex(int(M.settings().value("case_tab", 0) or 0))
            self._restoring = False
            if self.tabs.currentWidget() is self.board_tab:
                QTimer.singleShot(0, self.board_tab.activate)

    def on_select(self, it, _prev=None):
        if self.save_timer.isActive():
            self.save_timer.stop()
            self.save_info()
        self.laws_tab.save_now()
        self.board_tab.flush()
        if it is not None and (it.data(FU.FOLDER_ROLE) is not None or it.data(FU.HINT_ROLE)):
            self.cid = None                           # папка или подсказка — сводку покажет главное окно
            if it.data(FU.FOLDER_ROLE) is not None:
                self.cur_fid = it.data(FU.FOLDER_ROLE)
            return
        self.cid = it.data(Qt.UserRole) if it else None
        if self.cid:
            M.settings().setValue("last_case", self.cid)
        if not self.cid:
            self.stack.setCurrentIndex(0)
            return
        self.stack.setCurrentIndex(1)
        self.load_case()

    def load_case(self):
        c = db().case(self.cid)
        if not c:
            return
        self._loading = True
        self.h_title.setText(c["title"])
        self.b_arch.setText("Вернуть из архива" if c["archived"] else "В архив")
        for k, w in self.fields.items():
            v = c.get(k) or ""
            if isinstance(w, QComboBox):
                w.setCurrentText(str(v))
            elif isinstance(w, QPlainTextEdit):
                w.setPlainText(str(v))
            elif isinstance(w, QDoubleSpinBox):
                w.setValue(float(v or 0))
            else:
                w.setText(str(v))
        self._loading = False
        self.instances.set_case(self.cid)
        self.load_events()
        self.load_docs()
        self.sub_tab.set_case(self.cid)
        self.laws_tab.set_case(self.cid)
        self.board_tab.set_case(self.cid)
        if self.tabs.currentWidget() is self.board_tab:
            self.board_tab.activate()
        self.load_money()
        self.load_quotes()

    def save_info(self):
        if self._loading or not self.cid:
            return
        vals = {}
        for k, w in self.fields.items():
            if isinstance(w, QComboBox):
                vals[k] = w.currentText()
            elif isinstance(w, QPlainTextEdit):
                vals[k] = w.toPlainText()
            elif isinstance(w, QDoubleSpinBox):
                vals[k] = w.value()
            else:
                vals[k] = w.text()
        if not vals.get("title", "").strip():
            vals["title"] = "Без названия"
        cur = db().case(self.cid) or {}
        if not str(vals.get("folder", "")).strip() and cur.get("folder"):
            # папку дела программа завела сама после того, как форма открылась, — не стирать её пустым полем
            # (иначе при следующем сохранении появлялась вторая папка «<дело> (2)»)
            vals.pop("folder", None)
            w = self.fields.get("folder")
            if w is not None:
                w.blockSignals(True)
                w.setText(cur["folder"])
                w.blockSignals(False)
        db().update_case(self.cid, **vals)
        self.h_title.setText(vals["title"])
        it = self.list.currentItem()
        if it:
            c = db().case(self.cid) or {}
            sub = " · ".join(x for x in (c.get("number"), vals.get("client"), vals.get("stage")) if x)
            it.setText(vals["title"] + (f"\n{sub}" if sub else ""))

    def new_case(self, folder_id=None):
        """Новое дело. Открыта папка (или дело в папке) — дело создаётся в этой папке, доверитель
        подставляется из названия папки."""
        if folder_id is None:
            it = self.list.currentItem()
            if it is not None and it.data(FU.FOLDER_ROLE) is not None:
                folder_id = it.data(FU.FOLDER_ROLE)
            elif self.cid:
                folder_id = (db().case(self.cid) or {}).get("folder_id")
        f = db().folder(folder_id) if folder_id else None
        where = f"\n\nДело появится в папке «{f['name']}»." if f else ""
        name, ok = QInputDialog.getText(self, "Новое дело", "Название (например, «ООО Ромашка — взыскание долга»):"
                                        + where)
        if ok and name.strip():
            cid = db().add_case(title=name.strip(), stage=D.CASE_STAGES[1], **({"client": f["name"]} if f else {}))
            if f:
                db().set_case_folder(cid, f["id"])
                db().update_folder(f["id"], expanded=1)
            self.cid = cid
            sync_case_file(cid)                       # своя папка дела с файлом сведений — сразу
            self.show_arch.setChecked(False)
            self.reload()
            # reload() выбирает строку «молча» (сигналы выключены) — открываем новое дело явно
            self.on_select(self.list.currentItem())
            self.tabs.setCurrentIndex(0)
            if hasattr(self.main, "enter_case"):
                self.main.enter_case(cid)
                self.main.open_case_tab("overview")

    def delete_case(self, cid=None):
        """Дело — в корзину, сразу и без вопросов; внизу окна 10 секунд есть «Отменить»,
        а потом его можно вернуть из «🗑 Корзины» в течение 30 дней."""
        cid = cid or self.cid
        c = db().case(cid)
        if not c:
            return
        try:
            sync_case_file(cid)
        except Exception as e:
            M.log_error("Сведения дела перед удалением", e)
        tid = db().trash_case(cid)
        if cid == self.cid:
            self.cid = None
        self.reload()
        if hasattr(self.main, "show_home") and getattr(self.main, "mode_cid", None) == cid:
            self.main.show_home()
        refresh_trash_button(self.main)

        def undo():
            ok, why = db().trash_restore(tid)
            if not ok:
                self.main.toast(why)
                return
            self.reload()
            refresh_trash_button(self.main)
            if hasattr(self.main, "enter_case"):
                self.main.enter_case(cid)
        self.main.toast_undo(f"Дело «{c['title']}» — в корзине", undo)

    def toggle_archive(self):
        c = db().case(self.cid)
        db().update_case(self.cid, archived=0 if c["archived"] else 1)
        self.cid = None
        self.reload()

    def pick_folder(self):
        p = QFileDialog.getExistingDirectory(self, "Папка с документами дела", self.fields["folder"].text())
        if p:
            self.fields["folder"].setText(p)
            if QMessageBox.question(self, M.APP_NAME, "Добавить PDF и Word-файлы из этой папки в документы дела?") \
                    == QMessageBox.Yes:
                for f in sorted(Path(p).rglob("*")):
                    if f.suffix.lower() in (".pdf", ".docx", ".doc", ".jpg", ".png", ".tif", ".tiff"):
                        db().add_doc(self.cid, str(f))
                self.load_docs()

    # ------------------------------------------------------------ события
    def load_events(self):
        t = self.t_events
        t.setRowCount(0)
        for e in db().events(self.cid):
            r = t.rowCount()
            t.insertRow(r)
            for c, v in enumerate([CS.ru(e["date"]), e["time"], e["kind"], e["title"], e["place"],
                                   "✓" if e["done"] else ""]):
                it = QTableWidgetItem(v)
                it.setData(Qt.UserRole, e["id"])
                t.setItem(r, c, it)
        t.resizeColumnsToContents()
        t.horizontalHeader().setStretchLastSection(True)
        self.ev_stack.setCurrentIndex(0 if t.rowCount() else 1)

    def _event_dialog(self, e=None, kind="Срок"):
        e = e or {}
        dlg = QDialog(self)
        dlg.setWindowTitle(e.get("kind", kind))
        f = QFormLayout(dlg)
        de = date_edit(dt.date.fromisoformat(e["date"]) if e.get("date") else dt.date.today())
        te = QLineEdit(e.get("time", ""))
        te.setPlaceholderText("10:30")
        ke = QComboBox()
        ke.addItems(D.EVENT_KINDS)
        ke.setCurrentText(e.get("kind", kind))
        ti = QLineEdit(e.get("title", ""))
        ti.setPlaceholderText("Предварительное заседание / Отзыв на иск / …")
        pl = QLineEdit(e.get("place", ""))
        pl.setPlaceholderText("Каб. 305")
        f.addRow("Дата", de)
        f.addRow("Время", te)
        f.addRow("Вид", ke)
        f.addRow("Что", ti)
        f.addRow("Место", pl)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.accepted.connect(dlg.accept)
        bb.rejected.connect(dlg.reject)
        f.addRow(bb)
        if dlg.exec() != QDialog.Accepted:
            return None
        return dict(date=qdate(de), time=te.text().strip(), kind=ke.currentText(), title=ti.text().strip(),
                    place=pl.text().strip())

    def add_event(self, kind):
        v = self._event_dialog(kind=kind)
        if v:
            db().add_event(self.cid, v["date"], v["kind"], v["title"], v["time"], v["place"])
            self.load_events()
            self.reload_upcoming()

    def _cur_id(self, t):
        r = t.currentRow()
        return t.item(r, 0).data(Qt.UserRole) if r >= 0 and t.item(r, 0) else None

    def edit_event(self, row, _col):
        eid = self.t_events.item(row, 0).data(Qt.UserRole)
        e = next((x for x in db().events(self.cid) if x["id"] == eid), None)
        v = self._event_dialog(e)
        if v:
            db().update_event(eid, **v, notified=0)
            self.load_events()
            self.reload_upcoming()

    def toggle_done(self):
        eid = self._cur_id(self.t_events)
        if eid:
            e = next(x for x in db().events(self.cid) if x["id"] == eid)
            db().update_event(eid, done=0 if e["done"] else 1)
            self.load_events()
            self.reload_upcoming()

    def del_event(self):
        eid = self._cur_id(self.t_events)
        if eid:
            trash_event_with_undo(self.main, eid)

    # ------------------------------------------------------------ документы
    def refresh_docs_if(self, cid):
        if self.cid == cid:
            self.load_docs()

    def load_docs(self):
        docs = db().docs(self.cid) if self.cid else []
        self.l_docs.load(docs)
        self.docs_stack.setCurrentIndex(0 if docs else 1)
        if docs and hasattr(self.main, "prewarm_conversions"):
            # Word-файлы, которых ещё нет в PDF дела, заранее тихо превращаем в PDF — перетащить их потом мгновенно
            try:
                in_pdf = self.main.case_pdf_sources(self.cid)
            except Exception:
                in_pdf = set()
            todo = [d["path"] for d in docs if d["path"] and os.path.normcase(os.path.abspath(d["path"])) not in in_pdf]
            QTimer.singleShot(1500, lambda: self.main.prewarm_conversions(todo))

    def add_docs(self):
        paths, _ = QFileDialog.getOpenFileNames(self, "Документы дела", self.fields["folder"].text() or
                                                str(Path.home() / "Documents"), M.OPEN_FILTER)
        copy = copy_docs_enabled()
        added = []
        for p in paths:
            q = copy_into_case(self.cid, p) if copy else p
            db().add_doc(self.cid, q)
            added.append(q)
        self.load_docs()
        sync_case_file(self.cid)
        if added:                                   # и сразу в PDF дела
            self.main.show_main_ws(self.cid)
            self.main.open_paths(added)

    def add_doc_path(self, cid, path):
        db().add_doc(cid, path)
        if cid == self.cid:
            self.load_docs()

    def sel_doc_paths(self):
        return self.l_docs.paths(only_selected=True)

    def open_doc(self):
        for p in self.sel_doc_paths()[:1]:
            self.openFile.emit(p, 0)

    def package_from_docs(self):
        paths = self.sel_doc_paths() or self.l_docs.paths()
        if paths:
            PackageDialog(self.main, paths).exec()

    def search_docs(self):
        paths = self.l_docs.paths()
        SearchDialog(self.main, paths).exec()

    def _sel_docs(self):
        docs = [self.l_docs.row_doc(r) for r in self.l_docs.selected_rows()]
        return [d for d in docs if d[0] is not None]

    def del_doc(self):
        """Убрать выбранные документы из дела — сразу, без вопросов: в корзину программы (файлы на диске
        остаются), их страницы — из PDF дела. Внизу окна 10 секунд есть «Отменить»."""
        docs = self._sel_docs()
        if not docs:
            return
        cid, main = self.cid, self.main
        in_pdf = main.case_pdf_sources(cid)
        paths = [p for _i, p, _t in docs if p and os.path.normcase(os.path.abspath(p)) in in_pdf]
        snap, removed = None, 0
        if paths:                                       # один снимок для «Отменить» — на все документы сразу
            main.show_main_ws(cid)
            main.push_undo()
            snap = main.undo_stack[-1]
            removed = sum(main.remove_from_case_pdf(cid, p, undo=False) for p in paths)
            main.update_title()
        tid = db().trash_docs([d[0] for d in docs])
        self.load_docs()
        sync_case_file(cid)
        refresh_trash_button(main)

        def undo():
            ok, why = db().trash_restore(tid)
            if not ok:
                main.toast(why)
                return
            if removed:                                  # вернуть и страницы — отменой в PDF дела
                main.show_main_ws(cid)
                if main.undo_stack and main.undo_stack[-1] is snap:
                    main.undo()
                else:
                    main.open_paths([p for _i, p, _t in docs if p and os.path.exists(p)])
            self.refresh_docs_if(cid)
            sync_case_file(cid)
            refresh_trash_button(main)
        what = f"«{docs[0][2]}»" if len(docs) == 1 else f"Документы ({len(docs)})"
        extra = f", страниц из PDF дела: {removed}" if removed else ""
        main.toast_undo(f"{what} — убрано из дела{extra}. Файлы на диске целы", undo)

    def del_doc_files(self):
        """Удалить сами файлы с диска (в Корзину Windows) — это уже всерьёз, поэтому спрашиваем."""
        docs = self._sel_docs()
        if not docs:
            return
        what = f"файл «{docs[0][2]}»" if len(docs) == 1 else f"файлы выбранных документов ({len(docs)})"
        if QMessageBox.question(self, "Удалить с диска",
                                f"Удалить {what} с диска?\n\nФайлы уйдут в Корзину Windows (оттуда их можно "
                                "вернуть), документы уберутся из дела, а их страницы — из PDF дела.") \
                != QMessageBox.Yes:
            return
        gone = 0
        for _i, p, _t in docs:
            if p:
                self.main.remove_from_case_pdf(self.cid, p)
        for did, path, _t in docs:
            if path and os.path.exists(path):
                if not self.main.delete_file_completely(path):
                    continue
                db().forget_path(path)
                gone += 1
            db().delete_doc(did)
        self.load_docs()
        sync_case_file(self.cid)
        if gone:
            self.main.toast(f"🗑  В Корзину Windows: {gone} файл(ов)")

    def drop_missing_docs(self):
        n = 0
        for d in db().docs(self.cid):
            if not d["path"] or not os.path.exists(d["path"]):
                db().delete_doc(d["id"])
                n += 1
        self.load_docs()
        self.main.toast(f"Убрано из списка: {n}" if n else "Все файлы на месте")

    # ------------------------------------------------------------ время и деньги
    def load_money(self):
        t = self.t_time
        t.setRowCount(0)
        for e in db().time_entries(self.cid):
            r = t.rowCount()
            t.insertRow(r)
            s = e["amount"] or (e["hours"] or 0) * (e["rate"] or 0)
            for c, v in enumerate([CS.ru(e["date"]), f"{e['hours']:g}", L.money(e["rate"] or 0, False), L.money(s),
                                   e["description"]]):
                it = QTableWidgetItem(v)
                it.setData(Qt.UserRole, e["id"])
                t.setItem(r, c, it)
        p = self.t_pay
        p.setRowCount(0)
        for e in db().payments(self.cid):
            r = p.rowCount()
            p.insertRow(r)
            for c, v in enumerate([CS.ru(e["date"]), L.money(e["amount"]), e["note"]]):
                it = QTableWidgetItem(v)
                it.setData(Qt.UserRole, e["id"])
                p.setItem(r, c, it)
        b = db().balance(self.cid)
        nb = lambda x: f"{L.money(x)} ₽".replace(" ", "\u00a0")       # сумма не рвётся на две строки
        self.balance.setText(f"Часов:\u00a0{b['hours']:g} · Начислено:\u00a0{nb(b['billed'])} · "
                             f"Оплачено:\u00a0{nb(b['paid'])} · Долг:\u00a0{nb(b['due'])}")

    def _del_row(self, table, fn):
        i = self._cur_id(table)
        if i:
            fn(i)
            self.load_money()

    def add_time(self):
        dlg = QDialog(self)
        dlg.setWindowTitle("Работа по делу")
        f = QFormLayout(dlg)
        de = date_edit()
        hours = QDoubleSpinBox()
        hours.setRange(0, 1000)
        hours.setDecimals(2)
        hours.setValue(1)
        rate = money_edit(self.fields["rate"].value())
        fixed = money_edit(0)
        desc = QLineEdit()
        desc.setPlaceholderText("Подготовка отзыва на иск")
        f.addRow("Дата", de)
        f.addRow("Часы", hours)
        f.addRow("Ставка за час", rate)
        f.addRow("Или фикс. сумма", fixed)
        f.addRow("Работа", desc)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.accepted.connect(dlg.accept)
        bb.rejected.connect(dlg.reject)
        f.addRow(bb)
        if dlg.exec() == QDialog.Accepted:
            db().add_time(self.cid, qdate(de), hours.value(), desc.text(), rate.value(), fixed.value())
            self.load_money()

    def add_payment(self):
        dlg = QDialog(self)
        dlg.setWindowTitle("Оплата")
        f = QFormLayout(dlg)
        de, am, note = date_edit(), money_edit(0), QLineEdit()
        f.addRow("Дата", de)
        f.addRow("Сумма", am)
        f.addRow("Комментарий", note)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.accepted.connect(dlg.accept)
        bb.rejected.connect(dlg.reject)
        f.addRow(bb)
        if dlg.exec() == QDialog.Accepted:
            db().add_payment(self.cid, qdate(de), am.value(), note.text())
            self.load_money()

    def make_act(self):
        entries = [dict(date=CS.ru(e["date"]), description=e["description"], hours=e["hours"], rate=e["rate"],
                        amount=e["amount"]) for e in db().time_entries(self.cid)]
        if not entries:
            QMessageBox.information(self, M.APP_NAME, "Добавьте хотя бы одну работу.")
            return
        c = db().case(self.cid)
        executor, ok = QInputDialog.getText(self, "Акт", "Исполнитель (адвокат, ФИО):",
                                            text=M.settings().value("executor", ""))
        if not ok:
            return
        M.settings().setValue("executor", executor)
        p, _ = QFileDialog.getSaveFileName(self, "Акт", os.path.join(c.get("folder") or str(Path.home() / "Documents"),
                                           L.clean_filename(f"Акт — {c['title']}") + ".docx"), "Word (*.docx)")
        if p:
            L.act_docx(p, c, entries, executor=executor)
            db().add_doc(self.cid, p)
            self.load_docs()
            QDesktopServices.openUrl(QUrl.fromLocalFile(p))

    # ------------------------------------------------------------ выписки
    def load_quotes(self):
        self.l_quotes.clear()
        for q in db().quotes(self.cid):
            src = f"{q['source']}, л. {q['page']}" if q["page"] else q["source"]
            it = QListWidgetItem(f"«{q['text']}»\n— {src}")
            it.setData(Qt.UserRole, q["id"])
            it.setData(Qt.UserRole + 1, (q["text"], src))
            self.l_quotes.addItem(it)

    def _sel_quotes(self):
        items = self.l_quotes.selectedItems() or [self.l_quotes.item(i) for i in range(self.l_quotes.count())]
        return [it.data(Qt.UserRole + 1) for it in items]

    def copy_quotes(self):
        QApplication.clipboard().setText("\n\n".join(f"«{t}» ({s})" for t, s in self._sel_quotes()))
        self.main.msg("Выписки скопированы")

    def quotes_word(self):
        import docx
        qs = self._sel_quotes()
        if not qs:
            return
        c = db().case(self.cid)
        p, _ = QFileDialog.getSaveFileName(self, "Выписки", os.path.join(c.get("folder") or str(Path.home() / "Documents"),
                                           L.clean_filename(f"Выписки — {c['title']}") + ".docx"), "Word (*.docx)")
        if p:
            d = docx.Document()
            d.add_heading(f"Выписки по делу: {c['title']}", 1)
            for t, s in qs:
                d.add_paragraph(f"«{t}» ({s})")
            d.save(p)
            QDesktopServices.openUrl(QUrl.fromLocalFile(p))

    def del_quote(self):
        for it in self.l_quotes.selectedItems():
            db().delete_quote(it.data(Qt.UserRole))
        self.load_quotes()


# =============================================================================
#  Пакет в суд / на почту
# =============================================================================
class PackageDialog(QDialog):
    def __init__(self, main, paths=()):
        super().__init__(main)
        self.main = main
        self.setWindowTitle("Пакет в суд / на почту")
        self.resize(900, 640)
        v = QVBoxLayout(self)
        v.addLayout(title_row("Пакет документов", "package", big=True))
        f = QFormLayout()
        self.portal = QComboBox()
        self.keys = list(D.PORTALS)
        self.portal.addItems([D.PORTALS[k]["title"] for k in self.keys])
        self.pnote = QLabel()
        self.pnote.setWordWrap(True)
        self.pnote.setObjectName("hint")
        f.addRow("Куда", self.portal)
        f.addRow("", self.pnote)
        v.addLayout(f)
        self.t = QTableWidget(0, 4)
        self.t.setHorizontalHeaderLabels(["Документ (как в перечне приложений)", "Листов", "Экз.", "Файл"])
        self.t.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.t.horizontalHeader().setSectionResizeMode(3, QHeaderView.Interactive)
        self.t.verticalHeader().setVisible(False)
        self.t.setSelectionBehavior(QAbstractItemView.SelectRows)
        v.addWidget(self.t, 1)
        r = QHBoxLayout()
        for text, fn in (("+ Файлы…", self.add_files), ("+ Текущий документ", self.add_current), ("▲", lambda: self.move(-1)),
                         ("▼", lambda: self.move(1)), ("Убрать", lambda: self.t.removeRow(self.t.currentRow()))):
            b = QPushButton(text)
            b.clicked.connect(fn)
            r.addWidget(b)
        r.addStretch(1)
        hint = QLabel("Первая строка — основной документ (иск, жалоба, письмо)")
        hint.setObjectName("hint")
        r.addWidget(hint)
        v.addLayout(r)
        g = QGridLayout()
        self.c_prefix = QCheckBox("Номер по порядку в имени файла (01., 02., …)")
        self.c_prefix.setChecked(True)
        self.c_list = QCheckBox("Добавить перечень приложений последней страницей основного документа")
        self.c_duplex = QCheckBox("Двусторонняя печать (листов = страниц / 2)")
        self.c_combined = QCheckBox("Дополнительно — весь пакет одним файлом с закладками (для себя)")
        self.c_combined.setChecked(True)
        self.c_duplex.toggled.connect(self.recount)
        for i, c in enumerate((self.c_prefix, self.c_list, self.c_duplex, self.c_combined)):
            g.addWidget(c, i // 2, i % 2)
        v.addLayout(g)
        fo = QHBoxLayout()
        fo.addWidget(QLabel("Сохранить в папку:"))
        self.out = QLineEdit(os.path.join(main.default_dir(), "Пакет " + dt.date.today().strftime("%d.%m.%Y")))
        fo.addWidget(self.out, 1)
        b = _small_btn("…")
        b.setFixedWidth(34)
        b.clicked.connect(self.pick_out)
        fo.addWidget(b)
        v.addLayout(fo)
        bb = QHBoxLayout()
        b107 = QPushButton("Опись ф. 107 из списка…")
        b107.clicked.connect(self.to_f107)
        bl = QPushButton("Копировать перечень")
        bl.clicked.connect(self.copy_list)
        bb.addWidget(b107)
        bb.addWidget(bl)
        bb.addStretch(1)
        bc = QPushButton("Закрыть")
        bc.clicked.connect(self.reject)
        bg = QPushButton("Собрать пакет")
        bg.setObjectName("primary")
        bg.clicked.connect(self.build)
        bb.addWidget(bc)
        bb.addWidget(bg)
        v.addLayout(bb)
        self.portal.currentIndexChanged.connect(self.on_portal)
        self.portal.setCurrentIndex(int(M.settings().value("last_portal", 0) or 0))
        self.on_portal()
        for p in paths:
            self.add_path(p)

    def on_portal(self, *_):
        P = D.PORTALS[self.keys[self.portal.currentIndex()]]
        self.pnote.setText(P["note"] + f" Форматы: {P['formats']}.")
        M.settings().setValue("last_portal", self.portal.currentIndex())

    def pick_out(self):
        p = QFileDialog.getExistingDirectory(self, "Папка для пакета", self.out.text())
        if p:
            self.out.setText(os.path.join(p, "Пакет " + dt.date.today().strftime("%d.%m.%Y")))

    def add_path(self, p, title=None, pages=None):
        if pages is None:
            try:
                d = fitz.open(p) if p.lower().endswith(".pdf") else None
                pages = d.page_count if d else 1
            except Exception:
                pages = 1
        r = self.t.rowCount()
        self.t.insertRow(r)
        self.t.setItem(r, 0, QTableWidgetItem(title or Path(p).stem))
        sp = QSpinBox()
        sp.setRange(1, 100000)
        sp.setProperty("pages", pages)
        sp.setValue(L.count_sheets(pages, self.c_duplex.isChecked()))
        self.t.setCellWidget(r, 1, sp)
        cp = QSpinBox()
        cp.setRange(1, 100)
        self.t.setCellWidget(r, 2, cp)
        it = QTableWidgetItem(Path(p).name)
        it.setData(Qt.UserRole, p)
        it.setToolTip(p)
        it.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable)
        self.t.setItem(r, 3, it)

    def recount(self):
        for r in range(self.t.rowCount()):
            sp = self.t.cellWidget(r, 1)
            sp.setValue(L.count_sheets(int(sp.property("pages") or 1), self.c_duplex.isChecked()))

    def add_files(self):
        paths, _ = QFileDialog.getOpenFileNames(self, "Документы", self.main.default_dir(), M.OPEN_FILTER)
        for p in paths:
            self.add_path(p)

    def add_current(self):
        if not self.main.doc.page_count:
            return
        import tempfile
        p = os.path.join(tempfile.mkdtemp(prefix="pdfm_"), (self.main.base_name() or "Документ") + ".pdf")
        self.main.doc.save(p)
        self.add_path(p, self.main.base_name())

    def move(self, step):
        r = self.t.currentRow()
        n = r + step
        if r < 0 or not 0 <= n < self.t.rowCount():
            return
        rows = self.items()
        rows[r], rows[n] = rows[n], rows[r]
        self.t.setRowCount(0)
        for it in rows:
            self.add_path(it["path"], it["title"], it["pages"])
            self.t.cellWidget(self.t.rowCount() - 1, 1).setValue(it["sheets"])
            self.t.cellWidget(self.t.rowCount() - 1, 2).setValue(it["copies"])
        self.t.selectRow(n)

    def items(self):
        out = []
        for r in range(self.t.rowCount()):
            out.append(dict(title=self.t.item(r, 0).text().strip() or "Документ",
                            sheets=self.t.cellWidget(r, 1).value(), copies=self.t.cellWidget(r, 2).value(),
                            path=self.t.item(r, 3).data(Qt.UserRole),
                            pages=int(self.t.cellWidget(r, 1).property("pages") or 1)))
        return out

    def copy_list(self):
        lines = L.attachments_text(self.items())
        QApplication.clipboard().setText("Приложения:\n" + "\n".join(lines))
        QMessageBox.information(self, M.APP_NAME, "Перечень приложений скопирован:\n\n" + "\n".join(lines))

    def to_f107(self):
        rows = [(f"{it['title']} на {it['sheets']} л.", it["copies"], 1) for it in self.items()]
        F107Dialog(self, rows).exec()

    def build(self):
        items = self.items()
        if not items:
            return
        out = self.out.text().strip()
        if os.path.isdir(out) and os.listdir(out):
            if QMessageBox.question(self, M.APP_NAME, "Папка не пустая. Файлы с такими же именами будут "
                                    "перезаписаны. Продолжить?") != QMessageBox.Yes:
                return
        res = self.main.run("Сборка пакета…", L.build_package, items, self.keys[self.portal.currentIndex()], out,
                            number_prefix=self.c_prefix.isChecked(), add_list=self.c_list.isChecked(),
                            combined=self.c_combined.isChecked(), password_cb=self.main.ask_password)
        if res is M.FAILED:
            return
        lines = [f"• {os.path.basename(p)} — {n} стр., {s / 1048576:.1f} МБ" for p, n, s in res["files"]]
        text = "Готово. Файлы:\n" + "\n".join(lines)
        if res["warnings"]:
            text += "\n\nОбратите внимание:\n" + "\n".join("• " + w for w in res["warnings"])
        self.main.done(text, res["files"][0][0] if res["files"] else out, open_file=False)


# =============================================================================
#  Опись ф. 107
# =============================================================================
class F107Dialog(QDialog):
    def __init__(self, parent, rows=()):
        super().__init__(parent)
        self.setWindowTitle("Опись вложения ф. 107")
        self.resize(820, 620)
        v = QVBoxLayout(self)
        v.addLayout(title_row("Опись вложения ф. 107", "f107", big=True))
        s = M.settings()
        note = QLabel("Официальный бланк Почты России (как на pochta.ru): альбомный лист, два экземпляра рядом, "
                      "14 строк на листе, длинные названия переносятся целиком. Получатель и адрес в описи ф. 107 не указываются — они на конверте.")
        note.setObjectName("note")
        note.setWordWrap(True)
        v.addWidget(note)
        f = QFormLayout()
        self.sender = M.GrowEdit(s.value("f107_sender", "") or s.value("profile/Представитель", ""))
        self.sender.setPlaceholderText("ФИО или наименование юридического лица")
        self.spi = QLineEdit()
        self.spi.setPlaceholderText("необязательно — 14 цифр с чека или наклейки")
        self.spi.setMaxLength(20)
        f.addRow("Отправитель", self.sender)
        f.addRow("Почтовый идентификатор", self.spi)
        v.addLayout(f)
        self.t = QTableWidget(0, 3)
        self.t.setHorizontalHeaderLabels(["Наименование предметов", "Кол-во", "Ценность, руб."])
        self.t.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.t.setColumnWidth(1, 80)
        self.t.setColumnWidth(2, 130)
        self.t.setShowGrid(False)
        self.t.verticalHeader().setVisible(True)
        v.addWidget(self.t, 1)
        r = QHBoxLayout()
        for text, fn in (("+ Строка", lambda: self.add_row("", 1, 1)), ("Убрать", lambda: self.t.removeRow(self.t.currentRow()))):
            b = QPushButton(text)
            b.clicked.connect(fn)
            r.addWidget(b)
        r.addStretch(1)
        hint = QLabel("Пустая ценность — прочерк")
        hint.setObjectName("hint")
        r.addWidget(hint)
        v.addLayout(r)
        bb = QHBoxLayout()
        bb.addStretch(1)
        bc = QPushButton("Закрыть")
        bc.clicked.connect(self.reject)
        bs = QPushButton("Сохранить PDF…")
        bs.setObjectName("primary")
        bs.clicked.connect(self.save)
        bb.addWidget(bc)
        bb.addWidget(bs)
        v.addLayout(bb)
        for row in rows:
            self.add_row(*row)
        if not rows:
            self.add_row("", 1, 1)

    def add_row(self, name, q, val):
        r = self.t.rowCount()
        self.t.insertRow(r)
        self.t.setItem(r, 0, QTableWidgetItem(str(name)))
        self.t.setItem(r, 1, QTableWidgetItem(str(q)))
        self.t.setItem(r, 2, QTableWidgetItem("" if val in (None, "") else str(val)))

    def save(self):
        items = []
        for r in range(self.t.rowCount()):
            name = (self.t.item(r, 0) or QTableWidgetItem("")).text().strip()
            if not name:
                continue
            q = (self.t.item(r, 1) or QTableWidgetItem("1")).text().strip() or "1"
            val = (self.t.item(r, 2) or QTableWidgetItem("")).text().strip().replace(",", ".")
            try:
                items.append((name, int(q), float(val) if val else None))
            except ValueError:
                QMessageBox.warning(self, M.APP_NAME, f"Строка {r + 1}: количество и ценность — числа.")
                return
        if not items:
            return
        spi = "".join(ch for ch in self.spi.text() if not ch.isspace())
        if spi and (len(spi) != 14 or not spi.isdigit()):
            if QMessageBox.question(self, M.APP_NAME, f"Почтовый идентификатор обычно состоит из 14 цифр, а указано: "
                                    f"«{spi}». Всё равно продолжить?") != QMessageBox.Yes:
                return
        s = M.settings()
        s.setValue("f107_sender", self.sender.text())
        p, _ = QFileDialog.getSaveFileName(self, "Опись ф. 107", os.path.join(str(Path.home() / "Documents"),
                                           "Опись вложения ф107.pdf"), "PDF (*.pdf)")
        if not p:
            return
        if not p.lower().endswith(".pdf"):
            p += ".pdf"
        try:
            L.f107_pdf(items, sender=self.sender.text().strip(), spi=spi, out=p)
        except Exception as e:
            QMessageBox.warning(self, M.APP_NAME, str(e))
            return
        QDesktopServices.openUrl(QUrl.fromLocalFile(p))


# =============================================================================
#  Проверка перед подачей
# =============================================================================
class PreflightDialog(QDialog):
    ICON = {"error": "⛔", "warn": "⚠️", "info": "ℹ️", "ok": "✅"}

    def __init__(self, main):
        super().__init__(main)
        self.main = main
        self.setWindowTitle("Проверка перед подачей")
        self.resize(720, 460)
        v = QVBoxLayout(self)
        v.addLayout(title_row("Проверка перед подачей", "preflight", big=True))
        r = QHBoxLayout()
        r.addWidget(QLabel("Требования:"))
        self.portal = QComboBox()
        self.keys = [None] + list(D.PORTALS)
        self.portal.addItems(["Общие"] + [D.PORTALS[k]["title"] for k in self.keys[1:]])
        r.addWidget(self.portal, 1)
        v.addLayout(r)
        self.list = QListWidget()
        self.list.setWordWrap(True)
        self.list.itemDoubleClicked.connect(self.goto)
        v.addWidget(self.list, 1)
        self.portal.currentIndexChanged.connect(self.check)
        self.portal.setCurrentIndex(0)
        QTimer.singleShot(0, self.check)

    def check(self, *_):
        self.list.clear()
        size = os.path.getsize(self.main.path) if self.main.path and not self.main.modified else None
        res = L.preflight(self.main.doc, self.keys[self.portal.currentIndex()], size)
        for lvl, page, text in res:
            it = QListWidgetItem(f"{self.ICON[lvl]}  {text}")
            it.setData(Qt.UserRole, page)
            self.list.addItem(it)

    def goto(self, it):
        p = it.data(Qt.UserRole)
        if p:
            self.main.pages.clearSelection()
            self.main.pages.item(p - 1).setSelected(True)
            self.main.pages.scrollToItem(self.main.pages.item(p - 1))


# =============================================================================
#  Обезличивание
# =============================================================================
class AnonymizeDialog(QDialog):
    def __init__(self, main):
        super().__init__(main)
        self.main = main
        self.setWindowTitle("Обезличивание (152-ФЗ)")
        self.resize(760, 640)
        v = QVBoxLayout(self)
        v.addLayout(title_row("Обезличивание персональных данных", "anonymize", big=True))
        g = QGridLayout()
        self.checks = {}
        for i, (k, (label, _, _)) in enumerate(L.ANON_PATTERNS.items()):
            c = QCheckBox(label)
            c.setChecked(k not in ("name_patr",))
            self.checks[k] = c
            g.addWidget(c, i // 2, i % 2)
        v.addWidget(card(g))
        f = QFormLayout()
        self.custom = QLineEdit()
        self.custom.setPlaceholderText("Свои слова через запятую: Ромашка, Иванова, 1234567")
        self.mode = QComboBox()
        self.mode.addItems(L.ANON_MODES)
        self.meta = QCheckBox("Очистить свойства документа (автор, название)")
        self.meta.setChecked(True)
        f.addRow("Ещё скрыть", self.custom)
        f.addRow("Как скрыть", self.mode)
        f.addRow("", self.meta)
        v.addLayout(f)
        rb = QHBoxLayout()
        b = QPushButton("Найти")
        b.clicked.connect(self.find)
        rb.addWidget(b)
        self.cnt = QLabel()
        rb.addWidget(self.cnt, 1)
        v.addLayout(rb)
        self.list = QListWidget()
        v.addWidget(self.list, 1)
        bb = QHBoxLayout()
        bb.addStretch(1)
        bc = QPushButton("Отмена")
        bc.clicked.connect(self.reject)
        self.bgo = QPushButton("Скрыть отмеченное")
        self.bgo.setObjectName("primary")
        self.bgo.setEnabled(False)
        self.bgo.clicked.connect(self.apply)
        bb.addWidget(bc)
        bb.addWidget(self.bgo)
        v.addLayout(bb)

    def find(self):
        keys = [k for k, c in self.checks.items() if c.isChecked()]
        custom = [w.strip() for w in self.custom.text().split(",") if w.strip()]
        found = L.find_personal(self.main.doc, keys, custom)
        self.list.clear()
        seen = {}
        for pno, hits in found.items():
            for s, tag in hits:
                seen.setdefault((s, tag), []).append(pno + 1)
        for (s, tag), pages in sorted(seen.items(), key=lambda x: (x[0][1], x[0][0])):
            it = QListWidgetItem(f"{tag}  {s.replace(chr(10), ' ')}   — стр. {L.C_ranges(pages)}")
            it.setData(Qt.UserRole, s)
            it.setFlags(it.flags() | Qt.ItemIsUserCheckable)
            it.setCheckState(Qt.Checked)
            self.list.addItem(it)
        self.cnt.setText(f"Найдено: {len(seen)} уникальных фрагментов" if seen else "Ничего не найдено")
        self.bgo.setEnabled(bool(seen))

    def apply(self):
        words = [self.list.item(i).data(Qt.UserRole) for i in range(self.list.count())
                 if self.list.item(i).checkState() == Qt.Checked]
        if not words:
            return
        self.main.push_undo()
        n = self.main.run("Обезличивание…", L.anonymize, self.main.doc, [], words, self.mode.currentText(),
                          self.meta.isChecked())
        if n is M.FAILED:
            return
        self.main.after_change()
        self.accept()
        self.main.done(f"Скрыто фрагментов: {n}. Данные удалены из текста безвозвратно.\n"
                       "Просмотрите документ и сохраните его под новым именем (Файл → Сохранить как).")


# =============================================================================
#  Поиск по документам
# =============================================================================
class SearchDialog(QDialog):
    def __init__(self, main, paths=()):
        super().__init__(main)
        self.main = main
        self.paths = list(paths)
        self.setWindowTitle("Поиск по документам")
        self.resize(860, 560)
        v = QVBoxLayout(self)
        v.addLayout(title_row("Поиск по документам", "case_search", big=True))
        r = QHBoxLayout()
        self.q = QLineEdit()
        self.q.setPlaceholderText("Что ищем: «неустойка», «акт сверки», номер договора…")
        self.q.returnPressed.connect(self.search)
        r.addWidget(self.q, 1)
        b = QPushButton("Найти")
        b.setObjectName("primary")
        b.clicked.connect(self.search)
        r.addWidget(b)
        v.addLayout(r)
        r2 = QHBoxLayout()
        self.src = QLabel()
        r2.addWidget(self.src, 1)
        for text, fn in (("Файлы…", self.pick_files), ("Папка…", self.pick_folder)):
            bb = QPushButton(text)
            bb.clicked.connect(fn)
            r2.addWidget(bb)
        v.addLayout(r2)
        self.t = QTableWidget(0, 3)
        self.t.setHorizontalHeaderLabels(["Файл", "Где", "Фрагмент"])
        self.t.horizontalHeader().setSectionResizeMode(2, QHeaderView.Stretch)
        self.t.verticalHeader().setVisible(False)
        self.t.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.t.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.t.cellDoubleClicked.connect(self.open_hit)
        v.addWidget(self.t, 1)
        self.foot = QLabel()
        self.foot.setWordWrap(True)
        self.foot.setObjectName("hint")
        v.addWidget(self.foot)
        self.update_src()

    def update_src(self):
        self.src.setText(f"Документов для поиска: {len(self.paths)}")

    def pick_files(self):
        p, _ = QFileDialog.getOpenFileNames(self, "Файлы", self.main.default_dir(), M.OPEN_FILTER)
        if p:
            self.paths = p
            self.update_src()

    def pick_folder(self):
        d = QFileDialog.getExistingDirectory(self, "Папка", self.main.default_dir())
        if d:
            self.paths = [str(f) for f in sorted(Path(d).rglob("*")) if f.suffix.lower() in (".pdf", ".docx")]
            self.update_src()

    def search(self):
        q = self.q.text().strip()
        if not q or not self.paths:
            return
        L.SEARCH_CACHE = os.path.join(M.data_dir(), "search_cache.sqlite")
        res = self.main.run("Поиск…", L.search_files, [p for p in self.paths if os.path.exists(p)], q)
        if res is M.FAILED:
            return
        hits, no_text = res
        self.t.setRowCount(0)
        for p, pg, snip in hits:
            r = self.t.rowCount()
            self.t.insertRow(r)
            it = QTableWidgetItem(Path(p).name)
            it.setData(Qt.UserRole, p)
            self.t.setItem(r, 0, it)
            self.t.setItem(r, 1, QTableWidgetItem(str(pg)))
            self.t.setItem(r, 2, QTableWidgetItem(snip))
        self.t.resizeColumnToContents(0)
        foot = f"Найдено: {len(hits)}."
        if no_text:
            foot += " Без текстового слоя (нужен OCR): " + ", ".join(Path(p).name for p in no_text[:8])
        self.foot.setText(foot)

    def open_hit(self, row, _c):
        p = self.t.item(row, 0).data(Qt.UserRole)
        lab = self.t.item(row, 1).text()
        pg = int(lab) if lab.isdigit() else 1           # у Word — номер абзаца, открываем с начала
        self.main.open_external(p, pg - 1)


# =============================================================================
#  Шаблоны
# =============================================================================
def ensure_sample_templates():
    import templates_lib as TL
    try:
        TL.ensure_builtin(templates_dir())
    except Exception as e:
        M.log_error("Создание встроенных шаблонов", e)


PROFILE_FIELDS = [
    ("Представитель", "ФИО представителя", "Петров Пётр Петрович"),
    ("Контакты_представителя", "Контакты", "тел. +7 900 000-00-00, e-mail: lawyer@example.ru"),
    ("Реестровый_номер", "Рег. № в реестре адвокатов", "02/1234"),
    ("Субъект_реестра", "Субъект РФ реестра", "Республики …"),
    ("Адвокатское_образование", "Адвокатское образование", "Коллегия адвокатов «…»"),
]


def profile_values():
    s = M.settings()
    return {k: s.value(f"profile/{k}", "") or "" for k, _l, _p in PROFILE_FIELDS}


class ProfileDialog(QDialog):
    """«Мои реквизиты» — подставляются во все шаблоны."""

    def __init__(self, parent):
        super().__init__(parent)
        self.setWindowTitle("Мои реквизиты")
        self.resize(560, 330)
        v = QVBoxLayout(self)
        v.addLayout(title_row("Мои реквизиты"))
        hint = QLabel("Эти данные автоматически подставляются в шаблоны документов: шапку, подпись, адвокатский запрос.")
        hint.setObjectName("hint")
        hint.setWordWrap(True)
        v.addWidget(hint)
        f = QFormLayout()
        vals = profile_values()
        self.edits = {}
        for k, label, ph in PROFILE_FIELDS:
            e = M.GrowEdit(vals.get(k, ""))
            e.setPlaceholderText(ph)
            f.addRow(label, e)
            self.edits[k] = e
        v.addLayout(f)
        bb = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Save).setText("Сохранить")
        bb.button(QDialogButtonBox.Cancel).setText("Отмена")
        bb.accepted.connect(self.save)
        bb.rejected.connect(self.reject)
        v.addWidget(bb)

    def save(self):
        s = M.settings()
        for k, e in self.edits.items():
            s.setValue(f"profile/{k}", e.text().strip())
        self.accept()


def case_output_dir(cid):
    """Куда класть созданные документы: папка дела (та же, где PDF дела), без дела — «Дела/Без дела»."""
    if cid and db().case(cid):
        return CF.ensure_folder(db(), cid)
    c = None
    base = Path.home() / "Documents" / "LegalHelper" / "Дела"
    old = Path.home() / "Documents" / "PDF Мастер" / "Дела"     # папка версий до 1.6 — продолжаем её использовать
    if old.is_dir() and not base.exists():
        base = old
    name = L.clean_filename(c["title"]) if c else "Без дела"
    d = base / name
    d.mkdir(parents=True, exist_ok=True)
    return str(d)


PARTY_TYPES = {
    "Физическое лицо": [("ФИО", "Иванов Иван Иванович"), ("Дата_рождения", "01.01.1980"),
                        ("Место_рождения", "г. Москва"), ("Адрес", "Адрес места жительства (регистрации)"),
                        ("Паспорт", "серия и номер, кем и когда выдан"), ("СНИЛС", ""), ("ИНН", ""),
                        ("Телефон", ""), ("Email", "")],
    "Индивидуальный предприниматель": [("ФИО", "Иванов Иван Иванович"), ("ОГРНИП", ""), ("ИНН", ""),
                                       ("Дата_рождения", ""), ("Место_рождения", ""), ("Адрес", "Адрес места жительства"),
                                       ("Телефон", ""), ("Email", "")],
    "Организация (ООО, АО…)": [("Форма", "ООО"), ("Наименование", "Ромашка — без кавычек и формы"), ("ОГРН", ""),
                               ("ИНН", ""), ("КПП", ""), ("Адрес", "Адрес юридического лица"), ("Телефон", ""),
                               ("Email", "")],
    "Госорган / иное": [("Наименование", "Полное наименование"), ("Реквизиты", "ОГРН, ИНН и др."), ("Адрес", ""),
                        ("Телефон", ""), ("Email", "")],
}
ORG_FORMS = {"ООО": "Общество с ограниченной ответственностью", "АО": "Акционерное общество",
             "ПАО": "Публичное акционерное общество", "НАО": "Непубличное акционерное общество",
             "АНО": "Автономная некоммерческая организация", "ГУП": "Государственное унитарное предприятие",
             "МУП": "Муниципальное унитарное предприятие", "ТСЖ": "Товарищество собственников жилья",
             "СНТ": "Садоводческое некоммерческое товарищество"}
PARTY_LABELS = {"Истец": "Истец", "Ответчик": "Ответчик", "Третье_лицо": "Третье лицо",
                "Третье_лицо_2": "Третье лицо", "Заявитель": "Заявитель", "Заявитель_жалобы": "Заявитель жалобы"}


def initials(fio):
    parts = fio.split()
    if len(parts) < 2:
        return fio
    return parts[0] + " " + "".join(p[0] + "." for p in parts[1:3])


def guess_party(name):
    """Тип и поля по строке из карточки дела («ООО «Ромашка»», «ИП Сидоров С.С.», «Петров П.П.»)."""
    n = (name or "").strip()
    if not n:
        return {}
    import re as _re
    m = _re.match(r"^(ООО|АО|ПАО|НАО|АНО|ГУП|МУП|ТСЖ|СНТ)\s+[«\"]?(.+?)[»\"]?$", n)
    if m:
        return {"type": "Организация (ООО, АО…)", "Форма": m.group(1), "Наименование": m.group(2)}
    m = _re.match(r"^(?:ИП|Индивидуальный предприниматель)\s+(.+)$", n)
    if m:
        return {"type": "Индивидуальный предприниматель", "ФИО": m.group(1)}
    return {"type": "Физическое лицо", "ФИО": n}


class PartyEditor(QFrame):
    """Сторона для шапки документа: тип (физлицо / ИП / организация / иное) и реквизиты по ст. 125 АПК, 131 ГПК."""

    def __init__(self, role, data=None):
        super().__init__()
        self.role = role
        self.setObjectName("card")
        self.values = dict(data or {})
        v = QVBoxLayout(self)
        v.setContentsMargins(14, 10, 14, 10)
        h = QHBoxLayout()
        lab = QLabel(PARTY_LABELS.get(role, role.replace("_", " ")))
        lab.setObjectName("subtitle")
        h.addWidget(lab)
        h.addStretch(1)
        self.kind = QComboBox()
        self.kind.addItems(list(PARTY_TYPES))
        self.kind.setCurrentText(self.values.get("type") or "Физическое лицо")
        self.kind.currentTextChanged.connect(self._rebuild)
        h.addWidget(self.kind)
        v.addLayout(h)
        self.form = QFormLayout()
        self.form.setLabelAlignment(Qt.AlignRight | Qt.AlignVCenter)
        v.addLayout(self.form)
        self.edits = {}
        self._rebuild()

    def _rebuild(self, *_):
        self._collect()
        while self.form.rowCount():
            self.form.removeRow(0)
        self.edits = {}
        for key, hint in PARTY_TYPES[self.kind.currentText()]:
            if key == "Форма":
                e = QComboBox()
                e.setEditable(True)
                e.addItems(list(ORG_FORMS))
                e.setCurrentText(self.values.get(key, "ООО"))
            else:
                e = M.GrowEdit(self.values.get(key, ""), hint)
            self.form.addRow(key.replace("_", " ").replace("Email", "E-mail"), e)
            self.edits[key] = e

    def _collect(self):
        for k, e in getattr(self, "edits", {}).items():
            self.values[k] = (e.currentText() if isinstance(e, QComboBox) else e.text()).strip()
        if hasattr(self, "kind"):
            self.values["type"] = self.kind.currentText()

    def data(self):
        self._collect()
        return dict(self.values)

    def names(self):
        """(полное наименование для шапки, краткое — для текста)."""
        v = self.data()
        t = v.get("type")
        if t == "Организация (ООО, АО…)":
            name = v.get("Наименование", "").strip("«»\" ")
            if not name:
                return "", ""
            form = v.get("Форма", "")
            return f"{ORG_FORMS.get(form, form)} «{name}»", f"{form} «{name}»"
        if t == "Индивидуальный предприниматель":
            fio = v.get("ФИО", "")
            return (f"Индивидуальный предприниматель {fio}", f"ИП {initials(fio)}") if fio else ("", "")
        name = v.get("ФИО") or v.get("Наименование") or ""
        return name, name

    def block(self):
        """Текст блока для шапки: «Истец: …» и реквизиты — каждая строка отдельно."""
        full, _short = self.names()
        if not full:
            return ""
        v = self.data()
        lines = [f"{PARTY_LABELS.get(self.role, self.role.replace('_', ' '))}: {full}"]
        born = ", ".join(x for x in (v.get("Дата_рождения") and f"дата рождения: {v['Дата_рождения']}",
                                     v.get("Место_рождения") and f"место рождения: {v['Место_рождения']}") if x)
        ids = ", ".join(f"{k} {v[k]}" for k in ("ОГРН", "ОГРНИП", "ИНН", "КПП", "СНИЛС") if v.get(k))
        if born:
            lines.append(born[0].upper() + born[1:])
        if v.get("Паспорт"):
            lines.append(f"Паспорт: {v['Паспорт']}")
        if ids:
            lines.append(ids)
        if v.get("Реквизиты"):
            lines.append(v["Реквизиты"])
        if v.get("Адрес"):
            label = "Адрес" if v.get("type", "").startswith(("Организация", "Госорган")) else "Место жительства"
            lines.append(f"{label}: {v['Адрес']}")
        contacts = ", ".join(x for x in (v.get("Телефон") and f"тел.: {v['Телефон']}",
                                         v.get("Email") and f"e-mail: {v['Email']}") if x)
        if contacts:
            lines.append(contacts[0].upper() + contacts[1:])
        return "\n".join(lines)


def refresh_after_events(main):
    """Обновить всё, где видны сроки: карточку дела, «Обзор», «Главную»."""
    cp = main.cases_page
    try:
        if cp.cid:
            cp.load_events()
        cp.reload_upcoming()
        if getattr(main, "mode_cid", None) and hasattr(main, "overview"):
            main.overview.set_case(main.mode_cid)
        if hasattr(main, "home_page") and main.home_page.isVisible():
            main.home_page.refresh()
    except Exception as e:
        M.log_error("Обновление сроков", e)


def trash_event_with_undo(main, eid):
    e = next((x for x in db().events() if x["id"] == eid), None)
    tid = db().trash_events([eid])
    if not tid:
        return
    refresh_after_events(main)
    refresh_trash_button(main)

    def undo():
        ok, why = db().trash_restore(tid)
        if not ok:
            main.toast(why)
        refresh_after_events(main)
        refresh_trash_button(main)
    title = (e or {}).get("title") or (e or {}).get("kind") or "Событие"
    main.toast_undo(f"«{title}» — в корзине", undo)


TRASH_KINDS = {"case": "📁 Дело", "doc": "📄 Документ", "event": "📅 Срок / заседание"}


def refresh_trash_button(main):
    b = getattr(main, "b_trash", None)
    if b is None:
        return
    try:
        n = len(db().trash_items())
    except Exception:
        n = 0
    b.setText("🗑   Корзина" + (f" ({n})" if n else ""))


class TrashDialog(QDialog):
    """Корзина: удалённые дела, документы и сроки. Хранятся 30 дней, потом стираются сами."""

    def __init__(self, main):
        super().__init__(main)
        self.main = main
        self.setWindowTitle("Корзина")
        self.resize(720, 460)
        v = QVBoxLayout(self)
        v.setContentsMargins(20, 16, 20, 16)
        t = QLabel("🗑  Корзина")
        t.setObjectName("title")
        v.addWidget(t)
        hint = QLabel(f"Здесь лежит удалённое за последние {db().TRASH_DAYS} дней — потом оно стирается само. "
                      "Файлы на диске при удалении не трогаются.")
        hint.setObjectName("hint")
        hint.setWordWrap(True)
        v.addWidget(hint)
        self.stack = QStackedWidget()
        self.table = QTableWidget(0, 4)
        self.table.setHorizontalHeaderLabels(["Что", "Название", "Из дела", "Удалено"])
        self.table.verticalHeader().hide()
        self.table.setShowGrid(False)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        h = self.table.horizontalHeader()
        h.setSectionResizeMode(1, QHeaderView.Stretch)
        self.table.setColumnWidth(0, 150)
        self.table.setColumnWidth(2, 180)
        self.table.setColumnWidth(3, 150)
        self.table.verticalHeader().setDefaultSectionSize(36)
        self.table.cellDoubleClicked.connect(lambda *_: self.restore())
        self.stack.addWidget(self.table)
        self.empty = EmptyState("🗑", "Корзина пуста", "Удалённые дела, документы и сроки будут лежать здесь "
                                f"{db().TRASH_DAYS} дней — их можно вернуть.")
        self.stack.addWidget(self.empty)
        v.addWidget(self.stack, 1)
        row = QHBoxLayout()
        self.b_restore = QPushButton("↩  Восстановить")
        self.b_restore.setObjectName("primary")
        self.b_restore.clicked.connect(self.restore)
        self.b_forget = QPushButton("Удалить навсегда")
        self.b_forget.clicked.connect(self.forget)
        self.b_empty = QPushButton("Очистить корзину")
        self.b_empty.clicked.connect(self.empty_all)
        close = QPushButton("Закрыть")
        close.clicked.connect(self.accept)
        row.addWidget(self.b_restore)
        row.addWidget(self.b_forget)
        row.addStretch(1)
        row.addWidget(self.b_empty)
        row.addWidget(close)
        v.addLayout(row)
        self.table.itemSelectionChanged.connect(self._sync)
        self.load()

    def load(self):
        items = db().trash_items()
        self.table.setRowCount(0)
        for it in items:
            r = self.table.rowCount()
            self.table.insertRow(r)
            try:
                when = dt.datetime.fromisoformat(it["deleted"]).strftime("%d.%m.%Y %H:%M")
            except Exception:
                when = it["deleted"] or ""
            for c, val in enumerate((TRASH_KINDS.get(it["kind"], it["kind"]), it["title"],
                                     it["case_title"] if it["kind"] != "case" else "", when)):
                cell = QTableWidgetItem(val)
                cell.setData(Qt.UserRole, it["id"])
                self.table.setItem(r, c, cell)
        self.stack.setCurrentIndex(0 if items else 1)
        self.b_empty.setEnabled(bool(items))
        if items:
            self.table.selectRow(0)
        self._sync()
        refresh_trash_button(self.main)

    def _ids(self):
        return [self.table.item(r, 0).data(Qt.UserRole)
                for r in sorted({i.row() for i in self.table.selectedIndexes()})]

    def _sync(self):
        on = bool(self._ids()) and self.stack.currentIndex() == 0
        self.b_restore.setEnabled(on)
        self.b_forget.setEnabled(on)

    def restore(self):
        ids = self._ids()
        # дела — первыми: тогда вернутся и их документы и сроки, выбранные вместе с ними
        kinds = {i: k["kind"] for k in db().trash_items() for i in [k["id"]]}
        ids.sort(key=lambda i: kinds.get(i) != "case")
        done, fails = 0, []
        for tid in ids:
            ok, why = db().trash_restore(tid)
            if ok:
                done += 1
            elif why not in fails:
                fails.append(why)
        self.main.cases_page.reload()
        if self.main.cases_page.cid:
            self.main.cases_page.load_docs()
            self.main.cases_page.load_events()
        if hasattr(self.main, "overview") and getattr(self.main, "mode_cid", None):
            self.main.overview.set_case(self.main.mode_cid)
        self.load()
        if fails:
            QMessageBox.information(self, "Корзина", "\n".join(fails))
        elif done:
            self.main.toast(f"↩  Восстановлено: {done}")

    def forget(self):
        ids = self._ids()
        if not ids:
            return
        if QMessageBox.question(self, "Корзина", f"Удалить навсегда ({len(ids)})? Вернуть будет нельзя.") \
                != QMessageBox.Yes:
            return
        for tid in ids:
            db().trash_forget(tid)
        self.load()

    def empty_all(self):
        if QMessageBox.question(self, "Корзина", "Очистить корзину? Всё, что в ней, будет удалено навсегда.") \
                != QMessageBox.Yes:
            return
        db().trash_empty()
        self.load()


class TemplateDialog(QDialog):
    """Библиотека шаблонов + заполнение полями дела и реквизитами."""

    def __init__(self, main, cid=None, preselect=None):
        super().__init__(main)
        self.main = main
        self.setWindowTitle("Шаблоны документов")
        self.resize(1040, 700)
        ensure_sample_templates()
        root = QHBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)
        # ---- слева: библиотека
        left = QWidget()
        left.setObjectName("sidepanel")
        left.setFixedWidth(360)
        lv = QVBoxLayout(left)
        lv.setContentsMargins(16, 18, 12, 14)
        lv.addLayout(title_row("Шаблоны", "template", big=True))
        self.search = QLineEdit()
        self.search.setPlaceholderText("Найти шаблон: ознакомление, отложение…")
        self.search.textChanged.connect(self.filter)
        lv.addWidget(self.search)
        from PySide6.QtWidgets import QTreeWidget, QTreeWidgetItem
        self._TI = QTreeWidgetItem
        self.tree = QTreeWidget()
        self.tree.setHeaderHidden(True)
        self.tree.setObjectName("tpltree")
        self.tree.currentItemChanged.connect(self.on_pick)
        lv.addWidget(self.tree, 1)
        row = QHBoxLayout()
        b_add = QPushButton("+ Свой…")
        b_add.setToolTip("Добавить документ Word с полями в фигурных скобках, например {Суд}, {Номер_дела}")
        b_add.clicked.connect(self.add_own)
        b_dir = QPushButton("Папка")
        b_dir.clicked.connect(lambda: QDesktopServices.openUrl(QUrl.fromLocalFile(templates_dir())))
        self.b_word = QPushButton("Изменить в Word")
        self.b_word.setToolTip("Открыть выбранный шаблон в Word, чтобы поправить текст")
        self.b_word.clicked.connect(self.edit_in_word)
        for b in (b_add, b_dir, self.b_word):
            b.setObjectName("compact")
            row.addWidget(b)
        lv.addLayout(row)
        root.addWidget(left)
        # ---- справа: заполнение
        right = QWidget()
        rv = QVBoxLayout(right)
        rv.setContentsMargins(22, 18, 22, 16)
        self.t_name = QLabel("Выберите шаблон слева")
        self.t_name.setObjectName("subtitle")
        self.t_name.setWordWrap(True)
        rv.addWidget(self.t_name)
        top = QFormLayout()
        self.case = QComboBox()
        self.case.addItem("— без дела —", None)
        for c in db().cases():
            self.case.addItem(c["title"] + (f"  ({c['number']})" if c["number"] else ""), c["id"])
        cid = cid or main.last_case
        if cid:
            self.case.setCurrentIndex(max(0, self.case.findData(cid)))
        self.case.currentIndexChanged.connect(lambda *_: self.build_form())
        pr = QHBoxLayout()
        pr.addWidget(self.case, 1)
        b_prof = QPushButton("Мои реквизиты…")
        b_prof.clicked.connect(self.edit_profile)
        pr.addWidget(b_prof)
        top.addRow("Дело", pr)
        rv.addLayout(top)
        self.note = QLabel()
        self.note.setObjectName("hint")
        self.note.setWordWrap(True)
        rv.addWidget(self.note)
        self.form_box = QWidget()
        self.form = QFormLayout(self.form_box)
        self.form.setLabelAlignment(Qt.AlignRight | Qt.AlignVCenter)
        sc = QScrollArea()
        sc.setWidget(self.form_box)
        sc.setWidgetResizable(True)
        sc.setFrameShape(QFrame.NoFrame)
        rv.addWidget(sc, 1)
        where = QLabel("Документ Word сохранится в папку дела и сразу добавится справа, в «PDF дела».")
        where.setObjectName("hint")
        where.setWordWrap(True)
        rv.addWidget(where)
        self.c_word = QCheckBox("Открыть в Word, чтобы дописать (правки сами появятся в PDF дела)")
        self.c_word.setChecked(M.settings().value("tpl/word", "true") == "true")
        rv.addWidget(self.c_word)
        bb = QHBoxLayout()
        bb.addStretch(1)
        bc = QPushButton("Закрыть")
        bc.clicked.connect(self.reject)
        self.b_make = QPushButton("Создать документ")
        self.b_make.setObjectName("primary")
        self.b_make.clicked.connect(self.make)
        bb.addWidget(bc)
        bb.addWidget(self.b_make)
        rv.addLayout(bb)
        root.addWidget(right, 1)
        self.edits = {}
        self.path = None
        self.load_templates(preselect)

    # -------------------------------------------------------------- список
    def load_templates(self, select=None):
        self.tree.clear()
        base = Path(templates_dir())
        groups = {}
        for p in sorted(base.rglob("*.docx")):
            if p.name.startswith("~$"):
                continue
            rel = p.parent.relative_to(base)
            parts = [x for x in rel.parts if x != "Встроенные"]
            cat = " / ".join(parts) if parts else "Мои шаблоны"
            groups.setdefault(cat, []).append(p)
        order = sorted(groups, key=lambda c: (c.startswith("Мои"), c))
        cid = self.case.currentData() if hasattr(self, "case") else None
        court = ((db().case(cid) or {}).get("court", "") if cid else "").lower()
        prefer = "АПК" if "арбитраж" in court else ("ГПК" if court else "")
        if select:
            hits = [p for c in order for p in groups[c] if select.lower() in p.stem.lower()]
            best = next((p for p in hits if prefer and prefer in str(p.parent)), hits[0] if hits else None)
            select = str(best) if best else None
        first = None
        for cat in order:
            top = self._TI([cat])
            f = top.font(0)
            f.setBold(True)
            top.setFont(0, f)
            top.setFlags(Qt.ItemIsEnabled)
            self.tree.addTopLevelItem(top)
            for p in groups[cat]:
                it = self._TI([p.stem])
                it.setData(0, Qt.UserRole, str(p))
                it.setToolTip(0, str(p))
                top.addChild(it)
                if first is None or (select and str(p) == select):
                    first = it
            top.setExpanded(True)
        if first is not None:
            self.tree.setCurrentItem(first)
        self.b_word.setEnabled(first is not None)

    def filter(self, text):
        t = text.lower().strip()
        for i in range(self.tree.topLevelItemCount()):
            top = self.tree.topLevelItem(i)
            vis = 0
            for j in range(top.childCount()):
                ch = top.child(j)
                show = not t or t in ch.text(0).lower()
                ch.setHidden(not show)
                vis += show
            top.setHidden(vis == 0)

    def on_pick(self, cur, _prev=None):
        p = cur.data(0, Qt.UserRole) if cur else None
        if not p:
            return
        self.path = p
        self.t_name.setText(Path(p).stem)
        self.build_form()

    # -------------------------------------------------------------- поля
    def values_known(self):
        cid = self.case.currentData()
        vals = {}
        try:
            vals.update(json.loads(M.settings().value(f"tpl_vals/{cid or 0}", "{}") or "{}"))
        except Exception:
            pass
        if cid:
            vals.update({k: v for k, v in db().template_values(cid).items() if v})
        else:
            vals["Дата"] = dt.date.today().strftime("%d.%m.%Y")
        vals.update({k: v for k, v in profile_values().items() if v})
        vals.setdefault("Статус_доверителя", "Истец")
        return vals

    def build_form(self):
        while self.form.rowCount():
            self.form.removeRow(0)
        self.edits = {}
        if not self.path:
            return
        try:
            names = L.template_fields(self.path)
        except Exception as e:
            self.note.setText(f"Не удалось прочитать шаблон: {e}")
            return
        vals = self.values_known()
        empty = 0
        import templates_lib as TL
        self.parties = {}
        roles = [n[len("Сторона_"):] for n in names if n.startswith("Сторона_")]
        for role in roles:
            data = self.party_store.get(role) or guess_party(self._party_default(role, vals))
            pe = PartyEditor(role, data)
            self.form.addRow(pe)
            self.parties[role] = pe
        for n in names:
            if n.startswith("Сторона_") or n in roles:
                continue                               # краткое имя стороны берётся из её карточки
            if n.startswith(("Пункты_", "Блок_")):
                e = QPlainTextEdit(vals.get(n, ""))
                e.setPlaceholderText(TL.FIELD_HINTS.get(n, "Каждый пункт — с новой строки"))
                e.setFixedHeight(92 if n.startswith("Пункты_") else 78)
                e.text = e.toPlainText                  # единый способ прочитать значение
            else:
                e = M.GrowEdit(vals.get(n, ""), TL.FIELD_HINTS.get(n, ""))
            if not e.text():
                empty += 1
            self.form.addRow(n.replace("Пункты_", "").replace("Блок_", "").replace("_", " "), e)
            self.edits[n] = e
        cid = self.case.currentData()
        self.note.setText((f"Заполнено из дела и реквизитов: {len(names) - empty} из {len(names)}. "
                           if names else "В шаблоне нет полей. ") +
                          ("Пустые поля останутся в документе как есть — их можно дописать в Word."
                           if empty else "") + ("" if cid else "  Выберите дело, чтобы подставить его данные."))

    @property
    def party_store(self):
        try:
            store = json.loads(M.settings().value(f"tpl_vals/{self.case.currentData() or 0}", "{}") or "{}")
        except Exception:
            store = {}
        return {k[6:]: v for k, v in store.items() if k.startswith("party:") and isinstance(v, dict)}

    def _party_default(self, role, vals):
        """Кто в этой роли по карточке дела: доверитель — по «Статусу доверителя», иначе оппонент / третьи лица."""
        status = (vals.get("Статус_доверителя") or "Истец").lower()
        client, opp = vals.get("Доверитель", ""), vals.get("Оппонент", "")
        if role in ("Заявитель", "Заявитель_жалобы"):
            return client
        if role == "Истец":
            return client if status.startswith(("истец", "заявител")) else opp
        if role == "Ответчик":
            return opp if status.startswith(("истец", "заявител")) else client
        if role == "Третье_лицо":
            return vals.get("Третьи_лица", "")
        return ""

    def edit_profile(self):
        if ProfileDialog(self).exec():
            self.build_form()

    # -------------------------------------------------------------- действия
    def add_own(self):
        p, _ = QFileDialog.getOpenFileName(self, "Шаблон Word", str(Path.home()), "Word (*.docx)")
        if not p:
            return
        dst = Path(templates_dir()) / "Мои шаблоны" / Path(p).name
        dst.parent.mkdir(parents=True, exist_ok=True)
        import shutil
        shutil.copy2(p, dst)
        self.load_templates(dst.stem)
        QMessageBox.information(self, M.APP_NAME, "Шаблон добавлен в «Мои шаблоны».\n\nПоля пишутся в фигурных скобках: "
                                "{Суд}, {Номер_дела}, {Судья}, {Доверитель}, {Оппонент}, {Представитель}, {Дата} — они "
                                "заполнятся из дела автоматически. Любое другое {Поле} появится в окне для заполнения.")

    def edit_in_word(self):
        if self.path:
            QDesktopServices.openUrl(QUrl.fromLocalFile(self.path))

    def make(self):
        if not self.path:
            return
        cid = self.case.currentData()
        vals = {k: e.text().strip() for k, e in self.edits.items()}
        blocks = {k: v for k, v in vals.items() if k.startswith(L.BLOCK_PREFIXES)}   # пустые блоки — убрать абзац
        for role, pe in getattr(self, "parties", {}).items():
            blocks[f"Сторона_{role}"] = pe.block()
            vals[role] = pe.names()[1]
        # запомнить введённое вручную для этого дела
        known = set(profile_values()) | (set(db().template_values(cid)) if cid else set())
        store = {}
        try:
            store = json.loads(M.settings().value(f"tpl_vals/{cid or 0}", "{}") or "{}")
        except Exception:
            pass
        store.update({k: v for k, v in vals.items() if v and k not in known})
        for role, pe in getattr(self, "parties", {}).items():
            store[f"party:{role}"] = pe.data()
        M.settings().setValue(f"tpl_vals/{cid or 0}", json.dumps(store, ensure_ascii=False))
        M.settings().setValue("tpl/word", "true" if self.c_word.isChecked() else "false")
        out_dir = case_output_dir(cid)
        base = f"{Path(self.path).stem} {dt.date.today().strftime('%d.%m.%Y')}"
        out = os.path.join(out_dir, L.clean_filename(base) + ".docx")
        k = 2
        while os.path.exists(out):
            out = os.path.join(out_dir, L.clean_filename(f"{base} ({k})") + ".docx")
            k += 1
        try:
            L.fill_template(self.path, {**{k: v for k, v in vals.items() if v}, **blocks}, out)
        except Exception as e:
            return self.main.error("Не удалось создать документ", e)
        if cid:
            db().add_doc(cid, out)
            self.main.last_case = cid
        self.accept()
        self.main.add_created_doc(cid, out, open_word=self.c_word.isChecked())


# =============================================================================
#  Выписка
# =============================================================================
def save_quote(main, text, page_no):
    if not text:
        QMessageBox.information(main, M.APP_NAME, "В выделенной области нет текста. Если это скан — сначала "
                                                  "распознайте его (OCR).")
        return
    dlg = QDialog(main)
    dlg.setWindowTitle("Выписка")
    dlg.resize(560, 360)
    v = QVBoxLayout(dlg)
    ed = QPlainTextEdit(text)
    v.addWidget(ed, 1)
    f = QFormLayout()
    src = QLineEdit(main.base_name() or "Документ")
    pg = QSpinBox()
    pg.setRange(0, 100000)
    pg.setValue(page_no)
    f.addRow("Источник", src)
    f.addRow("Лист / стр.", pg)
    v.addLayout(f)
    bb = QHBoxLayout()
    b_copy = QPushButton("Только копировать")
    b_save = QPushButton("Сохранить в дело…")
    b_save.setObjectName("primary")
    bb.addWidget(b_copy)
    bb.addStretch(1)
    bb.addWidget(b_save)
    v.addLayout(bb)

    def copy():
        QApplication.clipboard().setText(f"«{ed.toPlainText().strip()}» ({src.text()}, л. {pg.value()})")
        main.msg("Цитата скопирована")
        dlg.accept()

    def save():
        cid = main.last_case or pick_case(dlg)
        if not cid:
            return
        db().add_quote(cid, ed.toPlainText().strip(), src.text(), pg.value())
        main.last_case = cid
        main.refresh_cases()
        main.msg("Выписка сохранена в дело")
        dlg.accept()
    b_copy.clicked.connect(copy)
    b_save.clicked.connect(save)
    dlg.exec()


# =============================================================================
#  Инструменты главного окна (вызываются как tool_<ключ>(main))
# =============================================================================
def tool_package(main):
    PackageDialog(main).exec()


def tool_f107(main):
    F107Dialog(main).exec()


def tool_preflight(main):
    if main.need_doc():
        PreflightDialog(main).exec()


def tool_anonymize(main):
    if main.need_doc():
        AnonymizeDialog(main).exec()


def tool_case_search(main):
    paths = []
    if main.cases_page and main.cases_page.cid:
        paths = [d["path"] for d in db().docs(main.cases_page.cid)]
    SearchDialog(main, paths).exec()


def tool_template(main, cid=None, preselect=None):
    TemplateDialog(main, cid or (main.cases_page.cid if main.cases_page else None), preselect).exec()


def tool_profile(main):
    ProfileDialog(main).exec()


def tool_quote(main):
    if main.need_doc():
        main.open_editor(mode="quote")


def tool_compare_ed(main):
    v = M.OptionsDialog.ask(main, "Сравнение редакций", [
        ("old", "Было (старая редакция)", "file", main.path or "", "PDF, Word (*.pdf *.docx *.txt)"),
        ("new", "Стало (новая редакция)", "file", "", "PDF, Word (*.pdf *.docx *.txt)"),
        ("only", "Показывать только изменённые абзацы", "check", False),
    ], ok_text="Сравнить", help="compare_ed")
    if not v or not v["old"] or not v["new"]:
        return
    base = os.path.join(os.path.dirname(v["new"]), f"Сравнение — {Path(v['new']).stem}")
    res = main.run("Сравнение…", lambda progress=None: L.compare_versions(v["old"], v["new"], base, v["only"]))
    if res is M.FAILED:
        return
    hp, pp, st = res
    main.done(f"Изменённых фрагментов: {st['changed']}, добавлено слов: {st['added']}, удалено: {st['deleted']}.\n"
              f"Отчёт: {hp}", pp or hp)
    QDesktopServices.openUrl(QUrl.fromLocalFile(hp))


def tool_sheetnum(main):
    if not main.need_doc():
        return
    v = M.OptionsDialog.ask(main, "Нумерация листов дела", [
        ("start", "Первый номер", "int", 1, (1, 100000)),
        ("per", "Листов в томе (0 — без томов)", "int", 250, (0, 5000)),
        ("pos", "Где ставить номер", "combo", None, L.SHEET_POSITIONS),
        ("size", "Размер цифр", "int", 12, (6, 36)),
        ("duplex", "Двусторонние листы (номер только на лицевой стороне)", "check", False),
        ("restart", "В каждом томе нумерация с 1", "check", True),
        ("split", "Сохранить тома отдельными файлами", "check", False),
        ("cover", "Добавить титульный лист тома", "check", False),
        ("title", "Надпись на титуле", "text", "Дело № ", None),
    ], ok_text="Пронумеровать", help="sheetnum",
        depends={"restart": ("per", lambda t: t not in ("0", "")), "cover": ("split", lambda c: c),
                 "title": ("split", lambda c: c)})
    if not v:
        return
    main.push_undo()
    vols = main.run("Нумерация листов…", L.number_sheets, main.doc, v["start"], v["per"], v["pos"], v["size"], 8,
                    v["duplex"], v["restart"])
    if vols is M.FAILED:
        return
    main.after_change()
    text = f"Пронумеровано. Томов: {len(vols)}.\n" + "\n".join(
        f"Том {i}: стр. {a + 1}–{b + 1}, листы {s1}–{s2}" for i, (a, b, s1, s2) in enumerate(vols, 1))
    if v["split"] and len(vols) >= 1:
        folder = QFileDialog.getExistingDirectory(main, "Папка для томов", main.default_dir())
        if folder:
            for i, (a, b, s1, s2) in enumerate(vols, 1):
                d = fitz.open()
                if v["cover"]:
                    d.insert_pdf(L.volume_cover(v["title"], i, s1, s2))
                d.insert_pdf(main.doc, from_page=a, to_page=b)
                C.save_pdf(d, os.path.join(folder, f"{main.base_name()} — том {i}.pdf"))
            main.done(text + f"\n\nТома сохранены в папку:\n{folder}", folder, open_file=False)
            return
    main.done(text)


def _stamp_templates():
    try:
        return json.loads(M.settings().value("stamp_templates", "") or "{}")
    except Exception:
        return {}


def tool_certify(main):
    if not main.need_doc():
        return
    tpls = _stamp_templates()
    default = "Копия верна\nАдвокат\n____________ Фамилия И.О.\n" + dt.date.today().strftime("%d.%m.%Y")
    names = ["— новый —"] + list(tpls)
    last = M.settings().value("stamp_last", "")
    v = M.OptionsDialog.ask(main, "Штамп заверения", [
        ("tpl", "Шаблон", "combo", last if last in tpls else names[0], names),
        ("lines", "Текст штампа", "multiline", tpls.get(last, default)),
        ("mode", "Где", "combo", None, L.STAMP_MODES),
        ("pos", "Положение", "combo", None, L.STAMP_POSITIONS),
        ("size", "Размер шрифта", "int", 10, (6, 20)),
        ("width", "Ширина штампа, мм", "int", 75, (30, 180)),
        ("frame", "Рамка", "check", True),
        ("color", "Цвет", "combo", None, ["Синий", "Чёрный"]),
        ("sig", "Подпись (PNG, необязательно)", "file", "", "Изображения (*.png *.jpg)"),
        ("bound", "На последней странице — «Прошито, пронумеровано N листов»", "check", False),
        ("save", "Сохранить как шаблон с именем", "text", ""),
    ], ok_text="Поставить штамп", help="certify")
    if not v:
        return
    if v["tpl"] != names[0] and v["lines"] == default:
        v["lines"] = tpls.get(v["tpl"], v["lines"])
    if v["save"].strip():
        tpls[v["save"].strip()] = v["lines"]
        M.settings().setValue("stamp_templates", json.dumps(tpls, ensure_ascii=False))
        M.settings().setValue("stamp_last", v["save"].strip())
    pages = L.stamp_pages(main.doc.page_count, v["mode"])
    sig = None
    if v["sig"]:
        try:
            sig = C.image_bytes(v["sig"])[0]
        except Exception as e:
            return main.error("Не удалось открыть картинку", e)
    color = (0.05, 0.1, 0.55) if v["color"] == "Синий" else (0, 0, 0)
    main.push_undo()

    def work(progress=None):
        L.certify_stamp(main.doc, pages, v["lines"].splitlines(), v["pos"], v["size"], v["width"], 10, v["frame"],
                        color, sig, progress)
        if v["bound"]:
            last_p = main.doc.page_count - 1
            lines = L.bound_stamp_lines(main.doc.page_count, "", "")
            pos = "Внизу слева" if "справа" in v["pos"] else "Внизу справа"
            L.certify_stamp(main.doc, [last_p], lines, pos, v["size"], 70, 10, True, color)
    if main.run("Штамп…", work) is M.FAILED:
        return
    main.after_change()
    main.msg(f"Штамп поставлен на {len(pages)} стр.")


# =============================================================================
#  Напоминания
# =============================================================================
class Reminders:
    def __init__(self, main):
        self.main = main
        self.tray = None
        if QSystemTrayIcon.isSystemTrayAvailable():
            self.tray = QSystemTrayIcon(main.windowIcon(), main)
            self.tray.setToolTip(M.APP_NAME)
            self.tray.messageClicked.connect(lambda: main.show_cases())
            self.tray.show()
        self.timer = QTimer(main)
        self.timer.setInterval(5 * 60 * 1000)
        self.timer.timeout.connect(self.check)
        self.timer.start()
        QTimer.singleShot(1500, lambda: self.check(startup=True))
        # напоминания — в точное время: проверка каждые 20 секунд
        self.popups = {}
        self.exact = QTimer(main)
        self.exact.setInterval(20 * 1000)
        self.exact.timeout.connect(self.check_exact)
        self.exact.start()
        QTimer.singleShot(2500, self.check_exact)

    def check_exact(self):
        """Показать напоминания, время которых наступило (и пропущенные, пока программа была закрыта)."""
        try:
            due = [e for e in db().reminders() if not e["notified"] and event_dt(e) <= now()]
        except Exception:
            return
        for e in due[:5]:
            if e["id"] in self.popups:
                continue
            db().update_event(e["id"], notified=1)
            p = ReminderPopup(self.main, e)
            self.popups[e["id"]] = p
            p.destroyed.connect(lambda *_, i=e["id"]: self.popups.pop(i, None))
            p.show()
            p.raise_()
            if self.tray:
                self.tray.showMessage("Напоминание", e["title"], QSystemTrayIcon.Information, 20000)
        if due:
            QApplication.beep()
            QApplication.alert(self.main)
            try:
                self.main.overview.load_reminders()
            except Exception:
                pass

    def check(self, startup=False):
        try:
            events = [e for e in db().events(upcoming_days=1, include_done=False) if e["kind"] != CS.CaseDB.REMINDER]
        except Exception:
            return
        today = now().date()
        due = [e for e in events if dt.date.fromisoformat(e["date"]) <= today + dt.timedelta(days=1)]
        if not due:
            self.main.set_banner("")
            return
        overdue = [e for e in due if dt.date.fromisoformat(e["date"]) < today]
        todays = [e for e in due if dt.date.fromisoformat(e["date"]) == today]
        tomorrow = [e for e in due if dt.date.fromisoformat(e["date"]) > today]
        first = (todays or tomorrow or overdue)[0]
        when = "сегодня" if dt.date.fromisoformat(first["date"]) == today else (
            "завтра" if dt.date.fromisoformat(first["date"]) > today else "просрочено")
        more = len(due) - 1
        self.main.set_banner(f"⚠  {when.capitalize()}{(' в ' + first['time']) if first['time'] else ''}: "
                             f"{first['kind'].lower()} «{first['title']}» · {first['case_title'] or 'без дела'}"
                             f"{f'  (и ещё {more})' if more > 0 else ''}   →",
                             first.get("case_id"))
        new = [e for e in due if not e["notified"]]
        if new and self.tray:
            text = "\n".join(f"{CS.ru(e['date'])} {e['time']} {e['kind']}: {e['title']} ({e['case_title'] or ''})"
                             for e in new[:5])
            self.tray.showMessage("Сроки и заседания", text, QSystemTrayIcon.Information, 15000)
            for e in new:
                db().update_event(e["id"], notified=1)


# =============================================================================
#  Сохранность данных: папка дела, перенос, резервные копии
# =============================================================================
def sync_case_file(cid):
    """Обновить LegalHelper-дело.json в папке дела. Ошибки — только в журнал: это страховка, а не работа."""
    if not cid:
        return
    try:
        if db().case(cid):
            CF.save_case_file(db(), cid)
    except Exception as e:
        M.log_error("Файл сведений в папке дела", e)


def copy_docs_enabled():
    return str(M.settings().value("copy_docs", "0")) == "1"      # по умолчанию файлы не копируются


def copy_into_case(cid, path):
    """Файл вне папки дела — копия в «<папка>/Документы», иначе тот же путь."""
    try:
        folder = CF.ensure_folder(db(), cid)
        if "rel" in CF._rel(path, folder) or not os.path.isfile(path):
            return path
        dest = os.path.join(folder, CF.DOCS_SUBDIR)
        os.makedirs(dest, exist_ok=True)
        stem, ext = os.path.splitext(os.path.basename(path))
        new, n = os.path.join(dest, stem + ext), 2
        while os.path.exists(new):
            new, n = os.path.join(dest, f"{stem} ({n}){ext}"), n + 1
        import shutil
        shutil.copy2(path, new)
        return new
    except Exception as e:
        M.log_error("Копирование файла в папку дела", e)
        return path


def open_case_folder(cid):
    if cid:
        sync_case_file(cid)
        QDesktopServices.openUrl(QUrl.fromLocalFile(CF.ensure_folder(db(), cid)))


def collect_case_files(main, cid):
    if not cid:
        return
    c = db().case(cid)
    folder = CF.ensure_folder(db(), cid)
    if QMessageBox.question(main, M.APP_NAME,
                            f"Скопировать все файлы дела «{c['title']}» в его папку?\n\n{folder}\n\n"
                            "Оригиналы останутся на месте, а дело будет работать с копиями в папке. После этого "
                            "папку можно целиком перенести на другой компьютер (флешкой, облаком) и открыть там: "
                            "«☰ → Файл → Открыть дело из папки…».") != QMessageBox.Yes:
        return
    try:
        copied, missing = CF.collect_files(db(), cid)
    except Exception as e:
        return main.error("Не удалось собрать файлы дела", e)
    main.cases_page.refresh_docs_if(cid) if hasattr(main.cases_page, "refresh_docs_if") else None
    text = f"Готово. Скопировано файлов: {copied}."
    if missing:
        text += "\n\nНе найдены (возможно, перемещены или удалены):\n" + "\n".join(missing[:10])
    box = QMessageBox(main)
    box.setWindowTitle(M.APP_NAME)
    box.setText(text)
    op = box.addButton("Открыть папку", QMessageBox.AcceptRole)
    box.addButton("OK", QMessageBox.RejectRole)
    box.exec()
    if box.clickedButton() is op:
        QDesktopServices.openUrl(QUrl.fromLocalFile(folder))


def open_case_from_folder(main):
    folder = QFileDialog.getExistingDirectory(main, "Папка дела (с файлом «LegalHelper-дело.json»)",
                                              CF.cases_root() if os.path.isdir(CF.cases_root()) else str(Path.home()))
    if not folder:
        return
    try:
        data = CF.read_case_file(folder)
    except Exception as e:
        QMessageBox.warning(main, M.APP_NAME, str(e))
        return
    title = data["case"].get("title", "")
    existing = CF.find_by_uid(db(), data["case"].get("uid"))
    replace = None
    if existing:
        box = QMessageBox(main)
        box.setWindowTitle(M.APP_NAME)
        box.setText(f"Дело «{title}» уже есть в программе.")
        box.setInformativeText("Заменить его сведениями из папки (например, более свежими с другого компьютера) "
                               "или просто открыть то, что уже есть? Перед заменой делается резервная копия.")
        b_rep = box.addButton("Заменить из папки", QMessageBox.AcceptRole)
        b_open = box.addButton("Открыть имеющееся", QMessageBox.RejectRole)
        box.addButton("Отмена", QMessageBox.RejectRole)
        box.exec()
        if box.clickedButton() is b_open:
            main.enter_case(existing)
            main.open_case_tab("overview")
            return
        if box.clickedButton() is not b_rep:
            return
        replace = existing
    try:
        main.cases_page.flush()
        BK.make_backup(M.data_dir(), "import", M.APP_VERSION)
        cid = CF.import_case(db(), folder, replace_cid=replace)
    except Exception as e:
        return main.error("Не удалось открыть дело из папки", e)
    main.cases_page.show_arch.setChecked(False)
    main.cases_page.reload()
    main.cases_page.on_select(None)
    main.enter_case(cid)
    main.open_case_tab("overview")
    missing = [d["path"] for d in db().docs(cid) if not os.path.exists(d["path"])]
    msg = f"Дело «{title}» открыто."
    if missing:
        msg += (f"\n\nФайлов не нашлось: {len(missing)} — они лежали вне папки дела. Перед переносом используйте "
                "«Документы → Ещё → Собрать все файлы в папку дела».")
    QMessageBox.information(main, M.APP_NAME, msg)


class BackupsDialog(QDialog):
    """Резервные копии: список, «сделать сейчас», «восстановить»."""

    def __init__(self, main):
        super().__init__(main)
        self.main = main
        self.setWindowTitle("Резервные копии")
        self.resize(720, 480)
        v = QVBoxLayout(self)
        v.addLayout(title_row("Резервные копии", big=True))
        note = QLabel("Программа сама сохраняет копию дел, настроек и своих шаблонов раз в день, перед каждым "
                      "обновлением, удалением дела и загрузкой дела из папки. Хранятся последние "
                      f"{BK.KEEP} копий в папке «Документы\\LegalHelper\\Резервные копии». Файлы документов лежат в "
                      "папках дел — их копия делается вместе с папкой дела.")
        note.setObjectName("note")
        note.setWordWrap(True)
        v.addWidget(note)
        self.t = QTableWidget(0, 4)
        self.t.setHorizontalHeaderLabels(["Когда", "Почему", "Дел", "Версия"])
        self.t.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
        self.t.setColumnWidth(0, 170)
        self.t.setColumnWidth(2, 60)
        self.t.setColumnWidth(3, 80)
        self.t.verticalHeader().hide()
        self.t.setShowGrid(False)
        self.t.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.t.setSelectionMode(QAbstractItemView.SingleSelection)
        self.t.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.t.itemSelectionChanged.connect(lambda: self.b_restore.setEnabled(bool(self.t.selectedItems())))
        v.addWidget(self.t, 1)
        r = QHBoxLayout()
        b = QPushButton("Сделать копию сейчас")
        b.clicked.connect(self.make_now)
        r.addWidget(b)
        b = QPushButton("Открыть папку с копиями")
        b.clicked.connect(lambda: QDesktopServices.openUrl(QUrl.fromLocalFile(BK.backups_dir())))
        r.addWidget(b)
        r.addStretch(1)
        self.b_restore = QPushButton("Восстановить выбранную…")
        self.b_restore.setObjectName("primary")
        self.b_restore.setEnabled(False)
        self.b_restore.clicked.connect(self.restore)
        r.addWidget(self.b_restore)
        v.addLayout(r)
        self.load()

    def load(self):
        self.items = BK.list_backups()
        self.t.setRowCount(0)
        for b in self.items:
            r = self.t.rowCount()
            self.t.insertRow(r)
            for c, text in enumerate((b["created"].strftime("%d.%m.%Y  %H:%M"), b["reason_text"],
                                      "" if b["cases"] is None else str(b["cases"]), b["version"])):
                self.t.setItem(r, c, QTableWidgetItem(text))
        if not self.items:
            self.t.insertRow(0)
            self.t.setItem(0, 1, QTableWidgetItem("Копий пока нет — нажмите «Сделать копию сейчас»"))

    def make_now(self):
        try:
            self.main.cases_page.flush()
            BK.make_backup(M.data_dir(), "manual", M.APP_VERSION)
        except Exception as e:
            return self.main.error("Не удалось сделать резервную копию", e)
        self.load()
        self.main.statusBar().showMessage("Резервная копия сохранена", 6000)

    def restore(self):
        rows = sorted({i.row() for i in self.t.selectedItems()})
        if not rows or rows[0] >= len(self.items):
            return
        b = self.items[rows[0]]
        if QMessageBox.question(self, M.APP_NAME, f"Вернуть дела, настройки и шаблоны к состоянию на "
                                f"{b['created']:%d.%m.%Y %H:%M}?\n\nВсё, что изменено после этого, пропадёт из "
                                "программы, но текущее состояние сначала сохранится отдельной копией («перед "
                                "восстановлением») — его тоже можно будет вернуть.\n\nПрограмма перезапустится."
                                ) != QMessageBox.Yes:
            return
        self.main.restore_backup(b["path"])


# =============================================================================
#  Главная (HUB): вся работа юриста на одном экране
# =============================================================================
WEEKDAYS = ["понедельник", "вторник", "среда", "четверг", "пятница", "суббота", "воскресенье"]
MONTHS_G = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа", "сентября", "октября",
            "ноября", "декабря"]
KIND_ICON = {"Заседание": "⚖️", "Срок": "⏳", "Задача": "✅", "Встреча": "🤝", "Напоминание": "⏰"}


class StatTile(QPushButton):
    def __init__(self, icon, caption, slot):
        super().__init__()
        self.setObjectName("actioncard")
        self.setCursor(Qt.PointingHandCursor)
        self.setMinimumHeight(86)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        h = QHBoxLayout(self)
        h.setContentsMargins(16, 10, 14, 10)
        ic = QLabel(icon)
        ic.setObjectName("cardicon")
        f = ic.font()
        f.setPointSize(17)
        ic.setFont(f)
        ic.setFixedSize(42, 42)
        ic.setAlignment(Qt.AlignCenter)
        ic.setAttribute(Qt.WA_TransparentForMouseEvents)
        h.addWidget(ic)
        v = QVBoxLayout()
        v.setSpacing(0)
        self.value = QLabel("—")
        self.value.setObjectName("bigresult")
        self.value.setAttribute(Qt.WA_TransparentForMouseEvents)
        cap = QLabel(caption)
        cap.setObjectName("carddesc")
        cap.setWordWrap(True)
        cap.setAttribute(Qt.WA_TransparentForMouseEvents)
        v.addWidget(self.value)
        v.addWidget(cap)
        h.addLayout(v, 1)
        self.clicked.connect(slot)
        anim.hover_lift(self)

    def set(self, text, color=None):
        """Число в карточке «отсчитывается» до нового значения (если оно изменилось)."""
        self.value.setStyleSheet(f"color: {color};" if color else "")
        m = re.match(r"^(\d(?:[\d  ]*\d)?)(.*)$", text)
        if not m:
            self.value.setText(text)
            self._num = None
            return
        new = int(re.sub(r"\D", "", m.group(1)) or 0)
        tail = m.group(2)
        sep = " " if " " in m.group(1) else " "

        def render(n):
            return f"{n:,}".replace(",", sep) + tail
        if not self.value.isVisible() and not getattr(self, "_seen", False):
            self.value.setText(render(new))         # Главную ещё не показывали — отсчёт будет при показе
            return
        self._seen = True
        prev = getattr(self, "_num", None)
        self._num = new
        anim.count_up(self.value, 0 if prev is None else prev, new, render)


def greeting(h):
    """Приветствие по времени суток: утро 5–11, день 12–17, вечер 18–22, ночь 23–4."""
    if 5 <= h < 12:
        return "Доброе утро"
    if 12 <= h < 18:
        return "Добрый день"
    if 18 <= h < 23:
        return "Добрый вечер"
    return "Доброй ночи"


class HomePage(QWidget):
    """Главная: что сегодня и на неделе, горящие сроки, напоминания, недавние дела, быстрые действия."""

    def __init__(self, main):
        super().__init__()
        self.main = main
        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        sc = QScrollArea()
        sc.setWidgetResizable(True)
        sc.setFrameShape(QFrame.NoFrame)
        outer.addWidget(sc)
        body = QWidget()
        sc.setWidget(body)
        v = QVBoxLayout(body)
        v.setContentsMargins(26, 20, 26, 20)
        v.setSpacing(14)
        self.hello = QLabel()
        self.hello.setObjectName("title")
        v.addWidget(self.hello)
        self._clock = QTimer(self)                  # приветствие и «Сегодня …» меняются, пока окно открыто
        self._clock.setInterval(60 * 1000)
        self._clock.timeout.connect(lambda: self.isVisible() and self.refresh())
        self._clock.start()
        self.today_lbl = QLabel()
        self.today_lbl.setObjectName("hint")
        v.addWidget(self.today_lbl)
        tiles = QHBoxLayout()
        tiles.setSpacing(12)
        m = main
        self.t_cases = StatTile("📁", "дел в работе", lambda: m.show_cases())
        self.t_today = StatTile("📅", "событий сегодня", lambda: self.focus_list())
        self.t_over = StatTile("🔥", "просрочено", lambda: self.focus_list())
        self.t_money = StatTile("💰", "к оплате доверителями", lambda: m.show_cases())
        for t in (self.t_cases, self.t_today, self.t_over, self.t_money):
            tiles.addWidget(t)
        v.addLayout(tiles)
        cols = QHBoxLayout()
        cols.setSpacing(14)
        self.l_week = self._card(cols, "🗓  Сегодня и ближайшие 7 дней", 3)
        self.l_recent = self._card(cols, "🕘  Недавние дела", 2)
        v.addLayout(cols)
        cols2 = QHBoxLayout()
        cols2.setSpacing(14)
        self.l_over = self._card(cols2, "🔥  Горящие и просроченные", 3)
        self.l_rem = self._card(cols2, "⏰  Напоминания", 2)
        v.addLayout(cols2)
        lab = QLabel("Быстрые действия")
        lab.setObjectName("subtitle")
        v.addWidget(lab)
        v.addLayout(card_grid([
            ActionCard("➕", "Новое дело", "Создать дело: папка, сроки, документы", lambda: m.cases_page.new_case()),
            ActionCard("📝", "Документ по шаблону", "Иски, жалобы, ходатайства, выступление",
                       lambda: tool_template(m, m.last_case)),
            ActionCard("📂", "Без дела — просто PDF", "Открыть, собрать, подписать, сжать PDF", lambda: m.enter_loose()),
            ActionCard("⏱️", "Посчитать срок", "Процессуальные сроки с праздниками", lambda: m.show_calc(0)),
            ActionCard("📥", "Открыть дело из папки", "Перенесённое с другого компьютера", lambda: open_case_from_folder(m)),
            ActionCard("🛟", "Резервные копии", "Сделать копию или вернуть данные", lambda: BackupsDialog(m).exec()),
        ]))
        v.addStretch(1)
        self.timer = QTimer(self)
        self.timer.setInterval(5 * 60 * 1000)
        self.timer.timeout.connect(lambda: self.isVisible() and self.refresh())
        self.timer.start()

    def _card(self, row, title, stretch):
        box = QFrame()
        box.setObjectName("card")
        bv = QVBoxLayout(box)
        bv.setContentsMargins(14, 12, 14, 12)
        lab = QLabel(title)
        lab.setObjectName("subtitle")
        bv.addWidget(lab)
        lst = QListWidget()
        lst.setObjectName("overlist")
        lst.setMinimumHeight(200)
        lst.setWordWrap(True)
        lst.itemClicked.connect(self.open_item)
        bv.addWidget(lst, 1)
        row.addWidget(box, stretch)
        return lst

    def focus_list(self):
        self.l_over.setFocus() if self.l_over.count() else self.l_week.setFocus()

    def open_item(self, it):
        cid = it.data(Qt.UserRole)
        if cid:
            self.main.enter_case(cid)
            self.main.open_case_tab("overview")

    @staticmethod
    def _empty(lst, text):
        it = QListWidgetItem(text)
        it.setFlags(Qt.NoItemFlags)
        lst.addItem(it)

    def refresh(self):
        today = now().date()
        name = (profile_values().get("Представитель") or "").split()
        first = name[1] if len(name) > 1 else ""
        self.hello.setText(f"{greeting(now().hour)}{', ' + first if first else ''}!")
        self.today_lbl.setText(f"Сегодня {WEEKDAYS[today.weekday()]}, {today.day} {MONTHS_G[today.month - 1]} "
                               f"{today.year} г.")
        cases = db().cases()
        self.t_cases.set(str(len(cases)))
        events = db().events(include_done=False)
        active = [e for e in events if e["case_id"] is None or any(c["id"] == e["case_id"] for c in cases)]
        todays = [e for e in active if e["date"] == today.isoformat()]
        over = [e for e in active if e["date"] < today.isoformat() and e["kind"] != CS.CaseDB.REMINDER]
        self.t_today.set(str(len(todays)))
        self.t_over.set(str(len(over)), M.T["danger"] if over else None)
        due = sum(max(0.0, db().balance(c["id"])["due"]) for c in cases)
        self.t_money.set(f"{L.money(due, cents=False)} ₽" if due else "0 ₽")

        def line(e, with_date=True):
            d = dt.date.fromisoformat(e["date"])
            left = (d - today).days
            when = "сегодня" if left == 0 else "завтра" if left == 1 else (
                f"просрочено {CS.ru(e['date'])}" if left < 0 else f"{WEEKDAYS[d.weekday()][:2].capitalize()}, {CS.ru(e['date'])}")
            t = f" {e['time']}" if e["time"] else ""
            return (f"{KIND_ICON.get(e['kind'], '•')}  {e['title'] or e['kind']}\n"
                    f"     {when if with_date else ''}{t} · {e['case_title'] or 'без дела'}")
        self.l_week.clear()
        week = [e for e in active if today.isoformat() <= e["date"] <= (today + dt.timedelta(days=7)).isoformat()]
        for e in sorted(week, key=lambda e: (e["date"], e["time"] or "")):
            it = QListWidgetItem(line(e))
            it.setData(Qt.UserRole, e["case_id"])
            if e["date"] == today.isoformat():
                it.setForeground(QColor(M.T["accent"]))
            self.l_week.addItem(it)
        if not week:
            self._empty(self.l_week, "На неделе ничего не запланировано 🎉")
        self.l_over.clear()
        soon = [e for e in active if e["kind"] in ("Срок", "Заседание") and
                e["date"] <= (today + dt.timedelta(days=2)).isoformat()]
        for e in sorted(soon, key=lambda e: (e["date"], e["time"] or "")):
            it = QListWidgetItem(line(e))
            it.setData(Qt.UserRole, e["case_id"])
            it.setForeground(QColor(M.T["danger"]))
            self.l_over.addItem(it)
        if not soon:
            self._empty(self.l_over, "Горящих сроков нет — всё под контролем")
        self.l_rem.clear()
        rems = sorted(db().reminders(), key=event_dt)[:8]
        for e in rems:
            w = event_dt(e)
            it = QListWidgetItem(f"⏰  {e['title']}\n     {w:%d.%m %H:%M} · {e['case_title'] or 'без дела'}")
            it.setData(Qt.UserRole, e["case_id"])
            self.l_rem.addItem(it)
        if not rems:
            self._empty(self.l_rem, "Напоминаний нет. Их можно поставить в обзоре дела.")
        self.l_recent.clear()
        nexts = {}
        for e in sorted(active, key=lambda e: (e["date"], e["time"] or "")):
            if e["case_id"] and e["date"] >= today.isoformat():
                nexts.setdefault(e["case_id"], e)
        for c in cases[:8]:
            sub = " · ".join(x for x in (c["number"], c["stage"]) if x)
            nx = nexts.get(c["id"])
            if nx:
                sub += ("\n     " if sub else "") + f"ближайшее: {CS.ru(nx['date'])} {nx['time'] or ''} {nx['kind'].lower()}".rstrip()
            it = QListWidgetItem(f"📁  {c['title']}" + (f"\n     {sub}" if sub else ""))
            it.setData(Qt.UserRole, c["id"])
            self.l_recent.addItem(it)
        if not cases:
            self._empty(self.l_recent, "Дел пока нет — «➕ Новое дело» ниже")
